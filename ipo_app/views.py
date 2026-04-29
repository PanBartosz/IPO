import csv
import io
import math
import mimetypes
import os
import re
import secrets
import json
import hashlib
import unicodedata
from collections import defaultdict
from datetime import datetime, timedelta
from html import unescape
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from typing import Iterable, Callable, Any, Optional

from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    get_flashed_messages,
    make_response,
    redirect,
    render_template,
    render_template_string,
    request,
    send_file,
    send_from_directory,
    session,
    url_for,
)
from sqlalchemy import or_, cast, String, func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload
from sqlalchemy.engine.url import make_url
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from PyPDF2 import PdfReader, PdfWriter
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from weasyprint import HTML
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from sendgrid import SendGridAPIClient
from sendgrid.helpers.mail import Mail

from .extensions import db, limiter
from .forms import (
    AssignmentBatchForm,
    AdminLoginForm,
    ReadingGroupForm,
    ReadingGroupMemberForm,
    StudentLoginForm,
    StudentRegistrationForm,
    StudentUploadForm,
    TeacherRegistrationForm,
    DelegationInviteForm,
    DelegationLeaderInviteForm,
    AdminDelegationInvitationForm,
    AdminInvitationForm,
    AdminInvitationBatchUploadForm,
    AdminInvitationBatchExecuteForm,
    AdminInvitationBatchDiscardForm,
    DashboardSettingsForm,
    TeacherMessageForm,
    TeacherMessageTemplateForm,
    QuickTeacherMessageForm,
    TeacherBoardLoginForm,
    TeacherPortalLoginForm,
    TeacherScoreForm,
    SurveyForm,
    SurveyQuestionForm,
    SurveyPublishForm,
    SurveyRecipientUpdateForm,
    SurveyResponseForm,
    PasswordResetRequestForm,
    PasswordResetConfirmForm,
    CommunicationComposeForm,
    CommunicationSendForm,
    CommunicationRetryForm,
    StaticPageForm,
    Stage3PacketForm,
    GradesExportForm,
    LANGUAGE_CHOICES,
    INVITATION_BATCH_TYPE_CHOICES,
)
from .models import (
    ActivityLog,
    AssignmentBatch,
    EssayAssignment,
    EssayScore,
    EssayAssignmentExtraReader,
    GroupMemberRole,
    PlannerAlgorithm,
    PlannerStatus,
    ReadingGroup,
    ReadingGroupMember,
    Stage,
    Student,
    Teacher,
    TeacherLanguage,
    AppSetting,
    TeacherMessage,
    TeacherMessagePriority,
    TeacherMessageStatus,
    TeacherMessageTemplate,
    Survey,
    SurveyAssignment,
    SurveyAssignmentStatus,
    SurveyAudience,
    SurveyQuestion,
    SurveyQuestionType,
    SurveyResponse,
    SurveyStatus,
    DelegationInvite,
    DelegationInviteRole,
    InvitationBatchDraft,
    PasswordResetAudience,
    PasswordResetToken,
    CommunicationCampaign,
    CommunicationStatus,
    CommunicationAudience,
    CommunicationRecipient,
    CommunicationRecipientRole,
    CommunicationRecipientStatus,
    EssayUpload,
    StaticPage,
    StaticPageAttachment,
)
from .scoring import (
    SCORE_VALIDATION_MESSAGE,
    format_score_value,
    parse_score_value,
    score_choices,
)

web_bp = Blueprint("web", __name__)
teacher_portal_bp = Blueprint("teacher_portal", __name__, url_prefix="/teacher")


ADMIN_MODULES = [
    ("Dashboard", "web.admin_dashboard"),
    ("Planner", "web.admin_planner"),
    ("Students", "web.admin_students"),
    ("Teachers", "web.admin_teachers"),
    ("Invitation", "web.admin_invitations"),
    ("Stage 1", "web.admin_uploads"),
    ("Stage 2", "web.admin_stage2_overview"),
    ("Messages", "web.admin_messages"),
    ("Communication", "web.admin_communications"),
    ("Static pages", "web.admin_static_pages"),
    ("Grades export", "web.admin_grades_export"),
    ("Surveys", "web.admin_surveys"),
    ("Logs", "web.admin_logs"),
    ("Exports", "web.admin_exports"),
    ("Help", "web.admin_help"),
    ("Settings", "web.admin_settings"),
]


LOG_PAGE_SIZE = 100
RESET_TOKEN_EXPIRY_HOURS = 24

LOGIN_RATE_LIMIT_PER_USERNAME = "10 per minute"
LOGIN_RATE_LIMIT_PER_IP = "120 per minute"
TEACHER_BOARD_LOGIN_RATE_LIMIT_PER_IP = "10 per minute"

INVITATION_BATCH_MAX_FILE_SIZE = 2 * 1024 * 1024
INVITATION_BATCH_MAX_ROWS = 1000
INVITATION_BATCH_MAX_CELL_LENGTH = 500

INVITATION_BATCH_TYPE_VALUES = {value for value, _ in INVITATION_BATCH_TYPE_CHOICES}
INVITATION_BATCH_REQUIRED_COLUMNS = {
    "delegation_leader": ["first_name", "last_name", "email", "country"],
    "delegation": ["first_name", "last_name", "email", "role", "delegation_leader_email"],
    "normal": ["first_name", "last_name", "email", "role", "country"],
}
INVITATION_BATCH_TEMPLATE_SAMPLE_ROWS = {
    "delegation_leader": {
        "first_name": "Alex",
        "last_name": "Leader",
        "email": "alex.leader@example.org",
        "country": "Poland",
    },
    "delegation": {
        "first_name": "Taylor",
        "last_name": "Teacher",
        "email": "taylor.teacher@example.org",
        "role": "teacher",
        "delegation_leader_email": "alex.leader@example.org",
    },
    "normal": {
        "first_name": "Jordan",
        "last_name": "Student",
        "email": "jordan.student@example.org",
        "role": "student",
        "country": "Poland",
    },
}
INVITATION_BATCH_EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


SETTING_EMAIL_AUTOMATION_ENABLED = "email_auto_enabled"
SETTING_EMAIL_STUDENT_SUBJECT = "email_student_subject"
SETTING_EMAIL_STUDENT_BODY = "email_student_body"
SETTING_EMAIL_TEACHER_SUBJECT = "email_teacher_subject"
SETTING_EMAIL_TEACHER_BODY = "email_teacher_body"
SETTING_EMAIL_MESSAGE_SUBJECT = "email_message_subject"
SETTING_EMAIL_MESSAGE_BODY = "email_message_body"
SETTING_EMAIL_STUDENT_RESET_LINK_SUBJECT = "email_student_reset_link_subject"
SETTING_EMAIL_STUDENT_RESET_LINK_BODY = "email_student_reset_link_body"
SETTING_EMAIL_STUDENT_RESET_SUBJECT = "email_student_reset_subject"
SETTING_EMAIL_STUDENT_RESET_BODY = "email_student_reset_body"
SETTING_EMAIL_TEACHER_RESET_LINK_SUBJECT = "email_teacher_reset_link_subject"
SETTING_EMAIL_TEACHER_RESET_LINK_BODY = "email_teacher_reset_link_body"
SETTING_EMAIL_TEACHER_RESET_SUBJECT = "email_teacher_reset_subject"
SETTING_EMAIL_TEACHER_RESET_BODY = "email_teacher_reset_body"
SETTING_EMAIL_SURVEY_PUBLISH_SUBJECT = "email_survey_publish_subject"
SETTING_EMAIL_SURVEY_PUBLISH_BODY = "email_survey_publish_body"
SETTING_EMAIL_SURVEY_REMINDER_SUBJECT = "email_survey_reminder_subject"
SETTING_EMAIL_SURVEY_REMINDER_BODY = "email_survey_reminder_body"
SETTING_TEACHER_PORTAL_ENABLED = "teacher_portal_enabled"
SETTING_REGISTRATION_INVITE_ONLY = "registration_invite_only"
SETTING_LOGIN_FOLDED_ALIAS_ENABLED = "login_folded_alias_enabled"
SETTING_LOGIN_EMAIL_ENABLED = "login_email_enabled"
SETTING_EMAIL_STUDENT_INVITE_SUBJECT = "email_student_invite_subject"
SETTING_EMAIL_STUDENT_INVITE_BODY = "email_student_invite_body"
SETTING_EMAIL_TEACHER_INVITE_SUBJECT = "email_teacher_invite_subject"
SETTING_EMAIL_TEACHER_INVITE_BODY = "email_teacher_invite_body"

DEFAULT_EMAIL_SUBJECTS = {
    SETTING_EMAIL_STUDENT_SUBJECT: "IPO registration confirmed for {{ first_name }} {{ last_name }}",
    SETTING_EMAIL_TEACHER_SUBJECT: "IPO teacher registration confirmed for {{ first_name }} {{ last_name }}",
    SETTING_EMAIL_STUDENT_INVITE_SUBJECT: "Complete your IPO student registration",
    SETTING_EMAIL_TEACHER_INVITE_SUBJECT: "Complete your IPO teacher registration",
    SETTING_EMAIL_MESSAGE_SUBJECT: "New IPO teacher board message: {{ title }}",
    SETTING_EMAIL_STUDENT_RESET_LINK_SUBJECT: "Reset your IPO portal password",
    SETTING_EMAIL_STUDENT_RESET_SUBJECT: "IPO portal password reset",
    SETTING_EMAIL_TEACHER_RESET_LINK_SUBJECT: "Reset your IPO teacher portal password",
    SETTING_EMAIL_TEACHER_RESET_SUBJECT: "IPO teacher portal password reset",
    SETTING_EMAIL_SURVEY_PUBLISH_SUBJECT: "New survey: {{ survey_title }}",
    SETTING_EMAIL_SURVEY_REMINDER_SUBJECT: "Reminder: please complete {{ survey_title }}",
}

DEFAULT_EMAIL_BODIES = {
    SETTING_EMAIL_STUDENT_BODY: """
<p>Hi {{ first_name }} {{ last_name }},</p>
<p>Thank you for registering for the International Philosophy Olympiad. Below is a summary of the details we received:</p>
<ul>
  <li><strong>Country:</strong> {{ country }}</li>
  <li><strong>School:</strong> {{ school }}</li>
  <li><strong>Essay language:</strong> {{ essay_language }}</li>
  <li><strong>Keyboard preference:</strong> {{ keyboard_preference }}</li>
  <li><strong>Meal preference:</strong> {{ meal_preference }}</li>
  {% if special_needs %}<li><strong>Special needs:</strong> {{ special_needs }}</li>{% endif %}
</ul>
<p>Your essay upload portal credentials:</p>
<ul>
  <li><strong>Username:</strong> {{ portal_username }}</li>
  <li><strong>Password:</strong> {{ portal_password }}</li>
</ul>
<p>You can access the portal here: <a href="{{ portal_link }}">{{ portal_link }}</a><br>
Login directly: <a href="{{ portal_login_link }}">{{ portal_login_link }}</a></p>
<p>If any information changes, please contact the organizing team as soon as possible.</p>
<p>We are looking forward to seeing you at the Olympiad!</p>
""",
    SETTING_EMAIL_TEACHER_BODY: """
<p>Hi {{ first_name }} {{ last_name }},</p>
<p>Your teacher registration for the International Philosophy Olympiad has been received. Here is a summary of your submission:</p>
<ul>
  <li><strong>Country:</strong> {{ country }}</li>
  <li><strong>School:</strong> {{ school }}</li>
  <li><strong>Languages:</strong> {{ languages }}</li>
  <li><strong>Meal preference:</strong> {{ meal_preference }}</li>
  <li><strong>T-shirt size:</strong> {{ tshirt_size }}</li>
  {% if special_needs %}<li><strong>Special needs:</strong> {{ special_needs }}</li>{% endif %}
  <li><strong>Delegation leader:</strong> {{ 'Yes' if delegation_leader else 'No' }}</li>
  <li><strong>Prior IPO assessor:</strong> {% if prior_ipo_assessor_level %}Yes (Level {{ prior_ipo_assessor_level }}){% else %}No{% endif %}</li>
</ul>
<p>Your teacher portal credentials:</p>
<ul>
  <li><strong>Username:</strong> {{ portal_username }}</li>
  <li><strong>Password:</strong> {{ portal_password }}</li>
</ul>
<p>Access the portal: <a href="{{ portal_link }}">{{ portal_link }}</a><br>
Login page: <a href="{{ portal_login_link }}">{{ portal_login_link }}</a></p>
<p>We appreciate your support and participation. Please reach out if anything needs to be updated.</p>
""",
    SETTING_EMAIL_STUDENT_INVITE_BODY: """
<p>Hi {{ first_name }} {{ last_name }},</p>
<p>{{ inviter_name }} has started your IPO student registration for the {{ country }} delegation.</p>
<p>Please complete your registration using this link:</p>
<p><a href="{{ invite_link }}">{{ invite_link }}</a></p>
<p>If you did not expect this invitation, please ignore this email.</p>
""",
    SETTING_EMAIL_TEACHER_INVITE_BODY: """
<p>Hi {{ first_name }} {{ last_name }},</p>
<p>{{ inviter_name }} has started your IPO teacher registration for the {{ country }} delegation.</p>
<p>Please complete your registration using this link:</p>
<p><a href="{{ invite_link }}">{{ invite_link }}</a></p>
<p>If you did not expect this invitation, please ignore this email.</p>
""",
    SETTING_EMAIL_MESSAGE_BODY: """
<p>Hello {{ teacher_first_name }} {{ teacher_last_name }},</p>
<p>A new message has been posted on the IPO teacher board.</p>
<p><strong>{{ title }}</strong><br>{{ body|replace('\\n', '<br>') }}</p>
<p>Priority: {{ priority|title }}</p>
<p>You can also read this message online: <a href="{{ board_url }}">Teacher Board</a>.</p>
""",
    SETTING_EMAIL_STUDENT_RESET_LINK_BODY: """
<p>Hello {{ first_name }} {{ last_name }},</p>
<p>We received a request to reset your IPO essay upload portal password. Please use the link below within the next 24 hours to continue:</p>
<p><a href="{{ reset_link }}">Reset your password</a></p>
<p>If you did not request this change, you can safely ignore this email — your password will remain the same.</p>
<p>Upload login page: <a href="{{ portal_login_link }}">{{ portal_login_link }}</a></p>
<p>This link expires automatically after 24 hours for security.</p>
""",
    SETTING_EMAIL_STUDENT_RESET_BODY: """
<p>Hello {{ first_name }} {{ last_name }},</p>
<p>Your IPO essay upload portal password has been reset.</p>
<ul>
  <li><strong>Username:</strong> {{ portal_username }}</li>
  <li><strong>New password:</strong> {{ portal_password }}</li>
</ul>
<p>Access the portal: <a href="{{ portal_link }}">{{ portal_link }}</a><br>
Login page: <a href="{{ portal_login_link }}">{{ portal_login_link }}</a></p>
""",
    SETTING_EMAIL_TEACHER_RESET_LINK_BODY: """
<p>Hello {{ first_name }} {{ last_name }},</p>
<p>We received a request to reset your IPO teacher portal password. Use the link below within the next 24 hours to continue:</p>
<p><a href="{{ reset_link }}">Reset your password</a></p>
<p>If you did not request this, simply ignore this email and your password will stay the same.</p>
<p>Teacher portal login: <a href="{{ portal_login_link }}">{{ portal_login_link }}</a></p>
<p>This link expires automatically after 24 hours.</p>
""",
    SETTING_EMAIL_TEACHER_RESET_BODY: """
<p>Hello {{ first_name }} {{ last_name }},</p>
<p>Your IPO teacher portal password has been reset.</p>
<ul>
  <li><strong>Username:</strong> {{ portal_username }}</li>
  <li><strong>New password:</strong> {{ portal_password }}</li>
</ul>
<p>Visit the portal: <a href="{{ portal_link }}">{{ portal_link }}</a><br>
Login page: <a href="{{ portal_login_link }}">{{ portal_login_link }}</a></p>
""",
    SETTING_EMAIL_SURVEY_PUBLISH_BODY: """
<p>Hello {{ recipient_name }},</p>
<p>A new survey has been published for you: <strong>{{ survey_title }}</strong>.</p>
{% if survey_description %}<p>{{ survey_description }}</p>{% endif %}
{% if survey_deadline %}<p><strong>Deadline:</strong> {{ survey_deadline }}</p>{% endif %}
<p><a href="{{ survey_link }}">Open this survey</a> or visit your portal: <a href="{{ portal_link }}">{{ portal_link }}</a>.</p>
""",
    SETTING_EMAIL_SURVEY_REMINDER_BODY: """
<p>Hello {{ recipient_name }},</p>
<p>This is a friendly reminder to complete the survey <strong>{{ survey_title }}</strong>.</p>
{% if survey_description %}<p>{{ survey_description }}</p>{% endif %}
{% if survey_deadline %}<p><strong>Deadline:</strong> {{ survey_deadline }}</p>{% endif %}
<p><a href="{{ survey_link }}">Resume the survey</a> or visit your portal: <a href="{{ portal_link }}">{{ portal_link }}</a>.</p>
""",
}

STAGE_LABELS = {
    Stage.STAGE_1: "Stage 1",
    Stage.STAGE_2: "Stage 2",
}

PORTAL_STATUS_LABELS = {
    "finalized": "Teacher finalized",
    "in_progress": "Draft saved",
    "awaiting_finalization": "Awaiting student finalization",
    "waiting_upload": "Awaiting essay upload",
    "pending": "Not started",
    "admin_entered": "Admin entered",
    "admin_override": "Admin override",
    "locked": "Locked by admin",
}

SETTING_UPLOAD_DEADLINE = "dashboard_upload_deadline"
SETTING_STAGE2_DEADLINE = "dashboard_stage2_deadline"
SETTING_STAGE1_SLA = "dashboard_stage1_sla_hours"
SETTING_STAGE2_SLA = "dashboard_stage2_sla_hours"
SETTING_STAGE1_LOCK = "grading_stage1_locked"
SETTING_STAGE2_LOCK = "grading_stage2_locked"
SETTING_UPLOAD_MODULE_ENABLED = "essay_upload_enabled"
SETTING_UPLOAD_LOCK = "upload_locked"
SETTING_EVENT_END = "event_end"
SETTING_EVENT_END_OVERRIDE = "event_end_override"
SETTING_PACKET_PRINT_MODE = "packet_print_mode"
SETTING_ESSAY_WORD_LIMIT = "essay_word_limit_words"
SETTING_ESSAY_WORD_LIMIT_MODE = "essay_word_limit_mode"
SETTING_STAGE1_TRIM_EXTREME_SCORES = "stage1_trim_extreme_scores"
SETTING_ETHICS_CONSENT_HTML = "ethics_consent_html"
SETTING_GDPR_NOTICE_HTML = "gdpr_notice_html"
SETTING_PHOTO_CONSENT_HTML = "photo_consent_html"
SETTING_TEACHER_COMPETENCE_HTML = "teacher_competence_html"
SETTING_ACCOMMODATION_INFO_HTML = "accommodation_info_html"

DEFAULT_ESSAY_WORD_LIMIT = 1500
DEFAULT_ESSAY_WORD_LIMIT_MODE = "soft"
DEFAULT_ETHICS_CONSENT_HTML = (
    '<p>Please review the <a href="https://www.philosophy-olympiad.org/?page_id=1710" target="_blank" rel="noopener">IPO Code of Ethics</a> before accepting this statement.</p>'
)
DEFAULT_GDPR_NOTICE_HTML = (
    "<p>Please review the data processing information before confirming your registration.</p>"
)
DEFAULT_PHOTO_CONSENT_HTML = (
    "<p>We may publish photos from the event for communication purposes. You can consent or decline below.</p>"
)
DEFAULT_TEACHER_COMPETENCE_HTML = (
    "<p>By checking the box below, you confirm that you are qualified and competent to assess philosophical essays.</p>"
)
DEFAULT_ACCOMMODATION_INFO_HTML = (
    "<p><strong>Important information</strong></p>"
    "<p>Participants who require individual consideration regarding accommodation arrangements may submit a request to the National Organizing Committee. Requests will be reviewed individually.</p>"
    '<p>Please contact:<br><strong>Magdalena Gawin</strong>, the Head of NOC, <a href="mailto:m.gawin@uw.edu.pl">m.gawin@uw.edu.pl</a></p>'
)

UPLOAD_ALLOWED_EXTENSIONS = {"pdf"}
MAX_UPLOAD_SIZE = 10 * 1024 * 1024
STATIC_PAGE_SLUG_RE = re.compile(r"[a-z0-9]+(?:-[a-z0-9]+)*")


def _essay_word_limit() -> int:
    raw = _get_setting(SETTING_ESSAY_WORD_LIMIT)
    if raw is None or str(raw).strip() == "":
        return DEFAULT_ESSAY_WORD_LIMIT
    try:
        limit = int(str(raw).strip())
    except ValueError:
        return DEFAULT_ESSAY_WORD_LIMIT
    return max(1, limit)


def _essay_word_limit_mode() -> str:
    raw = (_get_setting(SETTING_ESSAY_WORD_LIMIT_MODE) or "").strip().lower()
    if raw in {"off", "soft", "hard"}:
        return raw
    return DEFAULT_ESSAY_WORD_LIMIT_MODE


def _stage1_trim_extreme_scores_enabled() -> bool:
    return _parse_bool(_get_setting(SETTING_STAGE1_TRIM_EXTREME_SCORES, "false"))


def _ethics_consent_html() -> str:
    return _get_setting(SETTING_ETHICS_CONSENT_HTML) or DEFAULT_ETHICS_CONSENT_HTML


def _gdpr_notice_html() -> str:
    return _get_setting(SETTING_GDPR_NOTICE_HTML) or DEFAULT_GDPR_NOTICE_HTML


def _photo_consent_html() -> str:
    return _get_setting(SETTING_PHOTO_CONSENT_HTML) or DEFAULT_PHOTO_CONSENT_HTML


def _assessor_competence_html() -> str:
    return _get_setting(SETTING_TEACHER_COMPETENCE_HTML) or DEFAULT_TEACHER_COMPETENCE_HTML


def _accommodation_info_html() -> str:
    return _get_setting(SETTING_ACCOMMODATION_INFO_HTML) or DEFAULT_ACCOMMODATION_INFO_HTML


def _count_pdf_words(pdf_path: Path) -> int | None:
    try:
        from PyPDF2 import PdfReader
    except Exception:
        return None
    try:
        reader = PdfReader(str(pdf_path))
        chunks: list[str] = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text:
                chunks.append(text)
        full_text = " ".join(chunks).strip()
        if not full_text:
            return 0
        return len(full_text.split())
    except Exception:
        return None


@web_bp.before_app_request
def protect_admin_routes():
    if request.blueprint != "web":
        return
    path = request.path or ""
    if not path.startswith("/admin"):
        return
    allowed_endpoints = {
        "web.admin_login",
        "web.admin_logout",
    }
    if request.endpoint in allowed_endpoints:
        return
    if _current_admin():
        return
    return redirect(url_for("web.admin_login", next=request.full_path))


@web_bp.app_context_processor
def inject_globals():
    return {
        "current_year": datetime.utcnow().year,
        "current_admin": _current_admin(),
        "ADMIN_MODULES": ADMIN_MODULES,
        "admin_modules": ADMIN_MODULES,
        "teacher_portal_enabled": _teacher_portal_enabled(),
        "essay_upload_enabled": _is_upload_module_enabled(),
        "registration_invite_only": _registration_invite_required(),
    }


def _rate_limit_ip() -> str:
    return request.remote_addr or "unknown"


def _rate_limit_username() -> str:
    identifier = (request.form.get("username") or "").strip()
    normalized_identifier = identifier.lower()
    if identifier and _login_email_enabled() and _login_identifier_looks_like_email(identifier):
        normalized_email = _normalize_login_email(identifier)
        if normalized_email:
            return normalized_email
    if identifier and _login_folded_alias_enabled():
        folded = _fold_login_username(identifier)
        if folded:
            return folded
    return normalized_identifier or _rate_limit_ip()


def _ensure_language_entries(teacher: Teacher) -> None:
    if teacher.language_entries:
        return
    if not teacher.languages:
        return
    languages = [lang.strip() for lang in teacher.languages.split(",") if lang.strip()]
    existing = {entry.language for entry in teacher.language_entries}
    for lang in languages:
        if lang not in existing:
            teacher.language_entries.append(TeacherLanguage(language=lang))
    db.session.flush()


def _normalize_language(value: str | None) -> str | None:
    if not value:
        return None
    return value.strip().lower()


def _teacher_languages(teacher: Teacher) -> set[str]:
    _ensure_language_entries(teacher)
    languages: set[str] = set()
    for entry in teacher.language_entries:
        if entry.language:
            languages.add(_normalize_language(entry.language))
    if not languages and teacher.languages:
        for lang in teacher.languages.split(","):
            normalized = _normalize_language(lang)
            if normalized:
                languages.add(normalized)
    return languages


def _teacher_board_authorized() -> bool:
    return bool(session.get("teacher_board") or session.get("teacher_portal_id"))


def _prepare_message_form(form: TeacherMessageForm, selected_teacher_id: int | None = None) -> None:
    teachers = Teacher.query.order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()
    choices = [(0, "All teachers")]
    choices.extend((teacher.id, f"{teacher.first_name} {teacher.last_name}") for teacher in teachers)
    form.teacher_id.choices = choices
    if form.teacher_id.data is None:
        form.teacher_id.data = 0
    if selected_teacher_id is not None:
        form.teacher_id.data = selected_teacher_id if selected_teacher_id in {choice[0] for choice in choices} else 0


def _quick_message_context() -> dict:
    form = QuickTeacherMessageForm()
    if not form.priority.data:
        form.priority.data = "normal"
    templates = TeacherMessageTemplate.query.order_by(TeacherMessageTemplate.name.asc()).all()
    templates_data = [
        {
            "id": template.id,
            "name": template.name,
            "title": template.title,
            "body": template.body,
            "category": template.category or "",
        }
        for template in templates
    ]
    return {
        "quick_message_form": form,
        "quick_message_templates": templates,
        "quick_message_templates_data": templates_data,
    }


def _parse_datetime(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return None


def _log_filters_from_request() -> dict:
    return {
        "actor": request.args.get("actor", "").strip() or None,
        "action": request.args.get("action", "").strip() or None,
        "object_type": request.args.get("object_type", "").strip() or None,
        "q": request.args.get("q", "").strip() or None,
        "start": _parse_datetime(request.args.get("start") or None),
        "end": _parse_datetime(request.args.get("end") or None),
        "page": request.args.get("page", default=1, type=int),
    }


def _filtered_activity_logs(filters: dict):
    query = ActivityLog.query
    if filters.get("actor"):
        query = query.filter(ActivityLog.actor.ilike(f"%{filters['actor']}%"))
    if filters.get("action"):
        query = query.filter(ActivityLog.action == filters["action"])
    if filters.get("object_type"):
        query = query.filter(ActivityLog.object_type == filters["object_type"])
    if filters.get("q"):
        search = f"%{filters['q']}%"
        query = query.filter(
            or_(
                ActivityLog.description.ilike(search),
                cast(ActivityLog.details, String).ilike(search),
            )
        )
    if filters.get("start"):
        query = query.filter(ActivityLog.created_at >= filters["start"])
    if filters.get("end"):
        query = query.filter(ActivityLog.created_at <= filters["end"])
    return query


def _communication_country_choices() -> list[tuple[str, str]]:
    countries: set[str] = set()
    teacher_rows = (
        db.session.query(Teacher.country)
        .filter(Teacher.country.isnot(None))
        .distinct()
    )
    for (country,) in teacher_rows:
        if country:
            countries.add(country.strip())
    student_rows = (
        db.session.query(Student.country)
        .filter(Student.country.isnot(None))
        .distinct()
    )
    for (country,) in student_rows:
        if country:
            countries.add(country.strip())
    return sorted((country, country) for country in countries if country)


def _build_campaign_recipients(
    audience: CommunicationAudience,
    countries: list[str],
    include_absent: bool,
) -> list[CommunicationRecipient]:
    selected_countries = {country.strip() for country in countries or [] if country and country.strip()}
    filter_by_country = bool(selected_countries)
    recipients: list[CommunicationRecipient] = []
    seen_emails: set[str] = set()

    def add_recipient(recipient: CommunicationRecipient) -> None:
        email = (recipient.email or "").strip()
        if not email:
            return
        normalized_email = email.lower()
        if normalized_email in seen_emails:
            return
        seen_emails.add(normalized_email)
        recipient.email = email
        recipients.append(recipient)

    include_teachers = audience in {CommunicationAudience.TEACHER, CommunicationAudience.BOTH}
    include_students = audience in {CommunicationAudience.STUDENT, CommunicationAudience.BOTH}

    if include_teachers:
        teacher_query = Teacher.query
        if filter_by_country:
            teacher_query = teacher_query.filter(Teacher.country.in_(selected_countries))
        if not include_absent:
            teacher_query = teacher_query.filter(Teacher.is_absent.is_(False))
        teacher_query = teacher_query.order_by(Teacher.country.asc(), Teacher.last_name.asc())
        for teacher in teacher_query:
            recipient = CommunicationRecipient(
                role=CommunicationRecipientRole.TEACHER,
                teacher_id=teacher.id,
                email=teacher.email,
                first_name=teacher.first_name,
                last_name=teacher.last_name,
                country=teacher.country,
                status=CommunicationRecipientStatus.PENDING,
            )
            add_recipient(recipient)

    if include_students:
        student_query = Student.query
        if filter_by_country:
            student_query = student_query.filter(Student.country.in_(selected_countries))
        if not include_absent:
            student_query = student_query.filter(Student.is_absent.is_(False))
        student_query = student_query.order_by(Student.country.asc(), Student.last_name.asc())
        for student in student_query:
            recipient = CommunicationRecipient(
                role=CommunicationRecipientRole.STUDENT,
                student_id=student.id,
                email=student.email,
                first_name=student.first_name,
                last_name=student.last_name,
                country=student.country,
                status=CommunicationRecipientStatus.PENDING,
            )
            add_recipient(recipient)

    return recipients


def _campaign_recipient_stats(campaign_ids: Iterable[int]) -> dict[int, dict[str, int]]:
    ids = [int(campaign_id) for campaign_id in campaign_ids]
    if not ids:
        return {}
    rows = (
        db.session.query(
            CommunicationRecipient.campaign_id,
            CommunicationRecipient.status,
            func.count(CommunicationRecipient.id),
        )
        .filter(CommunicationRecipient.campaign_id.in_(ids))
        .group_by(CommunicationRecipient.campaign_id, CommunicationRecipient.status)
        .all()
    )
    stats: dict[int, dict[str, int]] = {campaign_id: {} for campaign_id in ids}
    for campaign_id, status, count in rows:
        stats.setdefault(campaign_id, {})
        stats[campaign_id][status.value if hasattr(status, "value") else status] = count
    return stats


def _campaign_role_counts(campaign: CommunicationCampaign) -> dict[str, int]:
    counts = {
        CommunicationRecipientRole.TEACHER.value: 0,
        CommunicationRecipientRole.STUDENT.value: 0,
    }
    for recipient in campaign.recipients:
        counts[recipient.role.value] = counts.get(recipient.role.value, 0) + 1
    return counts


def _campaign_sample_recipients(
    campaign: CommunicationCampaign,
    limit: int = 10,
) -> list[CommunicationRecipient]:
    return campaign.recipients[:limit]


def _deliver_campaign(
    campaign: CommunicationCampaign,
    *,
    only_failed: bool = False,
) -> dict[str, int]:
    if not _sendgrid_configured():
        return {"sent": 0, "failed": 0, "skipped": len(campaign.recipients)}

    target_status = (
        CommunicationRecipientStatus.FAILED
        if only_failed
        else CommunicationRecipientStatus.PENDING
    )
    recipients = [
        recipient
        for recipient in campaign.recipients
        if recipient.status == target_status
    ]
    if not recipients:
        return {"sent": 0, "failed": 0, "skipped": 0}

    sent = 0
    failed = 0
    for recipient in recipients:
        success, message_id, error = _send_email(
            recipient.email,
            campaign.subject,
            campaign.html_body,
            plain_text_content=campaign.text_body or None,
        )
        if success:
            recipient.status = CommunicationRecipientStatus.SENT
            recipient.sent_at = datetime.utcnow()
            recipient.message_id = message_id
            recipient.error = None
            sent += 1
        else:
            recipient.status = CommunicationRecipientStatus.FAILED
            recipient.sent_at = None
            recipient.message_id = None
            recipient.error = error or "unknown-error"
            failed += 1

    campaign.sent_count = sum(
        1
        for recipient in campaign.recipients
        if recipient.status == CommunicationRecipientStatus.SENT
    )
    campaign.failed_count = sum(
        1
        for recipient in campaign.recipients
        if recipient.status == CommunicationRecipientStatus.FAILED
    )
    campaign.total_recipients = len(campaign.recipients)

    if campaign.sent_count and not campaign.sent_at:
        campaign.sent_at = datetime.utcnow()

    if campaign.failed_count and campaign.sent_count:
        campaign.status = CommunicationStatus.FAILED
    elif campaign.failed_count:
        campaign.status = CommunicationStatus.FAILED
    else:
        campaign.status = CommunicationStatus.SENT

    return {"sent": sent, "failed": failed, "skipped": 0}


def _teacher_can_grade(student: Student, teacher: Teacher, language: str | None) -> bool:
    if not teacher or teacher.is_absent or not teacher.declared_assessor_competence:
        return False
    if student and teacher.country == student.country:
        return False
    normalized_language = _normalize_language(language)
    if normalized_language is None:
        return True
    teacher_langs = _teacher_languages(teacher)
    return normalized_language in teacher_langs


def _assessment_eligible_teachers_query():
    return Teacher.query.filter(
        Teacher.is_absent.is_(False),
        Teacher.declared_assessor_competence.is_(True),
    )


def _published_batch_guard_response(batch: AssignmentBatch):
    if batch.status != PlannerStatus.PUBLISHED:
        return None
    flash("Cannot modify a published batch. Unpublish it first.", "error")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


def _empty_load_entry() -> SimpleNamespace:
    return SimpleNamespace(
        stage1_expected=0,
        stage1_completed=0,
        stage2_assigned=0,
        stage2_completed=0,
        last_stage2_activity=None,
    )


def _teacher_load_snapshot(assignments: list[EssayAssignment]) -> dict[int, SimpleNamespace]:
    stats: dict[int, SimpleNamespace] = {}
    for assignment in assignments:
        stage1_scores = {
            score.teacher_id: score
            for score in assignment.scores
            if score.stage == Stage.STAGE_1 and score.score is not None
        }
        stage2_scores = {
            score.teacher_id: score
            for score in assignment.scores
            if score.stage == Stage.STAGE_2 and score.score is not None
        }

        for assessor in _stage_assessors(assignment, Stage.STAGE_1, include_group=True):
            teacher = assessor.teacher
            if not teacher:
                continue
            entry = stats.setdefault(teacher.id, _empty_load_entry())
            entry.stage1_expected += 1
            if stage1_scores.get(teacher.id):
                entry.stage1_completed += 1

        for assessor in _stage_assessors(assignment, Stage.STAGE_2, include_group=False):
            teacher = assessor.teacher
            if not teacher:
                continue
            entry = stats.setdefault(teacher.id, _empty_load_entry())
            entry.stage2_assigned += 1
            score = stage2_scores.get(teacher.id)
            if score:
                entry.stage2_completed += 1
                timestamp = score.submitted_at or score.updated_at
                if timestamp and (entry.last_stage2_activity is None or timestamp > entry.last_stage2_activity):
                    entry.last_stage2_activity = timestamp

    return stats


def _random_username(first_name: str, last_name: str, student_id: int | None = None) -> str:
    base = f"{_slugify(first_name[0])}{_slugify(last_name)}"
    suffix = f"{student_id or ''}{secrets.randbelow(1000)}"
    return f"{base}{suffix}"[:30]


def _generate_unique_username(first_name: str | None, last_name: str | None, max_attempts: int = 20) -> str:
    safe_first = (first_name or "student").strip() or "student"
    safe_last = (last_name or "candidate").strip() or "candidate"

    for attempt in range(max_attempts):
        candidate = _random_username(safe_first, safe_last)
        if not Student.query.filter_by(upload_username=candidate).first():
            return candidate

    fallback = f"{_slugify(safe_first[0])}{secrets.token_hex(4)}"
    return fallback[:30]


def _generate_teacher_portal_username(teacher: Teacher, max_attempts: int = 20) -> str:
    safe_first = (teacher.first_name or "teacher").strip() or "teacher"
    safe_last = (teacher.last_name or "mentor").strip() or "mentor"
    base = f"tch{_slugify(safe_last)[:8]}"
    for _ in range(max_attempts):
        suffix = f"{teacher.id or ''}{secrets.randbelow(1000)}"
        candidate = f"{base}{suffix}"[:30]
        exists = Teacher.query.filter(Teacher.portal_username == candidate).first()
        if not exists:
            return candidate
    return f"tch{secrets.token_hex(4)}"[:30]


def _generate_unique_teacher_token(max_attempts: int = 20) -> str:
    for _ in range(max_attempts):
        token = secrets.token_urlsafe(16)
        if not Teacher.query.filter(Teacher.portal_token == token).first():
            return token
    return secrets.token_hex(16)


def _generate_unique_invite_token(max_attempts: int = 20) -> str:
    for _ in range(max_attempts):
        token = secrets.token_urlsafe(24)
        if not DelegationInvite.query.filter(DelegationInvite.token == token).first():
            return token
    return secrets.token_hex(24)


def _ensure_teacher_portal_credentials(teacher: Teacher, *, reset_password: bool = False) -> tuple[str, str | None]:
    username_generated = False
    if not teacher.portal_username:
        teacher.portal_username = _generate_teacher_portal_username(teacher)
        username_generated = True

    plain_password: str | None = None
    if reset_password or not teacher.portal_password_hash:
        plain_password = secrets.token_urlsafe(8)
        teacher.portal_password_hash = generate_password_hash(plain_password)
        teacher.portal_token = _generate_unique_teacher_token()
        teacher.__dict__["_plain_portal_password"] = plain_password
    elif "_plain_portal_password" in teacher.__dict__:
        teacher.__dict__.pop("_plain_portal_password", None)

    if username_generated or plain_password is not None:
        db.session.flush()

    return teacher.portal_username, plain_password


def _slugify(value: str) -> str:
    return "".join(ch for ch in value.lower() if ch.isalnum())


_LOGIN_FOLD_REPLACEMENTS = {
    "ß": "ss",
    "ẞ": "ss",
    "Æ": "ae",
    "æ": "ae",
    "Œ": "oe",
    "œ": "oe",
    "Ø": "o",
    "ø": "o",
    "Đ": "d",
    "đ": "d",
    "Ð": "d",
    "ð": "d",
    "Þ": "th",
    "þ": "th",
    "Ł": "l",
    "ł": "l",
    "Ħ": "h",
    "ħ": "h",
    "ı": "i",
}


def _fold_login_username(value: str | None) -> str:
    raw = (value or "").strip()
    if not raw:
        return ""

    replaced = "".join(_LOGIN_FOLD_REPLACEMENTS.get(ch, ch) for ch in raw)
    normalized = unicodedata.normalize("NFKD", replaced)
    stripped = "".join(ch for ch in normalized if not unicodedata.combining(ch))
    return stripped.casefold()


def _normalize_login_email(value: str | None) -> str:
    return (value or "").strip().lower()


def _login_identifier_looks_like_email(value: str | None) -> bool:
    normalized = _normalize_login_email(value)
    return "@" in normalized if normalized else False


def _folded_username_conflicts(records: Iterable[Any], username_attr: str) -> list[dict[str, Any]]:
    folded_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        username = getattr(record, username_attr, None)
        if not username:
            continue
        folded_groups[_fold_login_username(username)].append(
            {
                "id": record.id,
                "username": username,
                "display_name": " ".join(
                    part for part in [getattr(record, "first_name", None), getattr(record, "last_name", None)] if part
                ),
                "is_absent": bool(getattr(record, "is_absent", False)),
            }
        )

    conflicts: list[dict[str, Any]] = []
    for folded_username, entries in sorted(folded_groups.items()):
        if len(entries) < 2:
            continue
        entries.sort(key=lambda item: (item["username"].casefold(), item["id"]))
        conflicts.append(
            {
                "folded_username": folded_username,
                "entries": entries,
                "entry_count": len(entries),
            }
        )
    return conflicts


def _login_alias_diagnostics() -> dict[str, Any]:
    student_records = Student.query.filter(Student.upload_username.isnot(None)).order_by(Student.id.asc()).all()
    teacher_records = Teacher.query.filter(Teacher.portal_username.isnot(None)).order_by(Teacher.id.asc()).all()

    student_conflicts = _folded_username_conflicts(student_records, "upload_username")
    teacher_conflicts = _folded_username_conflicts(teacher_records, "portal_username")

    student_report = {
        "label": "Students",
        "checked_count": len(student_records),
        "safe": not student_conflicts,
        "conflict_count": len(student_conflicts),
        "conflicts": student_conflicts,
    }
    teacher_report = {
        "label": "Teachers",
        "checked_count": len(teacher_records),
        "safe": not teacher_conflicts,
        "conflict_count": len(teacher_conflicts),
        "conflicts": teacher_conflicts,
    }
    enabled = _login_folded_alias_enabled()
    all_safe = student_report["safe"] and teacher_report["safe"]
    return {
        "enabled": enabled,
        "all_safe": all_safe,
        "enabled_but_unsafe": enabled and not all_safe,
        "students": student_report,
        "teachers": teacher_report,
    }


def _resolve_login_alias_match(
    model: Any,
    username_attr: str,
    email_attr: str,
    submitted_username: str,
) -> tuple[Any | None, dict[str, Any]]:
    username_column = getattr(model, username_attr)
    email_column = getattr(model, email_attr)

    query = model.query.filter(username_column == submitted_username)
    if hasattr(model, "is_absent"):
        query = query.filter(model.is_absent.is_(False))
    exact_match = query.first()
    if exact_match:
        return exact_match, {
            "match_type": "exact",
            "submitted_username": submitted_username,
            "canonical_username": getattr(exact_match, username_attr),
        }

    normalized_email = _normalize_login_email(submitted_username)
    if _login_email_enabled() and normalized_email:
        email_query = model.query.filter(email_column == normalized_email)
        if hasattr(model, "is_absent"):
            email_query = email_query.filter(model.is_absent.is_(False))
        email_match = email_query.first()
        if email_match:
            return email_match, {
                "match_type": "email",
                "submitted_username": submitted_username,
                "submitted_email": normalized_email,
                "canonical_username": getattr(email_match, username_attr),
                "canonical_email": getattr(email_match, email_attr),
            }

    if not _login_folded_alias_enabled():
        return None, {
            "match_type": "none",
            "submitted_username": submitted_username,
            "submitted_email": normalized_email,
            "folded_username": "",
            "match_count": 0,
        }

    folded_username = _fold_login_username(submitted_username)
    if not folded_username:
        return None, {
            "match_type": "none",
            "submitted_username": submitted_username,
            "submitted_email": normalized_email,
            "folded_username": folded_username,
            "match_count": 0,
        }

    folded_matches = []
    candidate_query = model.query.filter(username_column.isnot(None))
    if hasattr(model, "is_absent"):
        candidate_query = candidate_query.filter(model.is_absent.is_(False))
    for candidate in candidate_query.all():
        if _fold_login_username(getattr(candidate, username_attr)) == folded_username:
            folded_matches.append(candidate)

    if len(folded_matches) == 1:
        return folded_matches[0], {
            "match_type": "folded",
            "submitted_username": submitted_username,
            "submitted_email": normalized_email,
            "folded_username": folded_username,
            "canonical_username": getattr(folded_matches[0], username_attr),
            "match_count": 1,
        }

    return None, {
        "match_type": "ambiguous" if folded_matches else "none",
        "submitted_username": submitted_username,
        "submitted_email": normalized_email,
        "folded_username": folded_username,
        "match_count": len(folded_matches),
    }


def _generate_reset_token_value() -> str:
    for _ in range(10):
        token_candidate = secrets.token_urlsafe(32)[:64]
        existing = PasswordResetToken.query.filter(PasswordResetToken.token == token_candidate).first()
        if not existing:
            return token_candidate
    raise RuntimeError("Unable to generate unique password reset token")


def _invalidate_open_reset_tokens(
    audience: PasswordResetAudience, *, student_id: int | None = None, teacher_id: int | None = None
) -> None:
    query = PasswordResetToken.query.filter(
        PasswordResetToken.audience == audience,
        PasswordResetToken.used_at.is_(None),
    )
    if student_id is not None:
        query = query.filter(PasswordResetToken.student_id == student_id)
    if teacher_id is not None:
        query = query.filter(PasswordResetToken.teacher_id == teacher_id)

    for entry in query:
        entry.used_at = datetime.utcnow()


def _create_student_reset_token(student: Student) -> PasswordResetToken:
    _invalidate_open_reset_tokens(PasswordResetAudience.STUDENT, student_id=student.id)
    reset_token = PasswordResetToken(
        audience=PasswordResetAudience.STUDENT,
        student=student,
        token=_generate_reset_token_value(),
        expires_at=datetime.utcnow() + timedelta(hours=RESET_TOKEN_EXPIRY_HOURS),
    )
    db.session.add(reset_token)
    return reset_token


def _create_teacher_reset_token(teacher: Teacher) -> PasswordResetToken:
    _invalidate_open_reset_tokens(PasswordResetAudience.TEACHER, teacher_id=teacher.id)
    reset_token = PasswordResetToken(
        audience=PasswordResetAudience.TEACHER,
        teacher=teacher,
        token=_generate_reset_token_value(),
        expires_at=datetime.utcnow() + timedelta(hours=RESET_TOKEN_EXPIRY_HOURS),
    )
    db.session.add(reset_token)
    return reset_token


def _load_reset_token(token_value: str, audience: PasswordResetAudience) -> PasswordResetToken | None:
    return (
        PasswordResetToken.query.options(
            joinedload(PasswordResetToken.student),
            joinedload(PasswordResetToken.teacher),
        )
        .filter(
            PasswordResetToken.token == token_value,
            PasswordResetToken.audience == audience,
        )
        .first()
    )


def _reset_token_status(entry: PasswordResetToken | None) -> tuple[bool, str | None]:
    if not entry:
        return False, "invalid"
    if entry.used_at is not None:
        return False, "used"
    if entry.expires_at < datetime.utcnow():
        return False, "expired"
    if entry.audience == PasswordResetAudience.STUDENT and not entry.student:
        return False, "invalid"
    if entry.audience == PasswordResetAudience.TEACHER and not entry.teacher:
        return False, "invalid"
    return True, None


def _load_delegation_invite(token_value: str | None, role: DelegationInviteRole) -> DelegationInvite | None:
    if not token_value:
        return None
    return (
        DelegationInvite.query.filter(
            DelegationInvite.token == token_value,
            DelegationInvite.role == role,
        ).first()
    )


def _delegation_invite_state(invite: DelegationInvite | None) -> SimpleNamespace:
    if not invite:
        return SimpleNamespace(status="invalid", label="Invalid", active=False)
    if invite.revoked_at is not None:
        return SimpleNamespace(status="revoked", label="Revoked", active=False)
    if invite.accepted_at is not None:
        return SimpleNamespace(status="completed", label="Completed", active=False)
    if invite.opened_at is not None:
        return SimpleNamespace(status="started", label="Started", active=True)
    return SimpleNamespace(status="pending", label="Pending", active=True)


def _get_upload_folder() -> Path:
    folder = Path(current_app.config.get("UPLOAD_FOLDER", Path(current_app.instance_path) / "uploads"))
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _normalize_static_page_slug(value: str | None) -> str:
    return (value or "").strip().lower()


def _get_static_page_attachment_folder() -> Path:
    folder = Path(
        current_app.config.get(
            "STATIC_PAGE_ATTACHMENT_FOLDER",
            Path(current_app.config.get("UPLOAD_FOLDER", Path(current_app.instance_path) / "uploads")) / "static-pages",
        )
    )
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _static_page_attachment_path(stored_filename: str | None) -> Path | None:
    if not stored_filename:
        return None
    return _get_static_page_attachment_folder() / stored_filename


def _save_static_page_attachment(page: StaticPage, upload) -> StaticPageAttachment:
    original_filename = upload.filename or ""
    safe_filename = secure_filename(original_filename)
    if not safe_filename:
        raise ValueError("Attachment filename is invalid.")

    upload.seek(0, 2)
    file_size = upload.tell()
    upload.seek(0)

    suffix = Path(safe_filename).suffix
    token = secrets.token_hex(6)
    stored_filename = f"static-page-{page.id}-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{token}{suffix}"
    file_path = _get_static_page_attachment_folder() / stored_filename
    upload.save(file_path)

    attachment = StaticPageAttachment(
        page_id=page.id,
        original_filename=safe_filename,
        stored_filename=stored_filename,
        content_type=(upload.mimetype or None),
        file_size=file_size,
        uploaded_by=_current_admin(),
    )
    db.session.add(attachment)
    db.session.flush()
    return attachment


def _delete_static_page_attachment_file(stored_filename: str | None) -> None:
    file_path = _static_page_attachment_path(stored_filename)
    if not file_path:
        return
    try:
        if file_path.exists():
            file_path.unlink()
    except OSError:
        current_app.logger.warning("Failed to remove static page attachment file: %s", file_path)



def _get_active_batch(stage: Stage) -> AssignmentBatch | None:
    return (
        AssignmentBatch.query.filter_by(stage=stage, status=PlannerStatus.PUBLISHED)
        .order_by(
            AssignmentBatch.published_at.is_(None),
            AssignmentBatch.published_at.desc(),
            AssignmentBatch.created_at.desc(),
        )
        .first()
    )

def _student_from_session() -> Student | None:
    student_id = session.get("student_upload_id")
    if not student_id:
        return None
    return db.session.get(Student, student_id)


def _teacher_portal_from_session() -> Teacher | None:
    teacher_id = session.get("teacher_portal_id")
    if not teacher_id:
        return None
    return db.session.get(Teacher, teacher_id)


def _portal_next_url(candidate: str | None) -> str | None:
    if not candidate:
        return None
    candidate = candidate.strip()
    if not candidate.startswith("/"):
        return None
    if candidate.startswith("//"):
        return None
    if "\r" in candidate or "\n" in candidate:
        return None
    return candidate


def _teacher_portal_assignment_ids(teacher: Teacher, stage: Stage) -> set[int]:
    assignment_ids: set[int] = set()
    if stage == Stage.STAGE_1:
        group_query = (
            db.session.query(EssayAssignment.id)
            .join(ReadingGroup, EssayAssignment.group_id == ReadingGroup.id)
            .join(ReadingGroupMember, ReadingGroupMember.group_id == ReadingGroup.id)
            .join(AssignmentBatch, AssignmentBatch.id == EssayAssignment.batch_id)
            .filter(
                ReadingGroupMember.teacher_id == teacher.id,
                ReadingGroupMember.active.is_(True),
                ReadingGroup.stage == Stage.STAGE_1,
                AssignmentBatch.status == PlannerStatus.PUBLISHED,
            )
        )
        assignment_ids.update(aid for (aid,) in group_query.all())

        extra_query = (
            db.session.query(EssayAssignmentExtraReader.assignment_id)
            .join(EssayAssignment, EssayAssignment.id == EssayAssignmentExtraReader.assignment_id)
            .join(AssignmentBatch, AssignmentBatch.id == EssayAssignment.batch_id)
            .filter(
                EssayAssignmentExtraReader.teacher_id == teacher.id,
                EssayAssignmentExtraReader.stage == Stage.STAGE_1.value,
                AssignmentBatch.status == PlannerStatus.PUBLISHED,
            )
        )
        assignment_ids.update(aid for (aid,) in extra_query.all())
    else:
        extra_query = (
            db.session.query(EssayAssignmentExtraReader.assignment_id)
            .join(EssayAssignment, EssayAssignment.id == EssayAssignmentExtraReader.assignment_id)
            .join(AssignmentBatch, AssignmentBatch.id == EssayAssignment.batch_id)
            .filter(
                EssayAssignmentExtraReader.teacher_id == teacher.id,
                EssayAssignmentExtraReader.stage == Stage.STAGE_2.value,
                AssignmentBatch.status == PlannerStatus.PUBLISHED,
                EssayAssignment.stage_two_candidate.is_(True),
            )
        )
        assignment_ids.update(aid for (aid,) in extra_query.all())
    return assignment_ids


def _compute_teacher_portal_status(student: Student, score_record: EssayScore | None, pdf_available: bool) -> SimpleNamespace:
    if not pdf_available:
        return SimpleNamespace(code="waiting_upload", label="Awaiting essay upload", css="warning", timestamp=None)
    if not student.essay_finalized:
        return SimpleNamespace(code="awaiting_finalization", label="Awaiting student finalization", css="warning", timestamp=None)
    if score_record and score_record.teacher_finalized:
        timestamp = score_record.teacher_finalized_at or score_record.updated_at
        return SimpleNamespace(code="finalized", label="Teacher finalized", css="success", timestamp=timestamp)
    if score_record and score_record.score is not None:
        timestamp = score_record.submitted_at or score_record.updated_at
        return SimpleNamespace(code="in_progress", label="Draft saved", css="", timestamp=timestamp)
    return SimpleNamespace(code="pending", label="Not started", css="", timestamp=None)


def _resolve_portal_status(
    student: Student,
    score_record: EssayScore | None,
    pdf_available: bool,
    *,
    locked: bool = False,
) -> SimpleNamespace:
    status_info = _compute_teacher_portal_status(student, score_record, pdf_available)
    status_code = status_info.code
    status_css = status_info.css
    status_timestamp = status_info.timestamp

    locked_flag = locked or (score_record.locked_by_admin if score_record else False)

    if locked_flag:
        status_code = "locked"
        status_css = "warning"
        if score_record:
            status_timestamp = score_record.admin_updated_at or status_timestamp or score_record.updated_at
    elif score_record:
        if score_record.admin_override:
            status_code = "admin_override"
            status_css = "warning"
            status_timestamp = score_record.admin_updated_at or score_record.updated_at or status_timestamp
        elif score_record.admin_entered:
            status_code = "admin_entered"
            status_css = ""
            status_timestamp = score_record.admin_updated_at or score_record.updated_at or status_timestamp
        elif score_record.teacher_finalized:
            status_code = "finalized"
            status_css = "success"
            status_timestamp = score_record.teacher_finalized_at or status_timestamp or score_record.updated_at
        elif score_record.score is not None and status_code not in {"awaiting_finalization", "waiting_upload"}:
            status_code = "in_progress"

    resolved_label = PORTAL_STATUS_LABELS.get(status_code, status_info.label or "Not started")
    return SimpleNamespace(code=status_code, label=resolved_label, css=status_css, timestamp=status_timestamp)


def _teacher_portal_assignments(teacher: Teacher, stage: Stage) -> list[SimpleNamespace]:
    assignment_ids = _teacher_portal_assignment_ids(teacher, stage)
    if not assignment_ids:
        return []

    assignments = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.student),
            joinedload(EssayAssignment.batch),
            joinedload(EssayAssignment.group).joinedload(ReadingGroup.members).joinedload(ReadingGroupMember.teacher),
            joinedload(EssayAssignment.extra_readers),
            joinedload(EssayAssignment.scores),
        )
        .filter(EssayAssignment.id.in_(assignment_ids))
        .order_by(EssayAssignment.id.asc())
        .all()
    )

    global_stage_locked = _is_stage_locked(stage)

    entries: list[SimpleNamespace] = []
    for assignment in assignments:
        student = assignment.student
        if not student or student.is_absent:
            continue
        score_record = next(
            (score for score in assignment.scores if score.teacher_id == teacher.id and score.stage == stage),
            None,
        )
        pdf_path = _essay_pdf_path(student)
        pdf_available = pdf_path is not None
        if stage == Stage.STAGE_2 and not assignment.stage_two_candidate:
            continue
        assignment_locked = assignment.stage1_locked if stage == Stage.STAGE_1 else assignment.stage2_locked
        status_info = _resolve_portal_status(
            student,
            score_record,
            pdf_available,
            locked=(assignment_locked or global_stage_locked),
        )
        status = status_info.code
        essay_available = status not in {"awaiting_finalization", "waiting_upload"}
        resolved_label = status_info.label
        if assignment.group:
            group_label = assignment.group.name or str(assignment.group.id)
        else:
            group_label = "—"
        essay_id = student.id if student else None
        display_label = f"Essay {essay_id}" if essay_id is not None else f"Assignment {assignment.id}"
        entry = SimpleNamespace(
            assignment_id=assignment.id,
            stage=stage,
            essay_id=essay_id,
            display_label=display_label,
            group_label=group_label,
            stage_label=STAGE_LABELS[stage],
            essay_available=essay_available,
            score_record=score_record,
            status=status,
            status_label=resolved_label,
            status_css=status_info.css,
            status_timestamp=status_info.timestamp,
            updated_at=score_record.updated_at if score_record else None,
            finalized=bool(score_record.teacher_finalized) if score_record else False,
            download_url=(
                url_for(
                    "teacher_portal.download_assignment",
                    assignment_id=assignment.id,
                    stage=stage.value,
                )
                if essay_available and not (assignment_locked or global_stage_locked)
                else None
            ),
        )
        entries.append(entry)

    entries.sort(key=lambda entry: (entry.display_label, entry.assignment_id))
    return entries


def _get_stage_one_assignment(student: Student) -> EssayAssignment | None:
    active_batch = _get_active_batch(Stage.STAGE_1)
    if not active_batch:
        return None
    return (
        EssayAssignment.query
        .filter_by(student_id=student.id, stage=Stage.STAGE_1, batch_id=active_batch.id)
        .order_by(EssayAssignment.id.desc())
        .first()
    )


def _stage_assessors(assignment: EssayAssignment | None, stage: Stage, include_group: bool = False) -> list[SimpleNamespace]:
    assessors: list[SimpleNamespace] = []
    if not assignment:
        return assessors

    seen: set[int] = set()
    if include_group and assignment.group:
        for member in assignment.group.members:
            teacher = member.teacher
            if not member.active or not teacher or teacher.is_absent:
                continue
            if teacher.id in seen:
                continue
            assessors.append(
                SimpleNamespace(teacher=teacher, teacher_id=teacher.id, source="Group", is_extra=False, extra_reader=None)
            )
            seen.add(teacher.id)

    for extra in assignment.extra_readers:
        if extra.stage != stage:
            continue
        teacher = extra.teacher
        if not teacher or teacher.is_absent or teacher.id in seen:
            continue
        assessors.append(
            SimpleNamespace(teacher=teacher, teacher_id=teacher.id, source="Extra", is_extra=True, extra_reader=extra)
        )
        seen.add(teacher.id)

    return assessors


GRADE_STAGE_SLOTS = {
    Stage.STAGE_1.value: 5,
    Stage.STAGE_2.value: 2,
}

STAGE_1_REQUIRED_READINGS = 4


def _score_stats(
    values: list[float],
    *,
    trim_extremes: bool = False,
) -> SimpleNamespace:
    raw_values = [float(value) for value in values if value is not None]
    count = len(raw_values)
    if not raw_values:
        return SimpleNamespace(
            count=0,
            used_count=0,
            average=None,
            raw_average=None,
            spread=None,
            used_values=[],
            trimmed=False,
            dropped_low=None,
            dropped_high=None,
        )

    used_values = list(raw_values)
    dropped_low = None
    dropped_high = None
    trimmed = False
    if trim_extremes and count > STAGE_1_REQUIRED_READINGS:
        sorted_values = sorted(raw_values)
        dropped_low = sorted_values[0]
        dropped_high = sorted_values[-1]
        used_values = sorted_values[1:-1]
        trimmed = True

    raw_average = sum(raw_values) / count
    average = sum(used_values) / len(used_values) if used_values else None
    spread = max(raw_values) - min(raw_values) if count >= 2 else 0.0
    return SimpleNamespace(
        count=count,
        used_count=len(used_values),
        average=average,
        raw_average=raw_average,
        spread=spread,
        used_values=used_values,
        trimmed=trimmed,
        dropped_low=dropped_low,
        dropped_high=dropped_high,
    )


def _stage_one_score_stats(values: list[float]) -> SimpleNamespace:
    return _score_stats(values, trim_extremes=_stage1_trim_extreme_scores_enabled())


def _grade_score_stats(values: list[float], *, trim_extremes: bool = False) -> dict[str, Any]:
    stats = _score_stats(values, trim_extremes=trim_extremes)
    return {
        "count": stats.count,
        "used_count": stats.used_count,
        "average": stats.average,
        "raw_average": stats.raw_average,
        "spread": stats.spread,
        "used_values": stats.used_values,
        "trimmed": stats.trimmed,
        "dropped_low": stats.dropped_low,
        "dropped_high": stats.dropped_high,
    }


def _grade_export_rows(
    *,
    selected_stages: set[Stage],
    stage3_only: bool,
    countries: set[str] | None,
    include_absent: bool,
    finalized_only: bool,
) -> list[dict]:
    query = (
        db.session.query(
            EssayScore,
            EssayAssignment,
            Student,
            Teacher,
        )
        .join(EssayAssignment, EssayScore.assignment_id == EssayAssignment.id)
        .join(Student, EssayAssignment.student_id == Student.id)
        .join(Teacher, EssayScore.teacher_id == Teacher.id)
        .filter(EssayScore.score.isnot(None))
    )

    if not include_absent:
        query = query.filter(Student.is_absent.is_(False))
    if finalized_only:
        query = query.filter(EssayScore.teacher_finalized.is_(True))
    if countries:
        query = query.filter(Student.country.in_(countries))

    rows: list[dict] = []
    for score_record, assignment, student, teacher in query:
        stage_enum: Stage = score_record.stage
        if stage_enum not in selected_stages:
            continue
        if stage_enum == Stage.STAGE_2 and not assignment.stage_two_candidate:
            continue
        rows.append(
            {
                "student_id": student.id,
                "student_name": f"{student.first_name} {student.last_name}",
                "student_country": student.country or "",
                "student_absent": bool(student.is_absent),
                "stage": stage_enum,
                "score": float(score_record.score) if score_record.score is not None else None,
                "finalized": bool(score_record.teacher_finalized),
                "submitted_at": score_record.submitted_at,
                "updated_at": score_record.updated_at,
                "teacher_id": teacher.id,
                "teacher_name": f"{teacher.first_name} {teacher.last_name}",
                "teacher_country": teacher.country or "",
                "assignment_id": assignment.id,
                "stage_two_candidate": bool(assignment.stage_two_candidate),
                "stage_three_recommended": bool(assignment.stage_three_recommended),
            }
        )

    if stage3_only:
        rows = [row for row in rows if row["stage_three_recommended"]]
    return rows


def _summarize_grade_rows(rows: list[dict]) -> list[dict]:
    students: dict[int, dict] = {}
    for row in rows:
        student_entry = students.setdefault(
            row["student_id"],
            {
                "student_id": row["student_id"],
                "student_name": row["student_name"],
                "student_country": row["student_country"],
                "recommended": False,
                "stages": {},
                "all_scores": [],
            },
        )
        stage_enum: Stage = row["stage"]
        stage_key = stage_enum.value
        stage_bucket = student_entry["stages"].setdefault(
            stage_key,
            {"scores": [], "values": []},
        )
        stage_bucket["scores"].append(row)
        if row["score"] is not None:
            stage_bucket["values"].append(row["score"])
            student_entry["all_scores"].append(row["score"])
        if row["stage_three_recommended"]:
            student_entry["recommended"] = True

    summarized: list[dict] = []
    for entry in students.values():
        stage_summary: dict[str, dict] = {}
        combined_average_values: list[float] = []
        for stage_key, data in entry["stages"].items():
            stage_enum = Stage(stage_key)
            stats = _grade_score_stats(
                data["values"],
                trim_extremes=stage_enum == Stage.STAGE_1 and _stage1_trim_extreme_scores_enabled(),
            )
            sorted_scores = sorted(
                data["scores"],
                key=lambda item: (
                    item["teacher_id"],
                    item["submitted_at"] or item["updated_at"] or datetime.min,
                ),
            )
            for index, score_entry in enumerate(sorted_scores, start=1):
                score_entry["reader_alias"] = f"R{index}"
            stage_summary[stage_key] = {
                "scores": sorted_scores,
                "count": stats["count"],
                "used_count": stats["used_count"],
                "average": stats["average"],
                "raw_average": stats["raw_average"],
                "spread": stats["spread"],
                "average_is_trimmed": stats["trimmed"],
            }
            combined_average_values.extend(stats["used_values"])
        combined_stats = _grade_score_stats(combined_average_values)
        summarized.append(
            {
                "student_id": entry["student_id"],
                "student_name": entry["student_name"],
                "student_country": entry["student_country"],
                "recommended": entry["recommended"],
                "stages": stage_summary,
                "combined_count": combined_stats["count"],
                "combined_average": combined_stats["average"],
                "combined_spread": combined_stats["spread"],
            }
        )

    return summarized


def _grades_sort_key(entry: dict, sort_key: str) -> tuple:
    if sort_key == "stage1_avg":
        value = entry["stages"].get(Stage.STAGE_1.value, {}).get("average")
    elif sort_key == "stage2_avg":
        value = entry["stages"].get(Stage.STAGE_2.value, {}).get("average")
    elif sort_key == "combined_avg":
        value = entry.get("combined_average")
    else:
        value = entry["student_id"]

    if value is None:
        value = float("-inf") if sort_key != "student_id" else entry["student_id"]
    return (value, entry["student_id"])


def _safe_sheet_title(base: str, existing: set[str]) -> str:
    clean = re.sub(r"[\\/*?:\[\]]", "", base).strip() or "Sheet"
    clean = clean[:28]
    candidate = clean
    index = 2
    while candidate in existing:
        suffix = f"-{index}"
        candidate = f"{clean[:28 - len(suffix)]}{suffix}"
        index += 1
    existing.add(candidate)
    return candidate


def _format_score_value(value: float | None) -> str:
    if value is None:
        return ""
    return f"{value:.2f}"


def _apply_column_bold(sheet, column_indexes: list[int]) -> None:
    if not column_indexes or sheet.max_row < 2:
        return
    bold_font = Font(bold=True)
    for row_idx in range(2, sheet.max_row + 1):
        for col_idx in column_indexes:
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.font = bold_font


def _grades_export_xlsx(students: list[dict], options: dict) -> Any:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"

    anonymize_students = options.get("anonymize_students", False)
    anonymize_teachers = options.get("anonymize_teachers", True)
    include_countries = options.get("include_countries", False)
    include_timestamps = options.get("include_timestamps", True)
    group_by_country = options.get("group_by_country", False)
    compact_view = options.get("compact_view", False)
    stage_order = options.get("stage_order", [Stage.STAGE_1.value, Stage.STAGE_2.value])

    summary_header_entries: list[tuple[str, str | None]] = [("Student ID", None)]
    if not anonymize_students:
        summary_header_entries.append(("Student name", None))
    if include_countries:
        summary_header_entries.append(("Country", None))
    for stage_key in stage_order:
        stage_label = STAGE_LABELS[Stage(stage_key)]
        summary_header_entries.append((f"{stage_label} scores", None))
        summary_header_entries.append((f"{stage_label} average", "avg"))
        summary_header_entries.append((f"{stage_label} spread", None))
    summary_header_entries.append(("All scores count", None))
    summary_header_entries.append(("All scores average", "avg"))
    summary_header_entries.append(("All scores spread", None))
    summary_header_entries.append(("Stage 3 recommended", None))

    summary_sheet.append([text for text, _ in summary_header_entries])
    summary_avg_columns = [idx + 1 for idx, (_, kind) in enumerate(summary_header_entries) if kind == "avg"]

    for student in students:
        row = [student["student_id"]]
        if not anonymize_students:
            row.append(student["student_name"])
        if include_countries:
            row.append(student["student_country"])
        for stage_key in stage_order:
            stage_info = student["stages"].get(stage_key, {})
            row.append(stage_info.get("count", 0))
            row.append(_format_score_value(stage_info.get("average")))
            row.append(_format_score_value(stage_info.get("spread")))
        row.append(student["combined_count"])
        row.append(_format_score_value(student["combined_average"]))
        row.append(_format_score_value(student["combined_spread"]))
        row.append("Yes" if student["recommended"] else "No")
        summary_sheet.append(row)

    _apply_column_bold(summary_sheet, summary_avg_columns)

    if group_by_country and include_countries:
        existing_titles = {summary_sheet.title}
        countries = sorted({student["student_country"] for student in students if student["student_country"]})
        for country in countries:
            sheet = workbook.create_sheet(_safe_sheet_title(country, existing_titles))
            sheet.append([text for text, _ in summary_header_entries])
            for student in students:
                if student["student_country"] != country:
                    continue
                row = [student["student_id"]]
                if not anonymize_students:
                    row.append(student["student_name"])
                if include_countries:
                    row.append(student["student_country"])
                for stage_key in stage_order:
                    stage_info = student["stages"].get(stage_key, {})
                    row.append(stage_info.get("count", 0))
                    row.append(_format_score_value(stage_info.get("average")))
                    row.append(_format_score_value(stage_info.get("spread")))
                row.append(student["combined_count"])
                row.append(_format_score_value(student["combined_average"]))
                row.append(_format_score_value(student["combined_spread"]))
                row.append("Yes" if student["recommended"] else "No")
                sheet.append(row)
            _apply_column_bold(sheet, summary_avg_columns)

    detail_headers = ["Student ID"]
    if not anonymize_students:
        detail_headers.append("Student name")
    if include_countries:
        detail_headers.append("Student country")
    detail_headers.extend(["Stage", "Score", "Finalized"])
    if include_timestamps:
        detail_headers.append("Submitted at")
        detail_headers.append("Last update")
    if anonymize_teachers:
        detail_headers.append("Reader")
    else:
        detail_headers.append("Teacher ID")
        detail_headers.append("Teacher name")
        if include_countries:
            detail_headers.append("Teacher country")

    detail_sheet = workbook.create_sheet("Details")
    detail_sheet.append(detail_headers)

    for student in students:
        for stage_key in stage_order:
            stage_info = student["stages"].get(stage_key)
            if not stage_info:
                continue
            for score_entry in stage_info["scores"]:
                row = [student["student_id"]]
                if not anonymize_students:
                    row.append(student["student_name"])
                if include_countries:
                    row.append(student["student_country"])
                row.append(STAGE_LABELS[Stage(stage_key)])
                row.append(_format_score_value(score_entry["score"]))
                row.append("Yes" if score_entry["finalized"] else "No")
                if include_timestamps:
                    submitted_at = score_entry["submitted_at"] or score_entry["updated_at"]
                    row.append(submitted_at.strftime("%Y-%m-%d %H:%M") if submitted_at else "")
                    row.append(score_entry["updated_at"].strftime("%Y-%m-%d %H:%M") if score_entry["updated_at"] else "")
                if anonymize_teachers:
                    row.append(score_entry.get("reader_alias", ""))
                else:
                    row.append(score_entry["teacher_id"])
                    row.append(score_entry["teacher_name"])
                    if include_countries:
                        row.append(score_entry["teacher_country"])
                detail_sheet.append(row)

    if compact_view:
        stage_slots = GRADE_STAGE_SLOTS
        compact_header_entries: list[tuple[str, str | None]] = [("Student ID", None)]
        if not anonymize_students:
            compact_header_entries.append(("Student name", None))
        if include_countries:
            compact_header_entries.append(("Country", None))
        for stage_key in stage_order:
            stage_label = STAGE_LABELS[Stage(stage_key)]
            slots = stage_slots.get(stage_key, 0)
            for index in range(slots):
                compact_header_entries.append((f"{stage_label} reader {index + 1}", None))
            compact_header_entries.append((f"{stage_label} average", "avg"))
            compact_header_entries.append((f"{stage_label} spread", None))
        compact_header_entries.append(("Combined average", "avg"))
        compact_header_entries.append(("Combined spread", None))
        compact_header_entries.append(("Stage 3 recommended", None))

        compact_sheet = workbook.create_sheet("Compact")
        compact_sheet.append([text for text, _ in compact_header_entries])
        compact_avg_columns = [idx + 1 for idx, (_, kind) in enumerate(compact_header_entries) if kind == "avg"]

        for student in students:
            row = [student["student_id"]]
            if not anonymize_students:
                row.append(student["student_name"])
            if include_countries:
                row.append(student["student_country"])
            for stage_key in stage_order:
                stage_info = student["stages"].get(stage_key)
                slots = stage_slots.get(stage_key, 0)
                scores = stage_info["scores"] if stage_info else []
                for index in range(slots):
                    value = ""
                    if stage_info and index < len(scores):
                        score_entry = scores[index]
                        formatted = _format_score_value(score_entry["score"])
                        if formatted:
                            if anonymize_teachers:
                                alias = score_entry.get("reader_alias")
                                formatted = f"{formatted} ({alias})" if alias else formatted
                            else:
                                formatted = f"{formatted} (T{score_entry['teacher_id']})"
                        value = formatted or ""
                    row.append(value)
                if stage_info:
                    row.append(_format_score_value(stage_info["average"]))
                    row.append(_format_score_value(stage_info["spread"]))
                else:
                    row.extend(["", ""])
            row.append(_format_score_value(student["combined_average"]))
            row.append(_format_score_value(student["combined_spread"]))
            row.append("Yes" if student["recommended"] else "No")
            compact_sheet.append(row)

        _apply_column_bold(compact_sheet, compact_avg_columns)
        workbook.active = workbook.sheetnames.index(compact_sheet.title)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="grades-export.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _grades_export_pdf(students: list[dict], options: dict) -> Any:
    stage_order = options.get("stage_order", [Stage.STAGE_1.value, Stage.STAGE_2.value])
    stage_order_labels = [(stage_key, STAGE_LABELS[Stage(stage_key)]) for stage_key in stage_order]
    context = {
        "students": students,
        "anonymize_students": options.get("anonymize_students", False),
        "anonymize_teachers": options.get("anonymize_teachers", True),
        "include_countries": options.get("include_countries", False),
        "include_timestamps": options.get("include_timestamps", True),
        "stage_order": stage_order_labels,
        "generated_at": datetime.utcnow(),
        "stage_slots": {
            Stage.STAGE_1.value: 5,
            Stage.STAGE_2.value: 2,
        },
    }
    template = "admin/grades/pdf/report_compact.html" if options.get("compact_view") else "admin/grades/pdf/report.html"
    return _render_pdf(template, "grades-report.pdf", **context)


def _compute_stage_one_status(assignment: EssayAssignment | None) -> SimpleNamespace:
    assessors = _stage_assessors(assignment, Stage.STAGE_1, include_group=True)
    assigned_readers = len(assessors)
    expected = max(assigned_readers, STAGE_1_REQUIRED_READINGS) if assignment else 0
    missing_reader_assignments = max(STAGE_1_REQUIRED_READINGS - assigned_readers, 0) if assignment else 0
    if not assignment:
        return SimpleNamespace(
            entered_scores=0,
            expected_scores=expected,
            assigned_readers=assigned_readers,
            required_readers=STAGE_1_REQUIRED_READINGS,
            missing_reader_assignments=missing_reader_assignments,
            missing_scores=0,
            average_score=None,
            raw_average_score=None,
            average_score_count=0,
            raw_score_count=0,
            average_is_trimmed=False,
            trimmed_low_score=None,
            trimmed_high_score=None,
            score_range=None,
            recommendation="No published Stage 1 batch.",
            needs_extra=False,
            stage_two=False,
            completed=False,
        )

    scores = [
        score
        for score in assignment.scores
        if score.stage == Stage.STAGE_1 and score.score is not None
    ]
    entered = len(scores)
    values = [s.score for s in scores if s.score is not None]
    score_stats = _stage_one_score_stats(values)
    average = score_stats.average
    score_range = score_stats.spread if entered else None

    if expected == 0:
        recommendation = "No active readers assigned."
    elif missing_reader_assignments > 0:
        recommendation = (
            f"Needs {missing_reader_assignments} additional reader assignment"
            f"{'s' if missing_reader_assignments != 1 else ''} to reach "
            f"{STAGE_1_REQUIRED_READINGS} Stage 1 readings."
        )
    elif entered < expected:
        missing = expected - entered
        recommendation = f"Awaiting {missing} reader score{'s' if missing != 1 else ''}."
    elif score_range is not None and score_range > 3 and not score_stats.trimmed:
        recommendation = "Score disparity >3: assign additional reader."
    elif average is not None and average >= 7:
        recommendation = "Qualifies for Stage 2."
    else:
        recommendation = "Average below 7: retain in Stage 1."

    return SimpleNamespace(
        entered_scores=entered,
        expected_scores=expected,
        assigned_readers=assigned_readers,
        required_readers=STAGE_1_REQUIRED_READINGS,
        missing_reader_assignments=missing_reader_assignments,
        missing_scores=max(expected - entered, 0),
        average_score=average,
        raw_average_score=score_stats.raw_average,
        average_score_count=score_stats.used_count,
        raw_score_count=score_stats.count,
        average_is_trimmed=score_stats.trimmed,
        trimmed_low_score=score_stats.dropped_low,
        trimmed_high_score=score_stats.dropped_high,
        score_range=score_range,
        recommendation=recommendation,
        needs_extra=assignment.needs_additional_reading,
        stage_two=assignment.stage_two_candidate,
        completed=missing_reader_assignments == 0 and expected > 0 and entered >= expected,
    )


def _stage_one_extra_reader_state(
    assignment: EssayAssignment | None,
    status: SimpleNamespace | None = None,
) -> SimpleNamespace:
    if not assignment:
        return SimpleNamespace(
            bucket="none",
            label="No extra reader needed",
            needs_action=False,
            extra_readers=0,
            extra_scores=0,
            missing_extra_scores=0,
            missing_reader_assignments=0,
        )

    status = status or _compute_stage_one_status(assignment)
    extra_readers = [
        extra
        for extra in assignment.extra_readers
        if extra.stage == Stage.STAGE_1.value
    ]
    extra_teacher_ids = {extra.teacher_id for extra in extra_readers}
    extra_score_teacher_ids = {
        score.teacher_id
        for score in assignment.scores
        if score.stage == Stage.STAGE_1
        and score.teacher_id in extra_teacher_ids
        and score.score is not None
    }
    extra_count = len(extra_readers)
    extra_scores = len(extra_score_teacher_ids)
    missing_extra_scores = max(extra_count - extra_scores, 0)
    needs_extra = bool(status.needs_extra or assignment.needs_additional_reading)
    missing_reader_assignments = getattr(status, "missing_reader_assignments", 0)

    if missing_reader_assignments > 0:
        bucket = "needs_assignment"
        label = "Needs extra reader assignment"
        needs_action = True
    elif extra_count > 0 and getattr(status, "missing_scores", 0) > 0 and missing_extra_scores > 0:
        bucket = "awaiting_score"
        label = "Extra reader assigned; awaiting score"
        needs_action = True
    elif needs_extra and extra_count == 0:
        bucket = "needs_assignment"
        label = "Needs extra reader assignment"
        needs_action = True
    elif needs_extra and missing_extra_scores > 0:
        bucket = "awaiting_score"
        label = "Extra reader assigned; awaiting score"
        needs_action = True
    elif needs_extra:
        bucket = "review_spread"
        label = "Extra score entered; review spread"
        needs_action = True
    elif extra_count > 0:
        bucket = "resolved"
        label = "Resolved after extra reading"
        needs_action = False
    else:
        bucket = "none"
        label = "No extra reader needed"
        needs_action = False

    return SimpleNamespace(
        bucket=bucket,
        label=label,
        needs_action=needs_action,
        extra_readers=extra_count,
        extra_scores=extra_scores,
        missing_extra_scores=missing_extra_scores,
        missing_reader_assignments=missing_reader_assignments,
    )



def _compute_stage_two_status(assignment: EssayAssignment | None) -> SimpleNamespace:
    assessors = _stage_assessors(assignment, Stage.STAGE_2, include_group=False)
    expected = len(assessors)
    if not assignment:
        return SimpleNamespace(
            entered_scores=0,
            expected_scores=expected,
            average_score=None,
            score_range=None,
            completed=False,
        )

    scores = [
        score
        for score in assignment.scores
        if score.stage == Stage.STAGE_2 and score.score is not None
    ]
    entered = len(scores)
    average = sum(s.score for s in scores) / entered if entered else None
    score_range = None
    if len(scores) >= 2:
        values = [s.score for s in scores]
        score_range = max(values) - min(values)

    return SimpleNamespace(
        entered_scores=entered,
        expected_scores=expected,
        average_score=average,
        score_range=score_range,
        completed=expected > 0 and entered >= expected,
    )


def _compute_combined_status(assignment: EssayAssignment | None) -> SimpleNamespace:
    if not assignment:
        return SimpleNamespace(count=0, average=None, score_range=None)

    stage1_values = [
        score.score
        for score in assignment.scores
        if score.stage == Stage.STAGE_1 and score.score is not None
    ]
    stage2_values = [
        score.score
        for score in assignment.scores
        if score.stage == Stage.STAGE_2 and score.score is not None
    ]
    stage1_stats = _stage_one_score_stats(stage1_values)
    values = list(stage1_stats.used_values) + stage2_values
    if not values:
        return SimpleNamespace(count=0, average=None, score_range=None)

    average = sum(values) / len(values)
    score_range = max(values) - min(values) if len(values) >= 2 else None
    return SimpleNamespace(count=len(values), average=average, score_range=score_range)

def _apply_stage_one_rules(assignment: EssayAssignment) -> None:
    status = _compute_stage_one_status(assignment)
    if not status.completed:
        return
    if status.score_range is not None and status.score_range > 3 and not status.average_is_trimmed:
        assignment.needs_additional_reading = True
        assignment.stage_two_candidate = False
    else:
        assignment.needs_additional_reading = False
        if status.average_score is not None and status.average_score >= 7:
            assignment.stage_two_candidate = True
        else:
            assignment.stage_two_candidate = False


def _recompute_stage_one_rules() -> int:
    assignments = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.scores),
            joinedload(EssayAssignment.extra_readers).joinedload(EssayAssignmentExtraReader.teacher),
            joinedload(EssayAssignment.group)
            .joinedload(ReadingGroup.members)
            .joinedload(ReadingGroupMember.teacher),
        )
        .filter(EssayAssignment.stage == Stage.STAGE_1)
        .all()
    )
    for assignment in assignments:
        _apply_stage_one_rules(assignment)
    return len(assignments)


def _ensure_student_credentials(student: Student, *, reset_password: bool = False) -> tuple[str, str | None]:
    username_generated = False
    if not student.upload_username:
        student.upload_username = _generate_unique_username(student.first_name, student.last_name)
        username_generated = True

    plain_password: str | None = None
    if reset_password or not student.upload_password_hash or not student.upload_token:
        plain_password = secrets.token_urlsafe(8)
        student.upload_password_hash = generate_password_hash(plain_password)
        student.upload_token = plain_password
        student.__dict__["_plain_upload_password"] = plain_password

    if username_generated or plain_password is not None:
        db.session.flush()

    return student.upload_username, plain_password


def _student_from_session() -> Student | None:
    student_id = session.get("student_upload_id")
    if not student_id:
        return None
    return db.session.get(Student, student_id)

def _current_admin() -> str | None:
    return session.get("admin_user")


def _parse_admin_users() -> dict[str, str]:
    raw = current_app.config.get("ADMIN_USERS", "")
    if isinstance(raw, dict):
        return {str(k): str(v) for k, v in raw.items()}
    if isinstance(raw, (list, tuple)):
        entries = [str(item) for item in raw]
    else:
        entries = [segment.strip() for segment in str(raw).split(",") if segment.strip()]
    credentials: dict[str, str] = {}
    for entry in entries:
        if ":" not in entry:
            continue
        username, password = entry.split(":", 1)
        credentials[username.strip()] = password.strip()
    return credentials


def _validate_admin_credentials(username: str, password: str) -> bool:
    credentials = _parse_admin_users()
    stored = credentials.get(username)
    if not stored:
        return False
    if stored.startswith("pbkdf2:") or stored.startswith("scrypt:"):
        return check_password_hash(stored, password)
    return secrets.compare_digest(stored, password)


def log_activity(
    action: str,
    description: str,
    *,
    actor: str | None = None,
    object_type: str | None = None,
    object_id: str | None = None,
    details: dict | None = None,
) -> None:
    actor_name = actor or _current_admin() or "system"
    entry = ActivityLog(
        actor=actor_name,
        action=action,
        description=description,
        object_type=object_type,
        object_id=str(object_id) if object_id is not None else None,
        details=details,
    )
    db.session.add(entry)


def _essay_pdf_path(student: Student) -> Path | None:
    if not student.uploaded_filename:
        return None
    path = _get_upload_folder() / student.uploaded_filename
    if path.exists():
        return path
    return None


def _essay_pdf_path_by_filename(filename: str | None) -> Path | None:
    if not filename:
        return None
    path = _get_upload_folder() / filename
    if path.exists():
        return path
    return None


def _current_essay_upload(student: Student) -> EssayUpload | None:
    if not student.uploaded_filename:
        return None
    return (
        EssayUpload.query
        .filter_by(student_id=student.id, filename=student.uploaded_filename)
        .order_by(EssayUpload.id.desc())
        .first()
    )


def _create_teacher_title_page(teacher: Teacher, group: ReadingGroup, essay_ids: list[int]) -> PdfReader:
    buffer = BytesIO()
    width, height = A4
    c = canvas.Canvas(buffer, pagesize=A4)
    margin = 40

    c.setFont("Helvetica-Bold", 22)
    c.drawString(margin, height - margin - 10, "IPO Essay Packet")
    c.setFont("Helvetica", 14)
    c.drawString(margin, height - margin - 36, f"Group: {group.name or group.id}")

    group_label = getattr(group, "name", None) or getattr(group, "id", "—")
    info_rows = [
        ("Teacher ID", teacher.id),
        ("Name", f"{teacher.first_name} {teacher.last_name}"),
        ("Group", group_label),
    ]

    table_top = height - margin - 70
    row_height = 26
    col_widths = [150, width - 2 * margin - 150]
    c.setLineWidth(1)
    y = table_top
    for label, value in info_rows:
        c.rect(margin, y - row_height, col_widths[0], row_height)
        c.rect(margin + col_widths[0], y - row_height, col_widths[1], row_height)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(margin + 6, y - row_height + 8, str(label))
        c.setFont("Helvetica", 12)
        c.drawString(margin + col_widths[0] + 6, y - row_height + 8, str(value))
        y -= row_height

    essay_table_top = y - 20
    header_height = 26
    essay_row_height = 24
    columns = [180, width - 2 * margin - 180]
    c.rect(margin, essay_table_top - header_height, columns[0], header_height)
    c.rect(margin + columns[0], essay_table_top - header_height, columns[1], header_height)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(margin + 6, essay_table_top - header_height + 8, "Essay ID")
    c.drawString(margin + columns[0] + 6, essay_table_top - header_height + 8, "Score")

    y = essay_table_top - header_height
    c.setFont("Helvetica", 12)
    entries = essay_ids or ["—"]
    for essay_id in entries:
        c.rect(margin, y - essay_row_height, columns[0], essay_row_height)
        c.rect(margin + columns[0], y - essay_row_height, columns[1], essay_row_height)
        c.drawString(margin + 6, y - essay_row_height + 8, str(essay_id))
        y -= essay_row_height
        if y < margin + 40:
            break

    c.save()
    buffer.seek(0)
    return PdfReader(buffer)


def _create_committee_title_page(member_name: str, committee_label: str, essay_ids: list[int]) -> PdfReader:
    buffer = BytesIO()
    width, height = A4
    c = canvas.Canvas(buffer, pagesize=A4)
    margin = 40

    c.setFont("Helvetica-Bold", 22)
    c.drawString(margin, height - margin - 10, "IPO Stage 3 Packet")
    c.setFont("Helvetica", 14)
    c.drawString(margin, height - margin - 36, f"Committee: {committee_label}")

    info_rows = [
        ("Member", member_name or "—"),
        ("Committee", committee_label),
    ]

    table_top = height - margin - 70
    row_height = 26
    col_widths = [150, width - 2 * margin - 150]
    c.setLineWidth(1)
    y = table_top
    for label, value in info_rows:
        c.rect(margin, y - row_height, col_widths[0], row_height)
        c.rect(margin + col_widths[0], y - row_height, col_widths[1], row_height)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(margin + 6, y - row_height + 8, str(label))
        c.setFont("Helvetica", 12)
        c.drawString(margin + col_widths[0] + 6, y - row_height + 8, str(value))
        y -= row_height

    essay_table_top = y - 20
    header_height = 26
    essay_row_height = 24
    columns = [180, width - 2 * margin - 180]
    c.rect(margin, essay_table_top - header_height, columns[0], header_height)
    c.rect(margin + columns[0], essay_table_top - header_height, columns[1], header_height)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(margin + 6, essay_table_top - header_height + 8, "Essay ID")
    c.drawString(margin + columns[0] + 6, essay_table_top - header_height + 8, "Comments")

    y = essay_table_top - header_height
    c.setFont("Helvetica", 12)
    entries = essay_ids or ["—"]
    for essay_id in entries:
        c.rect(margin, y - essay_row_height, columns[0], essay_row_height)
        c.rect(margin + columns[0], y - essay_row_height, columns[1], essay_row_height)
        c.drawString(margin + 6, y - essay_row_height + 8, str(essay_id))
        y -= essay_row_height
        if y < margin + 40:
            break

    c.save()
    buffer.seek(0)
    return PdfReader(buffer)


def _add_blank_page(writer: PdfWriter, template_page=None) -> None:
    if template_page is not None:
        width = float(template_page.mediabox.right - template_page.mediabox.left)
        height = float(template_page.mediabox.top - template_page.mediabox.bottom)
    else:
        width, height = A4
    writer.add_blank_page(width=width, height=height)


def _packets_double_sided() -> bool:
    return _get_setting(SETTING_PACKET_PRINT_MODE, "single") == "double"


def _auto_email_enabled() -> bool:
    flag = (_get_setting(SETTING_EMAIL_AUTOMATION_ENABLED, "false") or "false").strip().lower() == "true"
    if not flag:
        return False
    if current_app.config.get("SENDGRID_API_KEY") and current_app.config.get("SENDGRID_FROM_EMAIL"):
        return True
    return bool(current_app.debug and current_app.config.get("EMAIL_DEV_FAKE_SEND"))


def _teacher_portal_enabled() -> bool:
    flag = _get_setting(SETTING_TEACHER_PORTAL_ENABLED, "false") or "false"
    return flag.lower() == "true"


def _login_folded_alias_enabled() -> bool:
    return _parse_bool(_get_setting(SETTING_LOGIN_FOLDED_ALIAS_ENABLED, "false"))


def _login_email_enabled() -> bool:
    return _parse_bool(_get_setting(SETTING_LOGIN_EMAIL_ENABLED, "false"))


def _registration_invite_required() -> bool:
    flag = _get_setting(SETTING_REGISTRATION_INVITE_ONLY, "false") or "false"
    return flag.lower() == "true"


def _is_upload_module_enabled() -> bool:
    flag = _get_setting(SETTING_UPLOAD_MODULE_ENABLED, "false") or "false"
    return flag.lower() == "true"


def _sendgrid_configured() -> bool:
    return bool(current_app.config.get("SENDGRID_API_KEY") and current_app.config.get("SENDGRID_FROM_EMAIL"))


def _get_email_subject(key: str) -> str:
    return _get_setting(key) or DEFAULT_EMAIL_SUBJECTS[key]


def _get_email_body(key: str) -> str:
    return _get_setting(key) or DEFAULT_EMAIL_BODIES[key]


class _PlainTextHTMLParser(HTMLParser):
    _BREAK_TAGS = {"br"}
    _BLOCK_TAGS = {
        "p",
        "div",
        "li",
        "section",
        "article",
        "table",
        "tr",
        "td",
        "th",
        "header",
        "footer",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
    }

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str]]) -> None:
        if tag.lower() in self._BREAK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in self._BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if data:
            self.parts.append(data)


def _html_to_text(value: str | None) -> str:
    if not value:
        return ""
    parser = _PlainTextHTMLParser()
    parser.feed(value)
    parser.close()
    text = unescape("".join(parser.parts))
    text = text.replace("\r\n", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    cleaned: list[str] = []
    previous_blank = False
    for line in lines:
        stripped = line.strip()
        if not stripped:
            if not previous_blank and cleaned:
                cleaned.append("")
            previous_blank = True
            continue
        cleaned.append(stripped)
        previous_blank = False
    return "\n".join(cleaned).strip()


def _send_email(
    to_email: str,
    subject: str,
    html_content: str,
    *,
    plain_text_content: str | None = None,
) -> tuple[bool, str | None, str | None]:
    if not to_email:
        return False, None, "missing-recipient"
    api_key = current_app.config.get("SENDGRID_API_KEY")
    from_email = current_app.config.get("SENDGRID_FROM_EMAIL")
    if not api_key or not from_email:
        if current_app.debug and current_app.config.get("EMAIL_DEV_FAKE_SEND") and not current_app.testing:
            current_app.logger.info(
                "DEV email fake-send: to=%s subject=%s",
                to_email,
                subject,
            )
            return True, "dev-fake-send", None
        current_app.logger.debug("SendGrid not configured; skipping email to %s", to_email)
        return False, None, "sendgrid-not-configured"
    try:
        sg = SendGridAPIClient(api_key)
        message = Mail(
            from_email=from_email,
            to_emails=to_email,
            subject=subject,
            html_content=html_content,
            plain_text_content=plain_text_content,
        )
        response = sg.send(message)
        headers = getattr(response, "headers", {}) or {}
        message_id = headers.get("X-Message-Id") or headers.get("x-message-id")
        return True, message_id, None
    except Exception as exc:
        current_app.logger.warning("SendGrid error sending to %s: %s", to_email, exc)
        return False, None, str(exc)


def _send_templated_email(to_email: str, subject_key: str, body_key: str, context: dict) -> None:
    if not _auto_email_enabled():
        return
    subject_template = _get_email_subject(subject_key)
    body_template = _get_email_body(body_key)
    subject = render_template_string(subject_template, **context)
    body = render_template_string(body_template, **context)
    _send_email(to_email, subject, body)


def _send_delegation_invite_email(invite: DelegationInvite, inviter: Teacher | None) -> tuple[bool, str | None]:
    if not _auto_email_enabled():
        return False, "auto-email-disabled"
    if not invite.email:
        return False, "missing-recipient"

    role_label = "student" if invite.role == DelegationInviteRole.STUDENT else "teacher"
    invite_link = url_for(
        "web.register_student" if invite.role == DelegationInviteRole.STUDENT else "web.register_teacher",
        invite=invite.token,
        _external=True,
    )
    if inviter and (inviter.first_name or inviter.last_name):
        inviter_name = f"{inviter.first_name} {inviter.last_name}".strip()
    elif inviter:
        inviter_name = "The delegation leader"
    else:
        inviter_name = "The organizing team"
    context = {
        "first_name": invite.first_name or "",
        "last_name": invite.last_name or "",
        "email": invite.email,
        "country": invite.country,
        "invite_link": invite_link,
        "inviter_name": inviter_name,
        "role_label": role_label,
    }
    subject_key = (
        SETTING_EMAIL_STUDENT_INVITE_SUBJECT
        if invite.role == DelegationInviteRole.STUDENT
        else SETTING_EMAIL_TEACHER_INVITE_SUBJECT
    )
    body_key = (
        SETTING_EMAIL_STUDENT_INVITE_BODY
        if invite.role == DelegationInviteRole.STUDENT
        else SETTING_EMAIL_TEACHER_INVITE_BODY
    )
    subject_template = _get_email_subject(subject_key)
    body_template = _get_email_body(body_key)
    subject = render_template_string(subject_template, **context)
    body = render_template_string(body_template, **context)
    _, _, error = _send_email(invite.email, subject, body)
    return error is None, error


def _send_student_confirmation_email(student: Student) -> None:
    if not _auto_email_enabled() or not student.email:
        return
    portal_password = getattr(student, "_plain_upload_password", None)
    context = {
        "first_name": student.first_name,
        "last_name": student.last_name,
        "email": student.email,
        "country": student.country,
        "school": student.school,
        "essay_language": student.essay_language,
        "keyboard_preference": student.keyboard_preference,
        "meal_preference": student.meal_preference,
        "special_needs": student.special_needs or "",
        "portal_username": student.upload_username,
        "portal_password": portal_password or "",
        "portal_link": url_for("web.student_portal", _external=True),
        "portal_login_link": url_for("web.student_login", _external=True),
    }
    _send_templated_email(student.email, SETTING_EMAIL_STUDENT_SUBJECT, SETTING_EMAIL_STUDENT_BODY, context)


def _send_teacher_confirmation_email(teacher: Teacher) -> None:
    if not _auto_email_enabled() or not teacher.email:
        return
    portal_password = getattr(teacher, "_plain_portal_password", None)
    context = {
        "first_name": teacher.first_name,
        "last_name": teacher.last_name,
        "email": teacher.email,
        "country": teacher.country,
        "school": teacher.school,
        "languages": teacher.languages,
        "meal_preference": teacher.meal_preference,
        "tshirt_size": teacher.tshirt_size,
        "special_needs": teacher.special_needs or "",
        "delegation_leader": teacher.delegation_leader,
        "prior_ipo_assessor": teacher.prior_ipo_assessor,
        "prior_ipo_assessor_level": teacher.prior_ipo_assessor_level,
        "portal_username": teacher.portal_username,
        "portal_password": portal_password or "",
        "portal_link": url_for("teacher_portal.dashboard", _external=True),
        "portal_login_link": url_for("teacher_portal.login", _external=True),
    }
    _send_templated_email(teacher.email, SETTING_EMAIL_TEACHER_SUBJECT, SETTING_EMAIL_TEACHER_BODY, context)


def _send_student_reset_link_email(student: Student, reset_token: PasswordResetToken) -> None:
    if not _auto_email_enabled() or not student.email:
        return
    context = {
        "first_name": student.first_name,
        "last_name": student.last_name,
        "reset_link": url_for("web.student_reset_password_token", token=reset_token.token, _external=True),
        "portal_login_link": url_for("web.student_login", _external=True),
    }
    _send_templated_email(
        student.email,
        SETTING_EMAIL_STUDENT_RESET_LINK_SUBJECT,
        SETTING_EMAIL_STUDENT_RESET_LINK_BODY,
        context,
    )


def _send_student_reset_email(student: Student, plain_password: str) -> None:
    if not _auto_email_enabled() or not student.email:
        return
    context = {
        "first_name": student.first_name,
        "last_name": student.last_name,
        "portal_username": student.upload_username,
        "portal_password": plain_password,
        "portal_link": url_for("web.student_portal", _external=True),
        "portal_login_link": url_for("web.student_login", _external=True),
    }
    _send_templated_email(student.email, SETTING_EMAIL_STUDENT_RESET_SUBJECT, SETTING_EMAIL_STUDENT_RESET_BODY, context)


def _send_teacher_reset_link_email(teacher: Teacher, reset_token: PasswordResetToken) -> None:
    if not _auto_email_enabled() or not teacher.email:
        return
    context = {
        "first_name": teacher.first_name,
        "last_name": teacher.last_name,
        "reset_link": url_for("teacher_portal.reset_token", token=reset_token.token, _external=True),
        "portal_login_link": url_for("teacher_portal.login", _external=True),
    }
    _send_templated_email(
        teacher.email,
        SETTING_EMAIL_TEACHER_RESET_LINK_SUBJECT,
        SETTING_EMAIL_TEACHER_RESET_LINK_BODY,
        context,
    )


def _send_teacher_reset_email(teacher: Teacher, plain_password: str) -> None:
    if not _auto_email_enabled() or not teacher.email:
        return
    context = {
        "first_name": teacher.first_name,
        "last_name": teacher.last_name,
        "portal_username": teacher.portal_username,
        "portal_password": plain_password,
        "portal_link": url_for("teacher_portal.dashboard", _external=True),
        "portal_login_link": url_for("teacher_portal.login", _external=True),
    }
    _send_templated_email(teacher.email, SETTING_EMAIL_TEACHER_RESET_SUBJECT, SETTING_EMAIL_TEACHER_RESET_BODY, context)


def _send_teacher_message_notifications(message: TeacherMessage) -> None:
    if not _auto_email_enabled():
        return
    recipients: list[Teacher] = []
    if message.deliver_to_all:
        recipients = (
            Teacher.query.filter(
                Teacher.email.isnot(None),
                Teacher.is_absent.is_(False),
            ).all()
        )
    elif message.teacher:
        recipients = [message.teacher]
    unique = {}
    for teacher in recipients:
        if not teacher or not teacher.email:
            continue
        if teacher.id in unique:
            continue
        unique[teacher.id] = teacher

    board_url = url_for("web.teacher_board", _external=True)
    for teacher in unique.values():
        context = {
            "teacher_first_name": teacher.first_name,
            "teacher_last_name": teacher.last_name,
            "teacher_email": teacher.email,
            "title": message.title,
            "body": message.body,
            "priority": message.priority.value,
            "board_url": board_url,
        }
        _send_templated_email(teacher.email, SETTING_EMAIL_MESSAGE_SUBJECT, SETTING_EMAIL_MESSAGE_BODY, context)


def _create_essay_overlay_page(
    page,
    group: ReadingGroup | None,
    assignment: EssayAssignment,
    *,
    header_text: str = "IPO Essay",
    labels: list[tuple[str, str]] | None = None,
) -> PdfReader:
    buffer = BytesIO()
    width = float(page.mediabox.right - page.mediabox.left)
    height = float(page.mediabox.top - page.mediabox.bottom)
    c = canvas.Canvas(buffer, pagesize=(width, height))
    margin = 36
    table_height = 36
    table_width = width - 2 * margin
    top = height - margin

    c.setStrokeColor(colors.black)
    c.setLineWidth(1)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(margin, top + 4, header_text)

    c.rect(margin, top - table_height, table_width, table_height)
    col_width = table_width / 3
    for i in range(1, 3):
        c.line(margin + col_width * i, top - table_height, margin + col_width * i, top)

    if labels is None:
        group_label = getattr(group, "name", None) if group is not None else None
        if not group_label:
            group_label = getattr(group, "id", None) if group is not None else None
        group_label = group_label or "—"
        essay_id = assignment.student.id if assignment.student else "—"
        labels = [
            ("Group", group_label),
            ("Essay", essay_id),
            ("Score", "     /10"),
        ]

    for idx, (label, value) in enumerate(labels):
        x = margin + col_width * idx + 6
        c.setFont("Helvetica-Bold", 10)
        c.drawString(x, top - 14, str(label))
        c.setFont("Helvetica", 10)
        c.drawString(x, top - 28, str(value))

    c.save()
    buffer.seek(0)
    return PdfReader(buffer)


def _planner_teacher_choices(stage: Stage, language: str | None = None) -> list[tuple[int, str]]:
    teachers = (
        _assessment_eligible_teachers_query()
        .order_by(Teacher.last_name.asc(), Teacher.first_name.asc())
        .all()
    )
    choices: list[tuple[int, str]] = []
    for teacher in teachers:
        _ensure_language_entries(teacher)
        languages = {entry.language for entry in teacher.language_entries}
        if language and language not in languages:
            continue
        display_langs = ", ".join(sorted(languages)) if languages else teacher.languages
        label = f"{teacher.first_name} {teacher.last_name} · {teacher.country} ({display_langs})"
        choices.append((teacher.id, label))
    return choices


def _teacher_has_active_group(teacher_id: int, stage: Stage, exclude_group_id: int | None = None) -> bool:
    query = (
        ReadingGroupMember.query.join(ReadingGroup)
        .filter(
            ReadingGroupMember.teacher_id == teacher_id,
            ReadingGroup.stage == stage,
            ReadingGroupMember.active.is_(True),
        )
    )
    if exclude_group_id is not None:
        query = query.filter(ReadingGroupMember.group_id != exclude_group_id)
    return db.session.query(query.exists()).scalar()


def _load_batch_groups(batch: AssignmentBatch) -> list[ReadingGroup]:
    return (
        ReadingGroup.query.options(
            joinedload(ReadingGroup.members).joinedload(ReadingGroupMember.teacher),
            joinedload(ReadingGroup.assignments).joinedload(EssayAssignment.student),
            joinedload(ReadingGroup.assignments)
            .joinedload(EssayAssignment.extra_readers)
            .joinedload(EssayAssignmentExtraReader.teacher),
            joinedload(ReadingGroup.assignments).joinedload(EssayAssignment.scores),
        )
        .filter(ReadingGroup.batch_id == batch.id)
        .order_by(ReadingGroup.language.asc(), ReadingGroup.name.asc())
        .all()
    )


def _reader_score_is_complete(score_record: EssayScore | None) -> bool:
    if not score_record or score_record.score is None:
        return False
    return bool(score_record.teacher_finalized or score_record.admin_entered or score_record.admin_override)


def _teacher_stage_one_assignment_ids_for_batch(batch: AssignmentBatch, teacher_id: int) -> set[int]:
    assignment_ids: set[int] = set()
    group_query = (
        db.session.query(EssayAssignment.id)
        .join(ReadingGroup, EssayAssignment.group_id == ReadingGroup.id)
        .join(ReadingGroupMember, ReadingGroupMember.group_id == ReadingGroup.id)
        .filter(
            EssayAssignment.batch_id == batch.id,
            EssayAssignment.stage == Stage.STAGE_1,
            ReadingGroupMember.teacher_id == teacher_id,
            ReadingGroupMember.active.is_(True),
        )
    )
    assignment_ids.update(assignment_id for (assignment_id,) in group_query.all())

    extra_query = (
        db.session.query(EssayAssignmentExtraReader.assignment_id)
        .join(EssayAssignment, EssayAssignment.id == EssayAssignmentExtraReader.assignment_id)
        .filter(
            EssayAssignment.batch_id == batch.id,
            EssayAssignment.stage == Stage.STAGE_1,
            EssayAssignmentExtraReader.teacher_id == teacher_id,
            EssayAssignmentExtraReader.stage == Stage.STAGE_1.value,
        )
    )
    assignment_ids.update(assignment_id for (assignment_id,) in extra_query.all())
    return assignment_ids


def _teacher_stage_one_workload(batch: AssignmentBatch, teacher: Teacher) -> SimpleNamespace:
    assignment_ids = _teacher_stage_one_assignment_ids_for_batch(batch, teacher.id)
    score_map: dict[int, EssayScore] = {}
    if assignment_ids:
        scores = (
            EssayScore.query.filter(
                EssayScore.assignment_id.in_(assignment_ids),
                EssayScore.teacher_id == teacher.id,
                EssayScore.stage == Stage.STAGE_1,
            )
            .order_by(EssayScore.id.asc())
            .all()
        )
        score_map = {score.assignment_id: score for score in scores if score.assignment_id}

    completed_count = sum(
        1
        for assignment_id in assignment_ids
        if _reader_score_is_complete(score_map.get(assignment_id))
    )
    assigned_count = len(assignment_ids)
    pending_count = max(assigned_count - completed_count, 0)
    return SimpleNamespace(
        assignment_ids=assignment_ids,
        assigned_count=assigned_count,
        completed_count=completed_count,
        pending_count=pending_count,
        is_free=pending_count == 0,
    )


def _supplemental_reader_assignments_for_group(group: ReadingGroup, teacher: Teacher) -> list[EssayAssignment]:
    eligible_assignments: list[EssayAssignment] = []
    assignments = sorted(
        group.assignments,
        key=lambda assignment: (
            assignment.student.country if assignment.student else "",
            assignment.student.id if assignment.student else 0,
        ),
    )
    for assignment in assignments:
        if assignment.stage != Stage.STAGE_1 or not assignment.student:
            continue
        assessors = _stage_assessors(assignment, Stage.STAGE_1, include_group=True)
        if len(assessors) >= STAGE_1_REQUIRED_READINGS:
            continue
        if any(entry.teacher_id == teacher.id for entry in assessors):
            continue
        if not _teacher_can_grade(assignment.student, teacher, assignment.student.essay_language):
            continue
        eligible_assignments.append(assignment)
    return eligible_assignments


def _supplemental_reader_context(
    batch: AssignmentBatch,
    groups: list[ReadingGroup],
    teachers: list[Teacher],
) -> dict[int, SimpleNamespace]:
    workload_map = {teacher.id: _teacher_stage_one_workload(batch, teacher) for teacher in teachers}

    group_map: dict[int, SimpleNamespace] = {}
    for group in groups:
        missing_slots = 0
        assignments_needing_readers = 0
        assigned_supplemental_counts: dict[int, int] = defaultdict(int)
        assigned_supplemental_teachers: dict[int, Teacher] = {}
        for assignment in group.assignments:
            if assignment.stage != Stage.STAGE_1:
                continue
            status = _compute_stage_one_status(assignment)
            if status.missing_reader_assignments > 0:
                assignments_needing_readers += 1
                missing_slots += status.missing_reader_assignments
            for assessor in _stage_assessors(assignment, Stage.STAGE_1, include_group=True):
                if not assessor.is_extra:
                    continue
                assigned_supplemental_counts[assessor.teacher_id] += 1
                assigned_supplemental_teachers[assessor.teacher_id] = assessor.teacher

        assigned_readers: list[SimpleNamespace] = []
        for teacher_id, assignment_count in assigned_supplemental_counts.items():
            teacher = assigned_supplemental_teachers[teacher_id]
            workload = workload_map.get(teacher.id)
            if not workload:
                workload = _teacher_stage_one_workload(batch, teacher)
                workload_map[teacher.id] = workload
            assigned_readers.append(
                SimpleNamespace(
                    teacher=teacher,
                    assignment_count=assignment_count,
                    workload=workload,
                )
            )
        assigned_readers.sort(key=lambda entry: (entry.teacher.last_name, entry.teacher.first_name, entry.teacher.id))

        candidates: list[SimpleNamespace] = []
        if missing_slots:
            for teacher in teachers:
                workload = workload_map.get(teacher.id)
                if not workload or not workload.is_free:
                    continue
                assignments = _supplemental_reader_assignments_for_group(group, teacher)
                if not assignments:
                    continue
                candidates.append(
                    SimpleNamespace(
                        teacher=teacher,
                        workload=workload,
                        eligible_assignments=assignments,
                        eligible_count=len(assignments),
                    )
                )
            candidates.sort(
                key=lambda entry: (
                    -entry.eligible_count,
                    -entry.workload.completed_count,
                    entry.teacher.last_name,
                    entry.teacher.first_name,
                    entry.teacher.id,
                )
            )

        group_map[group.id] = SimpleNamespace(
            missing_slots=missing_slots,
            assignments_needing_readers=assignments_needing_readers,
            assigned_readers=assigned_readers,
            candidates=candidates,
        )

    return group_map


def _planner_algorithm_for_batch(batch: AssignmentBatch) -> PlannerAlgorithm:
    algorithm = getattr(batch, "planner_algorithm", None)
    if isinstance(algorithm, PlannerAlgorithm):
        return algorithm
    if algorithm:
        try:
            return PlannerAlgorithm(algorithm)
        except ValueError:
            try:
                return PlannerAlgorithm[str(algorithm)]
            except KeyError:
                pass
    return PlannerAlgorithm.STRICT_FOUR_WITH_BACKUPS


def _planner_algorithm_choices() -> list[tuple[str, str]]:
    return [
        (PlannerAlgorithm.STRICT_FOUR_WITH_BACKUPS.value, "Strict 4 + backups"),
        (PlannerAlgorithm.ALL_ASSESSORS_NO_BACKUPS.value, "All assessors, no backups"),
    ]


def _remove_auto_generated_groups(batch: AssignmentBatch) -> tuple[int, int]:
    removed_groups = 0
    removed_assignments = 0
    for group in list(batch.groups):
        if group.auto_generated:
            for assignment in list(group.assignments):
                db.session.delete(assignment)
                removed_assignments += 1
            db.session.delete(group)
            removed_groups += 1

    db.session.flush()
    return removed_groups, removed_assignments


def _auto_generate_groups_for_batch(batch: AssignmentBatch) -> dict[str, int | dict[str, int]]:
    if _planner_algorithm_for_batch(batch) == PlannerAlgorithm.ALL_ASSESSORS_NO_BACKUPS:
        return _auto_generate_all_assessors_groups_for_batch(batch)
    return _auto_generate_strict_groups_for_batch(batch)


def _auto_generate_strict_groups_for_batch(batch: AssignmentBatch) -> dict[str, int | dict[str, int]]:
    removed_groups, removed_assignments = _remove_auto_generated_groups(batch)

    manual_teacher_ids: set[int] = set()
    manual_teacher_counts: dict[str, int] = defaultdict(int)
    manual_group_counts: dict[str, int] = defaultdict(int)
    for group in batch.groups:
        if group.stage != batch.stage:
            continue
        if not group.auto_generated:
            manual_group_counts[group.language] += 1
        for member in group.members:
            if not member.teacher or member.teacher.is_absent:
                continue
            manual_teacher_ids.add(member.teacher_id)
            if not group.auto_generated and member.active:
                manual_teacher_counts[group.language] += 1

    student_counts: dict[str, int] = defaultdict(int)
    for language, in Student.query.filter_by(is_absent=False).with_entities(Student.essay_language):
        if language:
            student_counts[language] += 1

    teachers = _assessment_eligible_teachers_query().order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()

    teacher_lookup: dict[int, Teacher] = {}
    eligible_teachers: list[Teacher] = []
    teacher_languages_map: dict[int, set[str]] = {}
    language_capacity: dict[str, int] = defaultdict(int)

    for teacher in teachers:
        teacher_lookup[teacher.id] = teacher
        if teacher.id in manual_teacher_ids:
            continue
        _ensure_language_entries(teacher)
        if _teacher_has_active_group(teacher.id, batch.stage):
            continue
        languages_spoken = {entry.language for entry in teacher.language_entries if entry.language}
        if not languages_spoken:
            continue
        eligible_teachers.append(teacher)
        teacher_languages_map[teacher.id] = languages_spoken
        for language in languages_spoken:
            language_capacity[language] += 1

    total_eligible = len(eligible_teachers)
    full_groups = total_eligible // 4
    remainder = total_eligible % 4
    slots_to_assign = full_groups * 4

    summary: dict[str, int | dict[str, int]] = {"removed": removed_groups, "removed_assignments": removed_assignments}

    if full_groups == 0:
        summary["created"] = {}
        summary["backups"] = total_eligible
        _recalculate_group_targets(batch, student_counts)
        return summary

    min_per_teacher_load = max(current_app.config.get("PLANNER_MIN_PER_TEACHER_LOAD", 3), 1)
    languages = set(student_counts.keys()) | set(language_capacity.keys())

    language_group_capacity: dict[str, int] = {language: language_capacity.get(language, 0) // 4 for language in languages}
    language_min_groups: dict[str, int] = {}
    shortfalls: dict[str, int] = {}
    total_min_groups = 0

    for language in languages:
        capacity_groups = language_group_capacity.get(language, 0)
        if capacity_groups == 0:
            language_min_groups[language] = 0
            if student_counts.get(language, 0) > 0:
                required_slots = max(4, math.ceil(student_counts[language] / min_per_teacher_load))
                manual_slots = manual_teacher_counts.get(language, 0)
                needed_slots = max(required_slots - manual_slots, 0)
                if needed_slots > 0:
                    shortfalls[language] = needed_slots
            continue

        required_slots = 0
        if student_counts.get(language, 0) > 0:
            required_slots = max(4, math.ceil(student_counts[language] / min_per_teacher_load))
        manual_slots = manual_teacher_counts.get(language, 0)
        needed_slots = max(required_slots - manual_slots, 0)
        required_groups = math.ceil(needed_slots / 4) if needed_slots else 0
        min_groups = min(required_groups, capacity_groups)
        language_min_groups[language] = min_groups
        total_min_groups += min_groups
        if required_groups > capacity_groups:
            shortfalls[language] = (required_groups - capacity_groups) * 4

    if total_min_groups > full_groups:
        summary["created"] = {}
        summary["backups"] = total_eligible
        if shortfalls:
            summary["shortfall"] = shortfalls
        summary["matching_failed"] = True
        _recalculate_group_targets(batch, student_counts)
        return summary

    language_group_plan: dict[str, int] = {language: language_min_groups.get(language, 0) for language in languages}
    assigned_groups = sum(language_group_plan.values())
    remaining_groups = full_groups - assigned_groups

    def _best_language_for_extra_group() -> str | None:
        best_language = None
        best_score = None
        for language in languages:
            capacity = language_group_capacity.get(language, 0)
            current = language_group_plan.get(language, 0)
            if current >= capacity:
                continue
            manual_slots = manual_teacher_counts.get(language, 0)
            assigned_slots = current * 4
            score = student_counts.get(language, 0) / max(1, manual_slots + assigned_slots)
            if best_language is None or score > best_score or (
                math.isclose(score, best_score or 0.0) and student_counts.get(language, 0) > student_counts.get(best_language, 0)
            ):
                best_language = language
                best_score = score
        return best_language

    while remaining_groups > 0:
        language = _best_language_for_extra_group()
        if language is None:
            break
        language_group_plan[language] = language_group_plan.get(language, 0) + 1
        remaining_groups -= 1

    if sum(language_group_plan.values()) < full_groups:
        summary["created"] = {}
        summary["backups"] = total_eligible
        if shortfalls:
            summary["shortfall"] = shortfalls
        summary["matching_failed"] = True
        _recalculate_group_targets(batch, student_counts)
        return summary

    def _group_plan_to_slots(plan: dict[str, int]) -> dict[str, int]:
        return {language: count * 4 for language, count in plan.items() if count > 0}

    language_slot_plan = _group_plan_to_slots(language_group_plan)
    planned_slot_total = sum(language_slot_plan.values())
    if planned_slot_total != slots_to_assign:
        summary["created"] = {}
        summary["backups"] = total_eligible
        if shortfalls:
            summary["shortfall"] = shortfalls
        summary["matching_failed"] = True
        _recalculate_group_targets(batch, student_counts)
        return summary

    def _match_teachers_to_slots(plan: dict[str, int]) -> tuple[dict[str, list[int]] | None, dict[str, int]]:
        slots: list[str] = []
        slot_language: dict[str, str] = {}
        for language, count in plan.items():
            for idx in range(count):
                slot_id = f"{language}#{idx}"
                slots.append(slot_id)
                slot_language[slot_id] = language

        adjacency: dict[int, list[str]] = {}
        for teacher in eligible_teachers:
            langs = teacher_languages_map.get(teacher.id, set())
            adjacency[teacher.id] = [slot for slot in slots if slot_language[slot] in langs]

        match_slot: dict[str, int] = {}
        match_teacher: dict[int, str] = {}

        teacher_order = sorted(
            eligible_teachers,
            key=lambda t: (len(adjacency.get(t.id, [])), t.id),
        )

        def dfs(teacher_id: int, seen: set[str]) -> bool:
            for slot in adjacency.get(teacher_id, []):
                if slot in seen:
                    continue
                seen.add(slot)
                if slot not in match_slot or dfs(match_slot[slot], seen):
                    match_slot[slot] = teacher_id
                    match_teacher[teacher_id] = slot
                    return True
            return False

        for teacher in teacher_order:
            dfs(teacher.id, set())

        if len(match_slot) < len(slots):
            return None, {}

        assignments: dict[str, list[int]] = defaultdict(list)
        for slot, teacher_id in match_slot.items():
            assignments[slot_language[slot]].append(teacher_id)
        return assignments, {language: plan.get(language, 0) for language in plan}

    assignments, _ = _match_teachers_to_slots(language_slot_plan)
    adjustment_attempts = 0
    while assignments is None and adjustment_attempts < len(language_group_plan) + 5:
        reducible = [
            language
            for language, group_count in language_group_plan.items()
            if group_count > language_min_groups.get(language, 0)
        ]
        if not reducible:
            break
        reducible.sort(
            key=lambda language: (
                -(language_group_plan.get(language, 0) - language_min_groups.get(language, 0)),
                student_counts.get(language, 0),
                language,
            )
        )
        language_to_reduce = reducible[0]
        language_group_plan[language_to_reduce] -= 1

        while sum(language_group_plan.values()) < full_groups:
            candidate = _best_language_for_extra_group()
            if candidate is None:
                break
            if language_group_plan.get(candidate, 0) >= language_group_capacity.get(candidate, 0):
                break
            language_group_plan[candidate] = language_group_plan.get(candidate, 0) + 1

        if sum(language_group_plan.values()) < full_groups:
            break

        language_slot_plan = _group_plan_to_slots(language_group_plan)
        assignments, _ = _match_teachers_to_slots(language_slot_plan)
        adjustment_attempts += 1

    if assignments is None:
        summary["created"] = {}
        summary["backups"] = total_eligible
        if shortfalls:
            summary["shortfall"] = shortfalls
        summary["matching_failed"] = True
        _recalculate_group_targets(batch, student_counts)
        return summary

    created_count: dict[str, int] = defaultdict(int)
    used_teacher_ids: set[int] = set()

    def _build_groups_for_language(language: str, teacher_ids: list[int]) -> list[list[Teacher]]:
        teachers_in_language = [teacher_lookup[tid] for tid in teacher_ids]
        teachers_in_language.sort(key=lambda t: (t.country, t.last_name, t.first_name, t.id))
        groups: list[list[Teacher]] = []
        pool = teachers_in_language
        while len(pool) >= 4:
            selected: list[Teacher] = []
            used_countries: set[str] = set()
            remainder_list: list[Teacher] = []
            for teacher in pool:
                if len(selected) == 4:
                    remainder_list.append(teacher)
                    continue
                if teacher.country not in used_countries:
                    selected.append(teacher)
                    used_countries.add(teacher.country)
                else:
                    remainder_list.append(teacher)
            if len(selected) < 4:
                needed = 4 - len(selected)
                selected.extend(remainder_list[:needed])
                remainder_list = remainder_list[needed:]
            groups.append(selected)
            pool = remainder_list
        return groups

    for language, teacher_ids in assignments.items():
        if not teacher_ids:
            continue
        language_groups = _build_groups_for_language(language, teacher_ids)
        for members in language_groups:
            group_name = f"{language} Group {created_count[language] + manual_group_counts.get(language, 0) + 1}"
            group = ReadingGroup(
                name=group_name,
                language=language,
                stage=batch.stage,
                target_essay_load=0,
                allow_same_country=False,
                auto_generated=True,
                status=batch.status,
                batch=batch,
            )
            db.session.add(group)
            db.session.flush()
            for idx, teacher in enumerate(members):
                role = GroupMemberRole.LEADER if idx == 0 else GroupMemberRole.READER
                member = ReadingGroupMember(
                    group_id=group.id,
                    teacher_id=teacher.id,
                    role=role,
                    active=True,
                )
                db.session.add(member)
                used_teacher_ids.add(teacher.id)
            created_count[language] += 1

    matched_teacher_ids = used_teacher_ids.copy()
    backup_teachers = [teacher for teacher in eligible_teachers if teacher.id not in matched_teacher_ids]

    if remainder == 0:
        backup_teachers = []

    if backup_teachers:
        target_groups = [
            group
            for group in batch.groups
            if group.stage == batch.stage and group.auto_generated and not db.inspect(group).deleted
        ]
        target_groups.sort(key=lambda g: (g.language, g.id))
        if target_groups:
            for idx, teacher in enumerate(backup_teachers):
                group = target_groups[idx % len(target_groups)]
                member = ReadingGroupMember(
                    group_id=group.id,
                    teacher_id=teacher.id,
                    role=GroupMemberRole.BACKUP,
                    active=False,
                )
                db.session.add(member)
                used_teacher_ids.add(teacher.id)

    summary_created = {language: count for language, count in created_count.items() if count}
    summary["created"] = summary_created
    summary["backups"] = len(backup_teachers)
    if shortfalls:
        summary["shortfall"] = shortfalls

    _recalculate_group_targets(batch, student_counts)
    return summary


def _auto_generate_all_assessors_groups_for_batch(batch: AssignmentBatch) -> dict[str, int | dict[str, int]]:
    removed_groups, removed_assignments = _remove_auto_generated_groups(batch)

    manual_teacher_ids: set[int] = set()
    manual_teacher_counts: dict[str, int] = defaultdict(int)
    manual_group_counts: dict[str, int] = defaultdict(int)
    for group in batch.groups:
        if group.stage != batch.stage:
            continue
        if not group.auto_generated:
            manual_group_counts[group.language] += 1
        for member in group.members:
            if not member.teacher or member.teacher.is_absent:
                continue
            manual_teacher_ids.add(member.teacher_id)
            if not group.auto_generated and member.active:
                manual_teacher_counts[group.language] += 1

    student_counts: dict[str, int] = defaultdict(int)
    for language, in Student.query.filter_by(is_absent=False).with_entities(Student.essay_language):
        if language:
            student_counts[language] += 1

    teachers = _assessment_eligible_teachers_query().order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()

    eligible_teachers: list[Teacher] = []
    teacher_languages_map: dict[int, set[str]] = {}
    for teacher in teachers:
        if teacher.id in manual_teacher_ids:
            continue
        _ensure_language_entries(teacher)
        if _teacher_has_active_group(teacher.id, batch.stage):
            continue
        languages_spoken = {entry.language for entry in teacher.language_entries if entry.language}
        if not languages_spoken:
            continue
        eligible_teachers.append(teacher)
        teacher_languages_map[teacher.id] = languages_spoken

    summary: dict[str, int | dict[str, int]] = {
        "removed": removed_groups,
        "removed_assignments": removed_assignments,
        "backups": 0,
    }

    if not eligible_teachers:
        summary["created"] = {}
        summary["teachers_assigned"] = 0
        summary["underfilled"] = 0
        _recalculate_group_targets(batch, student_counts)
        return summary

    min_per_teacher_load = max(current_app.config.get("PLANNER_MIN_PER_TEACHER_LOAD", 3), 1)
    planned_teacher_counts: dict[str, int] = defaultdict(int, manual_teacher_counts)
    teachers_by_language: dict[str, list[Teacher]] = defaultdict(list)

    def _teacher_sort_key(teacher: Teacher) -> tuple[int, int, str, str, int]:
        languages = teacher_languages_map.get(teacher.id, set())
        demand_languages = [language for language in languages if student_counts.get(language, 0) > 0]
        useful_options = demand_languages or list(languages)
        max_demand = max((student_counts.get(language, 0) for language in useful_options), default=0)
        return (len(useful_options), -max_demand, teacher.last_name, teacher.first_name, teacher.id)

    def _language_sort_key(language: str) -> tuple[int, int, float, int, int, str]:
        students = student_counts.get(language, 0)
        current_teachers = planned_teacher_counts.get(language, 0)
        required_teachers = math.ceil(students / min_per_teacher_load) if students else 0
        deficit = max(required_teachers - current_teachers, 0)
        load_score = students / max(current_teachers + 1, 1)
        return (
            0 if students else 1,
            -deficit,
            -load_score,
            -students,
            current_teachers,
            language,
        )

    for teacher in sorted(eligible_teachers, key=_teacher_sort_key):
        languages = sorted(teacher_languages_map.get(teacher.id, set()))
        languages_with_students = [language for language in languages if student_counts.get(language, 0) > 0]
        candidate_languages = languages_with_students or languages
        chosen_language = sorted(candidate_languages, key=_language_sort_key)[0]
        teachers_by_language[chosen_language].append(teacher)
        planned_teacher_counts[chosen_language] += 1

    def _build_groups_for_language(teachers_in_language: list[Teacher]) -> list[list[Teacher]]:
        pool = sorted(teachers_in_language, key=lambda t: (t.country, t.last_name, t.first_name, t.id))
        groups: list[list[Teacher]] = []
        while pool:
            selected: list[Teacher] = []
            used_countries: set[str] = set()
            remainder_list: list[Teacher] = []
            for teacher in pool:
                if len(selected) == 4:
                    remainder_list.append(teacher)
                    continue
                if teacher.country not in used_countries:
                    selected.append(teacher)
                    used_countries.add(teacher.country)
                else:
                    remainder_list.append(teacher)
            if len(selected) < 4:
                needed = min(4 - len(selected), len(remainder_list))
                selected.extend(remainder_list[:needed])
                remainder_list = remainder_list[needed:]
            groups.append(selected)
            pool = remainder_list
        return groups

    created_count: dict[str, int] = defaultdict(int)
    underfilled_groups = 0
    no_student_language_groups = 0

    for language in sorted(teachers_by_language):
        language_groups = _build_groups_for_language(teachers_by_language[language])
        for members in language_groups:
            group_name = f"{language} Group {created_count[language] + manual_group_counts.get(language, 0) + 1}"
            group = ReadingGroup(
                name=group_name,
                language=language,
                stage=batch.stage,
                target_essay_load=0,
                allow_same_country=False,
                auto_generated=True,
                status=batch.status,
                batch=batch,
            )
            db.session.add(group)
            db.session.flush()
            for idx, teacher in enumerate(members):
                role = GroupMemberRole.LEADER if idx == 0 else GroupMemberRole.READER
                member = ReadingGroupMember(
                    group_id=group.id,
                    teacher_id=teacher.id,
                    role=role,
                    active=True,
                )
                db.session.add(member)
            if len(members) < 4:
                underfilled_groups += 1
            if student_counts.get(language, 0) == 0:
                no_student_language_groups += 1
            created_count[language] += 1

    summary["created"] = {language: count for language, count in created_count.items() if count}
    summary["teachers_assigned"] = len(eligible_teachers)
    summary["underfilled"] = underfilled_groups
    if no_student_language_groups:
        summary["no_student_language_groups"] = no_student_language_groups

    _recalculate_group_targets(batch, student_counts)
    return summary


def _recalculate_group_targets(batch: AssignmentBatch, student_counts: dict[str, int] | None = None) -> None:
    if student_counts is None:
        counts: dict[str, int] = defaultdict(int)
        for language, in Student.query.filter_by(is_absent=False).with_entities(Student.essay_language):
            if language:
                counts[language] += 1
        student_counts = counts

    min_per_teacher_load = max(current_app.config.get("PLANNER_MIN_PER_TEACHER_LOAD", 3), 1)

    groups_by_language: dict[str, list[ReadingGroup]] = defaultdict(list)
    for group in batch.groups:
        if group.stage != batch.stage:
            continue
        groups_by_language[group.language].append(group)

    for language, groups in groups_by_language.items():
        total_students = student_counts.get(language, 0)
        active_counts: dict[int, int] = {}
        total_active_teachers = 0

        for group in groups:
            active_members = [
                member
                for member in group.members
                if member.active and member.teacher and not member.teacher.is_absent
            ]
            active_count = len(active_members)
            if active_count == 0:
                group.target_essay_load = 0
                continue
            active_counts[group.id] = active_count
            total_active_teachers += active_count

        if total_active_teachers == 0:
            continue

        if total_students == 0:
            for group in groups:
                if active_counts.get(group.id):
                    group.target_essay_load = 0
            continue

        per_teacher_load = max(math.ceil(total_students / total_active_teachers), min_per_teacher_load)

        for group in groups:
            active_count = active_counts.get(group.id)
            if not active_count:
                continue
            group.target_essay_load = max(per_teacher_load * active_count, 1)


def _auto_assign_essays(batch: AssignmentBatch) -> tuple[int, list[Student]]:
    # remove existing non-manual assignments
    for assignment in list(batch.assignments):
        if assignment.manual_override and assignment.group_id is None:
            db.session.delete(assignment)
            continue
        if assignment.manual_override:
            continue
        db.session.delete(assignment)
    db.session.flush()
    _recalculate_group_targets(batch)

    groups = [group for group in batch.groups if group.members]
    if not groups:
        return 0, []

    min_active_members = 1
    if _planner_algorithm_for_batch(batch) == PlannerAlgorithm.STRICT_FOUR_WITH_BACKUPS:
        min_active_members = 4

    group_teacher_countries: dict[int, set[str]] = {}
    for group in groups:
        countries = {
            member.teacher.country
            for member in group.members
            if member.active and member.teacher and not member.teacher.is_absent
        }
        group_teacher_countries[group.id] = countries

    group_assignment_counts: dict[int, int] = defaultdict(int)
    for assignment in batch.assignments:
        if assignment.group_id:
            group_assignment_counts[assignment.group_id] += 1

    students = (
        Student.query.filter_by(is_absent=False)
        .order_by(Student.essay_language.asc(), Student.country.asc())
        .all()
    )
    skipped: list[Student] = []
    assigned = 0

    for student in students:
        manual_assignment = next(
            (a for a in batch.assignments if a.student_id == student.id and a.manual_override and a.group_id),
            None,
        )
        if manual_assignment:
            continue

        candidate_groups = []
        for group in groups:
            if group.language != student.essay_language:
                continue
            if group.locked:
                continue
            active_members = [
                member
                for member in group.members
                if member.active and member.teacher and not member.teacher.is_absent
            ]
            if len(active_members) < min_active_members:
                continue
            current_load = group_assignment_counts.get(group.id, 0)
            if current_load >= group.target_essay_load:
                continue
            countries = group_teacher_countries.get(group.id, set())
            if not group.allow_same_country and student.country in countries:
                continue
            if _teacher_has_active_group_conflict(group, student.country):
                continue
            candidate_groups.append(group)

        if not candidate_groups:
            skipped.append(student)
            continue

        candidate_groups.sort(
            key=lambda grp: (
                group_assignment_counts.get(grp.id, 0),
                grp.target_essay_load,
                grp.id,
            )
        )
        chosen_group = candidate_groups[0]
        assignment = EssayAssignment(
            student_id=student.id,
            group_id=chosen_group.id,
            batch=batch,
            stage=batch.stage,
            manual_override=False,
            status=PlannerStatus.DRAFT,
        )
        db.session.add(assignment)
        group_assignment_counts[chosen_group.id] += 1
        assigned += 1

    return assigned, skipped


def _teacher_has_active_group_conflict(group: ReadingGroup, student_country: str) -> bool:
    if group.allow_same_country:
        return False
    for member in group.members:
        if not member.active or not member.teacher or member.teacher.is_absent:
            continue
        if member.teacher.country == student_country:
            return True
    return False


def _render_pdf(template_name: str, filename: str, **context):
    html = render_template(template_name, **context)
    pdf_bytes = HTML(string=html, base_url=request.url_root).write_pdf()
    response = make_response(pdf_bytes)
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = f"attachment; filename={filename}"
    return response


def _run_mysqldump(output_path: Path) -> bool:
    database_url = _get_database_url()
    if not database_url or "mysql" not in database_url:
        return False

    url = make_url(database_url)
    env = os.environ.copy()
    if url.password:
        env["MYSQL_PWD"] = url.password

    cmd = [
        "mysqldump",
        f"--host={url.host or 'localhost'}",
        f"--port={url.port or 3306}",
        f"--user={url.username or ''}",
        url.database or "",
    ]

    import subprocess

    with output_path.open("wb") as fh:
        subprocess.check_call(cmd, env=env, stdout=fh)
    return True


def _get_database_url() -> str | None:
    return current_app.config.get("DATABASE_URL") or current_app.config.get("SQLALCHEMY_DATABASE_URI")


def _export_database() -> tuple[Path, str] | None:
    backups_dir = Path(current_app.config.get("BACKUP_FOLDER", Path(current_app.instance_path) / "backups"))
    backups_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    filename = f"ipo-db-{timestamp}.sql"
    output_path = backups_dir / filename

    database_url = _get_database_url() or ""
    if database_url.startswith("sqlite"):
        db_path = database_url.replace("sqlite://", "")
        db_file = Path(db_path)
        if not db_file.exists():
            return None
        import shutil

        shutil.copyfile(db_file, output_path)
        return output_path, filename

    if "mysql" in database_url:
        if _run_mysqldump(output_path):
            return output_path, filename

    return None


def _zip_essays(only_finalized: bool = False) -> tuple[BytesIO, str] | None:
    base_folder = Path(_get_upload_folder())
    if not base_folder.exists():
        return None

    timestamp = datetime.utcnow().strftime("%Y%m%d-%H%M%S")
    zip_buffer = BytesIO()
    import zipfile
    import csv as csv_module

    students_query = Student.query
    if only_finalized:
        students_query = students_query.filter(Student.essay_finalized.is_(True))
    students = students_query.order_by(Student.id.asc()).all()

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        manifest_rows = []
        for student in students:
            pdf_path = _essay_pdf_path(student)
            if not pdf_path or not pdf_path.exists():
                continue
            safe_name = f"{student.id}-{student.last_name}_{student.first_name}.pdf"
            manifest_rows.append(
                {
                    "student_id": student.id,
                    "first_name": student.first_name,
                    "last_name": student.last_name,
                    "country": student.country,
                    "uploaded_at": student.uploaded_at.isoformat() if student.uploaded_at else "",
                    "finalized": "yes" if student.essay_finalized else "no",
                    "filename": Path(pdf_path).name,
                }
            )
            archive.write(pdf_path, f"essays/{safe_name}")

        if not manifest_rows:
            return None

        manifest_bytes = io.StringIO()
        fieldnames = ["student_id", "first_name", "last_name", "country", "uploaded_at", "finalized", "filename"]
        writer = csv_module.DictWriter(manifest_bytes, fieldnames=fieldnames)
        writer.writeheader()
        for row in manifest_rows:
            writer.writerow(row)
        archive.writestr("manifest.csv", manifest_bytes.getvalue())

    zip_buffer.seek(0)
    filename = f"ipo-essays-{timestamp}.zip"
    return zip_buffer, filename


def _build_teacher_schedules(groups: list[ReadingGroup]) -> list[dict]:
    schedule_map: dict[int, dict] = {}
    for group in groups:
        assignments = sorted(
            group.assignments,
            key=lambda assignment: (
                assignment.student.essay_language if assignment.student else "",
                assignment.student.country if assignment.student else "",
            ),
        )
        for member in group.members:
            if not member.active or not member.teacher:
                continue
            teacher = member.teacher
            entry = schedule_map.setdefault(
                teacher.id,
                {
                    "teacher": teacher,
                    "group": group,
                    "role": member.role,
                    "assignments": [],
                },
            )
            entry["assignments"] = assignments
            entry["role"] = member.role

    schedules = list(schedule_map.values())
    schedules.sort(key=lambda item: (item["teacher"].last_name, item["teacher"].first_name))
    return schedules


def _teacher_stage_one_assignments_for_batch(batch: AssignmentBatch, teacher: Teacher) -> list[EssayAssignment]:
    assignment_ids = _teacher_stage_one_assignment_ids_for_batch(batch, teacher.id)
    if not assignment_ids:
        return []
    assignments = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.student),
            joinedload(EssayAssignment.group),
            joinedload(EssayAssignment.extra_readers).joinedload(EssayAssignmentExtraReader.teacher),
        )
        .filter(EssayAssignment.id.in_(assignment_ids))
        .all()
    )
    assignments.sort(
        key=lambda assignment: (
            assignment.group.language if assignment.group else "",
            assignment.group.name if assignment.group else "",
            assignment.student.country if assignment.student else "",
            assignment.student.id if assignment.student else 0,
        )
    )
    return assignments


def _build_teacher_reader_packet(teacher: Teacher, assignments: list[EssayAssignment]) -> tuple[BytesIO, list[int]]:
    writer = PdfWriter()
    double_sided = _packets_double_sided()
    missing_files: list[int] = []
    packet_group = SimpleNamespace(name="Individual packet", id="—")
    essay_ids = [assignment.student.id for assignment in assignments if assignment.student]

    title_reader = _create_teacher_title_page(teacher, packet_group, essay_ids)
    writer.add_page(title_reader.pages[0])
    if double_sided and assignments:
        _add_blank_page(writer, title_reader.pages[0])

    for assignment in assignments:
        student = assignment.student
        if not student:
            continue
        pdf_path = _essay_pdf_path(student)
        if not pdf_path:
            missing_files.append(student.id)
            continue
        try:
            reader = PdfReader(str(pdf_path))
        except Exception:
            missing_files.append(student.id)
            continue

        overlay_reader = _create_essay_overlay_page(reader.pages[0], assignment.group, assignment)
        first_page = reader.pages[0]
        first_page.merge_page(overlay_reader.pages[0])
        writer.add_page(first_page)
        for page in reader.pages[1:]:
            writer.add_page(page)

        if double_sided and len(reader.pages) % 2 != 0:
            _add_blank_page(writer, reader.pages[0])

    output = BytesIO()
    writer.write(output)
    output.seek(0)
    return output, missing_files


@web_bp.route("/")
def index():
    return render_template("index.html", invite_only=_registration_invite_required())


@web_bp.route("/manuals/<slug>/", defaults={"filename": "index.html"})
@web_bp.route("/manuals/<slug>/<path:filename>")
def manuals(slug: str, filename: str):
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]*", slug or ""):
        abort(404)
    manuals_dir = Path(current_app.static_folder) / "manuals" / slug
    if not manuals_dir.is_dir():
        abort(404)
    return send_from_directory(manuals_dir, filename)


@web_bp.route("/pages/<string:slug>")
def static_page_show(slug: str):
    normalized_slug = _normalize_static_page_slug(slug)
    if normalized_slug != slug:
        return redirect(url_for("web.static_page_show", slug=normalized_slug), code=301)

    page = (
        StaticPage.query.options(joinedload(StaticPage.attachments))
        .filter_by(slug=normalized_slug)
        .first()
    )
    if not page:
        abort(404)
    return render_template("static_pages/show.html", page=page)


@web_bp.route("/pages/<string:slug>/attachments/<int:attachment_id>")
def static_page_attachment_download(slug: str, attachment_id: int):
    normalized_slug = _normalize_static_page_slug(slug)
    page = (
        StaticPage.query.options(joinedload(StaticPage.attachments))
        .filter_by(slug=normalized_slug)
        .first()
    )
    if not page:
        abort(404)

    attachment = next((item for item in page.attachments if item.id == attachment_id), None)
    if not attachment:
        abort(404)

    file_path = _static_page_attachment_path(attachment.stored_filename)
    if not file_path or not file_path.exists():
        abort(404)

    guessed_type, _ = mimetypes.guess_type(attachment.original_filename or "")
    content_type = attachment.content_type or guessed_type or "application/octet-stream"
    response = send_file(
        file_path,
        mimetype=content_type,
        as_attachment=False,
        download_name=attachment.original_filename,
    )
    response.headers["Cache-Control"] = "public, max-age=3600"
    return response


@web_bp.route("/admin/login", methods=["GET", "POST"])
@limiter.limit(LOGIN_RATE_LIMIT_PER_IP, key_func=_rate_limit_ip, methods=["POST"])
@limiter.limit(LOGIN_RATE_LIMIT_PER_USERNAME, key_func=_rate_limit_username, methods=["POST"])
def admin_login():
    if _current_admin():
        next_url = _portal_next_url(request.args.get("next"))
        return redirect(next_url or url_for("web.admin_dashboard"))

    form = AdminLoginForm()
    next_url = _portal_next_url(request.args.get("next") or request.form.get("next"))
    credentials = _parse_admin_users()
    if not credentials:
        flash("No admin accounts configured. Please set ADMIN_USERS.", "error")

    if form.validate_on_submit() and credentials:
        username = form.username.data.strip()
        if _validate_admin_credentials(username, form.password.data):
            session["admin_user"] = username
            session.permanent = True
            log_activity("admin_login", f"{username} signed in", actor=username)
            db.session.commit()
            flash("Signed in successfully.", "success")
            return redirect(next_url or url_for("web.admin_dashboard"))
        else:
            log_activity("admin_login_failed", f"Failed login attempt for {username}", actor=username)
            db.session.commit()
            flash("Invalid credentials.", "error")

    return render_template("admin/login.html", form=form, next_url=next_url)


@web_bp.route("/admin/logout")
def admin_logout():
    username = _current_admin()
    if username:
        log_activity("admin_logout", f"{username} signed out", actor=username)
        db.session.commit()
    session.pop("admin_user", None)
    flash("Signed out.", "success")
    return redirect(url_for("web.admin_login"))


@web_bp.route("/participant/login", methods=["GET", "POST"])
@limiter.limit(LOGIN_RATE_LIMIT_PER_IP, key_func=_rate_limit_ip, methods=["POST"])
@limiter.limit(LOGIN_RATE_LIMIT_PER_USERNAME, key_func=_rate_limit_username, methods=["POST"])
def student_login():
    form = StudentLoginForm()
    if form.submit.name in request.form and form.validate_on_submit():
        username = form.username.data.strip()
        student, match_details = _resolve_login_alias_match(Student, "upload_username", "email", username)
        if not student or not student.upload_password_hash:
            flash("Invalid credentials.", "error")
            if match_details["match_type"] == "ambiguous":
                log_activity(
                    "student_upload_login_ambiguous",
                    f"Ambiguous folded login for {username}",
                    details={
                        "submitted_username": username,
                        "folded_username": match_details.get("folded_username"),
                        "match_count": match_details.get("match_count", 0),
                    },
                )
                db.session.commit()
        elif not check_password_hash(student.upload_password_hash, form.password.data.strip()):
            flash("Invalid credentials.", "error")
            log_activity(
                "student_upload_login_failed",
                f"Failed login for {username}",
                details={
                    "username": student.upload_username,
                    "email": student.email,
                    "submitted_username": username,
                    "login_match_type": match_details["match_type"],
                },
            )
            db.session.commit()
        else:
            session["student_upload_id"] = student.id
            log_activity(
                "student_upload_login",
                f"Student {student.first_name} {student.last_name} signed in",
                object_type="student",
                object_id=student.id,
                details={
                    "username": student.upload_username,
                    "email": student.email,
                    "submitted_username": username,
                    "login_match_type": match_details["match_type"],
                },
            )
            db.session.commit()
            next_param = _portal_next_url(request.args.get("next"))
            target = next_param or url_for("web.student_portal")
            return redirect(target)
    return render_template(
        "student/login.html",
        form=form,
        login_email_enabled=_login_email_enabled(),
        upload_module_enabled=_is_upload_module_enabled(),
        upload_locked=_is_upload_locked(),
    )


@web_bp.route("/participant/reset", methods=["GET", "POST"])
def student_reset_password_request():
    form = PasswordResetRequestForm()
    submitted = False
    email_sent = False
    token_generated = False
    if form.submit.name in request.form and form.validate_on_submit():
        email = form.email.data.lower().strip()
        student = Student.query.filter_by(email=email, is_absent=False).first()
        if student:
            reset_token = _create_student_reset_token(student)
            db.session.flush()
            log_activity(
                "student_password_reset_requested",
                f"Student {student.first_name} {student.last_name} requested portal reset",
                object_type="student",
                object_id=student.id,
                details={"email": student.email, "reset_token_id": reset_token.id},
            )
            db.session.commit()
            token_generated = True
            if _auto_email_enabled() and student.email:
                _send_student_reset_link_email(student, reset_token)
                email_sent = True
        submitted = True
    return render_template(
        "student/reset_request.html",
        form=form,
        submitted=submitted,
        email_sent=email_sent,
        auto_email_enabled=_auto_email_enabled(),
        token_generated=token_generated,
    )


@web_bp.route("/participant/reset/<string:token>", methods=["GET", "POST"])
def student_reset_password_token(token: str):
    form = PasswordResetConfirmForm()
    reset_token = _load_reset_token(token, PasswordResetAudience.STUDENT)
    is_valid, reason = _reset_token_status(reset_token)
    state = "ready" if is_valid else "invalid"
    plain_password: str | None = None
    email_sent = False
    auto_email_enabled = _auto_email_enabled()

    if form.validate_on_submit():
        if not is_valid or not reset_token or not reset_token.student:
            state = "invalid"
        else:
            _, plain_password = _ensure_student_credentials(reset_token.student, reset_password=True)
            reset_token.used_at = datetime.utcnow()
            db.session.flush()
            log_activity(
                "student_password_reset_completed",
                f"Student {reset_token.student.first_name} {reset_token.student.last_name} completed portal reset",
                object_type="student",
                object_id=reset_token.student.id,
                details={
                    "email": reset_token.student.email,
                    "upload_username": reset_token.student.upload_username,
                    "reset_token_id": reset_token.id,
                },
            )
            db.session.commit()
            if plain_password and auto_email_enabled and reset_token.student.email:
                _send_student_reset_email(reset_token.student, plain_password)
                email_sent = True
            state = "completed"
            reason = None
            is_valid = False

    return render_template(
        "student/reset_token.html",
        form=form,
        state=state,
        reason=reason,
        token_entry=reset_token,
        plain_password=plain_password,
        email_sent=email_sent,
        auto_email_enabled=auto_email_enabled,
    )


@web_bp.route("/participant/portal")
def student_portal():
    student = _student_from_session()
    if not student or student.is_absent:
        flash("Please log in to access the student portal.", "error")
        return redirect(url_for("web.student_login", next=request.full_path))

    is_htmx = _is_htmx_request()
    survey_assignments = _survey_assignments_for_user(student=student)
    survey_items, survey_summary = _prepare_survey_assignment_view(
        survey_assignments,
        lambda assignment_id: url_for("web.student_survey_detail", assignment_id=assignment_id),
    )
    survey_outstanding = [item for item in survey_items if not item.is_completed and not item.state.locked][:5]

    upload_status = SimpleNamespace(
        has_upload=bool(student.uploaded_filename),
        finalized=bool(student.essay_finalized),
        uploaded_at=student.uploaded_at,
        finalized_at=student.essay_finalized_at,
    )
    essay_access = _student_essay_access(student)
    upload_module_enabled = _is_upload_module_enabled()
    upload_locked = _is_upload_locked()
    upload_available = upload_module_enabled and not upload_locked

    template_name = "student/partials/portal_content.html" if is_htmx else "student/portal.html"
    context = {
        "student": student,
        "survey_summary": survey_summary,
        "survey_outstanding": survey_outstanding,
        "survey_items": survey_items,
        "upload_status": upload_status,
        "upload_module_enabled": upload_module_enabled,
        "upload_locked": upload_locked,
        "upload_available": upload_available,
        "essay_access": essay_access,
        "download_notice": _essay_download_notice(essay_access),
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/participant/logout")
def student_logout():
    student = _student_from_session()
    if student:
        log_activity(
            "student_upload_logout",
            f"Student {student.first_name} {student.last_name} signed out",
            object_type="student",
            object_id=student.id,
        )
        db.session.commit()
    session.pop("student_upload_id", None)
    flash("Signed out.", "success")
    return redirect(url_for("web.student_login"))


@web_bp.route("/participant/upload", methods=["GET", "POST"])
def student_upload():
    student = _student_from_session()
    if not student or student.is_absent:
        flash("Please log in before uploading your essay.", "error")
        return redirect(url_for("web.student_login"))
    form = StudentUploadForm()
    action = request.form.get("action", "upload").strip()
    module_enabled = _is_upload_module_enabled()
    upload_locked = _is_upload_locked()
    is_htmx = _is_htmx_request()
    word_limit = _essay_word_limit()
    word_limit_mode = _essay_word_limit_mode()

    def _history_entries(access: SimpleNamespace) -> list[SimpleNamespace]:
        uploads = (
            EssayUpload.query
            .filter_by(student_id=student.id)
            .order_by(EssayUpload.created_at.desc())
            .all()
        )
        entries: list[SimpleNamespace] = []
        for upload in uploads:
            file_path = _essay_pdf_path_by_filename(upload.filename)
            entries.append(
                SimpleNamespace(
                    upload=upload,
                    is_current=upload.filename == student.uploaded_filename,
                    file_missing=file_path is None,
                    download_url=(
                        url_for("web.student_upload_history_download", upload_id=upload.id)
                        if access.download_allowed and file_path
                        else None
                    ),
                )
            )
        return entries

    def _render_upload_response():
        essay_access = _student_essay_access(student)
        template_name = "student/partials/upload_content.html" if is_htmx else "student/upload.html"
        context = {
            "form": form,
            "student": student,
            "upload_locked": upload_locked,
            "upload_module_enabled": module_enabled,
            "word_limit": word_limit,
            "word_limit_mode": word_limit_mode,
            "essay_access": essay_access,
            "download_notice": _essay_download_notice(essay_access),
            "history_entries": _history_entries(essay_access),
            "is_htmx": is_htmx,
        }
        if not is_htmx:
            return render_template(template_name, **context)
        return _render_htmx_template(template_name, **context)

    if request.method == "POST":
        if not module_enabled:
            flash("Essay uploads are currently disabled by the organizing team.", "error")
            if is_htmx:
                return _render_upload_response()
            return redirect(url_for("web.student_upload"))
        if upload_locked:
            flash("Essay uploads are currently locked by the organizing team.", "error")
            if is_htmx:
                return _render_upload_response()
            return redirect(url_for("web.student_upload"))

    word_count = student.essay_word_count

    if request.method == "POST" and action == "finalize":
        if student.essay_finalized:
            flash("Your essay is already marked as finished.", "info")
        elif not student.uploaded_filename:
            flash("Upload your essay before marking it as finished.", "error")
        elif word_limit_mode == "hard" and word_count is not None and word_count > word_limit:
            flash(
                f"Your essay is above the word limit ({word_count} > {word_limit}). "
                "Please upload a shorter version before finalizing.",
                "error",
            )
        else:
            student.essay_finalized = True
            student.essay_finalized_at = datetime.utcnow()
            upload_entry = _current_essay_upload(student)
            if upload_entry and not upload_entry.finalized_at:
                upload_entry.finalized_at = student.essay_finalized_at
            db.session.flush()
            log_activity(
                "student_essay_finalized",
                f"Student {student.first_name} {student.last_name} finalized essay",
                object_type="student",
                object_id=student.id,
                details={"filename": student.uploaded_filename},
            )
            db.session.commit()
            flash("Your essay has been marked as finished. Further uploads are disabled.", "success")
        if is_htmx:
            return _render_upload_response()
        return redirect(url_for("web.student_upload"))

    if student.essay_finalized:
        return _render_upload_response()

    if form.validate_on_submit():
        file = form.essay.data
        if not file:
            flash("No file provided.", "error")
            if is_htmx:
                return _render_upload_response()
            return redirect(url_for("web.student_upload"))
        filename = secure_filename(file.filename or "")
        if not filename.lower().endswith(".pdf"):
            flash("Only PDF files are allowed.", "error")
            if is_htmx:
                return _render_upload_response()
            return redirect(url_for("web.student_upload"))
        file.seek(0, 2)
        size = file.tell()
        file.seek(0)
        if size > MAX_UPLOAD_SIZE:
            flash("File too large (max 10 MB).", "error")
            if is_htmx:
                return _render_upload_response()
            return redirect(url_for("web.student_upload"))
        upload_folder = _get_upload_folder()
        timestamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
        stored_name = f"student-{student.id}-{timestamp}.pdf"
        file_path = upload_folder / stored_name
        file.save(file_path)
        uploaded_at = datetime.utcnow()
        student.uploaded_filename = stored_name
        student.uploaded_at = uploaded_at
        student.essay_finalized = False
        student.essay_finalized_at = None
        student.essay_word_count = None
        student.essay_word_counted_at = None

        counted = _count_pdf_words(file_path)
        if counted is not None:
            student.essay_word_count = counted
            student.essay_word_counted_at = datetime.utcnow()
            if word_limit_mode == "soft" and counted > word_limit:
                flash(
                    f"Word limit warning: your PDF appears to have {counted} words (limit {word_limit}).",
                    "warning",
                )
            elif word_limit_mode == "hard" and counted > word_limit:
                flash(
                    f"Word limit exceeded: your PDF appears to have {counted} words (limit {word_limit}). "
                    "You will not be able to finalize until you upload a shorter version.",
                    "error",
                )
        upload_entry = EssayUpload(
            student_id=student.id,
            filename=stored_name,
            word_count=student.essay_word_count,
            word_counted_at=student.essay_word_counted_at,
            created_at=uploaded_at,
            updated_at=uploaded_at,
        )
        db.session.add(upload_entry)
        log_activity(
            "student_essay_uploaded",
            f"Student {student.first_name} {student.last_name} uploaded essay",
            object_type="student",
            object_id=student.id,
            details={"filename": stored_name},
        )
        db.session.commit()
        flash("Essay uploaded successfully.", "success")
        if is_htmx:
            return _render_upload_response()
        return redirect(url_for("web.student_upload"))
    return _render_upload_response()


@web_bp.route("/participant/upload/preview")
def student_upload_preview():
    student = _student_from_session()
    if not student or student.is_absent or not student.uploaded_filename:
        abort(404)
    access = _student_essay_access(student)
    if not access.preview_allowed:
        abort(403)
    pdf_path = _essay_pdf_path(student)
    if not pdf_path or not pdf_path.exists():
        abort(404)
    response = send_file(str(pdf_path), mimetype="application/pdf")
    response.headers["Cache-Control"] = "no-store"
    return response


@web_bp.route("/participant/upload/download")
def student_upload_download():
    student = _student_from_session()
    if not student or student.is_absent or not student.uploaded_filename:
        abort(404)
    access = _student_essay_access(student)
    if not access.download_allowed:
        abort(403)
    pdf_path = _essay_pdf_path(student)
    if not pdf_path:
        abort(404)
    filename = secure_filename(f"essay-{student.id}.pdf")
    log_activity(
        "student_essay_downloaded",
        f"Student {student.first_name} {student.last_name} downloaded essay",
        object_type="student",
        object_id=student.id,
        details={"filename": student.uploaded_filename},
    )
    db.session.commit()
    return send_file(
        pdf_path,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf",
    )


@web_bp.route("/participant/upload/history/<int:upload_id>/download")
def student_upload_history_download(upload_id: int):
    student = _student_from_session()
    if not student or student.is_absent:
        abort(404)
    access = _student_essay_access(student)
    if not access.download_allowed:
        abort(403)
    upload = db.session.get(EssayUpload, upload_id)
    if not upload or upload.student_id != student.id:
        abort(404)
    pdf_path = _essay_pdf_path_by_filename(upload.filename)
    if not pdf_path:
        abort(404)
    filename = secure_filename(f"essay-{student.id}-v{upload.id}.pdf")
    log_activity(
        "student_essay_history_downloaded",
        f"Student {student.first_name} {student.last_name} downloaded essay history",
        object_type="student",
        object_id=student.id,
        details={"filename": upload.filename, "upload_id": upload.id},
    )
    db.session.commit()
    return send_file(
        pdf_path,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf",
    )


@web_bp.route("/participant/surveys")
def student_surveys():
    student = _student_from_session()
    if not student or student.is_absent:
        flash("Please log in to access surveys.", "error")
        return redirect(url_for("web.student_login", next=request.full_path))

    assignments = _survey_assignments_for_user(student=student)
    items, summary = _prepare_survey_assignment_view(
        assignments,
        lambda assignment_id: url_for("web.student_survey_detail", assignment_id=assignment_id),
    )
    outstanding = [item for item in items if not item.is_completed and not item.state.locked]
    closed = [item for item in items if not item.is_completed and item.state.locked]
    completed = [item for item in items if item.is_completed]

    return render_template(
        "student/surveys.html",
        student=student,
        assignments=items,
        summary=summary,
        outstanding=outstanding,
        closed=closed,
        completed=completed,
    )


@web_bp.route("/participant/surveys/<int:assignment_id>", methods=["GET", "POST"])
def student_survey_detail(assignment_id: int):
    student = _student_from_session()
    if not student or student.is_absent:
        flash("Please log in to access surveys.", "error")
        return redirect(url_for("web.student_login", next=request.full_path))

    assignment = _load_survey_assignment(assignment_id, student=student)
    if not assignment or not assignment.survey:
        abort(404)
    survey = assignment.survey
    if survey.status == SurveyStatus.DRAFT:
        abort(404)

    questions = _ordered_survey_questions(survey)
    existing_data = _survey_display_data(assignment.response.response_data if assignment.response else {})
    state = _survey_assignment_state(assignment)
    form = SurveyResponseForm()
    errors: dict[int, str] = {}

    if form.validate_on_submit():
        if state.locked:
            flash("This survey is not available for editing.", "info")
        else:
            answers, errors = _parse_survey_submission(questions, request.form)
            if not errors:
                _save_survey_response(
                    assignment,
                    answers,
                    respondent_type="student",
                    respondent_id=student.id,
                    actor=student.upload_username or student.email or f"student-{student.id}",
                )
                db.session.commit()
                flash("Survey response saved.", "success")
                return redirect(url_for("web.student_survey_detail", assignment_id=assignment.id))
            existing_data = _survey_display_data(answers)
            flash("Please correct the errors below.", "error")
        state = _survey_assignment_state(assignment)

    button_label = "Update response" if assignment.status == SurveyAssignmentStatus.COMPLETED else "Submit response"

    return render_template(
        "student/survey_detail.html",
        student=student,
        assignment=assignment,
        survey=survey,
        questions=questions,
        form=form,
        existing_data=existing_data,
        errors=errors,
        state=state,
        button_label=button_label,
        SurveyQuestionType=SurveyQuestionType,
        SurveyAssignmentStatus=SurveyAssignmentStatus,
    )


@web_bp.route("/register/student", methods=["GET", "POST"])
def register_student():
    form = StudentRegistrationForm()
    invite_token = request.args.get("invite") or request.form.get("invite")
    invite = _load_delegation_invite(invite_token, DelegationInviteRole.STUDENT)
    invite_state = _delegation_invite_state(invite) if invite else None
    invite_active = bool(invite_state and invite_state.active)
    invite_required = _registration_invite_required()
    invite_message = None

    if invite_required and not invite_active:
        contact_hint = (
            "the organizing team"
            if invite and (invite.delegation_leader or invite.invited_by_id is None)
            else "your delegation leader"
        )
        if not invite:
            invite_message = "Registration is available by invitation only. Please use your invite link."
        elif invite_state and invite_state.status == "revoked":
            invite_message = f"This invite link has been revoked. Please contact {contact_hint}."
        elif invite_state and invite_state.status == "completed":
            invite_message = "This invite link has already been used."
        else:
            invite_message = f"This invite link is invalid. Please check with {contact_hint}."
    if invite and not invite_active and not invite_required:
        invite = None
        invite_state = None
        invite_active = False

    if not form.is_submitted():
        form.submit.label.text = "Submit registration"
        form.is_absent.data = False
    if invite and (not form.is_submitted() or invite_active):
        form.email.data = invite.email
        form.country.data = invite.country
        if not form.is_submitted():
            if invite.first_name:
                form.first_name.data = invite.first_name
            if invite.last_name:
                form.last_name.data = invite.last_name
    if invite_active and invite.opened_at is None:
        invite.opened_at = datetime.utcnow()
        db.session.commit()
    if invite_required and not invite_active:
        return render_template(
            "register_student.html",
            form=form,
            show_absence_toggle=False,
            ethics_consent_html=_ethics_consent_html(),
            gdpr_notice_html=_gdpr_notice_html(),
            photo_consent_html=_photo_consent_html(),
            accommodation_info_html=_accommodation_info_html(),
            invite_message=invite_message,
            invite_required=invite_required,
            invite_active=invite_active,
        )
    if form.validate_on_submit():
        now = datetime.utcnow()
        email = invite.email.lower().strip() if invite_active else form.email.data.lower().strip()
        country = invite.country.strip() if invite_active else form.country.data.strip()
        student = Student(
            first_name=form.first_name.data.strip(),
            last_name=form.last_name.data.strip(),
            email=email,
            country=country,
            school=form.school.data.strip(),
            sex_document=form.sex_document.data,
            accommodation_room_sharing=form.accommodation_room_sharing.data,
            date_of_birth=form.date_of_birth.data,
            essay_language=form.essay_language.data,
            keyboard_preference=form.keyboard_preference.data,
            meal_preference=form.meal_preference.data,
            special_needs=(form.special_needs.data.strip() or None)
            if form.special_needs.data
            else None,
            gdpr_acknowledged=bool(form.gdpr_acknowledged.data),
            gdpr_acknowledged_at=now if form.gdpr_acknowledged.data else None,
            photo_consent=form.photo_consent.data == "yes",
            photo_consent_at=now,
            accepted_ethics=bool(form.accepted_ethics.data),
            accepted_ethics_at=now if form.accepted_ethics.data else None,
            is_absent=bool(form.is_absent.data),
        )
        _ensure_student_credentials(student)
        db.session.add(student)
        try:
            db.session.flush()
            if invite_active:
                invite.first_name = student.first_name
                invite.last_name = student.last_name
                invite.email = student.email
                invite.accepted_at = now
                invite.student_id = student.id
                log_activity(
                    "delegation_invite_accepted",
                    f"Delegation invite accepted for {invite.email}",
                    actor=student.upload_username or student.email,
                    object_type="delegation_invite",
                    object_id=invite.id,
                    details={"role": invite.role.value, "country": invite.country},
                )
            log_activity(
    "student_registered",
    f"Student {student.first_name} {student.last_name} registered",
    object_type="student",
    object_id=student.id,
    details={
        "email": student.email,
        "country": student.country,
        "language": student.essay_language,
        "upload_username": student.upload_username,
        "invite_id": invite.id if invite_active else None,
    },
)
            db.session.commit()
            _send_student_confirmation_email(student)
            if invite_active:
                flash("Registration completed. Please check your email for portal credentials.", "success")
                return redirect(url_for("web.index"))
            flash("Student registered successfully.", "success")
            return redirect(url_for("web.register_student"))
        except IntegrityError:
            db.session.rollback()
            flash("This email address is already registered.", "error")

    return render_template(
        "register_student.html",
        form=form,
        show_absence_toggle=False,
        ethics_consent_html=_ethics_consent_html(),
        gdpr_notice_html=_gdpr_notice_html(),
        photo_consent_html=_photo_consent_html(),
        accommodation_info_html=_accommodation_info_html(),
        invite_message=invite_message,
        invite_required=invite_required,
        invite_active=invite_active,
        invite_token=invite.token if invite_active else None,
    )


@web_bp.route("/register/teacher", methods=["GET", "POST"])
def register_teacher():
    form = TeacherRegistrationForm()
    invite_token = request.args.get("invite") or request.form.get("invite")
    invite = _load_delegation_invite(invite_token, DelegationInviteRole.TEACHER)
    invite_state = _delegation_invite_state(invite) if invite else None
    invite_active = bool(invite_state and invite_state.active)
    invite_required = _registration_invite_required()
    invite_message = None

    if invite_required and not invite_active:
        contact_hint = (
            "the organizing team"
            if invite and (invite.delegation_leader or invite.invited_by_id is None)
            else "your delegation leader"
        )
        if not invite:
            invite_message = "Registration is available by invitation only. Please use your invite link."
        elif invite_state and invite_state.status == "revoked":
            invite_message = f"This invite link has been revoked. Please contact {contact_hint}."
        elif invite_state and invite_state.status == "completed":
            invite_message = "This invite link has already been used."
        else:
            invite_message = f"This invite link is invalid. Please check with {contact_hint}."
    if invite and not invite_active and not invite_required:
        invite = None
        invite_state = None
        invite_active = False

    if not form.is_submitted() and not form.languages.data:
        form.languages.data = ["English"]
    form.submit.label.text = "Submit registration"
    if not form.is_submitted():
        form.is_absent.data = False
        form.non_assessor.data = False
    if invite and (not form.is_submitted() or invite_active):
        form.email.data = invite.email
        form.country.data = invite.country
        if not form.is_submitted():
            if invite.first_name:
                form.first_name.data = invite.first_name
            if invite.last_name:
                form.last_name.data = invite.last_name
    if invite_active and invite.opened_at is None:
        invite.opened_at = datetime.utcnow()
        db.session.commit()
    if invite_required and not invite_active:
        return render_template(
            "register_teacher.html",
            form=form,
            show_absence_toggle=False,
            ethics_consent_html=_ethics_consent_html(),
            gdpr_notice_html=_gdpr_notice_html(),
            photo_consent_html=_photo_consent_html(),
            assessor_competence_html=_assessor_competence_html(),
            accommodation_info_html=_accommodation_info_html(),
            invite_message=invite_message,
            invite_required=invite_required,
            invite_active=invite_active,
            invite_delegation_leader=False,
        )
    if form.validate_on_submit():
        non_assessor = bool(form.non_assessor.data)
        if not non_assessor and not form.declared_assessor_competence.data:
            form.declared_assessor_competence.errors.append(
                "This confirmation is required unless you register as a non-assessing member."
            )
        else:
            now = datetime.utcnow()
            email = invite.email.lower().strip() if invite_active else form.email.data.lower().strip()
            country = invite.country.strip() if invite_active else form.country.data.strip()
            delegation_leader = bool(invite.delegation_leader) if invite_active else False
            prior_level = 0 if non_assessor else (form.prior_ipo_assessor_level.data or 0)
            declared_assessor_competence = False if non_assessor else bool(form.declared_assessor_competence.data)
            teacher = Teacher(
                first_name=form.first_name.data.strip(),
                last_name=form.last_name.data.strip(),
                email=email,
                country=country,
                school=form.school.data.strip(),
                sex_document=form.sex_document.data,
                accommodation_room_sharing=form.accommodation_room_sharing.data,
                languages=", ".join(form.languages.data),
                meal_preference=form.meal_preference.data,
                tshirt_size=form.tshirt_size.data,
                special_needs=(form.special_needs.data.strip() or None)
                if form.special_needs.data
                else None,
                delegation_leader=delegation_leader,
                prior_ipo_assessor_level=prior_level,
                prior_ipo_assessor=bool(prior_level),
                language_entries=[TeacherLanguage(language=lang) for lang in form.languages.data],
                gdpr_acknowledged=bool(form.gdpr_acknowledged.data),
                gdpr_acknowledged_at=now if form.gdpr_acknowledged.data else None,
                photo_consent=form.photo_consent.data == "yes",
                photo_consent_at=now,
                declared_assessor_competence=declared_assessor_competence,
                declared_assessor_competence_at=now if declared_assessor_competence else None,
                accepted_ethics=bool(form.accepted_ethics.data),
                accepted_ethics_at=now if form.accepted_ethics.data else None,
                is_absent=bool(form.is_absent.data),
            )
            db.session.add(teacher)
            try:
                _ensure_teacher_portal_credentials(teacher)
                db.session.flush()
                if invite_active:
                    invite.first_name = teacher.first_name
                    invite.last_name = teacher.last_name
                    invite.email = teacher.email
                    invite.accepted_at = now
                    invite.teacher_id = teacher.id
                    log_activity(
                        "delegation_invite_accepted",
                        f"Delegation invite accepted for {invite.email}",
                        actor=teacher.portal_username or teacher.email,
                        object_type="delegation_invite",
                        object_id=invite.id,
                        details={"role": invite.role.value, "country": invite.country},
                    )
                log_activity(
                    "teacher_registered",
                    f"Teacher {teacher.first_name} {teacher.last_name} registered",
                    object_type="teacher",
                    object_id=teacher.id,
                    details={
                        "email": teacher.email,
                        "country": teacher.country,
                        "languages": form.languages.data,
                        "portal_username": teacher.portal_username,
                        "invite_id": invite.id if invite_active else None,
                        "non_assessor": non_assessor,
                    },
                )
                db.session.commit()
                _send_teacher_confirmation_email(teacher)
                if invite_active:
                    flash("Registration completed. Please check your email for portal credentials.", "success")
                    return redirect(url_for("web.index"))
                flash("Teacher registered successfully.", "success")
                return redirect(url_for("web.register_teacher"))
            except IntegrityError:
                db.session.rollback()
                flash("This email address is already registered.", "error")

    return render_template(
        "register_teacher.html",
        form=form,
        show_absence_toggle=False,
        ethics_consent_html=_ethics_consent_html(),
        gdpr_notice_html=_gdpr_notice_html(),
        photo_consent_html=_photo_consent_html(),
        assessor_competence_html=_assessor_competence_html(),
        accommodation_info_html=_accommodation_info_html(),
        invite_message=invite_message,
        invite_required=invite_required,
        invite_active=invite_active,
        invite_token=invite.token if invite_active else None,
        invite_delegation_leader=bool(invite.delegation_leader) if invite_active else False,
    )


@web_bp.route("/admin")
def admin_dashboard():
    filters = {
        "essay_language": request.args.get("essay_language") or None,
        "alert_level": request.args.get("alert_level") or None,
    }

    metrics, alerts, tables, deadlines = _collect_dashboard_data(filters)
    language_options = [choice[0] for choice in LANGUAGE_CHOICES]
    alert_options = [
        {"value": "", "label": "All"},
        {"value": "warning", "label": "Warnings"},
        {"value": "info", "label": "Info"},
    ]
    stage1_locked = _is_stage_locked(Stage.STAGE_1)
    stage2_locked = _is_stage_locked(Stage.STAGE_2)
    upload_module_enabled = _is_upload_module_enabled()
    upload_locked = _is_upload_locked()
    upload_state = SimpleNamespace(enabled=upload_module_enabled, locked=upload_locked)

    is_htmx = _is_htmx_request()
    template_name = "admin/partials/dashboard_content.html" if is_htmx else "admin/dashboard.html"
    context = {
        "is_htmx": is_htmx,
        "metrics": metrics,
        "alerts": alerts,
        "tables": tables,
        "deadlines": deadlines,
        "filters": filters,
        "language_options": language_options,
        "alert_options": alert_options,
        "stage_lock_state": {
            "stage_1": stage1_locked,
            "stage_2": stage2_locked,
        },
        "upload_state": upload_state,
        "Stage": Stage,
        "last_refreshed": datetime.utcnow(),
    }
    if not is_htmx:
        context.update(_quick_message_context())
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/grading/lock", methods=["POST"])
def admin_toggle_grading_lock():
    if not _current_admin():
        abort(403)

    stage_value = request.form.get("stage")
    target_state = request.form.get("state")

    try:
        stage = Stage(stage_value)
    except (ValueError, TypeError):
        abort(400)

    if target_state not in {"lock", "unlock"}:
        abort(400)

    desired_lock = target_state == "lock"
    _set_stage_lock(stage, desired_lock)

    log_activity(
        "grading_stage_lock",
        f"Grading lock for {stage.value} set to {desired_lock}",
        object_type="grading",
        details={
            "stage": stage.value,
            "locked": desired_lock,
            "actor": _current_admin(),
        },
    )

    db.session.commit()

    stage_label = STAGE_LABELS.get(stage, stage.value.title())
    flash(f"{stage_label} grading {'locked' if desired_lock else 'unlocked'}.", "success")

    if _is_htmx_request():
        return admin_dashboard()

    redirect_target = request.form.get("redirect") or request.referrer or url_for("web.admin_dashboard")
    return redirect(redirect_target)


@web_bp.route("/admin/uploads/lock", methods=["POST"])
def admin_toggle_upload_lock():
    if not _current_admin():
        abort(403)

    target_state = request.form.get("state")
    if target_state not in {"lock", "unlock"}:
        abort(400)

    desired_lock = target_state == "lock"
    _set_upload_lock(desired_lock)

    log_activity(
        "essay_upload_lock",
        f"Essay uploads locked: {desired_lock}",
        object_type="uploads",
        details={
            "locked": desired_lock,
            "actor": _current_admin(),
        },
    )

    db.session.commit()

    flash(f"Essay uploads {'locked' if desired_lock else 'unlocked'}.", "success")

    if _is_htmx_request():
        return admin_dashboard()

    redirect_target = request.form.get("redirect") or request.referrer or url_for("web.admin_dashboard")
    return redirect(redirect_target)


def _is_htmx_request() -> bool:
    return request.headers.get("HX-Request") == "true"

def _render_htmx_template(template_name: str, **context):
    messages = get_flashed_messages(with_categories=True)
    body = render_template(template_name, **context)
    if not messages:
        return body
    toast_html = render_template("partials/toast_messages.html", messages=messages, oob=True)
    return f"{toast_html}{body}"





def _collect_dashboard_data(filters: dict | None = None):
    filters = filters or {}
    essay_language_filter = filters.get("essay_language") or None
    alert_level_filter = filters.get("alert_level") or None

    now = datetime.utcnow()

    # Student metrics
    total_students = Student.query.count()
    uploads_complete = Student.query.filter(Student.uploaded_filename.isnot(None)).count()
    uploads_pending = Student.query.filter(
        Student.uploaded_filename.is_(None),
        Student.is_absent.is_(False),
    ).count()
    absent_students = Student.query.filter_by(is_absent=True).count()
    students_without_uploads = (
        Student.query.filter(
            Student.uploaded_filename.is_(None),
            Student.is_absent.is_(False),
        )
        .order_by(Student.created_at.asc())
        .all()
    )

    # Teacher metrics and language map
    teachers = Teacher.query.order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()
    teacher_languages: dict[int, set[str]] = {}
    teacher_load: dict[int, dict] = {}
    active_teachers = 0
    absent_teachers = 0
    for teacher in teachers:
        _ensure_language_entries(teacher)
        languages = {entry.language for entry in teacher.language_entries if entry.language}
        teacher_languages[teacher.id] = languages
        teacher_load[teacher.id] = {
            "teacher": teacher,
            "assigned": 0,
            "completed": 0,
            "pending": 0,
            "is_absent": teacher.is_absent,
        }
        if teacher.is_absent:
            absent_teachers += 1
        else:
            active_teachers += 1

    # Group staffing metrics
    groups = (
        ReadingGroup.query.options(
            joinedload(ReadingGroup.members).joinedload(ReadingGroupMember.teacher)
        ).all()
    )
    understaffed_groups = []
    active_teacher_ids = set()
    for group in groups:
        active_members = [
            member
            for member in group.members
            if member.active and member.teacher and not member.teacher.is_absent
        ]
        for member in active_members:
            active_teacher_ids.add(member.teacher_id)
        if len(active_members) < 4:
            understaffed_groups.append(
                {
                    "group": group,
                    "active_count": len(active_members),
                    "language": group.language,
                }
            )

    # Stage metrics
    stage1_assignments = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.scores),
            joinedload(EssayAssignment.extra_readers),
            joinedload(EssayAssignment.student),
            joinedload(EssayAssignment.group)
            .joinedload(ReadingGroup.members)
            .joinedload(ReadingGroupMember.teacher),
        )
        .filter(EssayAssignment.stage == Stage.STAGE_1)
        .all()
    )

    stage1_total = len(stage1_assignments)
    stage1_completed = 0
    stage1_pending_count = 0
    stage1_expected_scores = 0
    stage1_entered_scores = 0
    tie_break_needed = 0
    stage1_pending_list = []
    extra_reader_summary = {
        "needs_assignment": 0,
        "awaiting_score": 0,
        "review_spread": 0,
        "resolved": 0,
        "action_total": 0,
    }
    stage1_extra_reader_queue = []

    stage2_candidates = []
    stage2_completed = 0
    stage2_pending_count = 0
    stage2_expected_scores = 0
    stage2_entered_scores = 0
    stage2_pending_list = []
    stage3_recommended = 0

    for assignment in stage1_assignments:
        status = _compute_stage_one_status(assignment)
        extra_state = _stage_one_extra_reader_state(assignment, status)
        expected = status.expected_scores or 0
        entered = status.entered_scores or 0
        stage1_expected_scores += expected
        stage1_entered_scores += entered
        if status.completed:
            stage1_completed += 1
        else:
            stage1_pending_count += 1
            stage1_pending_list.append(
                {
                    "assignment": assignment,
                    "student": assignment.student,
                    "missing": max(expected - entered, 0),
                }
            )
        if status.needs_extra or assignment.needs_additional_reading:
            tie_break_needed += 1
        if extra_state.bucket in extra_reader_summary:
            extra_reader_summary[extra_state.bucket] += 1
        if extra_state.needs_action:
            extra_reader_summary["action_total"] += 1
            stage1_extra_reader_queue.append(
                {
                    "assignment": assignment,
                    "student": assignment.student,
                    "state": extra_state,
                    "score_range": status.score_range,
                    "average_score": status.average_score,
                }
            )

        stage1_scores_map = {
            score.teacher_id: score
            for score in assignment.scores
            if score.stage == Stage.STAGE_1 and score.score is not None
        }
        for assessor in _stage_assessors(assignment, Stage.STAGE_1, include_group=True):
            teacher = assessor.teacher
            info = teacher_load.get(teacher.id)
            if not info:
                continue
            info["assigned"] += 1
            if stage1_scores_map.get(teacher.id):
                info["completed"] += 1

        if assignment.stage_two_candidate:
            stage2_candidates.append(assignment)
            stage2_status = _compute_stage_two_status(assignment)
            expected2 = stage2_status.expected_scores or 0
            entered2 = stage2_status.entered_scores or 0
            stage2_expected_scores += expected2
            stage2_entered_scores += entered2
            if expected2 > 0 and entered2 >= expected2:
                stage2_completed += 1
            else:
                stage2_pending_count += 1
                stage2_pending_list.append(
                    {
                        "assignment": assignment,
                        "student": assignment.student,
                        "missing": max(expected2 - entered2, 0),
                    }
                )
        if assignment.stage_three_recommended:
            stage3_recommended += 1

    for info in teacher_load.values():
        info["pending"] = max(info["assigned"] - info["completed"], 0)

    teacher_workload = [
        info
        for info in teacher_load.values()
        if info["assigned"] > 0 and not info["is_absent"]
    ]
    teacher_workload.sort(key=lambda entry: (entry["pending"], entry["assigned"]), reverse=True)
    teacher_workload = teacher_workload[:8]

    teacher_pending_count = sum(1 for info in teacher_load.values() if info["pending"] > 0 and not info["is_absent"])

    stage1_completion_rate = (stage1_completed / stage1_total * 100) if stage1_total else 100
    stage1_score_rate = (stage1_entered_scores / stage1_expected_scores * 100) if stage1_expected_scores else 100
    stage2_completion_rate = (stage2_completed / len(stage2_candidates) * 100) if stage2_candidates else 100
    stage2_score_rate = (stage2_entered_scores / stage2_expected_scores * 100) if stage2_expected_scores else 100

    metrics = {
        "students": {
            "total": total_students,
            "uploads_pending": uploads_pending,
            "uploads_complete": uploads_complete,
            "absent": absent_students,
        },
        "teachers": {
            "active": active_teachers,
            "absent": absent_teachers,
            "pending_scores": teacher_pending_count,
        },
        "groups": {
            "total": len(groups),
            "understaffed": len(understaffed_groups),
        },
        "stage": {
            "stage1": {
                "total": stage1_total,
                "completed": stage1_completed,
                "pending": stage1_pending_count,
                "expected_scores": stage1_expected_scores,
                "entered_scores": stage1_entered_scores,
                "tie_break": tie_break_needed,
                "extra_readers": extra_reader_summary,
                "completion_rate": stage1_completion_rate,
                "score_rate": stage1_score_rate,
            },
            "stage2": {
                "candidates": len(stage2_candidates),
                "completed": stage2_completed,
                "pending": stage2_pending_count,
                "expected_scores": stage2_expected_scores,
                "entered_scores": stage2_entered_scores,
                "completion_rate": stage2_completion_rate,
                "score_rate": stage2_score_rate,
            },
            "stage3": {
                "recommended": stage3_recommended,
            },
        },
    }

    alerts = []
    if uploads_pending:
        alerts.append(
            {
                "message": f"{uploads_pending} students have not uploaded their essays.",
                "level": "warning",
                "link": url_for("web.admin_students"),
            }
        )
    if len(understaffed_groups):
        alerts.append(
            {
                "message": f"{len(understaffed_groups)} reading group(s) have fewer than four active teachers.",
                "level": "warning",
                "link": url_for("web.admin_planner"),
            }
        )
    if stage1_pending_count:
        alerts.append(
            {
                "message": f"{stage1_pending_count} Stage 1 essay(s) still need scores.",
                "level": "warning",
                "link": url_for("web.admin_uploads"),
            }
        )
    if extra_reader_summary["action_total"]:
        alerts.append(
            {
                "message": f"{extra_reader_summary['action_total']} Stage 1 essay(s) need additional assessor action.",
                "level": "warning",
                "link": url_for("web.admin_uploads", status="needs_extra"),
            }
        )
    if stage2_pending_count:
        alerts.append(
            {
                "message": f"{stage2_pending_count} Stage 2 essay(s) awaiting assessment.",
                "level": "info",
                "link": url_for("web.admin_stage2_overview"),
            }
        )
    if teacher_pending_count:
        alerts.append(
            {
                "message": f"{teacher_pending_count} teacher(s) still need to submit scores.",
                "level": "info",
                "link": url_for("web.admin_teachers"),
            }
        )

    if essay_language_filter:
        students_filtered = [
            student
            for student in students_without_uploads
            if student.essay_language == essay_language_filter
        ]
        stage1_filtered = [
            item
            for item in stage1_pending_list
            if item["student"] and item["student"].essay_language == essay_language_filter
        ]
        extra_reader_filtered = [
            item
            for item in stage1_extra_reader_queue
            if item["student"] and item["student"].essay_language == essay_language_filter
        ]
    else:
        students_filtered = students_without_uploads
        stage1_filtered = stage1_pending_list
        extra_reader_filtered = stage1_extra_reader_queue

    if alert_level_filter:
        alerts_filtered = [alert for alert in alerts if alert["level"] == alert_level_filter]
    else:
        alerts_filtered = alerts

    tables = {
        "teacher_workload": teacher_workload,
        "students_without_uploads": students_filtered[:10],
        "stage1_pending": stage1_filtered[:8],
        "stage1_extra_reader_queue": extra_reader_filtered[:8],
    }

    alerts = alerts_filtered

    deadlines = _collect_dashboard_deadlines(now)

    return metrics, alerts, tables, deadlines


def _get_setting(key: str, default: str | None = None) -> str | None:
    setting = AppSetting.query.filter_by(key=key).first()
    if setting and setting.value is not None:
        return setting.value
    return default


def _set_setting(key: str, value: str | None) -> None:
    setting = AppSetting.query.filter_by(key=key).first()
    if setting:
        setting.value = value
    else:
        setting = AppSetting(key=key, value=value)
        db.session.add(setting)


def _parse_bool(raw: str | None) -> bool:
    if raw is None:
        return False
    return raw.strip().lower() in {"1", "true", "yes", "on"}


def _stage_lock_key(stage: Stage) -> str:
    return SETTING_STAGE1_LOCK if stage == Stage.STAGE_1 else SETTING_STAGE2_LOCK


def _is_stage_locked(stage: Stage) -> bool:
    return _parse_bool(_get_setting(_stage_lock_key(stage), "false"))


def _set_stage_lock(stage: Stage, locked: bool) -> None:
    _set_setting(_stage_lock_key(stage), "true" if locked else "false")


def _is_upload_locked() -> bool:
    return _parse_bool(_get_setting(SETTING_UPLOAD_LOCK, "false"))


def _set_upload_lock(locked: bool) -> None:
    _set_setting(SETTING_UPLOAD_LOCK, "true" if locked else "false")


def _collect_dashboard_deadlines(now: datetime) -> list[dict]:
    config = current_app.config
    deadlines: list[dict] = []

    upload_raw = _get_setting(SETTING_UPLOAD_DEADLINE, config.get("DASHBOARD_UPLOAD_DEADLINE"))
    stage2_raw = _get_setting(SETTING_STAGE2_DEADLINE, config.get("DASHBOARD_STAGE2_DEADLINE"))
    event_end_raw = _get_setting(SETTING_EVENT_END, config.get("DASHBOARD_EVENT_END"))

    upload_deadline = _parse_deadline(upload_raw)
    stage2_deadline = _parse_deadline(stage2_raw)
    event_end = _parse_deadline(event_end_raw)

    for label, deadline in (
        ("Essay upload deadline", upload_deadline),
        ("Stage 2 deadline", stage2_deadline),
        ("Event end (essay access unlock)", event_end),
    ):
        if deadline is None:
            deadlines.append({
                "label": label,
                "value": None,
                "display": "Not configured",
                "status": "info",
                "time_remaining": None,
            })
        else:
            delta = deadline - now
            deadlines.append({
                "label": label,
                "value": deadline,
                "display": deadline.strftime("%Y-%m-%d %H:%M UTC"),
                "status": "warning" if delta.total_seconds() < 0 else "info",
                "time_remaining": _humanize_timedelta(delta),
            })

    stage1_sla = _get_setting(SETTING_STAGE1_SLA, str(config.get("DASHBOARD_STAGE1_SCORE_SLA_HOURS", 24)))
    stage2_sla = _get_setting(SETTING_STAGE2_SLA, str(config.get("DASHBOARD_STAGE2_SCORE_SLA_HOURS", 24)))

    deadlines.append({
        "label": "Stage 1 score SLA",
        "value": None,
        "display": f"{stage1_sla} hours",
        "status": "info",
        "time_remaining": None,
    })
    deadlines.append({
        "label": "Stage 2 score SLA",
        "value": None,
        "display": f"{stage2_sla} hours",
        "status": "info",
        "time_remaining": None,
    })

    return deadlines


def _parse_deadline(raw: str | None) -> datetime | None:
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw)
    except ValueError:
        return None


def _event_end() -> datetime | None:
    config = current_app.config
    raw = _get_setting(SETTING_EVENT_END, config.get("DASHBOARD_EVENT_END"))
    return _parse_deadline(raw)


def _event_end_override() -> str:
    raw = (_get_setting(SETTING_EVENT_END_OVERRIDE, "auto") or "auto").strip().lower()
    if raw not in {"auto", "force_end", "force_open"}:
        return "auto"
    return raw


def _essay_event_state(now: datetime | None = None) -> SimpleNamespace:
    event_end = _event_end()
    override = _event_end_override()
    now = now or datetime.utcnow()
    if override == "force_end":
        event_over = True
    elif override == "force_open":
        event_over = False
    else:
        event_over = bool(event_end and now >= event_end)
    display = event_end.strftime("%Y-%m-%d %H:%M UTC") if event_end else None
    override_label = {
        "auto": "Auto",
        "force_end": "Force end",
        "force_open": "Force open",
    }[override]
    return SimpleNamespace(
        event_end=event_end,
        event_over=event_over,
        display=display,
        override=override,
        override_label=override_label,
    )


def _student_essay_access(student: Student, *, now: datetime | None = None) -> SimpleNamespace:
    state = _essay_event_state(now)
    has_upload = bool(student.uploaded_filename)
    preview_allowed = has_upload
    download_allowed = has_upload and state.event_over
    return SimpleNamespace(
        has_upload=has_upload,
        preview_allowed=preview_allowed,
        download_allowed=download_allowed,
        event_over=state.event_over,
        event_end=state.event_end,
        event_end_display=state.display,
        override=state.override,
        override_label=state.override_label,
    )


def _essay_download_notice(access: SimpleNamespace) -> str | None:
    if not access.has_upload or access.download_allowed:
        return None
    if access.override == "force_open":
        return "Downloads are temporarily locked by the organizing team."
    if access.event_end_display:
        return f"Downloads unlock after the event ends ({access.event_end_display})."
    return "Downloads unlock after the event ends. Ask the organizers to set the event end time."


def _humanize_timedelta(delta: timedelta) -> str:
    seconds = int(delta.total_seconds())
    prefix = 'Due in ' if seconds >= 0 else 'Overdue by '
    seconds = abs(seconds)
    days, remainder = divmod(seconds, 86400)
    hours, remainder = divmod(remainder, 3600)
    minutes, _ = divmod(remainder, 60)
    parts = []
    if days:
        parts.append(f"{days}d")
    if hours:
        parts.append(f"{hours}h")
    if minutes and not days:
        parts.append(f"{minutes}m")
    if not parts:
        parts.append('moments')
    return prefix + ' '.join(parts)


def _populate_admin_settings_form(form: DashboardSettingsForm, config: dict[str, Any]) -> None:
    upload_raw = _get_setting(SETTING_UPLOAD_DEADLINE, config.get("DASHBOARD_UPLOAD_DEADLINE"))
    stage2_raw = _get_setting(SETTING_STAGE2_DEADLINE, config.get("DASHBOARD_STAGE2_DEADLINE"))
    event_end_raw = _get_setting(SETTING_EVENT_END, config.get("DASHBOARD_EVENT_END"))
    event_override_raw = _get_setting(SETTING_EVENT_END_OVERRIDE, "auto")
    stage1_sla_raw = _get_setting(SETTING_STAGE1_SLA, str(config.get("DASHBOARD_STAGE1_SCORE_SLA_HOURS", 24)))
    stage2_sla_raw = _get_setting(SETTING_STAGE2_SLA, str(config.get("DASHBOARD_STAGE2_SCORE_SLA_HOURS", 24)))

    if upload_raw:
        dt = _parse_deadline(upload_raw)
        if dt:
            form.upload_deadline.data = dt
    if stage2_raw:
        dt = _parse_deadline(stage2_raw)
        if dt:
            form.stage2_deadline.data = dt
    if event_end_raw:
        dt = _parse_deadline(event_end_raw)
        if dt:
            form.event_end.data = dt
    override_value = (event_override_raw or "auto").strip().lower()
    if override_value not in {"auto", "force_end", "force_open"}:
        override_value = "auto"
    form.event_end_override.data = override_value
    if stage1_sla_raw:
        try:
            form.stage1_sla.data = int(float(stage1_sla_raw))
        except ValueError:
            form.stage1_sla.data = None
    if stage2_sla_raw:
        try:
            form.stage2_sla.data = int(float(stage2_sla_raw))
        except ValueError:
            form.stage2_sla.data = None

    packet_mode = _get_setting(SETTING_PACKET_PRINT_MODE, "single")
    form.packet_print_mode.data = packet_mode if packet_mode in {"single", "double"} else "single"

    essay_word_limit_raw = _get_setting(SETTING_ESSAY_WORD_LIMIT)
    try:
        form.essay_word_limit.data = int(essay_word_limit_raw) if essay_word_limit_raw else DEFAULT_ESSAY_WORD_LIMIT
    except (TypeError, ValueError):
        form.essay_word_limit.data = DEFAULT_ESSAY_WORD_LIMIT
    form.essay_word_limit_mode.data = _essay_word_limit_mode()
    form.stage1_trim_extreme_scores_enabled.data = _stage1_trim_extreme_scores_enabled()
    form.ethics_consent_html.data = _ethics_consent_html()
    form.gdpr_notice_html.data = _gdpr_notice_html()
    form.photo_consent_html.data = _photo_consent_html()
    form.assessor_competence_html.data = _assessor_competence_html()
    form.accommodation_info_html.data = _accommodation_info_html()

    form.essay_upload_enabled.data = _is_upload_module_enabled()
    form.teacher_portal_enabled.data = _teacher_portal_enabled()
    form.registration_invite_only.data = _registration_invite_required()
    form.login_email_enabled.data = _login_email_enabled()
    form.login_folded_alias_enabled.data = _login_folded_alias_enabled()
    form.email_auto_enabled.data = _parse_bool(_get_setting(SETTING_EMAIL_AUTOMATION_ENABLED, "false"))
    form.email_student_subject.data = _get_setting(SETTING_EMAIL_STUDENT_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_STUDENT_SUBJECT]
    form.email_student_body.data = _get_setting(SETTING_EMAIL_STUDENT_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_STUDENT_BODY]
    form.email_teacher_subject.data = _get_setting(SETTING_EMAIL_TEACHER_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_TEACHER_SUBJECT]
    form.email_teacher_body.data = _get_setting(SETTING_EMAIL_TEACHER_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_TEACHER_BODY]
    form.email_student_invite_subject.data = _get_setting(SETTING_EMAIL_STUDENT_INVITE_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_STUDENT_INVITE_SUBJECT]
    form.email_student_invite_body.data = _get_setting(SETTING_EMAIL_STUDENT_INVITE_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_STUDENT_INVITE_BODY]
    form.email_teacher_invite_subject.data = _get_setting(SETTING_EMAIL_TEACHER_INVITE_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_TEACHER_INVITE_SUBJECT]
    form.email_teacher_invite_body.data = _get_setting(SETTING_EMAIL_TEACHER_INVITE_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_TEACHER_INVITE_BODY]
    form.email_message_subject.data = _get_setting(SETTING_EMAIL_MESSAGE_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_MESSAGE_SUBJECT]
    form.email_message_body.data = _get_setting(SETTING_EMAIL_MESSAGE_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_MESSAGE_BODY]
    form.email_survey_publish_subject.data = _get_setting(SETTING_EMAIL_SURVEY_PUBLISH_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_SURVEY_PUBLISH_SUBJECT]
    form.email_survey_publish_body.data = _get_setting(SETTING_EMAIL_SURVEY_PUBLISH_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_SURVEY_PUBLISH_BODY]
    form.email_survey_reminder_subject.data = _get_setting(SETTING_EMAIL_SURVEY_REMINDER_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_SURVEY_REMINDER_SUBJECT]
    form.email_survey_reminder_body.data = _get_setting(SETTING_EMAIL_SURVEY_REMINDER_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_SURVEY_REMINDER_BODY]
    form.email_student_reset_link_subject.data = _get_setting(SETTING_EMAIL_STUDENT_RESET_LINK_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_STUDENT_RESET_LINK_SUBJECT]
    form.email_student_reset_link_body.data = _get_setting(SETTING_EMAIL_STUDENT_RESET_LINK_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_STUDENT_RESET_LINK_BODY]
    form.email_student_reset_subject.data = _get_setting(SETTING_EMAIL_STUDENT_RESET_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_STUDENT_RESET_SUBJECT]
    form.email_student_reset_body.data = _get_setting(SETTING_EMAIL_STUDENT_RESET_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_STUDENT_RESET_BODY]
    form.email_teacher_reset_link_subject.data = _get_setting(SETTING_EMAIL_TEACHER_RESET_LINK_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_TEACHER_RESET_LINK_SUBJECT]
    form.email_teacher_reset_link_body.data = _get_setting(SETTING_EMAIL_TEACHER_RESET_LINK_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_TEACHER_RESET_LINK_BODY]
    form.email_teacher_reset_subject.data = _get_setting(SETTING_EMAIL_TEACHER_RESET_SUBJECT) or DEFAULT_EMAIL_SUBJECTS[SETTING_EMAIL_TEACHER_RESET_SUBJECT]
    form.email_teacher_reset_body.data = _get_setting(SETTING_EMAIL_TEACHER_RESET_BODY) or DEFAULT_EMAIL_BODIES[SETTING_EMAIL_TEACHER_RESET_BODY]


@web_bp.route("/admin/settings", methods=["GET", "POST"])
def admin_settings():
    form = DashboardSettingsForm()
    config = current_app.config
    if request.method == "GET":
        _populate_admin_settings_form(form, config)

    diagnostics = _login_alias_diagnostics()
    submitted_with_enabled_alias = form.is_submitted() and bool(form.login_folded_alias_enabled.data)
    enabling_alias_from_off_to_on = submitted_with_enabled_alias and not diagnostics["enabled"]

    if form.is_submitted() and form.validate():
        if enabling_alias_from_off_to_on and not diagnostics["all_safe"]:
            form.login_folded_alias_enabled.errors.append(
                "Cannot enable folded login aliases until all folded student and teacher usernames are unique."
            )
        else:
            previous_stage1_trim_extreme_scores = _stage1_trim_extreme_scores_enabled()
            upload_value = form.upload_deadline.data.strftime("%Y-%m-%dT%H:%M") if form.upload_deadline.data else None
            stage2_value = form.stage2_deadline.data.strftime("%Y-%m-%dT%H:%M") if form.stage2_deadline.data else None
            event_end_value = form.event_end.data.strftime("%Y-%m-%dT%H:%M") if form.event_end.data else None
            event_override_value = form.event_end_override.data if form.event_end_override.data != "auto" else None
            stage1_sla_value = str(form.stage1_sla.data) if form.stage1_sla.data else None
            stage2_sla_value = str(form.stage2_sla.data) if form.stage2_sla.data else None
            stage1_trim_extreme_scores_value = bool(form.stage1_trim_extreme_scores_enabled.data)

            _set_setting(SETTING_UPLOAD_DEADLINE, upload_value)
            _set_setting(SETTING_STAGE2_DEADLINE, stage2_value)
            _set_setting(SETTING_EVENT_END, event_end_value)
            _set_setting(SETTING_EVENT_END_OVERRIDE, event_override_value)
            _set_setting(SETTING_STAGE1_SLA, stage1_sla_value)
            _set_setting(SETTING_STAGE2_SLA, stage2_sla_value)
            _set_setting(SETTING_PACKET_PRINT_MODE, form.packet_print_mode.data)
            _set_setting(
                SETTING_ESSAY_WORD_LIMIT,
                str(form.essay_word_limit.data).strip() if form.essay_word_limit.data else None,
            )
            _set_setting(SETTING_ESSAY_WORD_LIMIT_MODE, form.essay_word_limit_mode.data)
            _set_setting(
                SETTING_STAGE1_TRIM_EXTREME_SCORES,
                "true" if stage1_trim_extreme_scores_value else "false",
            )
            ethics_html = (form.ethics_consent_html.data or "").strip() or None
            gdpr_html = (form.gdpr_notice_html.data or "").strip() or None
            photo_html = (form.photo_consent_html.data or "").strip() or None
            competence_html = (form.assessor_competence_html.data or "").strip() or None
            accommodation_html = (form.accommodation_info_html.data or "").strip() or None
            _set_setting(SETTING_ETHICS_CONSENT_HTML, ethics_html)
            _set_setting(SETTING_GDPR_NOTICE_HTML, gdpr_html)
            _set_setting(SETTING_PHOTO_CONSENT_HTML, photo_html)
            _set_setting(SETTING_TEACHER_COMPETENCE_HTML, competence_html)
            _set_setting(SETTING_ACCOMMODATION_INFO_HTML, accommodation_html)
            _set_setting(SETTING_UPLOAD_MODULE_ENABLED, "true" if form.essay_upload_enabled.data else "false")
            _set_setting(SETTING_TEACHER_PORTAL_ENABLED, "true" if form.teacher_portal_enabled.data else "false")
            _set_setting(SETTING_REGISTRATION_INVITE_ONLY, "true" if form.registration_invite_only.data else "false")
            _set_setting(SETTING_LOGIN_EMAIL_ENABLED, "true" if form.login_email_enabled.data else "false")
            _set_setting(SETTING_LOGIN_FOLDED_ALIAS_ENABLED, "true" if form.login_folded_alias_enabled.data else "false")
            _set_setting(SETTING_EMAIL_AUTOMATION_ENABLED, "true" if form.email_auto_enabled.data else "false")
            _set_setting(SETTING_EMAIL_STUDENT_SUBJECT, form.email_student_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_STUDENT_BODY, form.email_student_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_TEACHER_SUBJECT, form.email_teacher_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_TEACHER_BODY, form.email_teacher_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_STUDENT_INVITE_SUBJECT, form.email_student_invite_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_STUDENT_INVITE_BODY, form.email_student_invite_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_TEACHER_INVITE_SUBJECT, form.email_teacher_invite_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_TEACHER_INVITE_BODY, form.email_teacher_invite_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_MESSAGE_SUBJECT, form.email_message_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_MESSAGE_BODY, form.email_message_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_SURVEY_PUBLISH_SUBJECT, form.email_survey_publish_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_SURVEY_PUBLISH_BODY, form.email_survey_publish_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_SURVEY_REMINDER_SUBJECT, form.email_survey_reminder_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_SURVEY_REMINDER_BODY, form.email_survey_reminder_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_STUDENT_RESET_LINK_SUBJECT, form.email_student_reset_link_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_STUDENT_RESET_LINK_BODY, form.email_student_reset_link_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_STUDENT_RESET_SUBJECT, form.email_student_reset_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_STUDENT_RESET_BODY, form.email_student_reset_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_TEACHER_RESET_LINK_SUBJECT, form.email_teacher_reset_link_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_TEACHER_RESET_LINK_BODY, form.email_teacher_reset_link_body.data.strip() or None)
            _set_setting(SETTING_EMAIL_TEACHER_RESET_SUBJECT, form.email_teacher_reset_subject.data.strip() or None)
            _set_setting(SETTING_EMAIL_TEACHER_RESET_BODY, form.email_teacher_reset_body.data.strip() or None)
            if stage1_trim_extreme_scores_value != previous_stage1_trim_extreme_scores:
                _recompute_stage_one_rules()
            db.session.commit()
            flash("Dashboard settings updated.", "success")
            return redirect(url_for("web.admin_settings"))

    return render_template("admin/settings.html", form=form, config=config, login_alias_diagnostics=diagnostics)


@web_bp.route("/admin/exports")
def admin_exports():
    if not _current_admin():
        abort(403)

    database_url = _get_database_url() or ""
    database_available = database_url.startswith("sqlite") or "mysql" in database_url
    return render_template(
        "admin/exports.html",
        database_available=database_available,
        teacher_portal_enabled=_teacher_portal_enabled(),
    )


@web_bp.route("/admin/exports/database", methods=["POST"])
def admin_export_database():
    if not _current_admin():
        abort(403)
    result = _export_database()
    if not result:
        flash("Database export is not available for the current configuration.", "error")
        return redirect(url_for("web.admin_exports"))
    path, filename = result
    log_activity("database_exported", f"Database export {filename} generated", object_type="backup", object_id=filename)
    return send_file(path, as_attachment=True, download_name=filename, mimetype="application/sql")


@web_bp.route("/admin/exports/essays", methods=["POST"])
def admin_export_essays():
    if not _current_admin():
        abort(403)
    only_finalized = request.form.get("finalized_only") == "1"
    result = _zip_essays(only_finalized=only_finalized)
    if not result:
        flash("No essays available to export.", "warning")
        return redirect(url_for("web.admin_exports"))
    zip_buffer, filename = result
    log_activity(
        "essays_exported",
        f"Essay archive {filename} generated",
        object_type="essay_export",
        details={"finalized_only": only_finalized},
    )
    return send_file(zip_buffer, as_attachment=True, download_name=filename, mimetype="application/zip")


@web_bp.route("/admin/help")
def admin_help():
    if not _current_admin():
        abort(403)
    return render_template("admin/help.html")


@web_bp.route("/admin/messages/quick", methods=["POST"])
def admin_message_quick():
    if not _current_admin():
        abort(403)

    is_htmx = _is_htmx_request()
    success = False
    form = QuickTeacherMessageForm()
    if form.validate_on_submit():
        try:
            teacher_id = int(form.teacher_id.data)
        except (TypeError, ValueError):
            flash("Invalid teacher selected.", "error")
            if is_htmx:
                return render_template("partials/toast_oob.html")
            return redirect(request.referrer or url_for("web.admin_messages"))

        teacher = db.session.get(Teacher, teacher_id)
        if not teacher:
            flash("Teacher could not be found.", "error")
            if is_htmx:
                return render_template("partials/toast_oob.html")
            return redirect(request.referrer or url_for("web.admin_messages"))

        try:
            priority = TeacherMessagePriority(form.priority.data)
        except ValueError:
            flash("Unknown priority selected.", "error")
            if is_htmx:
                return render_template("partials/toast_oob.html")
            return redirect(request.referrer or url_for("web.admin_messages"))

        message = TeacherMessage(
            title=form.title.data.strip(),
            body=form.body.data.strip(),
            priority=priority,
            status=TeacherMessageStatus.ACTIVE,
            deliver_to_all=False,
            teacher_id=teacher.id,
            category=form.category.data.strip() or None,
            creator_id=_current_admin(),
        )
        db.session.add(message)
        db.session.flush()
        log_activity(
            "teacher_message_created",
            f"Message '{message.title}' created via quick compose",
            object_type="teacher_message",
            object_id=message.id,
            details={
                "priority": message.priority.value,
                "teacher_id": teacher.id,
                "deliver_to_all": False,
                "origin": request.referrer or request.path,
            },
        )
        db.session.commit()
        _send_teacher_message_notifications(message)
        flash(f"Message sent to {teacher.first_name} {teacher.last_name}.", "success")
        success = True
    else:
        for field_errors in form.errors.values():
            for error in field_errors:
                flash(error, "error")

    if is_htmx:
        response = make_response(render_template("partials/toast_oob.html"))
        if success:
            response.headers["HX-Trigger"] = json.dumps({"quick-message-sent": True})
        return response
    return redirect(request.referrer or url_for("web.admin_messages"))


@web_bp.route("/admin/messages", methods=["GET", "POST"])
def admin_messages():
    if not _current_admin():
        abort(403)

    is_htmx = _is_htmx_request()
    form = TeacherMessageForm()
    _prepare_message_form(form)

    templates = TeacherMessageTemplate.query.order_by(TeacherMessageTemplate.name.asc()).all()
    selected_template_id = request.args.get("template", type=int)
    selected_template = None
    if request.method == "GET" and selected_template_id:
        selected_template = next((tpl for tpl in templates if tpl.id == selected_template_id), None)
        if selected_template:
            form.title.data = selected_template.title
            form.body.data = selected_template.body
            form.category.data = selected_template.category or ""

    def _render_messages(current_form: TeacherMessageForm, template_id: int | None):
        messages = TeacherMessage.query.order_by(TeacherMessage.created_at.desc()).all()
        template_name = "admin/partials/messages_content.html" if is_htmx else "admin/messages.html"
        context = {
            "form": current_form,
            "templates": templates,
            "selected_template_id": template_id,
            "messages": messages,
            "priorities": list(TeacherMessagePriority),
            "statuses": list(TeacherMessageStatus),
            "TeacherMessageStatus": TeacherMessageStatus,
            "is_htmx": is_htmx,
        }
        if not is_htmx:
            return render_template(template_name, **context)
        return _render_htmx_template(template_name, **context)

    if form.validate_on_submit():
        deliver_to_all = form.teacher_id.data == 0
        teacher_id = None if deliver_to_all else form.teacher_id.data

        message = TeacherMessage(
            title=form.title.data.strip(),
            body=form.body.data.strip(),
            priority=TeacherMessagePriority(form.priority.data),
            status=TeacherMessageStatus.ACTIVE,
            deliver_to_all=deliver_to_all,
            teacher_id=teacher_id,
            category=(form.category.data.strip() or None),
            creator_id=_current_admin(),
        )
        db.session.add(message)
        db.session.flush()
        log_activity(
            "teacher_message_created",
            f"Message '{message.title}' created",
            object_type="teacher_message",
            object_id=message.id,
            details={
                "priority": message.priority.value,
                "teacher_id": teacher_id,
                "deliver_to_all": deliver_to_all,
                "origin": request.referrer or request.path,
            },
        )
        db.session.commit()
        _send_teacher_message_notifications(message)
        flash("Message posted.", "success")
        if is_htmx:
            fresh_form = TeacherMessageForm()
            _prepare_message_form(fresh_form)
            return _render_messages(fresh_form, None)
        return redirect(url_for("web.admin_messages"))

    return _render_messages(form, selected_template.id if selected_template else None)


@web_bp.route("/admin/message-templates")
def admin_message_templates():
    if not _current_admin():
        abort(403)

    is_htmx = _is_htmx_request()
    templates = (
        TeacherMessageTemplate.query.order_by(TeacherMessageTemplate.name.asc()).all()
    )
    template_name = "admin/partials/message_templates_content.html" if is_htmx else "admin/message_templates.html"
    context = {
        "templates": templates,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/message-templates/new", methods=["GET", "POST"])
def admin_message_template_new():
    if not _current_admin():
        abort(403)

    form = TeacherMessageTemplateForm()
    if form.validate_on_submit():
        template = TeacherMessageTemplate(
            name=form.name.data.strip(),
            title=form.title.data.strip(),
            body=form.body.data.strip(),
            category=form.category.data.strip() or None,
        )
        db.session.add(template)
        try:
            db.session.flush()
        except IntegrityError:
            db.session.rollback()
            form.name.errors.append("A template with this name already exists.")
        else:
            log_activity(
                "teacher_message_template_created",
                f"Template '{template.name}' created",
                object_type="teacher_message_template",
                object_id=template.id,
            )
            db.session.commit()
            flash("Template created.", "success")
            return redirect(url_for("web.admin_message_templates"))
    return render_template(
        "admin/message_template_form.html",
        form=form,
        template=None,
    )


@web_bp.route("/admin/message-templates/<int:template_id>/edit", methods=["GET", "POST"])
def admin_message_template_edit(template_id: int):
    if not _current_admin():
        abort(403)

    template = db.session.get(TeacherMessageTemplate, template_id)
    if not template:
        abort(404)

    form = TeacherMessageTemplateForm(obj=template)

    if form.validate_on_submit():
        template.name = form.name.data.strip()
        template.title = form.title.data.strip()
        template.body = form.body.data.strip()
        template.category = form.category.data.strip() or None
        try:
            db.session.flush()
        except IntegrityError:
            db.session.rollback()
            form.name.errors.append("A template with this name already exists.")
        else:
            log_activity(
                "teacher_message_template_updated",
                f"Template '{template.name}' updated",
                object_type="teacher_message_template",
                object_id=template.id,
            )
            db.session.commit()
            flash("Template updated.", "success")
            return redirect(url_for("web.admin_message_templates"))

    return render_template(
        "admin/message_template_form.html",
        form=form,
        template=template,
    )


@web_bp.route("/admin/message-templates/<int:template_id>/delete", methods=["POST"])
def admin_message_template_delete(template_id: int):
    if not _current_admin():
        abort(403)

    template = db.session.get(TeacherMessageTemplate, template_id)
    if not template:
        abort(404)

    db.session.delete(template)
    log_activity(
        "teacher_message_template_deleted",
        f"Template '{template.name}' deleted",
        object_type="teacher_message_template",
        object_id=template.id,
    )
    db.session.commit()
    flash("Template deleted.", "success")
    if _is_htmx_request():
        return admin_message_templates()
    return redirect(url_for("web.admin_message_templates"))


@web_bp.route("/admin/messages/<int:message_id>/edit", methods=["GET", "POST"])
def admin_message_edit(message_id: int):
    if not _current_admin():
        abort(403)

    message = db.session.get(TeacherMessage, message_id)
    if not message:
        abort(404)

    form = TeacherMessageForm(obj=message)
    if not form.is_submitted():
        _prepare_message_form(form, message.teacher_id or 0)
        form.priority.data = message.priority.value
    else:
        _prepare_message_form(form)

    if form.validate_on_submit():
        deliver_to_all = form.teacher_id.data == 0
        teacher_id = None if deliver_to_all else form.teacher_id.data

        message.title = form.title.data.strip()
        message.body = form.body.data.strip()
        message.priority = TeacherMessagePriority(form.priority.data)
        message.deliver_to_all = deliver_to_all
        message.teacher_id = teacher_id
        message.category = form.category.data.strip() or None
        message.updater_id = _current_admin()
        db.session.flush()
        log_activity(
            "teacher_message_updated",
            f"Message '{message.title}' updated",
            object_type="teacher_message",
            object_id=message.id,
            details={
                "teacher_id": teacher_id,
                "deliver_to_all": deliver_to_all,
            },
        )
        db.session.commit()
        flash("Message updated.", "success")
        return redirect(url_for("web.admin_messages"))

    return render_template(
        "admin/messages_edit.html",
        form=form,
        message=message,
    )


@web_bp.route("/admin/messages/<int:message_id>/archive", methods=["POST"])
def admin_message_archive(message_id: int):
    if not _current_admin():
        abort(403)
    message = db.session.get(TeacherMessage, message_id)
    if not message:
        abort(404)
    message.status = TeacherMessageStatus.ARCHIVED
    message.updater_id = _current_admin()
    db.session.flush()
    log_activity(
        "teacher_message_archived",
        f"Message '{message.title}' archived",
        object_type="teacher_message",
        object_id=message.id,
        details={"status": message.status.value},
    )
    db.session.commit()
    flash("Message archived.", "success")
    if _is_htmx_request():
        return admin_messages()
    return redirect(url_for("web.admin_messages"))


@web_bp.route("/admin/messages/<int:message_id>/restore", methods=["POST"])
def admin_message_restore(message_id: int):
    if not _current_admin():
        abort(403)
    message = db.session.get(TeacherMessage, message_id)
    if not message:
        abort(404)
    message.status = TeacherMessageStatus.ACTIVE
    message.updater_id = _current_admin()
    db.session.flush()
    log_activity(
        "teacher_message_restored",
        f"Message '{message.title}' restored",
        object_type="teacher_message",
        object_id=message.id,
        details={"status": message.status.value},
    )
    db.session.commit()
    flash("Message restored.", "success")
    if _is_htmx_request():
        return admin_messages()
    return redirect(url_for("web.admin_messages"))


@web_bp.route("/admin/messages/<int:message_id>/delete", methods=["POST"])
def admin_message_delete(message_id: int):
    if not _current_admin():
        abort(403)
    message = db.session.get(TeacherMessage, message_id)
    if not message:
        abort(404)
    db.session.delete(message)
    db.session.flush()
    log_activity(
        "teacher_message_deleted",
        f"Message '{message.title}' deleted",
        object_type="teacher_message",
        object_id=message.id,
    )
    db.session.commit()
    flash("Message deleted.", "success")
    if _is_htmx_request():
        return admin_messages()
    return redirect(url_for("web.admin_messages"))


@web_bp.route("/admin/static-pages")
def admin_static_pages():
    if not _current_admin():
        abort(403)

    pages = StaticPage.query.order_by(StaticPage.updated_at.desc(), StaticPage.id.desc()).all()
    return render_template("admin/static_pages/index.html", pages=pages)


@web_bp.route("/admin/static-pages/new", methods=["GET", "POST"])
def admin_static_page_new():
    if not _current_admin():
        abort(403)

    form = StaticPageForm()
    if form.validate_on_submit():
        slug = _normalize_static_page_slug(form.slug.data)
        form.slug.data = slug
        if not STATIC_PAGE_SLUG_RE.fullmatch(slug):
            form.slug.errors.append("Use lowercase letters, numbers, and single hyphens only.")
        else:
            existing = StaticPage.query.filter_by(slug=slug).first()
            if existing:
                form.slug.errors.append("A page with this slug already exists.")
            else:
                page = StaticPage(
                    title=form.title.data.strip(),
                    slug=slug,
                    html_body=form.html_body.data or "",
                    created_by=_current_admin(),
                    updated_by=_current_admin(),
                )
                db.session.add(page)
                db.session.flush()

                uploaded = form.attachment.data
                if uploaded and (uploaded.filename or "").strip():
                    try:
                        attachment = _save_static_page_attachment(page, uploaded)
                    except ValueError as exc:
                        db.session.rollback()
                        form.attachment.errors.append(str(exc))
                    else:
                        log_activity(
                            "static_page_attachment_added",
                            f"Attachment '{attachment.original_filename}' added to static page '{page.title}'",
                            object_type="static_page",
                            object_id=page.id,
                            details={
                                "slug": page.slug,
                                "attachment_id": attachment.id,
                                "filename": attachment.original_filename,
                            },
                        )

                if not form.errors:
                    log_activity(
                        "static_page_created",
                        f"Static page '{page.title}' created",
                        object_type="static_page",
                        object_id=page.id,
                        details={"slug": page.slug},
                    )
                    db.session.commit()
                    flash("Static page created.", "success")
                    return redirect(url_for("web.admin_static_pages"))

    return render_template(
        "admin/static_pages/form.html",
        form=form,
        page=None,
    )


@web_bp.route("/admin/static-pages/<int:page_id>/edit", methods=["GET", "POST"])
def admin_static_page_edit(page_id: int):
    if not _current_admin():
        abort(403)

    page = db.session.get(StaticPage, page_id)
    if not page:
        abort(404)

    form = StaticPageForm(obj=page)
    if request.method == "GET":
        form.slug.data = page.slug

    if form.validate_on_submit():
        slug = _normalize_static_page_slug(form.slug.data)
        form.slug.data = slug
        if not STATIC_PAGE_SLUG_RE.fullmatch(slug):
            form.slug.errors.append("Use lowercase letters, numbers, and single hyphens only.")
        else:
            existing = StaticPage.query.filter(StaticPage.slug == slug, StaticPage.id != page.id).first()
            if existing:
                form.slug.errors.append("A page with this slug already exists.")
            else:
                page.title = form.title.data.strip()
                page.slug = slug
                page.html_body = form.html_body.data or ""
                page.updated_by = _current_admin()

                uploaded = form.attachment.data
                if uploaded and (uploaded.filename or "").strip():
                    try:
                        attachment = _save_static_page_attachment(page, uploaded)
                    except ValueError as exc:
                        db.session.rollback()
                        form.attachment.errors.append(str(exc))
                    else:
                        log_activity(
                            "static_page_attachment_added",
                            f"Attachment '{attachment.original_filename}' added to static page '{page.title}'",
                            object_type="static_page",
                            object_id=page.id,
                            details={
                                "slug": page.slug,
                                "attachment_id": attachment.id,
                                "filename": attachment.original_filename,
                            },
                        )

                if not form.errors:
                    db.session.flush()
                    log_activity(
                        "static_page_updated",
                        f"Static page '{page.title}' updated",
                        object_type="static_page",
                        object_id=page.id,
                        details={"slug": page.slug},
                    )
                    db.session.commit()
                    flash("Static page updated.", "success")
                    return redirect(url_for("web.admin_static_pages"))

    return render_template(
        "admin/static_pages/form.html",
        form=form,
        page=page,
    )


@web_bp.route("/admin/static-pages/<int:page_id>/delete", methods=["POST"])
def admin_static_page_delete(page_id: int):
    if not _current_admin():
        abort(403)

    page = db.session.get(StaticPage, page_id)
    if not page:
        abort(404)

    file_names = [attachment.stored_filename for attachment in page.attachments]
    page_title = page.title
    page_slug = page.slug

    db.session.delete(page)
    db.session.flush()
    log_activity(
        "static_page_deleted",
        f"Static page '{page_title}' deleted",
        object_type="static_page",
        object_id=page_id,
        details={"slug": page_slug, "attachments_removed": len(file_names)},
    )
    db.session.commit()

    for file_name in file_names:
        _delete_static_page_attachment_file(file_name)

    flash("Static page deleted.", "success")
    return redirect(url_for("web.admin_static_pages"))


@web_bp.route("/admin/static-pages/<int:page_id>/attachments/<int:attachment_id>/delete", methods=["POST"])
def admin_static_page_attachment_delete(page_id: int, attachment_id: int):
    if not _current_admin():
        abort(403)

    page = db.session.get(StaticPage, page_id)
    if not page:
        abort(404)

    attachment = db.session.get(StaticPageAttachment, attachment_id)
    if not attachment or attachment.page_id != page.id:
        abort(404)

    file_name = attachment.stored_filename
    original_name = attachment.original_filename
    page_slug = page.slug

    db.session.delete(attachment)
    db.session.flush()
    log_activity(
        "static_page_attachment_deleted",
        f"Attachment '{original_name}' removed from static page '{page.title}'",
        object_type="static_page",
        object_id=page.id,
        details={"slug": page_slug, "attachment_id": attachment_id, "filename": original_name},
    )
    db.session.commit()

    _delete_static_page_attachment_file(file_name)

    flash("Attachment removed.", "success")
    return redirect(url_for("web.admin_static_page_edit", page_id=page.id))


@web_bp.route("/admin/communications")
def admin_communications():
    if not _current_admin():
        abort(403)

    campaigns = (
        CommunicationCampaign.query.order_by(CommunicationCampaign.created_at.desc()).all()
    )
    stats = _campaign_recipient_stats(campaign.id for campaign in campaigns)
    sendgrid_ready = _sendgrid_configured()
    if not sendgrid_ready:
        flash(
            "SendGrid is not configured. Drafts can be prepared, but sending is disabled until credentials are added in Settings.",
            "warning",
        )

    return render_template(
        "admin/communications/index.html",
        campaigns=campaigns,
        stats=stats,
        CommunicationStatus=CommunicationStatus,
        CommunicationRecipientStatus=CommunicationRecipientStatus,
        sendgrid_ready=sendgrid_ready,
    )


@web_bp.route("/admin/communications/new", methods=["GET", "POST"])
def admin_communications_new():
    if not _current_admin():
        abort(403)

    form = CommunicationComposeForm()
    country_choices = _communication_country_choices()
    form.countries.choices = country_choices
    sendgrid_ready = _sendgrid_configured()
    if not sendgrid_ready:
        flash(
            "SendGrid is not currently configured. You can prepare the message and audience, but sending will be blocked until credentials are added in Settings.",
            "warning",
        )

    if request.method == "GET" and not form.audience.data:
        form.audience.data = CommunicationAudience.BOTH.value

    if form.validate_on_submit():
        try:
            audience = CommunicationAudience(form.audience.data)
        except ValueError:
            form.audience.errors.append("Unknown audience selected.")
        else:
            selected_countries = [code for code in (form.countries.data or []) if code]
            recipients = _build_campaign_recipients(
                audience,
                selected_countries,
                bool(form.include_absent.data),
            )
            if not recipients:
                flash("No recipients match the selected filters.", "warning")
            else:
                subject = form.subject.data.strip()
                html_body = form.html_body.data or ""
                plain_text = _html_to_text(html_body) or subject
                campaign = CommunicationCampaign(
                    subject=subject,
                    html_body=html_body,
                    text_body=plain_text,
                    audience=audience,
                    country_filter=sorted(selected_countries) if selected_countries else None,
                    include_absent=bool(form.include_absent.data),
                    status=CommunicationStatus.DRAFT,
                    created_by=_current_admin(),
                    total_recipients=len(recipients),
                    sent_count=0,
                    failed_count=0,
                )
                for recipient in recipients:
                    campaign.recipients.append(recipient)
                db.session.add(campaign)
                db.session.flush()
                log_activity(
                    "communication_campaign_created",
                    f"Campaign '{campaign.subject}' drafted",
                    object_type="communication_campaign",
                    object_id=campaign.id,
                    details={
                        "audience": campaign.audience.value,
                        "total_recipients": campaign.total_recipients,
                        "countries": campaign.country_filter or [],
                        "include_absent": campaign.include_absent,
                    },
                )
                db.session.commit()
                flash(
                    f"Draft created with {campaign.total_recipients} recipient(s). Review the email before sending.",
                    "success",
                )
                return redirect(url_for("web.admin_communications_preview", campaign_id=campaign.id))

    return render_template(
        "admin/communications/new.html",
        form=form,
        campaign=None,
        sendgrid_ready=sendgrid_ready,
    )


@web_bp.route("/admin/communications/<int:campaign_id>/edit", methods=["GET", "POST"])
def admin_communications_edit(campaign_id: int):
    if not _current_admin():
        abort(403)
    campaign = db.session.get(CommunicationCampaign, campaign_id)
    if not campaign:
        abort(404)
    if campaign.status not in {CommunicationStatus.DRAFT, CommunicationStatus.FAILED}:
        flash("Only draft or failed campaigns can be edited.", "error")
        return redirect(url_for("web.admin_communications_detail", campaign_id=campaign.id))

    form = CommunicationComposeForm()
    country_choices = _communication_country_choices()
    existing_countries = set(country for country, _ in country_choices)
    selected_countries = campaign.country_filter or []
    sendgrid_ready = _sendgrid_configured()
    if not sendgrid_ready:
        flash(
            "SendGrid is not currently configured. You can prepare the message and audience, but sending will be blocked until credentials are added in Settings.",
            "warning",
        )
    for country in selected_countries:
        if country and country not in existing_countries:
            country_choices.append((country, country))
            existing_countries.add(country)
    country_choices.sort(key=lambda item: item[1])
    form.countries.choices = country_choices
    form.submit_preview.label.text = "Save changes"

    if request.method == "GET":
        form.subject.data = campaign.subject
        form.html_body.data = campaign.html_body
        form.audience.data = campaign.audience.value
        form.countries.data = selected_countries
        form.include_absent.data = bool(campaign.include_absent)

    if form.validate_on_submit():
        try:
            audience = CommunicationAudience(form.audience.data)
        except ValueError:
            form.audience.errors.append("Unknown audience selected.")
        else:
            selected_countries = [code for code in (form.countries.data or []) if code]
            recipients = _build_campaign_recipients(
                audience,
                selected_countries,
                bool(form.include_absent.data),
            )
            if not recipients:
                flash("No recipients match the selected filters.", "warning")
            else:
                subject = form.subject.data.strip()
                html_body = form.html_body.data or ""
                plain_text = _html_to_text(html_body) or subject

                # reset recipient list
                campaign.recipients.clear()
                db.session.flush()
                for recipient in recipients:
                    campaign.recipients.append(recipient)

                campaign.subject = subject
                campaign.html_body = html_body
                campaign.text_body = plain_text
                campaign.audience = audience
                campaign.country_filter = sorted(selected_countries) if selected_countries else None
                campaign.include_absent = bool(form.include_absent.data)
                campaign.total_recipients = len(recipients)
                campaign.sent_count = 0
                campaign.failed_count = 0
                campaign.sent_at = None
                campaign.sent_by = None
                campaign.status = CommunicationStatus.DRAFT

                db.session.flush()
                log_activity(
                    "communication_campaign_updated",
                    f"Campaign '{campaign.subject}' updated",
                    object_type="communication_campaign",
                    object_id=campaign.id,
                    details={
                        "audience": campaign.audience.value,
                        "total_recipients": campaign.total_recipients,
                        "countries": campaign.country_filter or [],
                        "include_absent": campaign.include_absent,
                    },
                )
                db.session.commit()
                flash("Draft updated. Review the email before sending.", "success")
                return redirect(url_for("web.admin_communications_preview", campaign_id=campaign.id))

    return render_template(
        "admin/communications/new.html",
        form=form,
        campaign=campaign,
        sendgrid_ready=sendgrid_ready,
    )


@web_bp.route("/admin/communications/<int:campaign_id>/preview", methods=["GET", "POST"])
def admin_communications_preview(campaign_id: int):
    if not _current_admin():
        abort(403)
    campaign = db.session.get(CommunicationCampaign, campaign_id)
    if not campaign:
        abort(404)

    form = CommunicationSendForm()
    sendgrid_ready = _sendgrid_configured()
    if not sendgrid_ready:
        flash("SendGrid is not configured. Sending is currently disabled.", "warning")
    recipients = list(campaign.recipients)
    role_counts = _campaign_role_counts(campaign)
    status_counts = {status.value: 0 for status in CommunicationRecipientStatus}
    for recipient in recipients:
        status_counts[recipient.status.value] = status_counts.get(recipient.status.value, 0) + 1

    if form.validate_on_submit():
        if campaign.status not in {CommunicationStatus.DRAFT, CommunicationStatus.FAILED}:
            flash("This campaign has already been processed.", "info")
            return redirect(url_for("web.admin_communications_detail", campaign_id=campaign.id))
        if not _sendgrid_configured():
            flash("SendGrid is not configured. Update the settings before sending.", "error")
            return redirect(url_for("web.admin_communications_preview", campaign_id=campaign.id))
        campaign.sent_by = _current_admin()
        campaign.status = CommunicationStatus.SENDING
        db.session.flush()
        result = _deliver_campaign(campaign, only_failed=False)
        if result["sent"]:
            campaign.sent_at = datetime.utcnow()
        db.session.flush()
        log_activity(
            "communication_campaign_sent",
            f"Campaign '{campaign.subject}' dispatched",
            object_type="communication_campaign",
            object_id=campaign.id,
            details={
                "sent": result["sent"],
                "failed": result["failed"],
                "total": campaign.total_recipients,
            },
        )
        db.session.commit()
        if result["sent"] and not result["failed"]:
            flash("Email sent to all recipients.", "success")
        elif result["sent"]:
            flash(
                f"Email sent, but {result['failed']} recipient(s) failed. Review the campaign log for details.",
                "warning",
            )
        else:
            flash("No emails were sent. All recipients may have already been processed.", "info")
        return redirect(url_for("web.admin_communications_detail", campaign_id=campaign.id))

    if request.method == "POST":
        flash("Confirm the send before continuing.", "error")

    return render_template(
        "admin/communications/preview.html",
        campaign=campaign,
        form=form,
        role_counts=role_counts,
        status_counts=status_counts,
        sample_recipients=_campaign_sample_recipients(campaign, limit=10),
        sendgrid_ready=sendgrid_ready,
        CommunicationRecipientStatus=CommunicationRecipientStatus,
        CommunicationRecipientRole=CommunicationRecipientRole,
    )


@web_bp.route("/admin/communications/<int:campaign_id>")
def admin_communications_detail(campaign_id: int):
    if not _current_admin():
        abort(403)
    campaign = db.session.get(CommunicationCampaign, campaign_id)
    if not campaign:
        abort(404)

    is_htmx = _is_htmx_request()
    sendgrid_ready = _sendgrid_configured()
    if not sendgrid_ready:
        flash(
            "SendGrid is not configured. Resending failed recipients is disabled until credentials are added.",
            "warning",
        )
    recipients = list(campaign.recipients)
    status_counts = {status.value: 0 for status in CommunicationRecipientStatus}
    for recipient in recipients:
        status_counts[recipient.status.value] = status_counts.get(recipient.status.value, 0) + 1
    role_counts = _campaign_role_counts(campaign)

    status_filter_value = request.args.get("status")
    status_filter = None
    if status_filter_value:
        try:
            status_filter = CommunicationRecipientStatus(status_filter_value)
        except ValueError:
            status_filter = None
    if status_filter:
        recipients = [recipient for recipient in recipients if recipient.status == status_filter]

    recipients.sort(
        key=lambda recipient: (
            recipient.role.value,
            (recipient.last_name or "").lower(),
            (recipient.first_name or "").lower(),
            recipient.email.lower(),
        )
    )

    retry_form = CommunicationRetryForm()

    template_name = "admin/partials/communications_detail_content.html" if is_htmx else "admin/communications/detail.html"
    context = {
        "campaign": campaign,
        "recipients": recipients,
        "status_counts": status_counts,
        "role_counts": role_counts,
        "status_filter": status_filter.value if status_filter else "all",
        "sendgrid_ready": sendgrid_ready,
        "retry_form": retry_form,
        "CommunicationRecipientStatus": CommunicationRecipientStatus,
        "CommunicationRecipientRole": CommunicationRecipientRole,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/communications/<int:campaign_id>/resend-failed", methods=["POST"])
def admin_communications_resend_failed(campaign_id: int):
    if not _current_admin():
        abort(403)
    form = CommunicationRetryForm()
    if not form.validate_on_submit():
        flash("Unable to process the request.", "error")
        return redirect(url_for("web.admin_communications_detail", campaign_id=campaign_id))
    campaign = db.session.get(CommunicationCampaign, campaign_id)
    if not campaign:
        abort(404)

    failed_recipients = [recipient for recipient in campaign.recipients if recipient.status == CommunicationRecipientStatus.FAILED]
    if not failed_recipients:
        flash("There are no failed recipients to resend.", "info")
        return redirect(url_for("web.admin_communications_detail", campaign_id=campaign.id))

    if not _sendgrid_configured():
        flash("SendGrid is not configured. Update the settings before resending.", "error")
        return redirect(url_for("web.admin_communications_detail", campaign_id=campaign.id))

    campaign.sent_by = _current_admin()
    campaign.status = CommunicationStatus.SENDING
    db.session.flush()
    result = _deliver_campaign(campaign, only_failed=True)
    if result["sent"]:
        campaign.sent_at = datetime.utcnow()
    db.session.flush()
    log_activity(
        "communication_campaign_resend_failed",
        f"Resent failed recipients for campaign '{campaign.subject}'",
        object_type="communication_campaign",
        object_id=campaign.id,
        details={
            "resent": result["sent"],
            "remaining_failures": campaign.failed_count,
        },
    )
    db.session.commit()

    if result["sent"] and not result["failed"]:
        flash("Failed recipients resent successfully.", "success")
    elif result["sent"]:
        flash(
            f"Retried {result['sent']} recipient(s), but {result['failed']} still failed.",
            "warning",
        )
    else:
        flash("No failed recipients were resent. Review the campaign details.", "info")
    if _is_htmx_request():
        return admin_communications_detail(campaign.id)
    return redirect(url_for("web.admin_communications_detail", campaign_id=campaign.id))


def _survey_stats_snapshot(survey: Survey) -> dict[str, int]:
    total = len(survey.assignments)
    completed = sum(
        1 for assignment in survey.assignments if assignment.status == SurveyAssignmentStatus.COMPLETED
    )
    responses = (
        db.session.query(func.count(SurveyResponse.id)).filter_by(survey_id=survey.id).scalar()
        or 0
    )
    return {
        "total": total,
        "completed": completed,
        "pending": max(total - completed, 0),
        "responses": responses,
    }


@web_bp.route("/admin/surveys")
def admin_surveys():
    if not _current_admin():
        abort(403)

    is_htmx = _is_htmx_request()
    surveys = (
        Survey.query.options(
            joinedload(Survey.assignments).joinedload(SurveyAssignment.teacher),
            joinedload(Survey.assignments).joinedload(SurveyAssignment.student),
        )
        .order_by(Survey.created_at.desc())
        .all()
    )

    survey_ids = [survey.id for survey in surveys]
    response_counts: dict[int, int] = {}
    if survey_ids:
        rows = (
            db.session.query(SurveyResponse.survey_id, func.count(SurveyResponse.id))
            .filter(SurveyResponse.survey_id.in_(survey_ids))
            .group_by(SurveyResponse.survey_id)
            .all()
        )
        response_counts = {survey_id: count for survey_id, count in rows}

    stats: dict[int, dict[str, int]] = {}
    for survey in surveys:
        total = len(survey.assignments)
        completed = sum(1 for assignment in survey.assignments if assignment.status == SurveyAssignmentStatus.COMPLETED)
        stats[survey.id] = {
            "total": total,
            "completed": completed,
            "pending": total - completed,
            "responses": response_counts.get(survey.id, 0),
        }

    template_name = "admin/partials/surveys_index_content.html" if is_htmx else "admin/surveys/index.html"
    context = {
        "surveys": surveys,
        "stats": stats,
        "SurveyStatus": SurveyStatus,
        "SurveyAudience": SurveyAudience,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/surveys/new", methods=["GET", "POST"])
def admin_new_survey():
    if not _current_admin():
        abort(403)

    form = SurveyForm()
    form.submit.label.text = "Create survey"

    if form.validate_on_submit():
        survey = Survey(
            title=form.title.data.strip(),
            description=(form.description.data.strip() or None) if form.description.data else None,
            audience=SurveyAudience(form.audience.data),
            allow_edit=bool(form.allow_edit.data),
            open_at=form.open_at.data,
            close_at=form.close_at.data,
            status=SurveyStatus.DRAFT,
            created_by=_current_admin(),
        )
        db.session.add(survey)
        db.session.flush()
        log_activity(
            "survey_created",
            f"Survey '{survey.title}' created",
            object_type="survey",
            object_id=survey.id,
            details={
                "audience": survey.audience.value,
                "allow_edit": survey.allow_edit,
            },
        )
        db.session.commit()
        flash("Survey created. Add questions before publishing.", "success")
        return redirect(url_for("web.admin_edit_survey", survey_id=survey.id))

    return render_template(
        "admin/surveys/edit.html",
        form=form,
        question_form=None,
        survey=None,
        questions=[],
        can_edit=True,
        SurveyStatus=SurveyStatus,
        responses_count=0,
    )


def _render_survey_edit_view(
    *,
    survey: Survey,
    form: SurveyForm,
    question_form: SurveyQuestionForm | None,
    can_edit: bool,
    responses_count: int,
    is_htmx: bool,
    set_dates: bool,
):
    if set_dates:
        form.open_at.data = survey.open_at
        form.close_at.data = survey.close_at
    questions = _ordered_survey_questions(survey)
    template_name = "admin/partials/surveys_edit_content.html" if is_htmx else "admin/surveys/edit.html"
    context = {
        "survey": survey,
        "form": form,
        "question_form": question_form,
        "questions": questions,
        "can_edit": can_edit,
        "SurveyQuestionType": SurveyQuestionType,
        "SurveyStatus": SurveyStatus,
        "responses_count": responses_count,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/surveys/<int:survey_id>/edit", methods=["GET", "POST"])
def admin_edit_survey(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)

    is_htmx = _is_htmx_request()
    form = SurveyForm(obj=survey)
    form.submit.label.text = "Save changes"
    question_form = SurveyQuestionForm()
    question_form.submit.label.text = "Add question"
    can_edit = survey.status == SurveyStatus.DRAFT
    responses_count = (
        db.session.query(func.count(SurveyResponse.id)).filter_by(survey_id=survey.id).scalar()
    )

    if request.method == "POST":
        form_name = request.form.get("form_name")
        if form_name == "survey":
            if not can_edit:
                flash("Published surveys cannot be edited.", "error")
            elif form.validate():
                survey.title = form.title.data.strip()
                survey.description = (form.description.data.strip() or None) if form.description.data else None
                survey.audience = SurveyAudience(form.audience.data)
                survey.allow_edit = bool(form.allow_edit.data)
                survey.open_at = form.open_at.data
                survey.close_at = form.close_at.data
                db.session.flush()
                log_activity(
                    "survey_updated",
                    f"Survey '{survey.title}' updated",
                    object_type="survey",
                    object_id=survey.id,
                    details={
                        "audience": survey.audience.value,
                        "allow_edit": survey.allow_edit,
                    },
                )
                db.session.commit()
                flash("Survey details updated.", "success")
                if is_htmx:
                    fresh_form = SurveyForm(obj=survey)
                    fresh_form.submit.label.text = "Save changes"
                    fresh_question_form = SurveyQuestionForm()
                    fresh_question_form.submit.label.text = "Add question"
                    return _render_survey_edit_view(
                        survey=survey,
                        form=fresh_form,
                        question_form=fresh_question_form,
                        can_edit=can_edit,
                        responses_count=responses_count,
                        is_htmx=is_htmx,
                        set_dates=True,
                    )
                return redirect(url_for("web.admin_edit_survey", survey_id=survey.id))
        elif form_name == "question":
            if not can_edit:
                flash("Questions cannot be modified once the survey is published.", "error")
            elif question_form.validate():
                question = SurveyQuestion(
                    survey=survey,
                    prompt=question_form.prompt.data.strip(),
                    help_text=(question_form.help_text.data.strip() or None) if question_form.help_text.data else None,
                    type=SurveyQuestionType(question_form.type.data),
                    required=bool(question_form.required.data),
                    options=_parse_survey_options(question_form.options_text.data, question_form.type.data),
                )
                question.position = _next_question_position(survey)
                db.session.add(question)
                db.session.flush()
                log_activity(
                    "survey_question_created",
                    f"Question added to survey '{survey.title}'",
                    object_type="survey_question",
                    object_id=question.id,
                    details={
                        "survey_id": survey.id,
                        "type": question.type.value,
                        "required": question.required,
                    },
                )
                db.session.commit()
                flash("Question added.", "success")
                if is_htmx:
                    fresh_form = SurveyForm(obj=survey)
                    fresh_form.submit.label.text = "Save changes"
                    fresh_question_form = SurveyQuestionForm()
                    fresh_question_form.submit.label.text = "Add question"
                    return _render_survey_edit_view(
                        survey=survey,
                        form=fresh_form,
                        question_form=fresh_question_form,
                        can_edit=can_edit,
                        responses_count=responses_count,
                        is_htmx=is_htmx,
                        set_dates=True,
                    )
                return redirect(url_for("web.admin_edit_survey", survey_id=survey.id, _anchor="questions"))

    return _render_survey_edit_view(
        survey=survey,
        form=form,
        question_form=question_form,
        can_edit=can_edit,
        responses_count=responses_count,
        is_htmx=is_htmx,
        set_dates=request.method == "GET",
    )


@web_bp.route("/admin/surveys/<int:survey_id>/preview")
def admin_preview_survey(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)
    questions = _ordered_survey_questions(survey)
    return render_template(
        "admin/surveys/preview.html",
        survey=survey,
        questions=questions,
        SurveyQuestionType=SurveyQuestionType,
    )


@web_bp.route("/admin/surveys/<int:survey_id>/questions/<int:question_id>/edit", methods=["GET", "POST"])
def admin_edit_survey_question(survey_id: int, question_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    question = db.session.get(SurveyQuestion, question_id)
    if not survey or not question or question.survey_id != survey.id:
        abort(404)
    if survey.status != SurveyStatus.DRAFT:
        flash("Published surveys cannot be edited.", "error")
        return redirect(url_for("web.admin_edit_survey", survey_id=survey.id))

    form = SurveyQuestionForm(obj=question)
    form.submit.label.text = "Save question"
    if request.method == "GET":
        form.options_text.data = "\n".join(question.options or [])

    if form.validate_on_submit():
        question.prompt = form.prompt.data.strip()
        question.help_text = (form.help_text.data.strip() or None) if form.help_text.data else None
        question.type = SurveyQuestionType(form.type.data)
        question.required = bool(form.required.data)
        question.options = _parse_survey_options(form.options_text.data, form.type.data)
        db.session.flush()
        log_activity(
            "survey_question_updated",
            f"Question '{question.prompt}' updated",
            object_type="survey_question",
            object_id=question.id,
            details={
                "survey_id": survey.id,
                "type": question.type.value,
                "required": question.required,
            },
        )
        db.session.commit()
        flash("Question updated.", "success")
        return redirect(url_for("web.admin_edit_survey", survey_id=survey.id, _anchor="questions"))

    return render_template(
        "admin/surveys/question_edit.html",
        survey=survey,
        question=question,
        form=form,
        SurveyQuestionType=SurveyQuestionType,
    )


@web_bp.route("/admin/surveys/<int:survey_id>/questions/<int:question_id>/delete", methods=["POST"])
def admin_delete_survey_question(survey_id: int, question_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    question = db.session.get(SurveyQuestion, question_id)
    if not survey or not question or question.survey_id != survey.id:
        abort(404)
    if survey.status != SurveyStatus.DRAFT:
        flash("Published surveys cannot be modified.", "error")
        if _is_htmx_request():
            fresh_form = SurveyForm(obj=survey)
            fresh_form.submit.label.text = "Save changes"
            fresh_question_form = SurveyQuestionForm()
            fresh_question_form.submit.label.text = "Add question"
            return _render_survey_edit_view(
                survey=survey,
                form=fresh_form,
                question_form=fresh_question_form,
                can_edit=False,
                responses_count=(
                    db.session.query(func.count(SurveyResponse.id))
                    .filter_by(survey_id=survey.id)
                    .scalar()
                ),
                is_htmx=True,
                set_dates=True,
            )
        return redirect(url_for("web.admin_edit_survey", survey_id=survey.id, _anchor="questions"))

    db.session.delete(question)
    _normalize_question_positions(survey)
    db.session.flush()
    log_activity(
        "survey_question_deleted",
        f"Question removed from survey '{survey.title}'",
        object_type="survey_question",
        object_id=question_id,
        details={"survey_id": survey.id},
    )
    db.session.commit()
    flash("Question deleted.", "success")
    if _is_htmx_request():
        fresh_form = SurveyForm(obj=survey)
        fresh_form.submit.label.text = "Save changes"
        fresh_question_form = SurveyQuestionForm()
        fresh_question_form.submit.label.text = "Add question"
        return _render_survey_edit_view(
            survey=survey,
            form=fresh_form,
            question_form=fresh_question_form,
            can_edit=True,
            responses_count=(
                db.session.query(func.count(SurveyResponse.id))
                .filter_by(survey_id=survey.id)
                .scalar()
            ),
            is_htmx=True,
            set_dates=True,
        )
    return redirect(url_for("web.admin_edit_survey", survey_id=survey.id, _anchor="questions"))


@web_bp.route(
    "/admin/surveys/<int:survey_id>/questions/<int:question_id>/move/<string:direction>",
    methods=["POST"],
)
def admin_move_survey_question(survey_id: int, question_id: int, direction: str):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    question = db.session.get(SurveyQuestion, question_id)
    if not survey or not question or question.survey_id != survey.id:
        abort(404)
    if survey.status != SurveyStatus.DRAFT:
        flash("Published surveys cannot be modified.", "error")
        if _is_htmx_request():
            fresh_form = SurveyForm(obj=survey)
            fresh_form.submit.label.text = "Save changes"
            fresh_question_form = SurveyQuestionForm()
            fresh_question_form.submit.label.text = "Add question"
            return _render_survey_edit_view(
                survey=survey,
                form=fresh_form,
                question_form=fresh_question_form,
                can_edit=False,
                responses_count=(
                    db.session.query(func.count(SurveyResponse.id))
                    .filter_by(survey_id=survey.id)
                    .scalar()
                ),
                is_htmx=True,
                set_dates=True,
            )
        return redirect(url_for("web.admin_edit_survey", survey_id=survey.id, _anchor="questions"))
    direction = direction.lower()
    if direction not in {"up", "down"}:
        abort(400)

    ordered = _ordered_survey_questions(survey)
    current_index = next((index for index, item in enumerate(ordered) if item.id == question.id), None)
    if current_index is None:
        abort(404)
    swap_index = current_index - 1 if direction == "up" else current_index + 1
    if swap_index < 0 or swap_index >= len(ordered):
        flash("Question already at the boundary.", "info")
        if _is_htmx_request():
            fresh_form = SurveyForm(obj=survey)
            fresh_form.submit.label.text = "Save changes"
            fresh_question_form = SurveyQuestionForm()
            fresh_question_form.submit.label.text = "Add question"
            return _render_survey_edit_view(
                survey=survey,
                form=fresh_form,
                question_form=fresh_question_form,
                can_edit=True,
                responses_count=(
                    db.session.query(func.count(SurveyResponse.id))
                    .filter_by(survey_id=survey.id)
                    .scalar()
                ),
                is_htmx=True,
                set_dates=True,
            )
        return redirect(url_for("web.admin_edit_survey", survey_id=survey.id, _anchor="questions"))

    ordered[current_index], ordered[swap_index] = ordered[swap_index], ordered[current_index]
    for idx, item in enumerate(ordered, start=1):
        item.position = idx
    db.session.flush()
    log_activity(
        "survey_question_reordered",
        f"Question reorder in survey '{survey.title}'",
        object_type="survey_question",
        object_id=question.id,
        details={"survey_id": survey.id, "direction": direction},
    )
    db.session.commit()
    flash("Question order updated.", "success")
    if _is_htmx_request():
        fresh_form = SurveyForm(obj=survey)
        fresh_form.submit.label.text = "Save changes"
        fresh_question_form = SurveyQuestionForm()
        fresh_question_form.submit.label.text = "Add question"
        return _render_survey_edit_view(
            survey=survey,
            form=fresh_form,
            question_form=fresh_question_form,
            can_edit=True,
            responses_count=(
                db.session.query(func.count(SurveyResponse.id))
                .filter_by(survey_id=survey.id)
                .scalar()
            ),
            is_htmx=True,
            set_dates=True,
        )
    return redirect(url_for("web.admin_edit_survey", survey_id=survey.id, _anchor="questions"))


@web_bp.route("/admin/surveys/<int:survey_id>/publish", methods=["GET", "POST"])
def admin_publish_survey(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)
    if survey.status != SurveyStatus.DRAFT:
        flash("Only draft surveys can be published.", "error")
        return redirect(url_for("web.admin_edit_survey", survey_id=survey.id))
    if not survey.questions:
        flash("Add at least one question before publishing.", "error")
        return redirect(url_for("web.admin_edit_survey", survey_id=survey.id))

    form = SurveyPublishForm()
    teacher_choices: list[tuple[int, str]] = []
    student_choices: list[tuple[int, str]] = []

    if survey.audience in (SurveyAudience.TEACHER, SurveyAudience.MIXED):
        teacher_choices = [
            (teacher.id, f"{teacher.first_name} {teacher.last_name} — {teacher.country}")
            for teacher in Teacher.query.order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()
        ]
    if survey.audience in (SurveyAudience.STUDENT, SurveyAudience.MIXED):
        student_choices = [
            (student.id, f"{student.first_name} {student.last_name} — {student.country}")
            for student in Student.query.order_by(Student.last_name.asc(), Student.first_name.asc()).all()
        ]

    form.teacher_ids.choices = teacher_choices
    form.student_ids.choices = student_choices
    if not form.is_submitted():
        form.notify_recipients.data = True

    if form.validate_on_submit():
        selected_teacher_ids: set[int] = set(form.teacher_ids.data or [])
        selected_student_ids: set[int] = set(form.student_ids.data or [])

        if survey.audience == SurveyAudience.TEACHER and (form.send_to_all_students.data or selected_student_ids):
            form.student_ids.errors.append("This survey targets teachers only.")
            return render_template(
                "admin/surveys/publish.html",
                survey=survey,
                form=form,
                teacher_choices=teacher_choices,
                student_choices=student_choices,
            )
        if survey.audience == SurveyAudience.STUDENT and (form.send_to_all_teachers.data or selected_teacher_ids):
            form.teacher_ids.errors.append("This survey targets students only.")
            return render_template(
                "admin/surveys/publish.html",
                survey=survey,
                form=form,
                teacher_choices=teacher_choices,
                student_choices=student_choices,
            )

        if form.send_to_all_teachers.data:
            selected_teacher_ids = {choice[0] for choice in teacher_choices}
        if form.send_to_all_students.data:
            selected_student_ids = {choice[0] for choice in student_choices}

        if not selected_teacher_ids and not selected_student_ids:
            flash("No recipients available for the selected audience.", "error")
            return render_template(
                "admin/surveys/publish.html",
                survey=survey,
                form=form,
                teacher_choices=teacher_choices,
                student_choices=student_choices,
            )

        new_assignments = _create_survey_assignments(
            survey,
            teacher_ids=selected_teacher_ids,
            student_ids=selected_student_ids,
        )
        if not new_assignments:
            flash("All selected recipients already have assignments.", "warning")
            return render_template(
                "admin/surveys/publish.html",
                survey=survey,
                form=form,
                teacher_choices=teacher_choices,
                student_choices=student_choices,
            )
        survey.status = SurveyStatus.PUBLISHED
        if not survey.open_at:
            survey.open_at = datetime.utcnow()
        db.session.flush()
        log_activity(
            "survey_published",
            f"Survey '{survey.title}' published",
            object_type="survey",
            object_id=survey.id,
            details={
                "assignments": len(new_assignments),
                "audience": survey.audience.value,
            },
        )
        db.session.commit()

        notify_sent = None
        if form.notify_recipients.data:
            notify_sent = _send_survey_emails(new_assignments, SETTING_EMAIL_SURVEY_PUBLISH_SUBJECT, SETTING_EMAIL_SURVEY_PUBLISH_BODY)
            if notify_sent > 0:
                log_activity(
                    "survey_notification_sent",
                    f"Survey '{survey.title}' notifications sent",
                    object_type="survey",
                    object_id=survey.id,
                    details={"recipients": notify_sent},
                )
                db.session.commit()

        message = f"Survey published to {len(new_assignments)} recipients."
        if notify_sent is not None:
            if notify_sent == -1:
                message += " Email notifications are disabled."
            elif notify_sent == 0:
                message += " No notification emails were sent."
            else:
                message += f" Notifications sent to {notify_sent} recipients."
        flash(message, "success")
        return redirect(url_for("web.admin_surveys"))

    return render_template(
        "admin/surveys/publish.html",
        survey=survey,
        form=form,
        teacher_choices=teacher_choices,
        student_choices=student_choices,
    )


@web_bp.route("/admin/surveys/<int:survey_id>/recipients/update", methods=["GET", "POST"])
def admin_update_survey_recipients(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)

    form = SurveyRecipientUpdateForm()
    if not form.is_submitted():
        form.notify_recipients.data = True

    if survey.status != SurveyStatus.PUBLISHED:
        flash("Only published surveys can have recipients updated.", "error")
        return redirect(url_for("web.admin_surveys"))

    plan = _survey_recipient_update_plan(survey)
    back_url = url_for("web.admin_survey_responses", survey_id=survey.id)

    if form.validate_on_submit():
        if plan.missing_total == 0:
            flash("No missing recipients were found for this survey audience.", "info")
            return redirect(back_url)

        new_assignments = _create_survey_assignments(
            survey,
            teacher_ids=plan.missing_teacher_ids,
            student_ids=plan.missing_student_ids,
        )
        if not new_assignments:
            flash("No missing recipients were found for this survey audience.", "info")
            return redirect(back_url)
        db.session.flush()

        student_count = sum(1 for assignment in new_assignments if assignment.student_id)
        teacher_count = sum(1 for assignment in new_assignments if assignment.teacher_id)
        log_activity(
            "survey_recipients_updated",
            f"Survey '{survey.title}' recipients updated",
            object_type="survey",
            object_id=survey.id,
            details={
                "new_assignments": len(new_assignments),
                "students": student_count,
                "teachers": teacher_count,
                "notify_recipients": bool(form.notify_recipients.data),
            },
        )
        db.session.commit()

        notify_sent = None
        if form.notify_recipients.data:
            notify_sent = _send_survey_emails(
                new_assignments,
                SETTING_EMAIL_SURVEY_PUBLISH_SUBJECT,
                SETTING_EMAIL_SURVEY_PUBLISH_BODY,
            )
            if notify_sent > 0:
                log_activity(
                    "survey_recipient_update_notification_sent",
                    f"Survey '{survey.title}' notifications sent to newly added recipients",
                    object_type="survey",
                    object_id=survey.id,
                    details={"recipients": notify_sent},
                )
                db.session.commit()

        message = (
            f"Added {len(new_assignments)} missing recipient(s): "
            f"{student_count} student(s), {teacher_count} teacher(s)."
        )
        if notify_sent is not None:
            if notify_sent == -1:
                message += " Email notifications are disabled."
            elif notify_sent == 0:
                message += " No notification emails were sent."
            else:
                message += f" Notifications sent to {notify_sent} recipient(s)."
        flash(message, "success")
        return redirect(back_url)

    return render_template(
        "admin/surveys/update_recipients.html",
        survey=survey,
        form=form,
        plan=plan,
        back_url=back_url,
        SurveyAudience=SurveyAudience,
    )


@web_bp.route("/admin/surveys/<int:survey_id>/close", methods=["POST"])
def admin_close_survey(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)
    if survey.status != SurveyStatus.PUBLISHED:
        flash("Only published surveys can be closed.", "error")
        return redirect(url_for("web.admin_surveys"))

    survey.status = SurveyStatus.CLOSED
    if not survey.close_at:
        survey.close_at = datetime.utcnow()
    db.session.flush()
    log_activity(
        "survey_closed",
        f"Survey '{survey.title}' closed",
        object_type="survey",
        object_id=survey.id,
    )
    db.session.commit()
    flash("Survey closed.", "success")
    if _is_htmx_request():
        summary = _survey_stats_snapshot(survey)
        return _render_htmx_template(
            "admin/partials/survey_card.html",
            survey=survey,
            summary=summary,
            SurveyStatus=SurveyStatus,
            SurveyAudience=SurveyAudience,
        )
    return redirect(url_for("web.admin_surveys"))


@web_bp.route("/admin/surveys/<int:survey_id>/reopen", methods=["POST"])
def admin_reopen_survey(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)
    if survey.status != SurveyStatus.CLOSED:
        flash("Only closed surveys can be reopened.", "error")
        return redirect(url_for("web.admin_surveys"))

    survey.status = SurveyStatus.PUBLISHED
    survey.close_at = None
    if not survey.open_at:
        survey.open_at = datetime.utcnow()
    db.session.flush()
    log_activity(
        "survey_reopened",
        f"Survey '{survey.title}' reopened",
        object_type="survey",
        object_id=survey.id,
    )
    db.session.commit()
    flash("Survey reopened.", "success")
    if _is_htmx_request():
        summary = _survey_stats_snapshot(survey)
        return _render_htmx_template(
            "admin/partials/survey_card.html",
            survey=survey,
            summary=summary,
            SurveyStatus=SurveyStatus,
            SurveyAudience=SurveyAudience,
        )
    return redirect(url_for("web.admin_surveys"))


@web_bp.route("/admin/surveys/<int:survey_id>/unpublish", methods=["POST"])
def admin_unpublish_survey(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)
    if survey.status != SurveyStatus.PUBLISHED:
        flash("Only published surveys can be unpublished.", "error")
        return redirect(url_for("web.admin_surveys"))

    response_exists = (
        db.session.query(SurveyResponse.id).filter_by(survey_id=survey.id).limit(1).first() is not None
    )
    if response_exists:
        flash("This survey already has responses and cannot be unpublished.", "error")
        return redirect(url_for("web.admin_surveys"))

    for assignment in list(survey.assignments):
        db.session.delete(assignment)

    survey.status = SurveyStatus.DRAFT
    survey.open_at = None
    survey.close_at = None
    db.session.flush()
    log_activity(
        "survey_unpublished",
        f"Survey '{survey.title}' reverted to draft",
        object_type="survey",
        object_id=survey.id,
    )
    db.session.commit()
    flash("Survey moved back to draft. Assignments were removed.", "success")
    if _is_htmx_request():
        summary = _survey_stats_snapshot(survey)
        return _render_htmx_template(
            "admin/partials/survey_card.html",
            survey=survey,
            summary=summary,
            SurveyStatus=SurveyStatus,
            SurveyAudience=SurveyAudience,
        )
    return redirect(url_for("web.admin_edit_survey", survey_id=survey.id))


@web_bp.route("/admin/surveys/<int:survey_id>/delete", methods=["POST"])
def admin_delete_survey(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)
    if survey.status != SurveyStatus.DRAFT:
        flash("Only draft surveys can be deleted.", "error")
        return redirect(url_for("web.admin_surveys"))

    title = survey.title
    db.session.delete(survey)
    db.session.flush()
    log_activity(
        "survey_deleted",
        f"Survey '{title}' deleted",
        object_type="survey",
        object_id=survey_id,
    )
    db.session.commit()
    flash("Survey deleted.", "success")
    if _is_htmx_request():
        return _render_htmx_template("partials/empty.html")
    return redirect(url_for("web.admin_surveys"))


@web_bp.route("/admin/surveys/<int:survey_id>/responses")
def admin_survey_responses(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = (
        Survey.query.options(joinedload(Survey.questions)).filter(Survey.id == survey_id).first()
    )
    if not survey:
        abort(404)

    questions = _ordered_survey_questions(survey)
    rows = _survey_assignment_rows(survey)
    audience_filter = request.args.get("audience", "all").lower()
    if audience_filter not in {"all", "teacher", "student"}:
        audience_filter = "all"
    status_filter = request.args.get("status", "all").lower()
    if status_filter not in {"all", "completed", "outstanding"}:
        status_filter = "all"

    filtered_rows = _filter_survey_rows(rows, audience_filter, status_filter)
    summary = _survey_summary_from_rows(filtered_rows)
    export_url = url_for(
        "web.admin_survey_responses_export",
        survey_id=survey.id,
        audience=audience_filter,
        status=status_filter,
    )
    is_htmx = _is_htmx_request()
    template_name = "admin/partials/surveys_responses_content.html" if is_htmx else "admin/surveys/responses.html"
    context = {
        "survey": survey,
        "questions": questions,
        "rows": filtered_rows,
        "summary": summary,
        "audience_filter": audience_filter,
        "status_filter": status_filter,
        "export_url": export_url,
        "SurveyAssignmentStatus": SurveyAssignmentStatus,
        "SurveyStatus": SurveyStatus,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/surveys/<int:survey_id>/responses/<int:assignment_id>")
def admin_survey_response_detail(survey_id: int, assignment_id: int):
    if not _current_admin():
        abort(403)
    survey = db.session.get(Survey, survey_id)
    if not survey:
        abort(404)

    assignment = (
        SurveyAssignment.query.options(
            joinedload(SurveyAssignment.response),
            joinedload(SurveyAssignment.teacher),
            joinedload(SurveyAssignment.student),
            joinedload(SurveyAssignment.survey).joinedload(Survey.questions),
        )
        .filter(
            SurveyAssignment.id == assignment_id,
            SurveyAssignment.survey_id == survey_id,
        )
        .first()
    )
    if not assignment:
        abort(404)

    questions = _ordered_survey_questions(survey)
    answers = _survey_display_data(assignment.response.response_data if assignment.response else {})
    state = _survey_assignment_state(assignment)

    if assignment.teacher:
        respondent_type = "teacher"
        respondent_name = f"{assignment.teacher.first_name} {assignment.teacher.last_name}"
        respondent_email = assignment.teacher.email
        respondent_country = assignment.teacher.country
    elif assignment.student:
        respondent_type = "student"
        respondent_name = f"{assignment.student.first_name} {assignment.student.last_name}"
        respondent_email = assignment.student.email
        respondent_country = assignment.student.country
    else:
        respondent_type = "unknown"
        respondent_name = "—"
        respondent_email = None
        respondent_country = None

    audience_filter = request.args.get("audience", "all").lower()
    if audience_filter not in {"all", "teacher", "student"}:
        audience_filter = "all"
    status_filter = request.args.get("status", "all").lower()
    if status_filter not in {"all", "completed", "outstanding"}:
        status_filter = "all"
    back_url = url_for(
        "web.admin_survey_responses",
        survey_id=survey.id,
        audience=audience_filter,
        status=status_filter,
    )

    return render_template(
        "admin/surveys/response_detail.html",
        survey=survey,
        assignment=assignment,
        questions=questions,
        answers=answers,
        state=state,
        respondent_type=respondent_type,
        respondent_name=respondent_name,
        respondent_email=respondent_email,
        respondent_country=respondent_country,
        audience_filter=audience_filter,
        status_filter=status_filter,
        back_url=back_url,
        SurveyAssignmentStatus=SurveyAssignmentStatus,
    )


@web_bp.route("/admin/surveys/<int:survey_id>/responses.xlsx")
def admin_survey_responses_export(survey_id: int):
    if not _current_admin():
        abort(403)
    survey = (
        Survey.query.options(joinedload(Survey.questions)).filter(Survey.id == survey_id).first()
    )
    if not survey:
        abort(404)

    questions = _ordered_survey_questions(survey)
    audience_filter = request.args.get("audience", "all").lower()
    if audience_filter not in {"all", "teacher", "student"}:
        audience_filter = "all"
    status_filter = request.args.get("status", "all").lower()
    if status_filter not in {"all", "completed", "outstanding"}:
        status_filter = "all"

    rows = _filter_survey_rows(_survey_assignment_rows(survey), audience_filter, status_filter)
    workbook = _survey_responses_workbook(survey, rows, questions)
    filename = secure_filename(f"survey-{survey.id}-responses.xlsx") or "survey-responses.xlsx"
    return send_file(
        workbook,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@web_bp.route("/admin/surveys/<int:survey_id>/remind", methods=["POST"])
def admin_survey_send_reminder(survey_id: int):
    if not _current_admin():
        abort(403)
    audience_filter = request.form.get("audience", "all").lower()
    if audience_filter not in {"all", "teacher", "student"}:
        audience_filter = "all"
    status_filter = request.form.get("status", "outstanding").lower()
    if status_filter not in {"all", "completed", "outstanding"}:
        status_filter = "outstanding"

    survey = (
        Survey.query.options(joinedload(Survey.questions)).filter(Survey.id == survey_id).first()
    )
    if not survey:
        abort(404)

    back_url = url_for(
        "web.admin_survey_responses",
        survey_id=survey_id,
        audience=audience_filter,
        status=status_filter,
    )

    rows = _filter_survey_rows(_survey_assignment_rows(survey), audience_filter, status_filter)
    assignments = [
        row.assignment
        for row in rows
        if row.assignment.status != SurveyAssignmentStatus.COMPLETED and not row.state.locked
    ]

    if not assignments:
        flash("No outstanding recipients found for the selected filters.", "info")
        return redirect(back_url)

    sent = _send_survey_emails(assignments, SETTING_EMAIL_SURVEY_REMINDER_SUBJECT, SETTING_EMAIL_SURVEY_REMINDER_BODY)
    if sent == -1:
        flash("Automatic emails are disabled. Enable them in Settings to send reminders.", "warning")
        return redirect(back_url)
    if sent == 0:
        flash("No reminder emails were sent.", "info")
        return redirect(back_url)

    log_activity(
        "survey_reminder_sent",
        f"Survey '{survey.title}' reminders sent",
        object_type="survey",
        object_id=survey.id,
        details={
            "recipients": sent,
            "audience_filter": audience_filter,
            "status_filter": status_filter,
        },
    )
    db.session.commit()
    flash(f"Reminder emails sent to {sent} recipients.", "success")
    if _is_htmx_request():
        summary = _survey_summary_from_rows(rows)
        export_url = url_for(
            "web.admin_survey_responses_export",
            survey_id=survey.id,
            audience=audience_filter,
            status=status_filter,
        )
        return _render_htmx_template(
            "admin/partials/surveys_responses_content.html",
            survey=survey,
            questions=_ordered_survey_questions(survey),
            rows=rows,
            summary=summary,
            audience_filter=audience_filter,
            status_filter=status_filter,
            export_url=export_url,
            SurveyAssignmentStatus=SurveyAssignmentStatus,
            SurveyStatus=SurveyStatus,
        )
    return redirect(back_url)


@web_bp.route("/admin/logs")
def admin_logs():
    if not _current_admin():
        abort(403)

    is_htmx = _is_htmx_request()
    raw_filters = _log_filters_from_request()
    filters = raw_filters.copy()
    filters["page"] = max(filters["page"], 1)
    query = _filtered_activity_logs(filters).order_by(ActivityLog.created_at.desc())
    pagination = query.paginate(page=filters["page"], per_page=LOG_PAGE_SIZE, error_out=False)

    actions = [row[0] for row in db.session.query(ActivityLog.action).distinct().order_by(ActivityLog.action.asc())]
    object_types = [row[0] for row in db.session.query(ActivityLog.object_type).filter(ActivityLog.object_type.isnot(None)).distinct().order_by(ActivityLog.object_type.asc())]

    filters_display = raw_filters.copy()
    filters_display["start"] = filters["start"].isoformat(timespec="seconds") if filters.get("start") else ""
    filters_display["end"] = filters["end"].isoformat(timespec="seconds") if filters.get("end") else ""

    query_args = request.args.to_dict()
    base_args = {key: value for key, value in query_args.items() if key != "page"}
    prev_url = url_for("web.admin_logs", **base_args, page=pagination.prev_num) if pagination.has_prev else None
    next_url = url_for("web.admin_logs", **base_args, page=pagination.next_num) if pagination.has_next else None
    export_url = url_for("web.admin_logs_export", **query_args)

    template_name = "admin/partials/logs_content.html" if is_htmx else "admin/logs.html"
    context = {
        "logs": pagination.items,
        "pagination": pagination,
        "filters": filters_display,
        "actions": actions,
        "object_types": object_types,
        "prev_url": prev_url,
        "next_url": next_url,
        "export_url": export_url,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/logs/export")
def admin_logs_export():
    if not _current_admin():
        abort(403)

    filters = _log_filters_from_request()
    query = _filtered_activity_logs(filters).order_by(ActivityLog.created_at.desc())
    limit = current_app.config.get("LOG_EXPORT_LIMIT", 5000)
    logs = query.limit(limit).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["timestamp", "actor", "action", "object_type", "object_id", "description", "details"])
    for entry in logs:
        writer.writerow([
            entry.created_at.isoformat(timespec="seconds") if entry.created_at else "",
            entry.actor,
            entry.action,
            entry.object_type or "",
            entry.object_id or "",
            entry.description or "",
            json.dumps(entry.details, ensure_ascii=False, indent=2) if entry.details else "",
        ])

    response = make_response(output.getvalue())
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = "attachment; filename=activity-logs.csv"
    return response


@web_bp.route("/board/login", methods=["GET", "POST"])
@limiter.limit(TEACHER_BOARD_LOGIN_RATE_LIMIT_PER_IP, key_func=_rate_limit_ip, methods=["POST"])
def teacher_board_login():
    if _current_admin():
        return redirect(url_for("web.teacher_board"))

    portal_teacher = _teacher_portal_from_session()
    if _teacher_portal_enabled() and portal_teacher:
        return redirect(url_for("web.teacher_board"))

    form = TeacherBoardLoginForm()
    password_config = current_app.config.get("TEACHER_BOARD_PASSWORD")
    expected = (password_config or "").strip()
    if not expected:
        flash("Teacher board is disabled until TEACHER_BOARD_PASSWORD is configured.", "error")

    if form.validate_on_submit():
        if not expected:
            return render_template("teacher_board/login.html", form=form), 503
        supplied = form.password.data.strip()
        if not secrets.compare_digest(supplied, expected):
            flash("Invalid password.", "error")
            log_activity(
                "teacher_board_login_failed",
                "Teacher board login failed",
                actor="teacher_board",
                details={"remote_addr": request.remote_addr},
            )
            db.session.commit()
        else:
            session["teacher_board"] = True
            flash("Access granted.", "success")
            log_activity(
                "teacher_board_login",
                "Teacher board access granted",
                actor="teacher_board",
                details={"remote_addr": request.remote_addr},
            )
            db.session.commit()
            return redirect(url_for("web.teacher_board"))

    return render_template("teacher_board/login.html", form=form)


@web_bp.route("/board/logout")
def teacher_board_logout():
    session.pop("teacher_board", None)
    flash("Logged out of teacher board.", "success")
    log_activity(
        "teacher_board_logout",
        "Teacher board session ended",
        actor=_current_admin() or "teacher_board",
        details={"remote_addr": request.remote_addr},
    )
    db.session.commit()
    return redirect(url_for("web.teacher_board_login"))


@web_bp.route("/board")
def teacher_board():
    is_admin = bool(_current_admin())
    portal_teacher = _teacher_portal_from_session()
    if not is_admin and not _teacher_board_authorized():
        return redirect(url_for("web.teacher_board_login"))

    is_htmx = _is_htmx_request()
    priority_param = request.args.get("priority") or ""
    teacher_param = request.args.get("teacher_id", type=int)
    category_param = request.args.get("category") or ""
    search_param = request.args.get("q") or ""

    query = TeacherMessage.query.filter(TeacherMessage.status == TeacherMessageStatus.ACTIVE)

    if portal_teacher and not is_admin:
        query = query.filter(
            or_(
                TeacherMessage.deliver_to_all.is_(True),
                TeacherMessage.teacher_id == portal_teacher.id,
            )
        )
        teacher_param = 0

    if priority_param:
        try:
            query = query.filter(TeacherMessage.priority == TeacherMessagePriority(priority_param))
        except ValueError:
            flash("Unknown priority selected.", "error")

    if teacher_param:
        if teacher_param == 0:
            pass
        else:
            query = query.filter(TeacherMessage.teacher_id == teacher_param)

    if category_param:
        query = query.filter(TeacherMessage.category.ilike(f"%{category_param}%"))

    if search_param:
        query = query.filter(
            or_(
                TeacherMessage.title.ilike(f"%{search_param}%"),
                TeacherMessage.body.ilike(f"%{search_param}%"),
            )
        )

    messages = query.order_by(TeacherMessage.created_at.desc()).all()
    if portal_teacher and not is_admin:
        teachers = [portal_teacher]
    else:
        teachers = Teacher.query.order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()

    template_name = "teacher_board/partials/index_content.html" if is_htmx else "teacher_board/index.html"
    context = {
        "messages": messages,
        "teachers": teachers,
        "priorities": list(TeacherMessagePriority),
        "filters": {
            "priority": priority_param,
            "teacher_id": teacher_param or 0,
            "category": category_param,
            "q": search_param,
        },
        "is_admin": is_admin,
        "portal_teacher": portal_teacher,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@teacher_portal_bp.before_request
def _teacher_portal_guard():
    if _teacher_portal_enabled():
        return None
    flash("Teacher portal is currently disabled. Please contact the organizing team for assistance.", "error")
    return redirect(url_for("web.teacher_board_login"))


@teacher_portal_bp.route("/login", methods=["GET", "POST"])
@limiter.limit(LOGIN_RATE_LIMIT_PER_IP, key_func=_rate_limit_ip, methods=["POST"])
@limiter.limit(LOGIN_RATE_LIMIT_PER_USERNAME, key_func=_rate_limit_username, methods=["POST"])
def login():
    if _teacher_portal_from_session():
        return redirect(url_for("teacher_portal.dashboard"))

    form = TeacherPortalLoginForm()
    next_candidate = request.args.get("next")

    if form.validate_on_submit():
        username = form.username.data.strip()
        password = form.password.data.strip()
        teacher, match_details = _resolve_login_alias_match(Teacher, "portal_username", "email", username)
        if not teacher or not teacher.portal_password_hash:
            flash("Invalid credentials.", "error")
            action = "teacher_portal_login_ambiguous" if match_details["match_type"] == "ambiguous" else "teacher_portal_login_failed"
            description = (
                f"Teacher portal login ambiguous for {username}"
                if match_details["match_type"] == "ambiguous"
                else f"Teacher portal login failed for {username}"
            )
            log_activity(
                action,
                description,
                actor="teacher_portal",
                details={
                    "username": username,
                    "email": match_details.get("submitted_email"),
                    "submitted_username": username,
                    "folded_username": match_details.get("folded_username"),
                    "match_count": match_details.get("match_count", 0),
                    "remote_addr": request.remote_addr,
                },
            )
            db.session.commit()
        elif not check_password_hash(teacher.portal_password_hash, password):
            flash("Invalid credentials.", "error")
            log_activity(
                "teacher_portal_login_failed",
                f"Teacher portal login failed for {username}",
                actor="teacher_portal",
                details={
                    "username": teacher.portal_username,
                    "email": teacher.email,
                    "submitted_username": username,
                    "login_match_type": match_details["match_type"],
                    "remote_addr": request.remote_addr,
                },
            )
            db.session.commit()
        else:
            session["teacher_portal_id"] = teacher.id
            session["teacher_board"] = True
            teacher.portal_last_login = datetime.utcnow()
            log_activity(
                "teacher_portal_login",
                f"Teacher {teacher.first_name} {teacher.last_name} signed in",
                actor=teacher.portal_username or teacher.email,
                object_type="teacher",
                object_id=teacher.id,
                details={
                    "username": teacher.portal_username,
                    "email": teacher.email,
                    "submitted_username": username,
                    "login_match_type": match_details["match_type"],
                    "remote_addr": request.remote_addr,
                },
            )
            db.session.commit()
            target = _portal_next_url(request.form.get("next") or next_candidate)
            return redirect(target or url_for("teacher_portal.dashboard"))

    return render_template(
        "teacher_portal/login.html",
        form=form,
        login_email_enabled=_login_email_enabled(),
        next_value=_portal_next_url(next_candidate),
    )


@teacher_portal_bp.route("/reset", methods=["GET", "POST"])
def reset_request():
    form = PasswordResetRequestForm()
    submitted = False
    email_sent = False
    token_generated = False
    if form.validate_on_submit():
        email = form.email.data.lower().strip()
        teacher = Teacher.query.filter_by(email=email, is_absent=False).first()
        if teacher:
            reset_token = _create_teacher_reset_token(teacher)
            db.session.flush()
            log_activity(
                "teacher_password_reset_requested",
                f"Teacher {teacher.first_name} {teacher.last_name} requested portal reset",
                object_type="teacher",
                object_id=teacher.id,
                details={"email": teacher.email, "portal_username": teacher.portal_username, "reset_token_id": reset_token.id},
            )
            db.session.commit()
            token_generated = True
            if _auto_email_enabled() and teacher.email:
                _send_teacher_reset_link_email(teacher, reset_token)
                email_sent = True
        submitted = True
    return render_template(
        "teacher_portal/reset_request.html",
        form=form,
        submitted=submitted,
        email_sent=email_sent,
        auto_email_enabled=_auto_email_enabled(),
        token_generated=token_generated,
    )


@teacher_portal_bp.route("/reset/<string:token>", methods=["GET", "POST"])
def reset_token(token: str):
    form = PasswordResetConfirmForm()
    reset_token = _load_reset_token(token, PasswordResetAudience.TEACHER)
    is_valid, reason = _reset_token_status(reset_token)
    state = "ready" if is_valid else "invalid"
    plain_password: str | None = None
    email_sent = False
    auto_email_enabled = _auto_email_enabled()

    if form.validate_on_submit():
        if not is_valid or not reset_token or not reset_token.teacher:
            state = "invalid"
        else:
            _, plain_password = _ensure_teacher_portal_credentials(reset_token.teacher, reset_password=True)
            reset_token.used_at = datetime.utcnow()
            db.session.flush()
            log_activity(
                "teacher_password_reset_completed",
                f"Teacher {reset_token.teacher.first_name} {reset_token.teacher.last_name} completed portal reset",
                object_type="teacher",
                object_id=reset_token.teacher.id,
                details={
                    "email": reset_token.teacher.email,
                    "portal_username": reset_token.teacher.portal_username,
                    "reset_token_id": reset_token.id,
                },
            )
            db.session.commit()
            if plain_password and auto_email_enabled and reset_token.teacher.email:
                _send_teacher_reset_email(reset_token.teacher, plain_password)
                email_sent = True
            state = "completed"
            reason = None
            is_valid = False

    return render_template(
        "teacher_portal/reset_token.html",
        form=form,
        state=state,
        reason=reason,
        token_entry=reset_token,
        plain_password=plain_password,
        email_sent=email_sent,
        auto_email_enabled=auto_email_enabled,
    )

@teacher_portal_bp.route("/logout")
def logout():
    teacher = _teacher_portal_from_session()
    if teacher:
        log_activity(
            "teacher_portal_logout",
            f"Teacher {teacher.first_name} {teacher.last_name} signed out",
            actor=teacher.portal_username or teacher.email,
            object_type="teacher",
            object_id=teacher.id,
        )
        db.session.commit()
    session.pop("teacher_portal_id", None)
    session.pop("teacher_board", None)
    flash("Signed out of teacher portal.", "success")
    return redirect(url_for("teacher_portal.login"))


@teacher_portal_bp.route("/")
def dashboard():
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))

    is_htmx = _is_htmx_request()
    stage1_entries = _teacher_portal_assignments(teacher, Stage.STAGE_1)
    stage2_entries = _teacher_portal_assignments(teacher, Stage.STAGE_2)

    def _summary(entries: list[SimpleNamespace]) -> SimpleNamespace:
        total = len(entries)
        finalized = sum(1 for entry in entries if entry.status == "finalized")
        in_progress = sum(1 for entry in entries if entry.status == "in_progress")
        awaiting_finalization = sum(1 for entry in entries if entry.status == "awaiting_finalization")
        waiting_upload = sum(1 for entry in entries if entry.status == "waiting_upload")
        pending = sum(1 for entry in entries if entry.status == "pending")
        return SimpleNamespace(
            total=total,
            finalized=finalized,
            awaiting_finalization=awaiting_finalization,
            waiting_upload=waiting_upload,
            in_progress=in_progress,
            pending=pending,
        )

    stage1_summary = _summary(stage1_entries)
    stage2_summary = _summary(stage2_entries)
    stage1_locked = _is_stage_locked(Stage.STAGE_1)
    stage2_locked = _is_stage_locked(Stage.STAGE_2)
    teacher_survey_assignments = _survey_assignments_for_user(teacher=teacher)
    teacher_survey_items, teacher_survey_summary = _prepare_survey_assignment_view(
        teacher_survey_assignments,
        lambda assignment_id: url_for("teacher_portal.survey_detail", assignment_id=assignment_id),
    )
    teacher_survey_outstanding = [
        item for item in teacher_survey_items if not item.is_completed and not item.state.locked
    ][:3]

    messages = (
        TeacherMessage.query.filter(TeacherMessage.status == TeacherMessageStatus.ACTIVE)
        .filter(
            or_(
                TeacherMessage.deliver_to_all.is_(True),
                TeacherMessage.teacher_id == teacher.id,
            )
        )
        .order_by(
            TeacherMessage.priority.desc(),
            TeacherMessage.created_at.desc(),
        )
        .limit(10)
        .all()
    )

    template_name = "teacher_portal/partials/dashboard_content.html" if is_htmx else "teacher_portal/dashboard.html"
    context = {
        "teacher": teacher,
        "stage1_entries": stage1_entries,
        "stage2_entries": stage2_entries,
        "stage1_summary": stage1_summary,
        "stage2_summary": stage2_summary,
        "stage1_locked": stage1_locked,
        "stage2_locked": stage2_locked,
        "survey_assignments": teacher_survey_items,
        "survey_summary": teacher_survey_summary,
        "survey_outstanding": teacher_survey_outstanding,
        "messages": messages,
        "stage_labels": STAGE_LABELS,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


def _delegation_invite_context(
    teacher: Teacher,
    *,
    student_form: DelegationInviteForm | None = None,
    teacher_form: DelegationInviteForm | None = None,
) -> dict:
    invites = (
        DelegationInvite.query.filter(DelegationInvite.country == teacher.country)
        .order_by(DelegationInvite.created_at.desc())
        .all()
    )
    student_form = student_form or DelegationInviteForm(prefix="student")
    teacher_form = teacher_form or DelegationInviteForm(prefix="teacher")
    student_form.submit.label.text = "Send student invite"
    teacher_form.submit.label.text = "Send teacher invite"
    auto_email_enabled = _auto_email_enabled()

    def _summary(role: DelegationInviteRole) -> SimpleNamespace:
        total = 0
        pending = 0
        started = 0
        completed = 0
        revoked = 0
        for invite in invites:
            if invite.role != role:
                continue
            total += 1
            state = _delegation_invite_state(invite)
            if state.status == "pending":
                pending += 1
            elif state.status == "started":
                started += 1
            elif state.status == "completed":
                completed += 1
            elif state.status == "revoked":
                revoked += 1
        return SimpleNamespace(
            total=total,
            pending=pending,
            started=started,
            completed=completed,
            revoked=revoked,
        )

    invite_entries: list[SimpleNamespace] = []
    for invite in invites:
        state = _delegation_invite_state(invite)
        invite_link = url_for(
            "web.register_student" if invite.role == DelegationInviteRole.STUDENT else "web.register_teacher",
            invite=invite.token,
            _external=True,
        )
        if invite.email_sent_at:
            email_status = "Sent"
            email_detail = invite.email_sent_at.strftime("%Y-%m-%d %H:%M")
        elif invite.email_error:
            email_status = "Failed"
            email_detail = invite.email_error
        elif auto_email_enabled:
            email_status = "Not sent"
            email_detail = None
        else:
            email_status = "Auto-email off"
            email_detail = None
        invite_entries.append(
            SimpleNamespace(
                invite=invite,
                role_label="Student" if invite.role == DelegationInviteRole.STUDENT else "Teacher",
                full_name=" ".join(part for part in [invite.first_name, invite.last_name] if part).strip() or "-",
                state=state,
                invite_link=invite_link,
                email_status=email_status,
                email_detail=email_detail,
                can_resend=state.active and auto_email_enabled,
                can_revoke=state.active,
            )
        )

    return {
        "teacher": teacher,
        "delegation_country": teacher.country,
        "student_form": student_form,
        "teacher_form": teacher_form,
        "invite_entries": invite_entries,
        "student_summary": _summary(DelegationInviteRole.STUDENT),
        "teacher_summary": _summary(DelegationInviteRole.TEACHER),
        "auto_email_enabled": auto_email_enabled,
    }


def _render_delegation_invites(
    teacher: Teacher,
    *,
    student_form: DelegationInviteForm | None = None,
    teacher_form: DelegationInviteForm | None = None,
):
    is_htmx = _is_htmx_request()
    context = _delegation_invite_context(
        teacher,
        student_form=student_form,
        teacher_form=teacher_form,
    )
    context["is_htmx"] = is_htmx
    template_name = "teacher_portal/partials/delegation_content.html" if is_htmx else "teacher_portal/delegation.html"
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


def _delegation_essay_context(teacher: Teacher) -> dict:
    students = (
        Student.query
        .filter_by(country=teacher.country, is_absent=False)
        .order_by(Student.last_name.asc(), Student.first_name.asc())
        .all()
    )
    event_state = _essay_event_state()
    entries: list[SimpleNamespace] = []
    total = len(students)
    uploaded = 0
    finalized = 0
    downloadable = 0

    for student in students:
        has_upload = bool(student.uploaded_filename)
        pdf_path = _essay_pdf_path(student)
        file_missing = has_upload and pdf_path is None
        if has_upload:
            uploaded += 1
        if student.essay_finalized:
            finalized += 1
        download_allowed = bool(
            event_state.event_over
            and student.essay_finalized
            and pdf_path is not None
        )
        if download_allowed:
            downloadable += 1

        if file_missing:
            status_label = "File missing"
            status_css = "warning"
        elif student.essay_finalized:
            status_label = "Finalized"
            status_css = "success"
        elif has_upload:
            status_label = "Uploaded"
            status_css = ""
        else:
            status_label = "No upload"
            status_css = "warning"

        entries.append(
            SimpleNamespace(
                student=student,
                status_label=status_label,
                status_css=status_css,
                download_url=(
                    url_for("teacher_portal.delegation_download_essay", student_id=student.id)
                    if download_allowed
                    else None
                ),
            )
        )

    summary = SimpleNamespace(
        total=total,
        uploaded=uploaded,
        finalized=finalized,
        pending=total - uploaded,
        downloadable=downloadable,
    )

    return {
        "teacher": teacher,
        "delegation_country": teacher.country,
        "entries": entries,
        "summary": summary,
        "event_state": event_state,
    }


def _render_delegation_essays(teacher: Teacher):
    is_htmx = _is_htmx_request()
    context = _delegation_essay_context(teacher)
    context["is_htmx"] = is_htmx
    template_name = (
        "teacher_portal/partials/delegation_essays_content.html"
        if is_htmx
        else "teacher_portal/delegation_essays.html"
    )
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@teacher_portal_bp.route("/delegation")
def delegation():
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))
    if not teacher.delegation_leader:
        flash("Delegation invites are available to delegation leaders only.", "error")
        return redirect(url_for("teacher_portal.dashboard"))
    return _render_delegation_invites(teacher)


@teacher_portal_bp.route("/delegation/essays")
def delegation_essays():
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))
    if not teacher.delegation_leader:
        flash("Delegation essays are available to delegation leaders only.", "error")
        return redirect(url_for("teacher_portal.dashboard"))
    return _render_delegation_essays(teacher)


@teacher_portal_bp.route("/delegation/essays/<int:student_id>/download")
def delegation_download_essay(student_id: int):
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))
    if not teacher.delegation_leader:
        abort(403)

    student = db.session.get(Student, student_id)
    if not student or student.is_absent or student.country != teacher.country:
        abort(404)

    event_state = _essay_event_state()
    if not event_state.event_over:
        flash("Delegation essay downloads unlock after the event ends.", "error")
        return redirect(url_for("teacher_portal.delegation_essays"))

    if not student.essay_finalized:
        flash("This essay has not been finalized yet.", "error")
        return redirect(url_for("teacher_portal.delegation_essays"))

    pdf_path = _essay_pdf_path(student)
    if not pdf_path:
        flash("Essay file is not available.", "error")
        return redirect(url_for("teacher_portal.delegation_essays"))

    log_activity(
        "delegation_leader_essay_downloaded",
        f"Delegation leader {teacher.first_name} {teacher.last_name} downloaded essay",
        actor=teacher.portal_username or teacher.email,
        object_type="student",
        object_id=student.id,
        details={"student_id": student.id, "country": student.country},
    )
    db.session.commit()

    filename = secure_filename(f"essay-{student.id}.pdf")
    return send_file(
        pdf_path,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf",
    )


@teacher_portal_bp.route("/delegation/invite/student", methods=["POST"])
def delegation_invite_student():
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))
    if not teacher.delegation_leader:
        abort(403)

    student_form = DelegationInviteForm(prefix="student")
    teacher_form = DelegationInviteForm(prefix="teacher")
    if not student_form.validate_on_submit():
        flash("Please correct the errors below.", "error")
        return _render_delegation_invites(teacher, student_form=student_form, teacher_form=teacher_form)

    email = student_form.email.data.lower().strip()
    if Student.query.filter(Student.email == email).first():
        flash("This email address is already registered as a student.", "error")
        return _render_delegation_invites(teacher, student_form=student_form, teacher_form=teacher_form)

    existing_invite = (
        DelegationInvite.query.filter(
            DelegationInvite.email == email,
            DelegationInvite.role == DelegationInviteRole.STUDENT,
            DelegationInvite.country == teacher.country,
            DelegationInvite.revoked_at.is_(None),
            DelegationInvite.accepted_at.is_(None),
        )
        .order_by(DelegationInvite.created_at.desc())
        .first()
    )
    if existing_invite:
        flash("An active student invite already exists for this email.", "warning")
        return _render_delegation_invites(teacher, student_form=student_form, teacher_form=teacher_form)

    invite = DelegationInvite(
        token=_generate_unique_invite_token(),
        role=DelegationInviteRole.STUDENT,
        country=teacher.country,
        first_name=student_form.first_name.data.strip(),
        last_name=student_form.last_name.data.strip(),
        email=email,
        invited_by_id=teacher.id,
    )
    db.session.add(invite)
    db.session.flush()

    sent, error = _send_delegation_invite_email(invite, teacher)
    now = datetime.utcnow()
    if sent:
        invite.email_sent_at = now
        invite.email_error = None
    elif error and error != "auto-email-disabled":
        invite.email_error = error

    log_activity(
        "delegation_invite_created",
        f"Delegation invite created for student {invite.email}",
        actor=teacher.portal_username or teacher.email,
        object_type="delegation_invite",
        object_id=invite.id,
        details={
            "role": invite.role.value,
            "email": invite.email,
            "country": invite.country,
            "email_sent": sent,
            "email_error": error,
        },
    )
    db.session.commit()

    if sent:
        flash("Student invite created and emailed.", "success")
    elif error == "auto-email-disabled":
        flash("Student invite created. Automatic emails are disabled; share the link manually.", "warning")
    else:
        flash("Student invite created, but email delivery failed. Share the link manually.", "warning")

    if _is_htmx_request():
        return delegation()
    return redirect(url_for("teacher_portal.delegation"))


@teacher_portal_bp.route("/delegation/invite/teacher", methods=["POST"])
def delegation_invite_teacher():
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))
    if not teacher.delegation_leader:
        abort(403)

    student_form = DelegationInviteForm(prefix="student")
    teacher_form = DelegationInviteForm(prefix="teacher")
    if not teacher_form.validate_on_submit():
        flash("Please correct the errors below.", "error")
        return _render_delegation_invites(teacher, student_form=student_form, teacher_form=teacher_form)

    email = teacher_form.email.data.lower().strip()
    if Teacher.query.filter(Teacher.email == email).first():
        flash("This email address is already registered as a teacher.", "error")
        return _render_delegation_invites(teacher, student_form=student_form, teacher_form=teacher_form)

    existing_invite = (
        DelegationInvite.query.filter(
            DelegationInvite.email == email,
            DelegationInvite.role == DelegationInviteRole.TEACHER,
            DelegationInvite.country == teacher.country,
            DelegationInvite.revoked_at.is_(None),
            DelegationInvite.accepted_at.is_(None),
        )
        .order_by(DelegationInvite.created_at.desc())
        .first()
    )
    if existing_invite:
        flash("An active teacher invite already exists for this email.", "warning")
        return _render_delegation_invites(teacher, student_form=student_form, teacher_form=teacher_form)

    invite = DelegationInvite(
        token=_generate_unique_invite_token(),
        role=DelegationInviteRole.TEACHER,
        country=teacher.country,
        first_name=teacher_form.first_name.data.strip(),
        last_name=teacher_form.last_name.data.strip(),
        email=email,
        invited_by_id=teacher.id,
    )
    db.session.add(invite)
    db.session.flush()

    sent, error = _send_delegation_invite_email(invite, teacher)
    now = datetime.utcnow()
    if sent:
        invite.email_sent_at = now
        invite.email_error = None
    elif error and error != "auto-email-disabled":
        invite.email_error = error

    log_activity(
        "delegation_invite_created",
        f"Delegation invite created for teacher {invite.email}",
        actor=teacher.portal_username or teacher.email,
        object_type="delegation_invite",
        object_id=invite.id,
        details={
            "role": invite.role.value,
            "email": invite.email,
            "country": invite.country,
            "email_sent": sent,
            "email_error": error,
        },
    )
    db.session.commit()

    if sent:
        flash("Teacher invite created and emailed.", "success")
    elif error == "auto-email-disabled":
        flash("Teacher invite created. Automatic emails are disabled; share the link manually.", "warning")
    else:
        flash("Teacher invite created, but email delivery failed. Share the link manually.", "warning")

    if _is_htmx_request():
        return delegation()
    return redirect(url_for("teacher_portal.delegation"))


@teacher_portal_bp.route("/delegation/invite/<int:invite_id>/revoke", methods=["POST"])
def delegation_invite_revoke(invite_id: int):
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))
    if not teacher.delegation_leader:
        abort(403)

    invite = db.session.get(DelegationInvite, invite_id)
    if not invite or invite.country != teacher.country:
        abort(404)
    if invite.revoked_at is None and invite.accepted_at is None:
        invite.revoked_at = datetime.utcnow()
        db.session.flush()
        log_activity(
            "delegation_invite_revoked",
            f"Delegation invite revoked for {invite.email}",
            actor=teacher.portal_username or teacher.email,
            object_type="delegation_invite",
            object_id=invite.id,
            details={"role": invite.role.value, "email": invite.email, "country": invite.country},
        )
        db.session.commit()
        flash("Invite revoked.", "success")
    else:
        flash("This invite can no longer be revoked.", "warning")

    if _is_htmx_request():
        return delegation()
    return redirect(url_for("teacher_portal.delegation"))


@teacher_portal_bp.route("/delegation/invite/<int:invite_id>/resend", methods=["POST"])
def delegation_invite_resend(invite_id: int):
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))
    if not teacher.delegation_leader:
        abort(403)

    invite = db.session.get(DelegationInvite, invite_id)
    if not invite or invite.country != teacher.country:
        abort(404)
    state = _delegation_invite_state(invite)
    if not state.active:
        flash("This invite is no longer active.", "warning")
        if _is_htmx_request():
            return delegation()
        return redirect(url_for("teacher_portal.delegation"))

    sent, error = _send_delegation_invite_email(invite, teacher)
    now = datetime.utcnow()
    if sent:
        invite.email_sent_at = now
        invite.email_error = None
    elif error and error != "auto-email-disabled":
        invite.email_error = error

    log_activity(
        "delegation_invite_resent",
        f"Delegation invite resent for {invite.email}",
        actor=teacher.portal_username or teacher.email,
        object_type="delegation_invite",
        object_id=invite.id,
        details={
            "role": invite.role.value,
            "email": invite.email,
            "country": invite.country,
            "email_sent": sent,
            "email_error": error,
        },
    )
    db.session.commit()

    if sent:
        flash("Invite email resent.", "success")
    elif error == "auto-email-disabled":
        flash("Automatic emails are disabled; share the link manually.", "warning")
    else:
        flash("Email delivery failed. Share the link manually.", "warning")

    if _is_htmx_request():
        return delegation()
    return redirect(url_for("teacher_portal.delegation"))


@teacher_portal_bp.route("/surveys")
def surveys():
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))

    assignments = _survey_assignments_for_user(teacher=teacher)
    items, summary = _prepare_survey_assignment_view(
        assignments,
        lambda assignment_id: url_for("teacher_portal.survey_detail", assignment_id=assignment_id),
    )
    outstanding = [item for item in items if not item.is_completed and not item.state.locked]
    closed = [item for item in items if not item.is_completed and item.state.locked]
    completed = [item for item in items if item.is_completed]

    return render_template(
        "teacher_portal/surveys.html",
        teacher=teacher,
        assignments=items,
        summary=summary,
        outstanding=outstanding,
        closed=closed,
        completed=completed,
    )


@teacher_portal_bp.route("/surveys/<int:assignment_id>", methods=["GET", "POST"])
def survey_detail(assignment_id: int):
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))

    assignment = _load_survey_assignment(assignment_id, teacher=teacher)
    if not assignment or not assignment.survey:
        abort(404)
    survey = assignment.survey
    if survey.status == SurveyStatus.DRAFT:
        abort(404)

    questions = _ordered_survey_questions(survey)
    existing_data = _survey_display_data(assignment.response.response_data if assignment.response else {})
    state = _survey_assignment_state(assignment)
    form = SurveyResponseForm()
    errors: dict[int, str] = {}

    if form.validate_on_submit():
        if state.locked:
            flash("This survey is not available for editing.", "info")
        else:
            answers, errors = _parse_survey_submission(questions, request.form)
            if not errors:
                _save_survey_response(
                    assignment,
                    answers,
                    respondent_type="teacher",
                    respondent_id=teacher.id,
                    actor=teacher.portal_username or teacher.email,
                )
                db.session.commit()
                flash("Survey response saved.", "success")
                return redirect(url_for("teacher_portal.survey_detail", assignment_id=assignment.id))
            existing_data = _survey_display_data(answers)
            flash("Please correct the errors below.", "error")
        state = _survey_assignment_state(assignment)

    button_label = "Update response" if assignment.status == SurveyAssignmentStatus.COMPLETED else "Submit response"

    return render_template(
        "teacher_portal/survey_detail.html",
        teacher=teacher,
        assignment=assignment,
        survey=survey,
        questions=questions,
        form=form,
        existing_data=existing_data,
        errors=errors,
        state=state,
        button_label=button_label,
        SurveyQuestionType=SurveyQuestionType,
        SurveyAssignmentStatus=SurveyAssignmentStatus,
    )


@teacher_portal_bp.route("/assignments/<int:assignment_id>/<stage>", methods=["GET", "POST"])
def assignment_detail(assignment_id: int, stage: str):
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))

    try:
        stage_enum = Stage(stage)
    except ValueError:
        abort(404)

    assignment = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.student),
            joinedload(EssayAssignment.batch),
            joinedload(EssayAssignment.group).joinedload(ReadingGroup.members).joinedload(ReadingGroupMember.teacher),
            joinedload(EssayAssignment.extra_readers).joinedload(EssayAssignmentExtraReader.teacher),
            joinedload(EssayAssignment.scores),
        )
        .filter_by(id=assignment_id)
        .first_or_404()
    )

    if assignment.batch and assignment.batch.status != PlannerStatus.PUBLISHED:
        abort(404)

    student = assignment.student
    if not student or student.is_absent:
        abort(404)

    if stage_enum == Stage.STAGE_2 and not assignment.stage_two_candidate:
        abort(404)

    assessors = _stage_assessors(assignment, stage_enum, include_group=stage_enum == Stage.STAGE_1)
    if not any(entry.teacher_id == teacher.id for entry in assessors):
        abort(403)

    is_htmx = _is_htmx_request()

    def _load_state() -> SimpleNamespace:
        score_record = (
            EssayScore.query.filter_by(
                assignment_id=assignment.id,
                teacher_id=teacher.id,
                stage=stage_enum,
            )
            .order_by(EssayScore.id.desc())
            .first()
        )
        essay_path = _essay_pdf_path(student)
        pdf_available = essay_path is not None
        stage_global_lock = _is_stage_locked(stage_enum)
        assignment_lock = assignment.stage1_locked if stage_enum == Stage.STAGE_1 else assignment.stage2_locked
        score_lock = bool(score_record.locked_by_admin) if score_record else False
        combined_lock = stage_global_lock or assignment_lock or score_lock
        status_info = _resolve_portal_status(student, score_record, pdf_available, locked=combined_lock)
        essay_available = status_info.code not in {"awaiting_finalization", "waiting_upload"}
        can_edit = essay_available and not combined_lock and (not score_record or not score_record.teacher_finalized)
        download_allowed = essay_available and not (stage_global_lock or assignment_lock)
        download_url = (
            url_for("teacher_portal.download_assignment", assignment_id=assignment.id, stage=stage_enum.value)
            if download_allowed
            else None
        )
        group_label = assignment.group.name or str(assignment.group.id) if assignment.group else "—"
        essay_id = student.id
        display_label = f"Essay {essay_id}" if essay_id is not None else f"Assignment {assignment.id}"
        return SimpleNamespace(
            score_record=score_record,
            status_info=status_info,
            essay_available=essay_available,
            can_edit=can_edit,
            download_url=download_url,
            group_label=group_label,
            display_label=display_label,
            pdf_available=pdf_available,
            grading_locked=combined_lock,
            assignment_locked=assignment_lock,
            global_lock=stage_global_lock,
            score_lock=score_lock,
        )

    def _render_response(form: TeacherScoreForm, state: SimpleNamespace):
        template_name = (
            "teacher_portal/partials/assignment_detail_content.html"
            if is_htmx
            else "teacher_portal/assignment_detail.html"
        )
        context = {
            "form": form,
            "stage": stage_enum,
            "stage_label": STAGE_LABELS[stage_enum],
            "score_record": state.score_record,
            "essay_available": state.essay_available,
            "can_edit": state.can_edit,
            "download_url": state.download_url,
            "group_label": state.group_label,
            "display_label": state.display_label,
            "status_info": state.status_info,
            "student_finalized": student.essay_finalized,
            "pdf_available": state.pdf_available,
            "grading_locked": state.grading_locked,
            "assignment_locked": state.assignment_locked,
            "global_lock": state.global_lock,
            "score_lock": state.score_lock,
            "is_htmx": is_htmx,
        }
        if not is_htmx:
            return render_template(template_name, **context)
        return _render_htmx_template(template_name, **context)

    state = _load_state()
    form = TeacherScoreForm()
    if request.method == "GET":
        if state.score_record and state.score_record.score is not None:
            form.score.data = format_score_value(state.score_record.score)
        if state.score_record and state.score_record.comments:
            form.comments.data = state.score_record.comments

    if form.validate_on_submit():
        if not state.can_edit:
            if state.status_info.code == "awaiting_finalization":
                flash("This essay is waiting for the student to finalize their upload before grading can begin.", "warning")
            elif state.status_info.code == "waiting_upload":
                flash("The essay file is not yet available for grading.", "warning")
            else:
                flash("Scores for this essay have been finalized. Contact the admin team for changes.", "info")
            if is_htmx:
                return _render_response(form, state)
            return redirect(url_for("teacher_portal.assignment_detail", assignment_id=assignment.id, stage=stage_enum.value))

        finalize_requested = bool(form.finalize.data)
        numeric_score, score_error = parse_score_value(form.score.data)
        if score_error:
            form.score.errors.append(score_error)
        if finalize_requested and numeric_score is None:
            form.score.errors.append("Score is required to finalize.")
            finalize_requested = False

        if form.errors:
            return _render_response(form, state)

        if state.score_record is None:
            state.score_record = EssayScore(
                student_id=student.id,
                teacher_id=teacher.id,
                assignment_id=assignment.id,
                stage=stage_enum,
            )
            db.session.add(state.score_record)

        state.score_record.score = float(numeric_score) if numeric_score is not None else None
        state.score_record.comments = (form.comments.data or "").strip() or None
        state.score_record.submitted_at = datetime.utcnow()
        state.score_record.admin_entered = False
        state.score_record.admin_override = False
        state.score_record.admin_updated_at = None

        if finalize_requested:
            state.score_record.teacher_finalized = True
            state.score_record.teacher_finalized_at = datetime.utcnow()
            flash("Score finalized. Thank you!", "success")
            log_activity(
                "teacher_portal_score_finalized",
                f"{teacher.first_name} {teacher.last_name} finalized {stage_enum.value} score",
                actor=teacher.portal_username or teacher.email,
                object_type="essay_assignment",
                object_id=assignment.id,
                details={
                    "stage": stage_enum.value,
                    "student_id": student.id,
                    "score": state.score_record.score,
                },
            )
        else:
            state.score_record.teacher_finalized = False
            state.score_record.teacher_finalized_at = None
            flash("Draft saved.", "success")
            log_activity(
                "teacher_portal_score_saved",
                f"{teacher.first_name} {teacher.last_name} saved {stage_enum.value} score draft",
                actor=teacher.portal_username or teacher.email,
                object_type="essay_assignment",
                object_id=assignment.id,
                details={
                    "stage": stage_enum.value,
                    "student_id": student.id,
                    "score": state.score_record.score,
                },
            )

        db.session.commit()
        if is_htmx:
            updated_state = _load_state()
            fresh_form = TeacherScoreForm()
            if updated_state.score_record and updated_state.score_record.score is not None:
                fresh_form.score.data = format_score_value(updated_state.score_record.score)
            if updated_state.score_record and updated_state.score_record.comments:
                fresh_form.comments.data = updated_state.score_record.comments
            return _render_response(fresh_form, updated_state)
        return redirect(url_for("teacher_portal.assignment_detail", assignment_id=assignment.id, stage=stage_enum.value))

    return _render_response(form, state)


@teacher_portal_bp.route("/assignments/<int:assignment_id>/<stage>/download")
def download_assignment(assignment_id: int, stage: str):
    teacher = _teacher_portal_from_session()
    if not teacher:
        flash("Please sign in to continue.", "error")
        return redirect(url_for("teacher_portal.login", next=request.full_path))

    try:
        stage_enum = Stage(stage)
    except ValueError:
        abort(404)

    assignment = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.student),
            joinedload(EssayAssignment.batch),
            joinedload(EssayAssignment.group).joinedload(ReadingGroup.members),
            joinedload(EssayAssignment.extra_readers),
        )
        .filter_by(id=assignment_id)
        .first_or_404()
    )

    if assignment.batch and assignment.batch.status != PlannerStatus.PUBLISHED:
        abort(404)
    student = assignment.student
    if not student or student.is_absent:
        abort(404)
    if stage_enum == Stage.STAGE_2 and not assignment.stage_two_candidate:
        abort(404)

    assessors = _stage_assessors(assignment, stage_enum, include_group=stage_enum == Stage.STAGE_1)
    if not any(entry.teacher_id == teacher.id for entry in assessors):
        abort(403)

    stage_global_lock = _is_stage_locked(stage_enum)
    assignment_lock = assignment.stage1_locked if stage_enum == Stage.STAGE_1 else assignment.stage2_locked
    if stage_global_lock or assignment_lock:
        flash("Essay downloads are currently locked by the admin team.", "error")
        return redirect(url_for("teacher_portal.assignment_detail", assignment_id=assignment.id, stage=stage_enum.value))

    if not student.essay_finalized:
        flash("This essay is awaiting student finalization and cannot be downloaded yet.", "error")
        return redirect(url_for("teacher_portal.assignment_detail", assignment_id=assignment.id, stage=stage_enum.value))

    pdf_path = _essay_pdf_path(student)
    if not pdf_path:
        flash("Essay file is not yet available.", "error")
        return redirect(url_for("teacher_portal.assignment_detail", assignment_id=assignment.id, stage=stage_enum.value))

    log_activity(
        "teacher_portal_download_essay",
        f"{teacher.first_name} {teacher.last_name} downloaded essay",
        actor=teacher.portal_username or teacher.email,
        object_type="essay_assignment",
        object_id=assignment.id,
        details={"stage": stage_enum.value, "student_id": student.id},
    )
    db.session.commit()

    filename = secure_filename(f"essay-{student.id}-{stage_enum.value}.pdf")
    return send_file(
        pdf_path,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf",
    )


@web_bp.route("/admin/students")
def admin_students():
    students = Student.query.order_by(Student.created_at.desc()).all()
    return render_template("admin/students.html", students=students)


@web_bp.route("/admin/teachers")
def admin_teachers():
    teachers = Teacher.query.order_by(Teacher.created_at.desc()).all()
    return render_template(
        "admin/teachers.html",
        teachers=teachers,
        **_quick_message_context(),
    )


def _delegation_leader_choices() -> list[tuple[int, str]]:
    leaders = (
        Teacher.query.filter(Teacher.delegation_leader.is_(True))
        .order_by(Teacher.country.asc(), Teacher.last_name.asc(), Teacher.first_name.asc())
        .all()
    )
    choices: list[tuple[int, str]] = []
    for leader in leaders:
        full_name = " ".join(part for part in [leader.first_name, leader.last_name] if part).strip() or leader.email
        label = f"{leader.country} — {full_name} ({leader.email})"
        if leader.is_absent:
            label = f"{label} [absent]"
        choices.append((leader.id, label))
    return choices


def _active_invite_for_email(email: str, role: DelegationInviteRole) -> DelegationInvite | None:
    return (
        DelegationInvite.query.filter(
            DelegationInvite.email == email,
            DelegationInvite.role == role,
            DelegationInvite.revoked_at.is_(None),
            DelegationInvite.accepted_at.is_(None),
        )
        .order_by(DelegationInvite.created_at.desc())
        .first()
    )


def _admin_invitation_conflict(
    *,
    email: str,
    role: DelegationInviteRole,
    delegation_leader: bool = False,
) -> SimpleNamespace | None:
    if delegation_leader:
        if Teacher.query.filter(Teacher.email == email).first():
            return SimpleNamespace(
                message=(
                    "This email address is already registered as a teacher. "
                    "Use Teachers -> Edit to mark them as a delegation leader."
                ),
                category="error",
            )
    else:
        if role == DelegationInviteRole.STUDENT and Student.query.filter(Student.email == email).first():
            return SimpleNamespace(message="This email address is already registered as a student.", category="error")
        if role == DelegationInviteRole.TEACHER and Teacher.query.filter(Teacher.email == email).first():
            return SimpleNamespace(message="This email address is already registered as a teacher.", category="error")

    existing_invite = _active_invite_for_email(email, role)
    if existing_invite:
        if delegation_leader and role == DelegationInviteRole.TEACHER:
            return SimpleNamespace(
                message="An active teacher invite already exists for this email address.",
                category="warning",
            )
        return SimpleNamespace(
            message="An active invite already exists for this email address.",
            category="warning",
        )
    return None


def _admin_create_invite_and_send(
    *,
    role: DelegationInviteRole,
    delegation_leader: bool,
    country: str,
    first_name: str,
    last_name: str,
    email: str,
    inviter: Teacher | None,
    category: str,
    actor: str | None,
    log_message: str,
) -> SimpleNamespace:
    invite = DelegationInvite(
        token=_generate_unique_invite_token(),
        role=role,
        delegation_leader=delegation_leader,
        country=country,
        first_name=first_name,
        last_name=last_name,
        email=email,
        invited_by_id=inviter.id if inviter else None,
    )
    db.session.add(invite)
    db.session.flush()

    sent, error = _send_delegation_invite_email(invite, inviter)
    now = datetime.utcnow()
    if sent:
        invite.email_sent_at = now
        invite.email_error = None
    elif error and error != "auto-email-disabled":
        invite.email_error = error

    details = {
        "role": invite.role.value,
        "email": invite.email,
        "country": invite.country,
        "category": category,
        "delegation_leader": bool(delegation_leader),
        "email_sent": sent,
        "email_error": error,
    }
    if inviter:
        details["invited_by_id"] = inviter.id

    log_activity(
        "invitation_created",
        log_message,
        actor=actor,
        object_type="delegation_invite",
        object_id=invite.id,
        details=details,
    )
    return SimpleNamespace(invite=invite, sent=sent, error=error)


def _normalize_invitation_batch_header(value: str) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "_", (value or "").strip().lower())
    return normalized.strip("_")


def _normalize_invitation_batch_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def _parse_invitation_batch_role(value: str) -> DelegationInviteRole | None:
    normalized = (value or "").strip().lower()
    if normalized == DelegationInviteRole.STUDENT.value:
        return DelegationInviteRole.STUDENT
    if normalized == DelegationInviteRole.TEACHER.value:
        return DelegationInviteRole.TEACHER
    return None


def _parse_invitation_batch_csv(file_bytes: bytes) -> tuple[list[str], list[dict], list[str], str | None]:
    warnings: list[str] = []
    text = None
    encoding_used = None
    for encoding in ("utf-8-sig", "utf-8", "iso-8859-1"):
        try:
            text = file_bytes.decode(encoding)
            encoding_used = encoding
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        return [], [], warnings, "Could not decode CSV file."

    if encoding_used == "iso-8859-1":
        warnings.append("CSV encoding fallback used (ISO-8859-1). Prefer UTF-8 for predictable results.")

    reader = csv.reader(io.StringIO(text))
    matrix = list(reader)
    if not matrix:
        return [], [], warnings, "The uploaded file is empty."

    header_raw = [_normalize_invitation_batch_cell(value) for value in matrix[0]]
    rows: list[dict] = []
    for row_number, raw_values in enumerate(matrix[1:], start=2):
        values = [_normalize_invitation_batch_cell(value) for value in raw_values]
        if not any(values):
            continue
        rows.append({"row_number": row_number, "values": values})
        if len(rows) > INVITATION_BATCH_MAX_ROWS:
            return [], [], warnings, f"Too many rows. Maximum allowed is {INVITATION_BATCH_MAX_ROWS}."

    return header_raw, rows, warnings, None


def _parse_invitation_batch_xlsx(file_bytes: bytes) -> tuple[list[str], list[dict], list[str], str | None]:
    warnings: list[str] = []
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception:
        return [], [], warnings, "Could not read XLSX file."

    try:
        if not workbook.worksheets:
            return [], [], warnings, "The uploaded workbook has no worksheets."
        worksheet = workbook.worksheets[0]
        matrix = []
        for row in worksheet.iter_rows(values_only=True):
            matrix.append([_normalize_invitation_batch_cell(value) for value in row])
    finally:
        workbook.close()

    if not matrix:
        return [], [], warnings, "The uploaded file is empty."

    header_raw = matrix[0]
    rows: list[dict] = []
    for row_number, values in enumerate(matrix[1:], start=2):
        if not any(values):
            continue
        rows.append({"row_number": row_number, "values": values})
        if len(rows) > INVITATION_BATCH_MAX_ROWS:
            return [], [], warnings, f"Too many rows. Maximum allowed is {INVITATION_BATCH_MAX_ROWS}."

    return header_raw, rows, warnings, None


def _parse_invitation_batch_file(file_bytes: bytes, filename: str) -> tuple[list[str], list[dict], list[str], str | None]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _parse_invitation_batch_csv(file_bytes)
    if suffix == ".xlsx":
        return _parse_invitation_batch_xlsx(file_bytes)
    return [], [], [], "Unsupported file type. Upload a .csv or .xlsx file."


def _validate_invitation_batch_row(row_values: dict[str, str], batch_type: str) -> tuple[dict | None, str | None]:
    required_columns = INVITATION_BATCH_REQUIRED_COLUMNS[batch_type]
    normalized: dict[str, str] = {}

    for key, value in row_values.items():
        cleaned = (value or "").strip()
        if len(cleaned) > INVITATION_BATCH_MAX_CELL_LENGTH:
            return None, f"Value in '{key}' is too long (max {INVITATION_BATCH_MAX_CELL_LENGTH} characters)."
        normalized[key] = cleaned

    for column in required_columns:
        if not normalized.get(column):
            return None, f"Missing required value in column '{column}'."

    first_name = normalized.get("first_name", "")
    last_name = normalized.get("last_name", "")
    if len(first_name) > 80:
        return None, "First name is too long (max 80 characters)."
    if len(last_name) > 80:
        return None, "Last name is too long (max 80 characters)."

    email = normalized.get("email", "").lower()
    if not INVITATION_BATCH_EMAIL_PATTERN.match(email):
        return None, "Invalid email address."
    if len(email) > 255:
        return None, "Email is too long (max 255 characters)."
    normalized["email"] = email

    if batch_type in {"normal", "delegation"}:
        role = _parse_invitation_batch_role(normalized.get("role", ""))
        if role is None:
            return None, "Role must be either 'student' or 'teacher'."
        normalized["role"] = role.value

    if batch_type in {"normal", "delegation_leader"}:
        country = normalized.get("country", "")
        if len(country) > 100:
            return None, "Country is too long (max 100 characters)."
        normalized["country"] = country

    if batch_type == "delegation":
        leader_email = normalized.get("delegation_leader_email", "").lower()
        if not INVITATION_BATCH_EMAIL_PATTERN.match(leader_email):
            return None, "Invalid delegation leader email address."
        if len(leader_email) > 255:
            return None, "Delegation leader email is too long (max 255 characters)."
        normalized["delegation_leader_email"] = leader_email

    return normalized, None


def _prepare_invitation_batch_request(normalized_row: dict[str, str], batch_type: str) -> tuple[dict | None, str | None]:
    email = normalized_row["email"]
    first_name = normalized_row["first_name"]
    last_name = normalized_row["last_name"]

    if batch_type == "delegation_leader":
        role = DelegationInviteRole.TEACHER
        conflict = _admin_invitation_conflict(email=email, role=role, delegation_leader=True)
        if conflict:
            return None, conflict.message
        return {
            "role": role.value,
            "delegation_leader": True,
            "country": normalized_row["country"],
            "first_name": first_name,
            "last_name": last_name,
            "email": email,
            "inviter_id": None,
            "category": "delegation",
            "log_message": f"Invitation created for delegation leader {email}",
        }, None

    role = _parse_invitation_batch_role(normalized_row.get("role", ""))
    if role is None:
        return None, "Role must be either 'student' or 'teacher'."

    if batch_type == "delegation":
        inviter = Teacher.query.filter(
            func.lower(Teacher.email) == normalized_row["delegation_leader_email"],
            Teacher.delegation_leader.is_(True),
        ).first()
        if not inviter:
            return None, "No delegation leader found for the provided delegation_leader_email."
        conflict = _admin_invitation_conflict(email=email, role=role, delegation_leader=False)
        if conflict:
            return None, conflict.message
        return {
            "role": role.value,
            "delegation_leader": False,
            "country": inviter.country,
            "first_name": first_name,
            "last_name": last_name,
            "email": email,
            "inviter_id": inviter.id,
            "category": "delegation",
            "log_message": f"Delegation invitation created for {email}",
        }, None

    conflict = _admin_invitation_conflict(email=email, role=role, delegation_leader=False)
    if conflict:
        return None, conflict.message
    return {
        "role": role.value,
        "delegation_leader": False,
        "country": normalized_row["country"],
        "first_name": first_name,
        "last_name": last_name,
        "email": email,
        "inviter_id": None,
        "category": "normal",
        "log_message": f"Normal invitation created for {email}",
    }, None


def _build_invitation_batch_preview(
    *,
    batch_type: str,
    file_bytes: bytes,
    source_filename: str,
    created_by: str,
) -> tuple[InvitationBatchDraft | None, str | None]:
    headers, raw_rows, warnings, parse_error = _parse_invitation_batch_file(file_bytes, source_filename)
    if parse_error:
        return None, parse_error

    if not headers:
        return None, "The uploaded file does not contain a header row."

    normalized_headers: list[str] = []
    seen_headers: set[str] = set()
    for index, header in enumerate(headers, start=1):
        normalized = _normalize_invitation_batch_header(header)
        if not normalized:
            return None, f"Column header #{index} is empty."
        if normalized in seen_headers:
            return None, f"Duplicate column detected after normalization: '{normalized}'."
        normalized_headers.append(normalized)
        seen_headers.add(normalized)

    missing_columns = [column for column in INVITATION_BATCH_REQUIRED_COLUMNS[batch_type] if column not in seen_headers]
    if missing_columns:
        return None, f"Missing required columns: {', '.join(missing_columns)}."

    preview_rows: list[dict] = []
    ready_rows = 0
    skip_rows = 0
    error_rows = 0
    seen_recipients: set[tuple[str, str]] = set()

    for raw_row in raw_rows:
        row_number = int(raw_row["row_number"])
        values = list(raw_row["values"])
        if len(values) < len(normalized_headers):
            values.extend([""] * (len(normalized_headers) - len(values)))
        row_data = {header: values[idx] if idx < len(values) else "" for idx, header in enumerate(normalized_headers)}

        normalized_row, row_error = _validate_invitation_batch_row(row_data, batch_type)
        if row_error:
            error_rows += 1
            preview_rows.append(
                {
                    "row_number": row_number,
                    "status": "error",
                    "reason": row_error,
                    "input": row_data,
                    "request": None,
                }
            )
            continue

        role_value = normalized_row.get("role", DelegationInviteRole.TEACHER.value)
        dedupe_key = (normalized_row["email"], role_value)
        if dedupe_key in seen_recipients:
            skip_rows += 1
            preview_rows.append(
                {
                    "row_number": row_number,
                    "status": "skip",
                    "reason": "Duplicate email/role in this uploaded file.",
                    "input": normalized_row,
                    "request": None,
                }
            )
            continue
        seen_recipients.add(dedupe_key)

        request_payload, skip_reason = _prepare_invitation_batch_request(normalized_row, batch_type)
        if skip_reason:
            skip_rows += 1
            preview_rows.append(
                {
                    "row_number": row_number,
                    "status": "skip",
                    "reason": skip_reason,
                    "input": normalized_row,
                    "request": None,
                }
            )
            continue

        ready_rows += 1
        preview_rows.append(
            {
                "row_number": row_number,
                "status": "ready",
                "reason": "Ready to create invite.",
                "input": normalized_row,
                "request": request_payload,
            }
        )

    total_rows = len(preview_rows)
    payload = {
        "batch_type": batch_type,
        "headers": normalized_headers,
        "warnings": warnings,
        "rows": preview_rows,
        "summary": {
            "total_rows": total_rows,
            "ready_rows": ready_rows,
            "skip_rows": skip_rows,
            "error_rows": error_rows,
            "would_send_emails": ready_rows if _auto_email_enabled() else 0,
            "would_require_manual_share": 0 if _auto_email_enabled() else ready_rows,
        },
    }
    draft = InvitationBatchDraft(
        created_by=created_by,
        batch_type=batch_type,
        source_filename=source_filename,
        file_sha256=hashlib.sha256(file_bytes).hexdigest(),
        status="preview",
        total_rows=total_rows,
        ready_rows=ready_rows,
        skip_rows=skip_rows,
        error_rows=error_rows,
        result_payload=payload,
    )
    return draft, None


def _find_admin_invitation_batch_draft(draft_id: int, admin_user: str | None) -> InvitationBatchDraft | None:
    if not admin_user:
        return None
    return InvitationBatchDraft.query.filter(
        InvitationBatchDraft.id == draft_id,
        InvitationBatchDraft.created_by == admin_user,
    ).first()


def _admin_invitations_context(
    *,
    leader_form: DelegationLeaderInviteForm | None = None,
    delegation_form: AdminDelegationInvitationForm | None = None,
    normal_form: AdminInvitationForm | None = None,
    batch_upload_form: AdminInvitationBatchUploadForm | None = None,
    batch_execute_form: AdminInvitationBatchExecuteForm | None = None,
    batch_discard_form: AdminInvitationBatchDiscardForm | None = None,
    batch_draft: InvitationBatchDraft | None = None,
) -> dict:
    auto_email_enabled = _auto_email_enabled()
    leader_form = leader_form or DelegationLeaderInviteForm(prefix="leader")
    delegation_form = delegation_form or AdminDelegationInvitationForm(prefix="delegation")
    normal_form = normal_form or AdminInvitationForm(prefix="normal")
    batch_upload_form = batch_upload_form or AdminInvitationBatchUploadForm(prefix="batch-upload")
    batch_execute_form = batch_execute_form or AdminInvitationBatchExecuteForm(prefix="batch-exec")
    batch_discard_form = batch_discard_form or AdminInvitationBatchDiscardForm(prefix="batch-discard")

    leader_form.submit.label.text = "Send delegation leader invite"
    delegation_form.submit.label.text = "Send delegation invite"
    normal_form.submit.label.text = "Send invite"
    batch_upload_form.submit_preview.label.text = "Validate and preview"
    batch_execute_form.submit_execute.label.text = "Execute batch"
    batch_discard_form.submit_discard.label.text = "Discard preview"

    if batch_draft:
        batch_execute_form.draft_id.data = str(batch_draft.id)
        batch_discard_form.draft_id.data = str(batch_draft.id)
    batch_payload = batch_draft.result_payload if batch_draft else None

    delegation_form.delegation_leader_id.choices = _delegation_leader_choices()
    delegation_leaders_available = bool(delegation_form.delegation_leader_id.choices)

    invites = (
        DelegationInvite.query.options(
            joinedload(DelegationInvite.invited_by),
        )
        .order_by(DelegationInvite.created_at.desc())
        .all()
    )

    delegation_summary = SimpleNamespace(total=0, pending=0, started=0, completed=0, revoked=0)
    normal_summary = SimpleNamespace(total=0, pending=0, started=0, completed=0, revoked=0)
    delegation_entries: list[SimpleNamespace] = []
    normal_entries: list[SimpleNamespace] = []

    def bump_summary(summary: SimpleNamespace, state: SimpleNamespace) -> None:
        summary.total += 1
        if state.status == "pending":
            summary.pending += 1
        elif state.status == "started":
            summary.started += 1
        elif state.status == "completed":
            summary.completed += 1
        elif state.status == "revoked":
            summary.revoked += 1

    for invite in invites:
        state = _delegation_invite_state(invite)
        is_delegation = bool(invite.delegation_leader or invite.invited_by_id)
        category_label = "Delegation" if is_delegation else "Normal"
        if invite.role == DelegationInviteRole.TEACHER and invite.delegation_leader:
            role_label = "Delegation leader"
        else:
            role_label = "Teacher" if invite.role == DelegationInviteRole.TEACHER else "Student"

        invite_link = url_for(
            "web.register_student" if invite.role == DelegationInviteRole.STUDENT else "web.register_teacher",
            invite=invite.token,
            _external=True,
        )

        if invite.email_sent_at:
            email_status = "Sent"
            email_detail = invite.email_sent_at.strftime("%Y-%m-%d %H:%M")
        elif invite.email_error:
            email_status = "Failed"
            email_detail = invite.email_error
        elif auto_email_enabled:
            email_status = "Not sent"
            email_detail = None
        else:
            email_status = "Auto-email off"
            email_detail = None

        sender_name = "Organizing team"
        if invite.invited_by:
            sender_name = (
                f"{invite.invited_by.first_name} {invite.invited_by.last_name}".strip()
                if invite.invited_by.first_name or invite.invited_by.last_name
                else invite.invited_by.email
            )

        record_url = None
        record_label = None
        if invite.teacher_id:
            record_url = url_for("web.edit_teacher", teacher_id=invite.teacher_id)
            record_label = f"Teacher #{invite.teacher_id}"
        elif invite.student_id:
            record_url = url_for("web.edit_student", student_id=invite.student_id)
            record_label = f"Student #{invite.student_id}"

        entry = SimpleNamespace(
            invite=invite,
            category_label=category_label,
            role_label=role_label,
            full_name=" ".join(part for part in [invite.first_name, invite.last_name] if part).strip() or "-",
            state=state,
            invite_link=invite_link,
            email_status=email_status,
            email_detail=email_detail,
            sender_name=sender_name,
            record_url=record_url,
            record_label=record_label,
            can_resend=state.active and auto_email_enabled,
            can_revoke=state.active,
        )

        if is_delegation:
            delegation_entries.append(entry)
            bump_summary(delegation_summary, state)
        else:
            normal_entries.append(entry)
            bump_summary(normal_summary, state)

    return {
        "delegation_leader_form": leader_form,
        "delegation_form": delegation_form,
        "normal_form": normal_form,
        "batch_upload_form": batch_upload_form,
        "batch_execute_form": batch_execute_form,
        "batch_discard_form": batch_discard_form,
        "batch_draft": batch_draft,
        "batch_payload": batch_payload,
        "batch_type_labels": {value: label for value, label in INVITATION_BATCH_TYPE_CHOICES},
        "delegation_entries": delegation_entries,
        "normal_entries": normal_entries,
        "delegation_summary": delegation_summary,
        "normal_summary": normal_summary,
        "auto_email_enabled": auto_email_enabled,
        "delegation_leaders_available": delegation_leaders_available,
    }


def _render_admin_invitations(
    *,
    leader_form: DelegationLeaderInviteForm | None = None,
    delegation_form: AdminDelegationInvitationForm | None = None,
    normal_form: AdminInvitationForm | None = None,
    batch_upload_form: AdminInvitationBatchUploadForm | None = None,
    batch_execute_form: AdminInvitationBatchExecuteForm | None = None,
    batch_discard_form: AdminInvitationBatchDiscardForm | None = None,
    batch_draft: InvitationBatchDraft | None = None,
):
    is_htmx = _is_htmx_request()
    context = _admin_invitations_context(
        leader_form=leader_form,
        delegation_form=delegation_form,
        normal_form=normal_form,
        batch_upload_form=batch_upload_form,
        batch_execute_form=batch_execute_form,
        batch_discard_form=batch_discard_form,
        batch_draft=batch_draft,
    )
    context["is_htmx"] = is_htmx
    template_name = "admin/partials/invitations_content.html" if is_htmx else "admin/invitations.html"
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/invitations")
def admin_invitations():
    if not _current_admin():
        abort(403)
    draft = None
    draft_id_raw = (request.args.get("batch_draft") or "").strip()
    if draft_id_raw.isdigit():
        draft = _find_admin_invitation_batch_draft(int(draft_id_raw), _current_admin())
        if not draft:
            flash("Batch preview not found.", "warning")
    return _render_admin_invitations(batch_draft=draft)

@web_bp.route("/admin/invitations/delegation-leader", methods=["POST"])
def admin_invitation_create_delegation_leader():
    if not _current_admin():
        abort(403)

    leader_form = DelegationLeaderInviteForm(prefix="leader")
    delegation_form = AdminDelegationInvitationForm(prefix="delegation")
    normal_form = AdminInvitationForm(prefix="normal")
    delegation_form.delegation_leader_id.choices = _delegation_leader_choices()

    if not leader_form.validate_on_submit():
        flash("Please correct the errors below.", "error")
        return _render_admin_invitations(
            leader_form=leader_form,
            delegation_form=delegation_form,
            normal_form=normal_form,
        )

    email = leader_form.email.data.lower().strip()
    country = leader_form.country.data.strip()
    conflict = _admin_invitation_conflict(
        email=email,
        role=DelegationInviteRole.TEACHER,
        delegation_leader=True,
    )
    if conflict:
        flash(conflict.message, conflict.category)
        return _render_admin_invitations(
            leader_form=leader_form,
            delegation_form=delegation_form,
            normal_form=normal_form,
        )

    result = _admin_create_invite_and_send(
        role=DelegationInviteRole.TEACHER,
        delegation_leader=True,
        country=country,
        first_name=leader_form.first_name.data.strip(),
        last_name=leader_form.last_name.data.strip(),
        email=email,
        inviter=None,
        category="delegation",
        actor=_current_admin(),
        log_message=f"Invitation created for delegation leader {email}",
    )
    sent, error = result.sent, result.error
    db.session.commit()

    if sent:
        flash("Delegation leader invite created and emailed.", "success")
    elif error == "auto-email-disabled":
        flash("Delegation leader invite created. Automatic emails are disabled; share the link manually.", "warning")
    else:
        flash("Delegation leader invite created, but email delivery failed. Share the link manually.", "warning")

    if _is_htmx_request():
        return _render_admin_invitations()
    return redirect(url_for("web.admin_invitations"))


@web_bp.route("/admin/invitations/delegation", methods=["POST"])
def admin_invitation_create_delegation():
    if not _current_admin():
        abort(403)

    leader_form = DelegationLeaderInviteForm(prefix="leader")
    delegation_form = AdminDelegationInvitationForm(prefix="delegation")
    normal_form = AdminInvitationForm(prefix="normal")
    delegation_form.delegation_leader_id.choices = _delegation_leader_choices()

    if not delegation_form.delegation_leader_id.choices:
        flash("No delegation leaders found. Create a delegation leader invite first.", "error")
        return _render_admin_invitations(
            leader_form=leader_form,
            delegation_form=delegation_form,
            normal_form=normal_form,
        )

    if not delegation_form.validate_on_submit():
        flash("Please correct the errors below.", "error")
        return _render_admin_invitations(
            leader_form=leader_form,
            delegation_form=delegation_form,
            normal_form=normal_form,
        )

    inviter = db.session.get(Teacher, int(delegation_form.delegation_leader_id.data))
    if not inviter or not inviter.delegation_leader:
        flash("Selected delegation leader is invalid.", "error")
        return _render_admin_invitations(
            leader_form=leader_form,
            delegation_form=delegation_form,
            normal_form=normal_form,
        )

    email = delegation_form.email.data.lower().strip()
    role = (
        DelegationInviteRole.STUDENT
        if delegation_form.role.data == DelegationInviteRole.STUDENT.value
        else DelegationInviteRole.TEACHER
    )

    conflict = _admin_invitation_conflict(email=email, role=role, delegation_leader=False)
    if conflict:
        flash(conflict.message, conflict.category)
        return _render_admin_invitations(leader_form=leader_form, delegation_form=delegation_form, normal_form=normal_form)

    result = _admin_create_invite_and_send(
        role=role,
        delegation_leader=False,
        country=inviter.country,
        first_name=delegation_form.first_name.data.strip(),
        last_name=delegation_form.last_name.data.strip(),
        email=email,
        inviter=inviter,
        category="delegation",
        actor=_current_admin(),
        log_message=f"Delegation invitation created for {email}",
    )
    sent, error = result.sent, result.error
    db.session.commit()

    if sent:
        flash("Delegation invite created and emailed.", "success")
    elif error == "auto-email-disabled":
        flash("Delegation invite created. Automatic emails are disabled; share the link manually.", "warning")
    else:
        flash("Delegation invite created, but email delivery failed. Share the link manually.", "warning")

    if _is_htmx_request():
        return _render_admin_invitations()
    return redirect(url_for("web.admin_invitations"))


@web_bp.route("/admin/invitations/normal", methods=["POST"])
def admin_invitation_create_normal():
    if not _current_admin():
        abort(403)

    leader_form = DelegationLeaderInviteForm(prefix="leader")
    delegation_form = AdminDelegationInvitationForm(prefix="delegation")
    normal_form = AdminInvitationForm(prefix="normal")
    delegation_form.delegation_leader_id.choices = _delegation_leader_choices()

    if not normal_form.validate_on_submit():
        flash("Please correct the errors below.", "error")
        return _render_admin_invitations(
            leader_form=leader_form,
            delegation_form=delegation_form,
            normal_form=normal_form,
        )

    email = normal_form.email.data.lower().strip()
    country = normal_form.country.data.strip()
    role = (
        DelegationInviteRole.STUDENT
        if normal_form.role.data == DelegationInviteRole.STUDENT.value
        else DelegationInviteRole.TEACHER
    )

    conflict = _admin_invitation_conflict(email=email, role=role, delegation_leader=False)
    if conflict:
        flash(conflict.message, conflict.category)
        return _render_admin_invitations(leader_form=leader_form, delegation_form=delegation_form, normal_form=normal_form)

    result = _admin_create_invite_and_send(
        role=role,
        delegation_leader=False,
        country=country,
        first_name=normal_form.first_name.data.strip(),
        last_name=normal_form.last_name.data.strip(),
        email=email,
        inviter=None,
        category="normal",
        actor=_current_admin(),
        log_message=f"Normal invitation created for {email}",
    )
    sent, error = result.sent, result.error
    db.session.commit()

    if sent:
        flash("Invite created and emailed.", "success")
    elif error == "auto-email-disabled":
        flash("Invite created. Automatic emails are disabled; share the link manually.", "warning")
    else:
        flash("Invite created, but email delivery failed. Share the link manually.", "warning")

    if _is_htmx_request():
        return _render_admin_invitations()
    return redirect(url_for("web.admin_invitations"))


def _init_admin_invitation_forms() -> SimpleNamespace:
    leader_form = DelegationLeaderInviteForm(prefix="leader")
    delegation_form = AdminDelegationInvitationForm(prefix="delegation")
    normal_form = AdminInvitationForm(prefix="normal")
    batch_upload_form = AdminInvitationBatchUploadForm(prefix="batch-upload")
    batch_execute_form = AdminInvitationBatchExecuteForm(prefix="batch-exec")
    batch_discard_form = AdminInvitationBatchDiscardForm(prefix="batch-discard")
    delegation_form.delegation_leader_id.choices = _delegation_leader_choices()
    return SimpleNamespace(
        leader_form=leader_form,
        delegation_form=delegation_form,
        normal_form=normal_form,
        batch_upload_form=batch_upload_form,
        batch_execute_form=batch_execute_form,
        batch_discard_form=batch_discard_form,
    )


@web_bp.route("/admin/invitations/batch/template/<string:batch_type>.<string:fmt>")
def admin_invitation_batch_template(batch_type: str, fmt: str):
    if not _current_admin():
        abort(403)
    if batch_type not in INVITATION_BATCH_TYPE_VALUES:
        abort(404)
    fmt = (fmt or "").strip().lower()
    if fmt not in {"csv", "xlsx"}:
        abort(404)

    headers = INVITATION_BATCH_REQUIRED_COLUMNS[batch_type]
    sample_row = INVITATION_BATCH_TEMPLATE_SAMPLE_ROWS.get(batch_type, {})

    if fmt == "csv":
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerow({header: sample_row.get(header, "") for header in headers})
        csv_bytes = output.getvalue().encode("utf-8")
        return send_file(
            BytesIO(csv_bytes),
            mimetype="text/csv; charset=utf-8",
            as_attachment=True,
            download_name=f"invitation-batch-{batch_type}-template.csv",
        )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Invites"
    worksheet.append(headers)
    worksheet.append([sample_row.get(header, "") for header in headers])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        download_name=f"invitation-batch-{batch_type}-template.xlsx",
    )


@web_bp.route("/admin/invitations/batch/preview", methods=["POST"])
def admin_invitation_batch_preview():
    if not _current_admin():
        abort(403)
    forms = _init_admin_invitation_forms()

    if not forms.batch_upload_form.validate_on_submit():
        flash("Please correct the batch upload form errors.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
        )

    batch_type = forms.batch_upload_form.batch_type.data
    if batch_type not in INVITATION_BATCH_TYPE_VALUES:
        flash("Unknown batch type selected.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
        )

    uploaded = forms.batch_upload_form.invite_file.data
    filename = secure_filename((uploaded.filename or "").strip())
    if not filename:
        flash("Uploaded file name is invalid.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
        )

    file_bytes = uploaded.read()
    if not file_bytes:
        flash("Uploaded file is empty.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
        )
    if len(file_bytes) > INVITATION_BATCH_MAX_FILE_SIZE:
        flash(f"File is too large. Maximum allowed size is {INVITATION_BATCH_MAX_FILE_SIZE // (1024 * 1024)} MB.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
        )

    draft, error = _build_invitation_batch_preview(
        batch_type=batch_type,
        file_bytes=file_bytes,
        source_filename=filename,
        created_by=_current_admin() or "",
    )
    if error or not draft:
        flash(error or "Could not build batch preview.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
        )

    db.session.add(draft)
    db.session.flush()
    log_activity(
        "invitation_batch_previewed",
        f"Invitation batch previewed ({batch_type})",
        actor=_current_admin(),
        object_type="invitation_batch_draft",
        object_id=draft.id,
        details={
            "batch_type": batch_type,
            "source_filename": filename,
            "total_rows": draft.total_rows,
            "ready_rows": draft.ready_rows,
            "skip_rows": draft.skip_rows,
            "error_rows": draft.error_rows,
        },
    )
    db.session.commit()
    flash(
        f"Preview created: {draft.ready_rows} ready, {draft.skip_rows} skipped, {draft.error_rows} errors.",
        "success",
    )

    payload = draft.result_payload or {}
    for warning in payload.get("warnings", []):
        flash(warning, "warning")

    forms.batch_execute_form.draft_id.data = str(draft.id)
    forms.batch_discard_form.draft_id.data = str(draft.id)
    if _is_htmx_request():
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
            batch_draft=draft,
        )
    return redirect(url_for("web.admin_invitations", batch_draft=draft.id))


@web_bp.route("/admin/invitations/batch/<int:draft_id>/execute", methods=["POST"])
def admin_invitation_batch_execute(draft_id: int):
    if not _current_admin():
        abort(403)
    forms = _init_admin_invitation_forms()
    execute_form = forms.batch_execute_form

    draft = _find_admin_invitation_batch_draft(draft_id, _current_admin())
    if not draft:
        abort(404)

    if not execute_form.validate_on_submit():
        flash("Confirm batch execution to continue.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=execute_form,
            batch_discard_form=forms.batch_discard_form,
            batch_draft=draft,
        )

    if str(execute_form.draft_id.data or "") != str(draft_id):
        flash("Batch draft mismatch.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=execute_form,
            batch_discard_form=forms.batch_discard_form,
            batch_draft=draft,
        )

    if draft.status != "preview":
        flash("This batch preview has already been processed.", "warning")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=execute_form,
            batch_discard_form=forms.batch_discard_form,
            batch_draft=draft,
        )

    payload = dict(draft.result_payload or {})
    rows = list(payload.get("rows", []))
    ready_rows = [row for row in rows if row.get("status") == "ready" and row.get("request")]
    if not ready_rows:
        flash("No rows are ready for execution.", "warning")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=execute_form,
            batch_discard_form=forms.batch_discard_form,
            batch_draft=draft,
        )

    draft.status = "executing"
    db.session.flush()

    created_count = 0
    email_sent_count = 0
    email_failed_count = 0
    execute_skip_count = 0
    execute_error_count = 0
    updated_rows: list[dict] = []

    for row in rows:
        row_out = dict(row)
        request_payload = row_out.get("request") or {}
        if row_out.get("status") != "ready" or not request_payload:
            updated_rows.append(row_out)
            continue

        role = _parse_invitation_batch_role(str(request_payload.get("role") or ""))
        if role is None:
            row_out["status"] = "error"
            row_out["reason"] = "Invalid role in prepared request."
            execute_error_count += 1
            updated_rows.append(row_out)
            continue

        delegation_leader = bool(request_payload.get("delegation_leader"))
        inviter = None
        inviter_id = request_payload.get("inviter_id")
        if inviter_id is not None:
            inviter = db.session.get(Teacher, int(inviter_id))
            if not inviter or not inviter.delegation_leader:
                row_out["status"] = "skip"
                row_out["reason"] = "Delegation leader no longer exists or is no longer marked as leader."
                execute_skip_count += 1
                updated_rows.append(row_out)
                continue

        conflict = _admin_invitation_conflict(
            email=str(request_payload.get("email") or ""),
            role=role,
            delegation_leader=delegation_leader,
        )
        if conflict:
            row_out["status"] = "skip"
            row_out["reason"] = conflict.message
            execute_skip_count += 1
            updated_rows.append(row_out)
            continue

        country_value = str(request_payload.get("country") or "").strip()
        if inviter:
            country_value = inviter.country

        try:
            with db.session.begin_nested():
                result = _admin_create_invite_and_send(
                    role=role,
                    delegation_leader=delegation_leader,
                    country=country_value,
                    first_name=str(request_payload.get("first_name") or "").strip(),
                    last_name=str(request_payload.get("last_name") or "").strip(),
                    email=str(request_payload.get("email") or "").strip(),
                    inviter=inviter,
                    category=str(request_payload.get("category") or "normal"),
                    actor=_current_admin(),
                    log_message=str(
                        request_payload.get("log_message")
                        or f"Batch invitation created for {str(request_payload.get('email') or '').strip()}"
                    ),
                )
            created_count += 1
            row_out["status"] = "created"
            row_out["reason"] = "Invite created."
            row_out["invite_id"] = result.invite.id
            if result.sent:
                email_sent_count += 1
            elif result.error and result.error != "auto-email-disabled":
                email_failed_count += 1
                row_out["email_error"] = result.error
        except IntegrityError:
            row_out["status"] = "error"
            row_out["reason"] = "Database conflict while creating invite."
            execute_error_count += 1
        updated_rows.append(row_out)

    payload["rows"] = updated_rows
    payload["execution"] = {
        "executed_at": datetime.utcnow().isoformat(),
        "created_rows": created_count,
        "email_sent": email_sent_count,
        "email_failed": email_failed_count,
        "skip_rows": execute_skip_count,
        "error_rows": execute_error_count,
    }

    draft.status = "executed"
    draft.executed_at = datetime.utcnow()
    draft.result_payload = payload
    draft.total_rows = len(updated_rows)
    draft.ready_rows = sum(1 for row in updated_rows if row.get("status") == "ready")
    draft.skip_rows = sum(1 for row in updated_rows if row.get("status") == "skip")
    draft.error_rows = sum(1 for row in updated_rows if row.get("status") == "error")

    log_activity(
        "invitation_batch_executed",
        f"Invitation batch executed ({draft.batch_type})",
        actor=_current_admin(),
        object_type="invitation_batch_draft",
        object_id=draft.id,
        details={
            "batch_type": draft.batch_type,
            "source_filename": draft.source_filename,
            "created_rows": created_count,
            "email_sent": email_sent_count,
            "email_failed": email_failed_count,
            "skip_rows": execute_skip_count,
            "error_rows": execute_error_count,
        },
    )
    db.session.commit()
    flash(
        f"Batch executed. Created {created_count} invites, sent {email_sent_count} emails, failed {email_failed_count} emails.",
        "success",
    )

    forms.batch_execute_form.draft_id.data = str(draft.id)
    forms.batch_discard_form.draft_id.data = str(draft.id)
    if _is_htmx_request():
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
            batch_draft=draft,
        )
    return redirect(url_for("web.admin_invitations", batch_draft=draft.id))


@web_bp.route("/admin/invitations/batch/<int:draft_id>/discard", methods=["POST"])
def admin_invitation_batch_discard(draft_id: int):
    if not _current_admin():
        abort(403)
    forms = _init_admin_invitation_forms()
    discard_form = forms.batch_discard_form

    draft = _find_admin_invitation_batch_draft(draft_id, _current_admin())
    if not draft:
        abort(404)

    if not discard_form.validate_on_submit():
        flash("Could not discard batch preview.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=discard_form,
            batch_draft=draft,
        )
    if str(discard_form.draft_id.data or "") != str(draft_id):
        flash("Batch draft mismatch.", "error")
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=discard_form,
            batch_draft=draft,
        )

    draft.status = "discarded"
    log_activity(
        "invitation_batch_discarded",
        f"Invitation batch discarded ({draft.batch_type})",
        actor=_current_admin(),
        object_type="invitation_batch_draft",
        object_id=draft.id,
        details={"batch_type": draft.batch_type, "source_filename": draft.source_filename},
    )
    db.session.commit()
    flash("Batch preview discarded.", "success")

    if _is_htmx_request():
        return _render_admin_invitations(
            leader_form=forms.leader_form,
            delegation_form=forms.delegation_form,
            normal_form=forms.normal_form,
            batch_upload_form=forms.batch_upload_form,
            batch_execute_form=forms.batch_execute_form,
            batch_discard_form=forms.batch_discard_form,
            batch_draft=None,
        )
    return redirect(url_for("web.admin_invitations"))


@web_bp.route("/admin/invitations/invite/<int:invite_id>/revoke", methods=["POST"])
def admin_invitation_revoke(invite_id: int):
    if not _current_admin():
        abort(403)

    invite = db.session.get(DelegationInvite, invite_id)
    if not invite:
        abort(404)

    if invite.revoked_at is None and invite.accepted_at is None:
        invite.revoked_at = datetime.utcnow()
        db.session.flush()
        log_activity(
            "invitation_revoked",
            f"Invitation revoked for {invite.email}",
            actor=_current_admin(),
            object_type="delegation_invite",
            object_id=invite.id,
            details={
                "role": invite.role.value,
                "email": invite.email,
                "country": invite.country,
                "delegation_leader": bool(invite.delegation_leader),
                "invited_by_id": invite.invited_by_id,
            },
        )
        db.session.commit()
        flash("Invite revoked.", "success")
    else:
        flash("This invite can no longer be revoked.", "warning")

    if _is_htmx_request():
        return _render_admin_invitations()
    return redirect(url_for("web.admin_invitations"))


@web_bp.route("/admin/invitations/invite/<int:invite_id>/resend", methods=["POST"])
def admin_invitation_resend(invite_id: int):
    if not _current_admin():
        abort(403)

    invite = db.session.get(DelegationInvite, invite_id)
    if not invite:
        abort(404)

    state = _delegation_invite_state(invite)
    if not state.active:
        flash("This invite is no longer active.", "warning")
        if _is_htmx_request():
            return _render_admin_invitations()
        return redirect(url_for("web.admin_invitations"))

    inviter = invite.invited_by if invite.invited_by_id else None
    sent, error = _send_delegation_invite_email(invite, inviter)
    now = datetime.utcnow()
    if sent:
        invite.email_sent_at = now
        invite.email_error = None
    elif error and error != "auto-email-disabled":
        invite.email_error = error

    log_activity(
        "invitation_resent",
        f"Invitation resent for {invite.email}",
        actor=_current_admin(),
        object_type="delegation_invite",
        object_id=invite.id,
        details={
            "role": invite.role.value,
            "email": invite.email,
            "country": invite.country,
            "delegation_leader": bool(invite.delegation_leader),
            "invited_by_id": invite.invited_by_id,
            "email_sent": sent,
            "email_error": error,
        },
    )
    db.session.commit()

    if sent:
        flash("Invite email resent.", "success")
    elif error == "auto-email-disabled":
        flash("Automatic emails are disabled; share the link manually.", "warning")
    else:
        flash("Email delivery failed. Share the link manually.", "warning")

    if _is_htmx_request():
        return _render_admin_invitations()
    return redirect(url_for("web.admin_invitations"))


@web_bp.route("/admin/students/<int:student_id>/edit", methods=["GET", "POST"])
def edit_student(student_id: int):
    student = db.session.get(Student, student_id)
    if not student:
        abort(404)

    form = StudentRegistrationForm(obj=student)
    form.accepted_ethics.validators = []
    form.gdpr_acknowledged.validators = []
    form.accommodation_room_sharing.validators = []
    form.submit.label.text = "Save changes"
    if not form.is_submitted():
        form.gdpr_acknowledged.data = bool(student.gdpr_acknowledged)
        form.photo_consent.data = "yes" if student.photo_consent else "no"

    if form.validate_on_submit():
        student.first_name = form.first_name.data.strip()
        student.last_name = form.last_name.data.strip()
        student.email = form.email.data.lower().strip()
        student.country = form.country.data.strip()
        student.school = form.school.data.strip()
        student.sex_document = form.sex_document.data
        student.accommodation_room_sharing = form.accommodation_room_sharing.data
        student.date_of_birth = form.date_of_birth.data
        student.essay_language = form.essay_language.data
        student.keyboard_preference = form.keyboard_preference.data
        student.meal_preference = form.meal_preference.data
        student.special_needs = (form.special_needs.data.strip() or None) if form.special_needs.data else None
        if form.gdpr_acknowledged.data:
            student.gdpr_acknowledged = True
            student.gdpr_acknowledged_at = student.gdpr_acknowledged_at or datetime.utcnow()
        else:
            student.gdpr_acknowledged = False
            student.gdpr_acknowledged_at = None
        new_photo_consent = form.photo_consent.data == "yes"
        if student.photo_consent != new_photo_consent:
            student.photo_consent = new_photo_consent
            student.photo_consent_at = datetime.utcnow()
        if form.accepted_ethics.data:
            student.accepted_ethics = True
            student.accepted_ethics_at = student.accepted_ethics_at or datetime.utcnow()
        else:
            student.accepted_ethics = False
            student.accepted_ethics_at = None
        student.is_absent = bool(form.is_absent.data)
        if not student.upload_username or not student.upload_password_hash:
            _ensure_student_credentials(student)

        try:
            db.session.flush()
            log_activity(
        "student_updated",
        f"Student {student.first_name} {student.last_name} updated",
        object_type="student",
        object_id=student.id,
        details={
            "email": student.email,
            "country": student.country,
            "language": student.essay_language,
            "upload_username": student.upload_username,
        },
    )
            db.session.commit()
            flash("Student details updated successfully.", "success")
            return redirect(url_for("web.admin_students"))
        except IntegrityError:
            db.session.rollback()
            flash("This email address is already registered with another student.", "error")

    return render_template(
        "register_student.html",
        form=form,
        form_title="Edit student",
        show_absence_toggle=True,
        ethics_consent_html=_ethics_consent_html(),
        gdpr_notice_html=_gdpr_notice_html(),
        photo_consent_html=_photo_consent_html(),
        accommodation_info_html=_accommodation_info_html(),
    )


@web_bp.route("/admin/students/<int:student_id>/reset-password", methods=["POST"])
def admin_student_reset_password(student_id: int):
    if not _current_admin():
        abort(403)
    student = db.session.get(Student, student_id)
    if not student:
        abort(404)

    _, plain_password = _ensure_student_credentials(student, reset_password=True)
    db.session.flush()
    log_activity(
        "student_password_reset",
        f"Student {student.first_name} {student.last_name} portal password reset",
        object_type="student",
        object_id=student.id,
        details={"upload_username": student.upload_username},
    )
    db.session.commit()

    if plain_password:
        if _auto_email_enabled() and student.email:
            _send_student_reset_email(student, plain_password)
            flash(
                f"Portal password reset and emailed to {student.email}. New password: {plain_password}",
                "success",
            )
        else:
            flash(
                f"Portal password reset. Automatic email is disabled; share this password manually: {plain_password}",
                "warning",
            )
    else:
        flash("Portal password reset, but no password was generated.", "error")

    return redirect(url_for("web.admin_students"))


@web_bp.route("/admin/students/<int:student_id>/delete", methods=["POST"])
def admin_student_delete(student_id: int):
    if not _current_admin():
        abort(403)
    student = db.session.get(Student, student_id)
    if not student:
        abort(404)

    expected_email = (student.email or "").strip().lower()
    confirm_email = (request.form.get("confirm_email") or "").strip().lower()
    if confirm_email != expected_email:
        flash("Confirmation email does not match.", "error")
        return redirect(url_for("web.admin_students"))

    student_name = " ".join(part for part in [student.first_name, student.last_name] if part).strip()
    student_email = student.email
    student_country = student.country

    db.session.delete(student)
    try:
        db.session.flush()
    except IntegrityError:
        db.session.rollback()
        flash("Student could not be deleted because related records exist.", "error")
        return redirect(url_for("web.admin_students"))

    log_activity(
        "student_deleted",
        f"Student {student_name or student_email} deleted",
        object_type="student",
        object_id=student_id,
        details={
            "email": student_email,
            "country": student_country,
        },
    )
    db.session.commit()
    flash("Student deleted.", "success")
    return redirect(url_for("web.admin_students"))


@web_bp.route("/admin/teachers/<int:teacher_id>/edit", methods=["GET", "POST"])
def edit_teacher(teacher_id: int):
    teacher = db.session.get(Teacher, teacher_id)
    if not teacher:
        abort(404)

    form = TeacherRegistrationForm(obj=teacher)
    form.accepted_ethics.validators = []
    form.gdpr_acknowledged.validators = []
    form.declared_assessor_competence.validators = []
    form.accommodation_room_sharing.validators = []
    if not form.is_submitted():
        languages = [lang.strip() for lang in teacher.languages.split(",")] if teacher.languages else ["English"]
        form.languages.data = [lang for lang in languages if lang]
        form.prior_ipo_assessor_level.data = teacher.prior_ipo_assessor_level or 0
        form.gdpr_acknowledged.data = bool(teacher.gdpr_acknowledged)
        form.photo_consent.data = "yes" if teacher.photo_consent else "no"
        form.delegation_leader.data = bool(teacher.delegation_leader)
        form.non_assessor.data = not bool(teacher.declared_assessor_competence)
    form.submit.label.text = "Save changes"

    if form.validate_on_submit():
        teacher.first_name = form.first_name.data.strip()
        teacher.last_name = form.last_name.data.strip()
        teacher.email = form.email.data.lower().strip()
        teacher.country = form.country.data.strip()
        teacher.school = form.school.data.strip()
        teacher.sex_document = form.sex_document.data
        teacher.accommodation_room_sharing = form.accommodation_room_sharing.data
        teacher.languages = ", ".join(form.languages.data)

        current_languages = {entry.language for entry in teacher.language_entries}
        desired_languages = set(form.languages.data)

        for entry in list(teacher.language_entries):
            if entry.language not in desired_languages:
                teacher.language_entries.remove(entry)
                db.session.delete(entry)

        for lang in desired_languages - current_languages:
            teacher.language_entries.append(TeacherLanguage(language=lang))
        teacher.meal_preference = form.meal_preference.data
        teacher.tshirt_size = form.tshirt_size.data
        teacher.special_needs = (form.special_needs.data.strip() or None) if form.special_needs.data else None
        teacher.delegation_leader = bool(form.delegation_leader.data)
        non_assessor = bool(form.non_assessor.data)
        prior_level = 0 if non_assessor else (form.prior_ipo_assessor_level.data or 0)
        teacher.prior_ipo_assessor_level = prior_level
        teacher.prior_ipo_assessor = bool(prior_level)
        if form.gdpr_acknowledged.data:
            teacher.gdpr_acknowledged = True
            teacher.gdpr_acknowledged_at = teacher.gdpr_acknowledged_at or datetime.utcnow()
        else:
            teacher.gdpr_acknowledged = False
            teacher.gdpr_acknowledged_at = None
        new_photo_consent = form.photo_consent.data == "yes"
        if teacher.photo_consent != new_photo_consent:
            teacher.photo_consent = new_photo_consent
            teacher.photo_consent_at = datetime.utcnow()
        if non_assessor:
            teacher.declared_assessor_competence = False
            teacher.declared_assessor_competence_at = None
        elif form.declared_assessor_competence.data:
            teacher.declared_assessor_competence = True
            teacher.declared_assessor_competence_at = teacher.declared_assessor_competence_at or datetime.utcnow()
        else:
            teacher.declared_assessor_competence = False
            teacher.declared_assessor_competence_at = None
        if form.accepted_ethics.data:
            teacher.accepted_ethics = True
            teacher.accepted_ethics_at = teacher.accepted_ethics_at or datetime.utcnow()
        else:
            teacher.accepted_ethics = False
            teacher.accepted_ethics_at = None
        teacher.is_absent = bool(form.is_absent.data)

        try:
            db.session.flush()
            log_activity(
                "teacher_updated",
                f"Teacher {teacher.first_name} {teacher.last_name} updated",
                object_type="teacher",
                object_id=teacher.id,
                details={
                    "email": teacher.email,
                    "country": teacher.country,
                    "languages": form.languages.data,
                },
            )
            db.session.commit()
            flash("Teacher details updated successfully.", "success")
            return redirect(url_for("web.admin_teachers"))
        except IntegrityError:
            db.session.rollback()
            flash("This email address is already registered with another teacher.", "error")

    return render_template(
        "register_teacher.html",
        form=form,
        form_title="Edit teacher",
        show_absence_toggle=True,
        ethics_consent_html=_ethics_consent_html(),
        gdpr_notice_html=_gdpr_notice_html(),
        photo_consent_html=_photo_consent_html(),
        assessor_competence_html=_assessor_competence_html(),
        accommodation_info_html=_accommodation_info_html(),
    )


@web_bp.route("/admin/teachers/<int:teacher_id>/reset-password", methods=["POST"])
def admin_teacher_reset_password(teacher_id: int):
    if not _current_admin():
        abort(403)
    teacher = db.session.get(Teacher, teacher_id)
    if not teacher:
        abort(404)

    username, plain_password = _ensure_teacher_portal_credentials(teacher, reset_password=True)
    db.session.flush()
    log_activity(
        "teacher_password_reset",
        f"Teacher {teacher.first_name} {teacher.last_name} portal password reset",
        object_type="teacher",
        object_id=teacher.id,
        details={"portal_username": username},
    )
    db.session.commit()

    if plain_password:
        if _auto_email_enabled() and teacher.email:
            _send_teacher_reset_email(teacher, plain_password)
            flash(
                f"Portal password reset and emailed to {teacher.email}. New password: {plain_password}",
                "success",
            )
        else:
            flash(
                f"Portal password reset. Automatic email is disabled; share this password manually: {plain_password}",
                "warning",
            )
    else:
        flash("Portal password reset, but no password was generated.", "error")

    return redirect(url_for("web.admin_teachers"))


@web_bp.route("/admin/teachers/<int:teacher_id>/delete", methods=["POST"])
def admin_teacher_delete(teacher_id: int):
    if not _current_admin():
        abort(403)
    teacher = db.session.get(Teacher, teacher_id)
    if not teacher:
        abort(404)

    expected_email = (teacher.email or "").strip().lower()
    confirm_email = (request.form.get("confirm_email") or "").strip().lower()
    if confirm_email != expected_email:
        flash("Confirmation email does not match.", "error")
        return redirect(url_for("web.admin_teachers"))

    teacher_name = " ".join(part for part in [teacher.first_name, teacher.last_name] if part).strip()
    teacher_email = teacher.email
    teacher_country = teacher.country
    delegation_leader = bool(teacher.delegation_leader)

    db.session.delete(teacher)
    try:
        db.session.flush()
    except IntegrityError:
        db.session.rollback()
        flash("Teacher could not be deleted because related records exist.", "error")
        return redirect(url_for("web.admin_teachers"))

    log_activity(
        "teacher_deleted",
        f"Teacher {teacher_name or teacher_email} deleted",
        object_type="teacher",
        object_id=teacher_id,
        details={
            "email": teacher_email,
            "country": teacher_country,
            "delegation_leader": delegation_leader,
        },
    )
    db.session.commit()
    flash("Teacher deleted.", "success")
    return redirect(url_for("web.admin_teachers"))


@web_bp.route("/admin/uploads")
def admin_uploads():
    active_batch = _get_active_batch(Stage.STAGE_1)
    if not active_batch:
        flash("No Stage 1 batch is published; draft assignments are ignored.", "warning")
    stage1_filter = request.args.get("status", "")
    all_students = (
        Student.query.filter_by(is_absent=False)
        .order_by(Student.last_name.asc(), Student.first_name.asc())
        .all()
    )
    status_map = {}
    extra_reader_map = {}
    students = []
    for student in all_students:
        assignment = _get_stage_one_assignment(student)
        status = _compute_stage_one_status(assignment)
        extra_state = _stage_one_extra_reader_state(assignment, status)
        if stage1_filter == "needs_extra" and not extra_state.needs_action:
            continue
        if stage1_filter == "extra_needs_assignment" and extra_state.bucket != "needs_assignment":
            continue
        if stage1_filter == "extra_awaiting_score" and extra_state.bucket != "awaiting_score":
            continue
        if stage1_filter == "extra_review_spread" and extra_state.bucket != "review_spread":
            continue
        students.append(student)
        status_map[student.id] = status
        extra_reader_map[student.id] = extra_state
    is_htmx = _is_htmx_request()
    template_name = "admin/partials/uploads_content.html" if is_htmx else "admin/uploads.html"
    context = {
        "students": students,
        "status_map": status_map,
        "extra_reader_map": extra_reader_map,
        "stage1_filter": stage1_filter,
        "active_batch": active_batch,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/essays/jump")
def admin_essay_jump():
    stage = request.args.get("stage", Stage.STAGE_1.value)
    raw_essay_id = (request.args.get("essay_id") or "").strip()
    fallback_endpoint = "web.admin_stage2_overview" if stage == Stage.STAGE_2.value else "web.admin_uploads"

    if not raw_essay_id:
        flash("Enter an Essay ID.", "error")
        return redirect(url_for(fallback_endpoint))
    if not raw_essay_id.isdigit():
        flash("Essay ID must be a number.", "error")
        return redirect(url_for(fallback_endpoint))

    essay_id = int(raw_essay_id)
    student = db.session.get(Student, essay_id)
    if not student or student.is_absent:
        flash(f"No essay with ID {essay_id} was found.", "error")
        return redirect(url_for(fallback_endpoint))

    assignment = _get_stage_one_assignment(student)
    if not assignment:
        flash(f"Essay {essay_id} has no published Stage 1 assignment.", "error")
        return redirect(url_for(fallback_endpoint))

    if stage == Stage.STAGE_2.value:
        if assignment.stage_two_candidate:
            return redirect(url_for("web.admin_stage2_score_student", student_id=student.id))
        flash(f"Essay {essay_id} is not marked for Stage 2. Opening Stage 1 scoring instead.", "warning")

    return redirect(url_for("web.admin_score_student", student_id=student.id))


@web_bp.route("/admin/uploads/<int:student_id>/download")
def admin_download_essay(student_id: int):
    student = db.session.get(Student, student_id)
    if not student or not student.uploaded_filename:
        abort(404)
    upload_folder = _get_upload_folder()
    file_path = upload_folder / student.uploaded_filename
    if not file_path.exists():
        abort(404)
    return send_from_directory(upload_folder, student.uploaded_filename, as_attachment=True)


@web_bp.route("/admin/students/cards.pdf")
def admin_student_cards_pdf():
    active_batch = _get_active_batch(Stage.STAGE_1)
    if not active_batch:
        flash("Publish a Stage 1 batch before generating cards.", "error")
        return redirect(url_for("web.admin_planner"))
    students = (
        Student.query.filter_by(is_absent=False)
        .order_by(Student.country.asc(), Student.last_name.asc())
        .all()
    )
    html = render_template(
        "admin/pdf/student_cards.html",
        students=students,
        current_year=datetime.utcnow().year,
        upload_url=url_for("web.student_login", _external=True),
        event_theme="Essay Competition",
    )
    pdf_bytes = HTML(string=html, base_url=request.url_root).write_pdf()
    response = make_response(pdf_bytes)
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = "attachment; filename=participant-cards.pdf"
    return response


@web_bp.route("/admin/teachers/cards.pdf")
def admin_teacher_cards_pdf():
    if not _current_admin():
        abort(403)
    if not _teacher_portal_enabled():
        flash("Enable teacher portal in Settings before generating teacher cards.", "error")
        return redirect(url_for("web.admin_settings"))

    teachers = (
        Teacher.query.filter_by(is_absent=False)
        .order_by(Teacher.country.asc(), Teacher.last_name.asc())
        .all()
    )
    card_entries: list[dict] = []
    for teacher in teachers:
        username, plain_password = _ensure_teacher_portal_credentials(teacher, reset_password=True)
        if plain_password is None:
            # Ensure a password is available for the card by forcing a reset
            username, plain_password = _ensure_teacher_portal_credentials(teacher, reset_password=True)
        card_entries.append({
            "teacher": teacher,
            "username": username,
            "password": plain_password,
        })

    db.session.commit()

    html = render_template(
        "admin/pdf/teacher_cards.html",
        teachers=card_entries,
        login_url=url_for("teacher_portal.login", _external=True),
        board_url=url_for("web.teacher_board", _external=True),
        current_year=datetime.utcnow().year,
    )
    pdf_bytes = HTML(string=html, base_url=request.url_root).write_pdf()
    response = make_response(pdf_bytes)
    response.headers["Content-Type"] = "application/pdf"
    response.headers["Content-Disposition"] = "attachment; filename=teacher-portal-cards.pdf"
    log_activity(
        "teacher_portal_cards_generated",
        "Teacher portal cards generated",
        object_type="teacher_portal",
    )
    return response


@web_bp.route("/admin/uploads/<int:student_id>/scores", methods=["GET", "POST"])
def admin_score_student(student_id: int):
    student = db.session.get(Student, student_id)
    if not student or student.is_absent:
        abort(404)

    assignment = _get_stage_one_assignment(student)
    if not assignment:
        flash("No published Stage 1 assignment found for this student.", "error")
        return redirect(url_for("web.admin_uploads"))

    is_htmx = _is_htmx_request()

    def _load_scoring_state() -> SimpleNamespace:
        assessors = _stage_assessors(assignment, Stage.STAGE_1, include_group=True)
        scores = EssayScore.query.filter_by(
            assignment_id=assignment.id,
            stage=Stage.STAGE_1,
        ).all()
        scores_map = {score.teacher_id: score for score in scores}
        pdf_available = _essay_pdf_path(student) is not None
        stage1_global_lock = _is_stage_locked(Stage.STAGE_1)
        assignment_stage1_locked = assignment.stage1_locked or stage1_global_lock

        reader_entries = []
        for assessor in assessors:
            teacher = assessor.teacher
            score_record = scores_map.get(teacher.id)
            status_info = _resolve_portal_status(student, score_record, pdf_available, locked=assignment_stage1_locked)
            reader_entries.append(
                SimpleNamespace(
                    teacher=teacher,
                    teacher_id=teacher.id,
                    source=assessor.source,
                    is_extra=assessor.is_extra,
                    score=score_record.score if score_record else None,
                    score_value=format_score_value(score_record.score) if score_record and score_record.score is not None else "",
                    updated_at=score_record.updated_at if score_record else None,
                    finalized=score_record.teacher_finalized if score_record else False,
                    finalized_at=score_record.teacher_finalized_at if score_record else None,
                    submitted_at=score_record.submitted_at if score_record else None,
                    portal_status_code=status_info.code,
                    portal_status_label=status_info.label,
                    portal_status_css=status_info.css,
                    portal_status_timestamp=status_info.timestamp,
                )
            )

        assigned_teacher_ids = {entry.teacher_id for entry in reader_entries}
        stage2_teacher_ids = {
            assessor.teacher.id
            for assessor in _stage_assessors(assignment, Stage.STAGE_2, include_group=False)
        }
        exclude_ids = assigned_teacher_ids.union(stage2_teacher_ids)
        available_query = _assessment_eligible_teachers_query()
        if exclude_ids:
            available_query = available_query.filter(~Teacher.id.in_(exclude_ids))
        available_teachers = [
            teacher
            for teacher in available_query.order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()
            if _teacher_can_grade(student, teacher, student.essay_language)
        ]

        db.session.expire(assignment, ["scores"])
        status = _compute_stage_one_status(assignment)
        return SimpleNamespace(
            scores_map=scores_map,
            reader_entries=reader_entries,
            assigned_teacher_ids=assigned_teacher_ids,
            available_teachers=available_teachers,
            stage1_global_lock=stage1_global_lock,
            status=status,
        )

    def _render_score_response():
        state = _load_scoring_state()
        template_name = "admin/partials/scoring_content.html" if is_htmx else "admin/scoring.html"
        context = {
            "student": student,
            "assignment": assignment,
            "reader_entries": state.reader_entries,
            "status": state.status,
            "available_teachers": state.available_teachers,
            "stage1_global_lock": state.stage1_global_lock,
            "score_choices": score_choices(include_blank=False),
            "is_htmx": is_htmx,
        }
        if not is_htmx:
            context.update(_quick_message_context())
            return render_template(template_name, **context)
        return _render_htmx_template(template_name, **context)

    state = _load_scoring_state()

    def _post_response():
        if is_htmx:
            return _render_score_response()
        return redirect(url_for("web.admin_score_student", student_id=student.id))

    remove_id = request.form.get("remove_reader")
    unlock_id = request.form.get("unlock_reader")
    action = request.form.get("action")

    if request.method == "POST" and action == "toggle_stage1_lock":
        assignment.stage1_locked = not assignment.stage1_locked
        now = datetime.utcnow()
        for score in EssayScore.query.filter_by(
            assignment_id=assignment.id,
            stage=Stage.STAGE_1,
        ).all():
            score.locked_by_admin = assignment.stage1_locked
            if assignment.stage1_locked:
                score.admin_updated_at = now
        log_activity(
            "stage1_grading_lock",
            f"Stage 1 grading lock set to {assignment.stage1_locked} for student {student.id}",
            object_type="student",
            object_id=student.id,
            details={"locked": assignment.stage1_locked},
        )
        db.session.commit()
        flash(
            "Stage 1 grading locked." if assignment.stage1_locked else "Stage 1 grading unlocked.",
            "success",
        )
        return _post_response()

    if request.method == "POST" and unlock_id:
        try:
            teacher_id = int(unlock_id)
        except ValueError:
            flash("Invalid teacher selection.", "error")
            return _post_response()
        score_record = state.scores_map.get(teacher_id)
        if not score_record or not score_record.teacher_finalized:
            flash("Score is not finalized for this teacher.", "warning")
            return _post_response()
        score_record.teacher_finalized = False
        score_record.teacher_finalized_at = None
        score_record.admin_updated_at = datetime.utcnow()
        log_activity(
            "stage1_score_unlocked",
            f"Unlocked Stage 1 score for teacher {teacher_id} on student {student.id}",
            object_type="student",
            object_id=student.id,
            details={"teacher_id": teacher_id},
        )
        db.session.commit()
        flash("Teacher score unlocked for editing.", "success")
        return _post_response()

    if request.method == "POST" and remove_id:
        try:
            teacher_id = int(remove_id)
        except ValueError:
            flash("Invalid teacher selection.", "error")
            return _post_response()
        extra = EssayAssignmentExtraReader.query.filter_by(assignment_id=assignment.id, teacher_id=teacher_id, stage=Stage.STAGE_1).first()
        if not extra:
            flash("Extra reader not found for this essay.", "error")
            return _post_response()
        score_record = state.scores_map.get(teacher_id)
        if score_record and score_record.stage == Stage.STAGE_1:
            db.session.delete(score_record)
            state.scores_map.pop(teacher_id, None)
        db.session.delete(extra)
        log_activity(
            "stage1_extra_reader_removed",
            f"Removed extra reader {teacher_id} for student {student.id}",
            object_type="student",
            object_id=student.id,
            details={"teacher_id": teacher_id},
        )
        db.session.commit()
        flash("Extra reader removed.", "success")
        return _post_response()

    if request.method == "POST" and action == "add_reader":
        teacher_id = request.form.get("teacher_id")
        if not teacher_id:
            flash("Select a teacher to add.", "error")
            return _post_response()
        try:
            teacher_id_int = int(teacher_id)
        except ValueError:
            flash("Invalid teacher selection.", "error")
            return _post_response()
        teacher = db.session.get(Teacher, teacher_id_int)
        if not teacher or teacher.is_absent:
            flash("Teacher not available.", "error")
            return _post_response()
        if teacher_id_int in state.assigned_teacher_ids:
            flash("Teacher already assigned.", "warning")
            return _post_response()
        if not _teacher_can_grade(student, teacher, student.essay_language):
            flash("Teacher is not eligible to grade this essay.", "error")
            return _post_response()
        db.session.add(EssayAssignmentExtraReader(assignment_id=assignment.id, teacher_id=teacher_id_int, stage=Stage.STAGE_1))
        log_activity(
            "stage1_extra_reader_added",
            f"Added extra reader {teacher_id_int} for student {student.id}",
            object_type="student",
            object_id=student.id,
            details={"teacher_id": teacher_id_int},
        )
        db.session.commit()
        flash("Additional reader assigned.", "success")
        return _post_response()

    if request.method == "POST" and action == "update_scores":
        updates = {}
        now = datetime.utcnow()
        for entry in state.reader_entries:
            field_name = f"score_{entry.teacher_id}"
            value = request.form.get(field_name, "").strip()
            if not value:
                continue
            score_value, score_error = parse_score_value(value)
            if score_error or score_value is None:
                flash(
                    f"Invalid score for {entry.teacher.first_name} {entry.teacher.last_name}. "
                    f"{SCORE_VALIDATION_MESSAGE}",
                    "error",
                )
                return _post_response()
            numeric = float(score_value)
            score_record = state.scores_map.get(entry.teacher_id)
            if not score_record:
                score_record = EssayScore(
                    student_id=student.id,
                    teacher_id=entry.teacher_id,
                    assignment_id=assignment.id,
                    stage=Stage.STAGE_1,
                )
                db.session.add(score_record)
                state.scores_map[entry.teacher_id] = score_record
            prev_admin_entered = score_record.admin_entered if score_record else False
            prev_admin_override = score_record.admin_override if score_record else False
            had_teacher_input = bool(score_record.submitted_at) or bool(score_record.teacher_finalized)
            existing_score = score_record.score if score_record and score_record.score is not None else None

            if existing_score is not None and math.isclose(existing_score, numeric, abs_tol=1e-9):
                continue

            score_record.score = numeric
            score_record.admin_updated_at = now
            if had_teacher_input and not prev_admin_override and not prev_admin_entered:
                score_record.admin_override = True
                score_record.admin_entered = False
            elif prev_admin_override:
                score_record.admin_override = True
                score_record.admin_entered = False
            else:
                score_record.admin_entered = True
                score_record.admin_override = False

            score_record.teacher_finalized = False
            score_record.teacher_finalized_at = None
            updates[entry.teacher_id] = numeric
        if not updates:
            flash("No score changes detected.", "info")
            return _post_response()

        _apply_stage_one_rules(assignment)
        log_activity(
            "stage1_scores_updated",
            f"Stage 1 scores updated for student {student.id}",
            object_type="student",
            object_id=student.id,
            details={"updates": updates},
        )
        db.session.commit()
        flash("Scores saved.", "success")
        return _post_response()

    if request.method == "POST" and action == "mark_stage2":
        assignment.stage_two_candidate = not assignment.stage_two_candidate
        if assignment.stage_two_candidate:
            assignment.needs_additional_reading = False
        log_activity(
            "stage1_mark_stage2",
            f"Stage 2 flag set to {assignment.stage_two_candidate} for student {student.id}",
            object_type="student",
            object_id=student.id,
        )
        db.session.commit()
        flash(
            "Marked as Stage 2 candidate."
            if assignment.stage_two_candidate
            else "Stage 2 recommendation removed.",
            "success",
        )
        return _post_response()

    if request.method == "POST" and action == "mark_extra":
        assignment.needs_additional_reading = not assignment.needs_additional_reading
        if assignment.needs_additional_reading:
            assignment.stage_two_candidate = False
        log_activity(
            "stage1_mark_extra",
            f"Extra reading flag set to {assignment.needs_additional_reading} for student {student.id}",
            object_type="student",
            object_id=student.id,
        )
        db.session.commit()
        flash(
            "Marked as needing extra reading."
            if assignment.needs_additional_reading
            else "Extra reading flag removed.",
            "success",
        )
        return _post_response()

    return _render_score_response()


@web_bp.route("/admin/planner", methods=["GET", "POST"])
def admin_planner():
    form = AssignmentBatchForm()
    form.stage.data = Stage.STAGE_1.value
    is_htmx = _is_htmx_request()

    def _render_planner_overview(current_form: AssignmentBatchForm):
        batches = AssignmentBatch.query.order_by(AssignmentBatch.created_at.desc()).all()
        active_batch = _get_active_batch(Stage.STAGE_1)
        template_name = "admin/partials/planner_overview_content.html" if is_htmx else "admin/planner/overview.html"
        context = {
            "form": current_form,
            "batches": batches,
            "active_batch": active_batch,
            "PlannerStatus": PlannerStatus,
            "PlannerAlgorithm": PlannerAlgorithm,
            "is_htmx": is_htmx,
        }
        if not is_htmx:
            return render_template(template_name, **context)
        return _render_htmx_template(template_name, **context)

    if form.validate_on_submit():
        batch = AssignmentBatch(
            name=form.name.data.strip(),
            stage=Stage(form.stage.data),
            status=PlannerStatus.DRAFT,
            planner_algorithm=PlannerAlgorithm(
                form.planner_algorithm.data or PlannerAlgorithm.STRICT_FOUR_WITH_BACKUPS.value
            ),
            notes=form.notes.data.strip() if form.notes.data else None,
        )
        db.session.add(batch)
        db.session.flush()
        log_activity(
            "batch_created",
            f"Created assignment batch {batch.name}",
            object_type="assignment_batch",
            object_id=batch.id,
            details={"stage": batch.stage.value},
        )
        db.session.commit()
        flash("Assignment batch created.", "success")
        if is_htmx:
            fresh_form = AssignmentBatchForm()
            fresh_form.stage.data = Stage.STAGE_1.value
            fresh_form.planner_algorithm.data = PlannerAlgorithm.STRICT_FOUR_WITH_BACKUPS.value
            return _render_planner_overview(fresh_form)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    return _render_planner_overview(form)


@web_bp.route("/admin/planner/batches/<int:batch_id>")
def admin_planner_batch(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)

    is_htmx = _is_htmx_request()

    group_form = ReadingGroupForm()
    if not group_form.target_essay_load.data:
        group_form.target_essay_load.data = 5
    group_form.language.data = group_form.language.data or "English"
    group_form.auto_generated.data = False

    member_form = ReadingGroupMemberForm()
    member_form.teacher_id.choices = _planner_teacher_choices(batch.stage)
    member_form.role.choices = [(role.value, role.name.title()) for role in GroupMemberRole]
    if "active" not in request.form:
        member_form.active.data = True

    groups = _load_batch_groups(batch)
    assigned_student_ids = {assignment.student_id for group in groups for assignment in group.assignments}

    for group in groups:
        group.members.sort(key=lambda member: (0 if member.role == GroupMemberRole.LEADER else 1, member.id))

    students_by_language: dict[str, list[Student]] = defaultdict(list)
    for student in Student.query.order_by(Student.essay_language.asc(), Student.country.asc(), Student.last_name.asc()).all():
        if student.id in assigned_student_ids:
            continue
        students_by_language[student.essay_language].append(student)

    eligible_students_map: dict[int, list[Student]] = {}
    for group in groups:
        active_countries = {
            member.teacher.country
            for member in group.members
            if member.active and member.teacher
        }
        candidates = [
            student
            for student in students_by_language.get(group.language, [])
            if group.allow_same_country or student.country not in active_countries
        ]
        eligible_students_map[group.id] = candidates[:100]

    teachers = _assessment_eligible_teachers_query().order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()
    teacher_languages_map: dict[int, set[str]] = {}
    for teacher in teachers:
        _ensure_language_entries(teacher)
        teacher_languages_map[teacher.id] = {
            entry.language for entry in teacher.language_entries if entry.language
        }

    active_teacher_ids = {
        member.teacher_id
        for group in groups
        for member in group.members
        if member.active and member.teacher
    }
    available_teachers_map: dict[int, list[Teacher]] = {}
    group_status_map: dict[int, dict[str, int | bool]] = {}

    for group in groups:
        group_teacher_ids = {member.teacher_id for member in group.members if member.teacher}
        active_count = sum(
            1
            for member in group.members
            if member.active and member.teacher and not member.teacher.is_absent
        )
        needs_teachers = active_count < 4
        group_status_map[group.id] = {
            "active_count": active_count,
            "needs_teachers": needs_teachers,
            "missing": max(0, 4 - active_count),
        }

        available_teachers: list[Teacher] = []
        for teacher in teachers:
            if teacher.id in group_teacher_ids:
                continue
            if teacher.id in active_teacher_ids:
                continue
            languages_spoken = teacher_languages_map.get(teacher.id, set())
            if group.language not in languages_spoken:
                continue
            available_teachers.append(teacher)
        available_teachers_map[group.id] = available_teachers

    supplemental_reader_map = _supplemental_reader_context(batch, groups, teachers)
    assignment_stage1_status_map = {
        assignment.id: _compute_stage_one_status(assignment)
        for group in groups
        for assignment in group.assignments
        if assignment.stage == Stage.STAGE_1
    }

    template_name = "admin/partials/planner_batch_content.html" if is_htmx else "admin/planner/batch_detail.html"
    context = {
        "batch": batch,
        "groups": groups,
        "group_form": group_form,
        "member_form": member_form,
        "eligible_students_map": eligible_students_map,
        "available_teachers_map": available_teachers_map,
        "supplemental_reader_map": supplemental_reader_map,
        "assignment_stage1_status_map": assignment_stage1_status_map,
        "group_status_map": group_status_map,
        "PlannerStatus": PlannerStatus,
        "PlannerAlgorithm": PlannerAlgorithm,
        "planner_algorithm_choices": _planner_algorithm_choices(),
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        context.update(_quick_message_context())
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route("/admin/planner/batches/<int:batch_id>/algorithm", methods=["POST"])
def admin_planner_update_algorithm(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response

    algorithm_value = request.form.get("planner_algorithm")
    try:
        algorithm = PlannerAlgorithm(algorithm_value)
    except Exception:
        flash("Invalid planner algorithm.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    batch.planner_algorithm = algorithm
    log_activity(
        "planner_algorithm_updated",
        f"Planner algorithm updated for batch {batch.name}",
        object_type="assignment_batch",
        object_id=batch.id,
        details={"planner_algorithm": algorithm.value},
    )
    db.session.commit()
    flash("Planner algorithm updated.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/groups", methods=["POST"])
def admin_planner_add_group(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response

    form = ReadingGroupForm()
    if form.validate_on_submit():
        target = form.target_essay_load.data or 5
        group = ReadingGroup(
            name=form.name.data.strip(),
            language=form.language.data,
            stage=batch.stage,
            target_essay_load=target,
            allow_same_country=bool(form.allow_same_country.data),
            auto_generated=bool(form.auto_generated.data),
            status=batch.status,
            batch=batch,
        )
        db.session.add(group)
        db.session.flush()
        log_activity(
            "group_created",
            f"Group {group.name} created for {group.language}",
            object_type="reading_group",
            object_id=group.id,
            details={"batch_id": batch.id, "language": group.language},
        )
        db.session.commit()
        flash("Reading group created.", "success")
    else:
        for errors in form.errors.values():
            for error in errors:
                flash(error, "error")

    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/groups/<int:group_id>/toggle-lock", methods=["POST"])
def admin_planner_toggle_lock(batch_id: int, group_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response
    group = ReadingGroup.query.filter_by(id=group_id, batch_id=batch.id).first_or_404()
    group.locked = not group.locked
    log_activity(
        "group_lock_toggled",
        f"Group {group.name} lock set to {group.locked}",
        object_type="reading_group",
        object_id=group.id,
        details={"locked": group.locked},
    )
    db.session.commit()
    flash("Group lock status updated.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/groups/<int:group_id>/delete", methods=["POST"])
def admin_planner_delete_group(batch_id: int, group_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response
    group = ReadingGroup.query.filter_by(id=group_id, batch_id=batch.id).first_or_404()
    if group.assignments:
        flash("Group cannot be deleted while assignments exist.", "error")
    else:
        group_name = group.name
        db.session.delete(group)
        log_activity(
            "group_deleted",
            f"Group {group_name} deleted",
            object_type="reading_group",
            object_id=group_id,
        )
        db.session.commit()
        flash("Reading group deleted.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/members", methods=["POST"])
def admin_planner_add_member(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response

    form = ReadingGroupMemberForm()
    form.teacher_id.choices = _planner_teacher_choices(batch.stage)
    form.role.choices = [(role.value, role.name.title()) for role in GroupMemberRole]
    group_id_value = request.form.get("group_id")
    if not group_id_value:
        flash("Group selection is required.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    group = ReadingGroup.query.filter_by(id=int(group_id_value), batch_id=batch.id).first()
    if not group:
        flash("Invalid group selection.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    if form.validate_on_submit():
        teacher_id = form.teacher_id.data
        if _teacher_has_active_group(teacher_id, batch.stage, exclude_group_id=group.id):
            flash("Teacher already assigned to an active group for this stage.", "error")
        else:
            already_member = any(member.teacher_id == teacher_id for member in group.members)
            if already_member:
                flash("Teacher is already part of this group.", "error")
            else:
                member = ReadingGroupMember(
                    group=group,
                    teacher_id=teacher_id,
                    role=GroupMemberRole(form.role.data),
                    active=bool(form.active.data),
                )
                db.session.add(member)
                db.session.flush()
                log_activity(
                    "group_member_added",
                    f"Teacher {member.teacher_id} added to group {group.name}",
                    object_type="reading_group",
                    object_id=group.id,
                    details={
                        "teacher_id": member.teacher_id,
                        "role": member.role.value,
                        "active": member.active,
                    },
                )
                db.session.commit()
                flash("Group member added.", "success")
    else:
        for errors in form.errors.values():
            for error in errors:
                flash(error, "error")

    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/members/<int:member_id>/remove", methods=["POST"])
def admin_planner_remove_member(batch_id: int, member_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response
    member = (
        ReadingGroupMember.query.join(ReadingGroup)
        .filter(ReadingGroupMember.id == member_id, ReadingGroup.batch_id == batch.id)
        .first_or_404()
    )
    group_name = member.group.name if member.group else str(batch.id)
    teacher_id = member.teacher_id
    db.session.delete(member)
    log_activity(
        "group_member_removed",
        f"Teacher {teacher_id} removed from group {group_name}",
        object_type="reading_group",
        object_id=member.group_id,
        details={"teacher_id": teacher_id},
    )
    db.session.commit()
    flash("Group member removed.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route(
    "/admin/planner/batches/<int:batch_id>/groups/<int:group_id>/members/<int:member_id>/role",
    methods=["POST"],
)
def admin_planner_update_member_role(batch_id: int, group_id: int, member_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response

    member = (
        ReadingGroupMember.query.join(ReadingGroup)
        .filter(
            ReadingGroupMember.id == member_id,
            ReadingGroupMember.group_id == group_id,
            ReadingGroup.batch_id == batch.id,
        )
        .first_or_404()
    )

    group = member.group
    target_role_value = request.form.get("role")
    try:
        target_role = GroupMemberRole(target_role_value)
    except Exception:
        flash("Invalid role change requested.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    if target_role not in {GroupMemberRole.LEADER, GroupMemberRole.READER}:
        flash("Invalid role change requested.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    if target_role == GroupMemberRole.LEADER:
        (
            ReadingGroupMember.query.filter(
                ReadingGroupMember.group_id == group.id,
                ReadingGroupMember.id != member.id,
                ReadingGroupMember.role == GroupMemberRole.LEADER,
            ).update({ReadingGroupMember.role: GroupMemberRole.READER}, synchronize_session=False)
        )
        member.role = GroupMemberRole.LEADER
        message = "Member promoted to leader."
    else:
        member.role = GroupMemberRole.READER
        message = "Member marked as reader."

    log_activity(
        "group_member_role_updated",
        f"Teacher {member.teacher_id} role updated in group {group.name}",
        object_type="reading_group",
        object_id=group.id,
        details={"member_id": member.id, "role": member.role.value},
    )
    db.session.commit()
    flash(message, "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/groups/<int:group_id>/assign", methods=["POST"])
def admin_planner_assign_manual(batch_id: int, group_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response
    group = ReadingGroup.query.filter_by(id=group_id, batch_id=batch.id).first_or_404()
    if group.locked:
        flash("Group is locked.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    student_id_value = request.form.get("student_id")
    if not student_id_value:
        flash("Student selection is required.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    student = db.session.get(Student, int(student_id_value))
    if not student:
        flash("Selected student does not exist.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    existing_assignment = EssayAssignment.query.filter_by(batch_id=batch.id, student_id=student.id).first()
    if existing_assignment:
        flash("Student already assigned in this batch.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    if student.essay_language != group.language:
        flash("Student language does not match the group language.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    active_countries = {
        member.teacher.country
        for member in group.members
        if member.active and member.teacher and not member.teacher.is_absent
    }
    if not group.allow_same_country and student.country in active_countries:
        flash("Group cannot read students from the same country.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    assignment = EssayAssignment(
        student_id=student.id,
        group_id=group.id,
        batch=batch,
        stage=batch.stage,
        manual_override=True,
        status=PlannerStatus.DRAFT,
    )
    db.session.add(assignment)
    db.session.flush()
    log_activity(
        "assignment_added",
        f"Student {student.first_name} {student.last_name} assigned to {group.name}",
        object_type="essay_assignment",
        object_id=assignment.id,
        details={
            "student_id": student.id,
            "group_id": group.id,
            "batch_id": batch.id,
        },
    )
    db.session.commit()
    flash("Essay assigned to group.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route(
    "/admin/planner/batches/<int:batch_id>/groups/<int:group_id>/supplemental-readers",
    methods=["POST"],
)
def admin_planner_assign_supplemental_reader(batch_id: int, group_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    group = ReadingGroup.query.filter_by(id=group_id, batch_id=batch.id).first_or_404()
    if batch.stage != Stage.STAGE_1 or group.stage != Stage.STAGE_1:
        flash("Supplemental readers are only supported for Stage 1 batches.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    teacher_id = request.form.get("teacher_id")
    if not teacher_id:
        flash("Select a teacher to assign.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))
    try:
        teacher_id_int = int(teacher_id)
    except ValueError:
        flash("Invalid teacher selection.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    teacher = db.session.get(Teacher, teacher_id_int)
    if not teacher or teacher.is_absent or not teacher.declared_assessor_competence:
        flash("Teacher is not available for assessment.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    workload = _teacher_stage_one_workload(batch, teacher)
    if not workload.is_free:
        flash(
            f"{teacher.first_name} {teacher.last_name} still has "
            f"{workload.pending_count} pending Stage 1 assignment"
            f"{'s' if workload.pending_count != 1 else ''}.",
            "error",
        )
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    assignments = _supplemental_reader_assignments_for_group(group, teacher)
    if not assignments:
        flash("No underfilled essays in this group are eligible for that teacher.", "warning")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    created_assignment_ids: list[int] = []
    for assignment in assignments:
        db.session.add(
            EssayAssignmentExtraReader(
                assignment_id=assignment.id,
                teacher_id=teacher.id,
                stage=Stage.STAGE_1.value,
            )
        )
        created_assignment_ids.append(assignment.id)

    log_activity(
        "stage1_supplemental_reader_assigned",
        f"Assigned supplemental reader {teacher.id} to group {group.name}",
        object_type="reading_group",
        object_id=group.id,
        details={
            "batch_id": batch.id,
            "teacher_id": teacher.id,
            "assignment_ids": created_assignment_ids,
        },
    )
    db.session.commit()
    flash(
        f"Assigned {teacher.first_name} {teacher.last_name} to "
        f"{len(created_assignment_ids)} supplemental Stage 1 reading"
        f"{'s' if len(created_assignment_ids) != 1 else ''}.",
        "success",
    )
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route(
    "/admin/planner/batches/<int:batch_id>/groups/<int:group_id>/supplemental-readers/<int:teacher_id>/remove",
    methods=["POST"],
)
def admin_planner_remove_supplemental_reader(batch_id: int, group_id: int, teacher_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    group = ReadingGroup.query.filter_by(id=group_id, batch_id=batch.id).first_or_404()
    if batch.stage != Stage.STAGE_1 or group.stage != Stage.STAGE_1:
        flash("Supplemental readers are only supported for Stage 1 batches.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    extra_rows = (
        EssayAssignmentExtraReader.query
        .join(EssayAssignment, EssayAssignment.id == EssayAssignmentExtraReader.assignment_id)
        .filter(
            EssayAssignment.batch_id == batch.id,
            EssayAssignment.group_id == group.id,
            EssayAssignment.stage == Stage.STAGE_1,
            EssayAssignmentExtraReader.teacher_id == teacher_id,
            EssayAssignmentExtraReader.stage == Stage.STAGE_1.value,
        )
        .all()
    )
    if not extra_rows:
        flash("Supplemental reader assignment not found for this group.", "warning")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    assignment_ids = [row.assignment_id for row in extra_rows]
    score_records = (
        EssayScore.query.filter(
            EssayScore.assignment_id.in_(assignment_ids),
            EssayScore.teacher_id == teacher_id,
            EssayScore.stage == Stage.STAGE_1,
        )
        .all()
    )
    score_ids = [score.id for score in score_records]
    for score in score_records:
        db.session.delete(score)
    for extra in extra_rows:
        db.session.delete(extra)

    teacher = db.session.get(Teacher, teacher_id)
    teacher_label = f"{teacher.first_name} {teacher.last_name}" if teacher else f"teacher {teacher_id}"
    log_activity(
        "stage1_supplemental_reader_removed",
        f"Removed supplemental reader {teacher_id} from group {group.name}",
        object_type="reading_group",
        object_id=group.id,
        details={
            "batch_id": batch.id,
            "teacher_id": teacher_id,
            "assignment_ids": assignment_ids,
            "score_ids": score_ids,
        },
    )
    db.session.commit()
    flash(
        f"Removed {teacher_label} from {len(assignment_ids)} supplemental Stage 1 reading"
        f"{'s' if len(assignment_ids) != 1 else ''}.",
        "success",
    )
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/assignments/<int:assignment_id>/remove", methods=["POST"])
def admin_planner_remove_assignment(batch_id: int, assignment_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response

    assignment = EssayAssignment.query.filter_by(id=assignment_id, batch_id=batch.id).first_or_404()

    db.session.delete(assignment)
    log_activity(
        "assignment_removed",
        f"Assignment {assignment_id} removed from batch {batch.name}",
        object_type="essay_assignment",
        object_id=assignment_id,
        details={"batch_id": batch.id, "group_id": assignment.group_id, "student_id": assignment.student_id},
    )
    db.session.commit()
    flash("Assignment removed.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/auto-generate", methods=["POST"])
def admin_planner_auto_generate(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response
    summary = _auto_generate_groups_for_batch(batch)
    if "created" in summary:
        created_total = sum(summary["created"].values())
    else:
        created_total = sum(
            count
            for key, count in summary.items()
            if key
            not in {
                "removed",
                "removed_assignments",
                "backups",
                "shortfall",
                "matching_failed",
                "teachers_assigned",
                "underfilled",
                "no_student_language_groups",
            }
            and isinstance(count, int)
        )
    log_activity(
        "groups_auto_generated",
        f"Auto-generated groups for batch {batch.name}",
        object_type="assignment_batch",
        object_id=batch.id,
        details=summary,
    )
    db.session.commit()
    removed_count = summary.get("removed", 0)
    removed_assignment_count = summary.get("removed_assignments", 0)
    backups_count = summary.get("backups", 0)
    message = f"Generated {created_total} groups (removed {removed_count} auto-generated groups)."
    if removed_assignment_count:
        message += f" Removed {removed_assignment_count} assignment(s) from replaced groups."
    if backups_count:
        message += f" {backups_count} teacher(s) assigned as backups."
    if summary.get("teachers_assigned") is not None:
        message += f" {summary['teachers_assigned']} teacher(s) assigned as active readers."
    if summary.get("matching_failed"):
        flash("Unable to auto-generate groups for all teachers. Please adjust teacher availability or language coverage.", "error")
    else:
        flash(message, "success")
        if summary.get("underfilled"):
            flash(f"{summary['underfilled']} generated group(s) have fewer than four active teachers.", "warning")
        if summary.get("no_student_language_groups"):
            flash(
                f"{summary['no_student_language_groups']} generated group(s) have no matching student essays yet.",
                "warning",
            )
        if summary.get("shortfall"):
            shortfall_items = ", ".join(f"{language}: {value}" for language, value in summary["shortfall"].items())
            flash(f"Not enough teachers in: {shortfall_items}. Some language coverage targets were reduced.", "warning")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/auto-assign", methods=["POST"])
def admin_planner_auto_assign(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    guard_response = _published_batch_guard_response(batch)
    if guard_response:
        return guard_response
    assigned, skipped = _auto_assign_essays(batch)
    metadata = {
        "assigned": assigned,
        "skipped": [student.id for student in skipped],
        "planner_algorithm": _planner_algorithm_for_batch(batch).value,
    }
    log_activity(
        "essays_auto_assigned",
        f"Auto-assigned essays for batch {batch.name}",
        object_type="assignment_batch",
        object_id=batch.id,
        details=metadata,
    )
    db.session.commit()
    if skipped:
        flash(
            f"Assigned {assigned} essays. {len(skipped)} essays need manual attention (no eligible group).",
            "error",
        )
    else:
        flash(f"Assigned {assigned} essays to reading groups.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/unpublish", methods=["POST"])
def admin_planner_unpublish(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    if batch.status != PlannerStatus.PUBLISHED:
        flash("Batch is not published.", "warning")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    batch.status = PlannerStatus.DRAFT
    batch.published_at = None

    for group in batch.groups:
        group.status = PlannerStatus.DRAFT
    for assignment in batch.assignments:
        assignment.status = PlannerStatus.DRAFT

    log_activity(
        "batch_unpublished",
        f"Unpublished batch {batch.name}",
        object_type="assignment_batch",
        object_id=batch.id,
        details={"stage": batch.stage.value},
    )
    db.session.commit()
    flash("Batch returned to draft state.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/publish", methods=["POST"])
def admin_planner_publish(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)

    algorithm = _planner_algorithm_for_batch(batch)
    unresolved_groups = []
    for group in batch.groups:
        active_count = sum(
            1
            for member in group.members
            if member.active and member.teacher and not member.teacher.is_absent
        )
        if algorithm == PlannerAlgorithm.STRICT_FOUR_WITH_BACKUPS and active_count < 4:
            unresolved_groups.append(group)
        elif algorithm == PlannerAlgorithm.ALL_ASSESSORS_NO_BACKUPS and active_count == 0:
            unresolved_groups.append(group)
    if unresolved_groups:
        if algorithm == PlannerAlgorithm.STRICT_FOUR_WITH_BACKUPS:
            flash("All groups must have at least four active teachers before publishing.", "error")
        else:
            flash("All groups must have at least one active teacher before publishing.", "error")
        if _is_htmx_request():
            return admin_planner_batch(batch.id)
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    now = datetime.utcnow()
    batch.status = PlannerStatus.PUBLISHED
    batch.published_at = now

    other_batches = (
        AssignmentBatch.query.filter(
            AssignmentBatch.id != batch.id,
            AssignmentBatch.stage == batch.stage,
            AssignmentBatch.status == PlannerStatus.PUBLISHED,
        )
        .all()
    )
    for other in other_batches:
        other.status = PlannerStatus.DRAFT
        other.published_at = None

    for group in batch.groups:
        group.status = PlannerStatus.PUBLISHED
    for assignment in batch.assignments:
        assignment.status = PlannerStatus.PUBLISHED

    log_activity(
        "batch_published",
        f"Published assignment batch {batch.name}",
        object_type="assignment_batch",
        object_id=batch.id,
        details={
            "groups": len(batch.groups),
            "assignments": len(batch.assignments),
            "stage": batch.stage.value,
            "planner_algorithm": algorithm.value,
        },
    )
    db.session.commit()
    flash("Assignment batch published.", "success")
    if _is_htmx_request():
        return admin_planner_batch(batch.id)
    return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))


@web_bp.route("/admin/planner/batches/<int:batch_id>/delete", methods=["POST"])
def admin_planner_delete_batch(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    if batch.status == PlannerStatus.PUBLISHED:
        flash("Cannot delete a published batch. Unpublish another batch first.", "error")
        return redirect(url_for("web.admin_planner"))
    name = batch.name
    stage = batch.stage.value
    log_activity(
        "batch_deleted",
        f"Deleted batch {name}",
        object_type="assignment_batch",
        object_id=batch_id,
        details={"stage": stage},
    )
    db.session.delete(batch)
    db.session.commit()
    flash("Batch deleted.", "success")
    if _is_htmx_request():
        return admin_planner()
    return redirect(url_for("web.admin_planner"))


@web_bp.route("/admin/planner/batches/<int:batch_id>/export/groups.pdf")
def admin_planner_export_groups_pdf(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    groups = _load_batch_groups(batch)
    slug = batch.name.lower().replace(" ", "-")
    filename = f"{slug}-groups.pdf"
    return _render_pdf("admin/planner/pdf/group_packet.html", filename, batch=batch, groups=groups)


@web_bp.route("/admin/planner/batches/<int:batch_id>/export/teachers.pdf")
def admin_planner_export_teachers_pdf(batch_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    if not batch:
        abort(404)
    groups = _load_batch_groups(batch)
    schedules = _build_teacher_schedules(groups)
    slug = batch.name.lower().replace(" ", "-")
    filename = f"{slug}-teachers.pdf"
    return _render_pdf(
        "admin/planner/pdf/teacher_packet.html",
        filename,
        batch=batch,
        schedules=schedules,
    )


@web_bp.route("/admin/planner/batches/<int:batch_id>/teachers/<int:teacher_id>/reader-packet.pdf")
def admin_teacher_reader_packet(batch_id: int, teacher_id: int):
    batch = db.session.get(AssignmentBatch, batch_id)
    teacher = db.session.get(Teacher, teacher_id)
    if not batch or not teacher:
        abort(404)

    assignments = _teacher_stage_one_assignments_for_batch(batch, teacher)
    if not assignments:
        flash("No Stage 1 assignments found for this teacher in the selected batch.", "warning")
        return redirect(url_for("web.admin_planner_batch", batch_id=batch.id))

    output, missing_files = _build_teacher_reader_packet(teacher, assignments)
    if missing_files:
        flash(
            "Some essays are missing uploads: " + ", ".join(str(sid) for sid in sorted(set(missing_files))),
            "warning",
        )

    filename_name = secure_filename(f"{teacher.last_name}-{teacher.first_name}") or f"teacher-{teacher.id}"
    filename = f"{filename_name}-stage1-packet.pdf"
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name=filename)


@web_bp.route("/admin/planner/groups/<int:group_id>/reader-packet.pdf")
def admin_group_reader_packet(group_id: int):
    group = (
        ReadingGroup.query.options(
            joinedload(ReadingGroup.members).joinedload(ReadingGroupMember.teacher),
            joinedload(ReadingGroup.assignments).joinedload(EssayAssignment.student),
        )
        .filter_by(id=group_id)
        .first()
    )
    if not group:
        abort(404)

    active_batch = _get_active_batch(group.stage)
    if not active_batch or group.batch_id != active_batch.id:
        flash("This group is not part of the active batch.", "error")
        return redirect(url_for("web.admin_planner_batch", batch_id=group.batch_id))

    assignments = [assignment for assignment in group.assignments if assignment.stage == Stage.STAGE_1]
    if not assignments:
        flash("No Stage 1 assignments found for this group.", "error")
        return redirect(url_for("web.admin_planner_batch", batch_id=group.batch_id))

    writer = PdfWriter()
    double_sided = _packets_double_sided()
    missing_files: list[int] = []

    for member in group.members:
        teacher = member.teacher
        if not member.active or not teacher or teacher.is_absent:
            continue

        essays = []
        for assignment in assignments:
            student = assignment.student
            if not student:
                continue
            pdf_path = _essay_pdf_path(student)
            if not pdf_path:
                missing_files.append(student.id)
                continue
            essays.append((assignment, student, pdf_path))

        title_reader = _create_teacher_title_page(teacher, group, [student.id for _, student, _ in essays])
        writer.add_page(title_reader.pages[0])

        if double_sided and essays:
            _add_blank_page(writer, title_reader.pages[0])

        for assignment, student, pdf_path in essays:
            try:
                reader = PdfReader(str(pdf_path))
            except Exception:
                missing_files.append(student.id)
                continue
            overlay_reader = _create_essay_overlay_page(reader.pages[0], group, assignment)
            first_page = reader.pages[0]
            first_page.merge_page(overlay_reader.pages[0])
            writer.add_page(first_page)
            for page in reader.pages[1:]:
                writer.add_page(page)

            if double_sided and len(reader.pages) % 2 != 0:
                _add_blank_page(writer, reader.pages[0])

    if not writer.pages:
        flash("No essays with uploads available for this group.", "error")
        return redirect(url_for("web.admin_planner_batch", batch_id=group.batch_id))

    output = BytesIO()
    writer.write(output)
    output.seek(0)

    if missing_files:
        flash(
            "Some essays are missing uploads: " + ", ".join(str(sid) for sid in sorted(set(missing_files))),
            "warning",
        )

    filename = f"group-{group.id}-reader-packet.pdf"
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name=filename)




@web_bp.route('/admin/stage2/<int:assignment_id>/packet.pdf')
def admin_stage2_packet(assignment_id: int):
    assignment = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.student),
            joinedload(EssayAssignment.extra_readers).joinedload(EssayAssignmentExtraReader.teacher),
        )
        .filter_by(id=assignment_id)
        .first()
    )
    if not assignment or not assignment.student:
        abort(404)

    if not assignment.stage_two_candidate:
        flash('Essay is not marked for Stage 2.', 'error')
        return redirect(url_for('web.admin_stage2_overview'))

    group = assignment.group or SimpleNamespace(name='Stage 1 group', id='—')
    assessors = _stage_assessors(assignment, Stage.STAGE_2, include_group=False)
    if not assessors:
        flash('No Stage 2 readers assigned.', 'warning')
        return redirect(url_for('web.admin_stage2_score_student', student_id=assignment.student_id))

    pdf_path = _essay_pdf_path(assignment.student)
    if not pdf_path:
        flash('Essay file missing.', 'error')
        return redirect(url_for('web.admin_stage2_score_student', student_id=assignment.student_id))

    writer = PdfWriter()
    double_sided = _packets_double_sided()

    for assessor in assessors:
        teacher = assessor.teacher
        title_reader = _create_teacher_title_page(teacher, assignment.group, [assignment.student.id] if assignment.student else [])
        writer.add_page(title_reader.pages[0])
        if double_sided:
            _add_blank_page(writer, title_reader.pages[0])

        reader = PdfReader(str(pdf_path))
        overlay_reader = _create_essay_overlay_page(reader.pages[0], assignment.group, assignment)
        first_page = reader.pages[0]
        first_page.merge_page(overlay_reader.pages[0])
        writer.add_page(first_page)
        for page in reader.pages[1:]:
            writer.add_page(page)

        if double_sided and len(reader.pages) % 2 != 0:
            _add_blank_page(writer, reader.pages[0])

    output = BytesIO()
    writer.write(output)
    output.seek(0)
    filename = f'stage2-packet-assignment-{assignment.id}.pdf'
    return send_file(output, mimetype='application/pdf', as_attachment=True, download_name=filename)


@web_bp.route("/admin/stage3/packets", methods=["GET", "POST"])
def admin_stage3_packets():
    if not _current_admin():
        abort(403)

    assignments = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.student),
            joinedload(EssayAssignment.group),
        )
        .filter(EssayAssignment.stage_three_recommended.is_(True))
        .order_by(EssayAssignment.student_id.asc())
        .all()
    )

    assignment_choices: list[tuple[int, str]] = []
    assignment_lookup: dict[int, EssayAssignment] = {}
    for assignment in assignments:
        student = assignment.student
        if not student:
            continue
        label = f"Essay {student.id} – {student.first_name} {student.last_name} ({student.country})"
        assignment_choices.append((assignment.id, label))
        assignment_lookup[assignment.id] = assignment
    if not assignment_choices:
        flash(
            "No essays are currently marked for Stage 3. Mark recommendations in Stage 2 scoring before creating packets.",
            "warning",
        )

    form = Stage3PacketForm()
    form.essay_ids.choices = assignment_choices
    if request.method == "GET":
        form.essay_ids.data = [assignment_id for assignment_id, _ in assignment_choices]

    if form.validate_on_submit():
        raw_names = form.member_names.data.replace(",", "\n")
        member_names = [name.strip() for name in raw_names.splitlines() if name.strip()]
        if not member_names:
            form.member_names.errors.append("Provide at least one committee member.")
        else:
            selected_ids = [int(essay_id) for essay_id in form.essay_ids.data or []]
            selected_assignments = [assignment_lookup.get(essay_id) for essay_id in selected_ids]
            selected_assignments = [assignment for assignment in selected_assignments if assignment and assignment.student]
            if not selected_assignments:
                form.essay_ids.errors.append("Select at least one essay to include.")
            else:
                committee_label = form.committee_label.data.strip() or "Steering Committee"
                writer = PdfWriter()
                double_sided = _packets_double_sided()
                missing_files: list[int] = []
                essay_ids_for_table = [
                    assignment.student.id
                    for assignment in selected_assignments
                    if assignment.student
                ]

                for member_name in member_names:
                    title_reader = _create_committee_title_page(member_name, committee_label, essay_ids_for_table)
                    writer.add_page(title_reader.pages[0])
                    if double_sided and selected_assignments:
                        _add_blank_page(writer, title_reader.pages[0])

                    for assignment in selected_assignments:
                        student = assignment.student
                        if not student:
                            continue
                        pdf_path = _essay_pdf_path(student)
                        if not pdf_path:
                            missing_files.append(student.id)
                            continue
                        try:
                            reader = PdfReader(str(pdf_path))
                        except Exception:
                            missing_files.append(student.id)
                            continue

                        overlay_labels = [
                            ("Committee", committee_label),
                            ("Essay", student.id),
                            ("Comments", ""),
                        ]
                        overlay_reader = _create_essay_overlay_page(
                            reader.pages[0],
                            None,
                            assignment,
                            header_text="IPO Stage 3 Review",
                            labels=overlay_labels,
                        )
                        first_page = reader.pages[0]
                        first_page.merge_page(overlay_reader.pages[0])
                        writer.add_page(first_page)
                        for page in reader.pages[1:]:
                            writer.add_page(page)

                        if double_sided and len(reader.pages) % 2 != 0:
                            _add_blank_page(writer, reader.pages[0])

                if not writer.pages:
                    flash("No packets could be generated. Verify that the selected essays have uploads.", "error")
                else:
                    output = BytesIO()
                    writer.write(output)
                    output.seek(0)
                    filename_label = committee_label.lower().replace(" ", "-")
                    filename = f"stage3-{filename_label}-packets.pdf"
                    log_activity(
                        "stage3_packets_generated",
                        f"Stage 3 packets generated for {len(member_names)} committee member(s).",
                        details={
                            "committee": committee_label,
                            "members": member_names,
                            "essay_ids": essay_ids_for_table,
                        },
                    )
                    if missing_files:
                        flash(
                            "Some essays are missing uploads: " + ", ".join(str(sid) for sid in sorted(set(missing_files))),
                            "warning",
                        )
                    return send_file(
                        output,
                        mimetype="application/pdf",
                        as_attachment=True,
                        download_name=filename,
                    )

    return render_template(
        "admin/stage2/stage3_packets.html",
        form=form,
        assignments=assignments,
    )


@web_bp.route("/admin/grades-export", methods=["GET", "POST"])
def admin_grades_export():
    if not _current_admin():
        abort(403)

    form = GradesExportForm()
    country_rows = (
        db.session.query(Student.country)
        .filter(Student.country.isnot(None))
        .distinct()
        .order_by(Student.country.asc())
        .all()
    )
    form.countries.choices = [(country, country) for (country,) in country_rows if country]

    if not form.is_submitted():
        form.stages.data = ["stage1", "stage2"]
        form.anonymize_teachers.data = True
        form.include_timestamps.data = True

    if form.validate_on_submit():
        country_filter = {value for value in form.countries.data if value}
        if not country_filter:
            country_filter = None

        stage_codes = set(form.stages.data or [])
        stage3_only = "stage3" in stage_codes
        stage_codes.discard("stage3")

        selected_stage_enums: set[Stage] = set()
        if "stage1" in stage_codes:
            selected_stage_enums.add(Stage.STAGE_1)
        if "stage2" in stage_codes:
            selected_stage_enums.add(Stage.STAGE_2)
        if not selected_stage_enums:
            selected_stage_enums = {Stage.STAGE_1, Stage.STAGE_2}
        if stage3_only and not stage_codes:
            selected_stage_enums = {Stage.STAGE_1, Stage.STAGE_2}

        rows = _grade_export_rows(
            selected_stages=selected_stage_enums,
            stage3_only=stage3_only,
            countries=country_filter,
            include_absent=bool(form.include_absent.data),
            finalized_only=bool(form.finalized_only.data),
        )
        if not rows:
            flash("No scores found for the selected filters.", "warning")
        else:
            students = _summarize_grade_rows(rows)
            if not students:
                flash("No scores found for the selected filters.", "warning")
            else:
                stage_order = [stage.value for stage in sorted(selected_stage_enums, key=lambda item: item.value)]
                sort_key = form.sort_by.data or "student_id"
                students.sort(key=lambda entry: _grades_sort_key(entry, sort_key), reverse=sort_key != "student_id")
                group_by_country = (
                    bool(form.group_by_country.data)
                    and form.export_type.data == "xlsx"
                    and bool(form.include_countries.data)
                )
                options = {
                    "anonymize_students": bool(form.anonymize_students.data),
                    "anonymize_teachers": bool(form.anonymize_teachers.data),
                    "include_countries": bool(form.include_countries.data),
                    "include_timestamps": bool(form.include_timestamps.data),
                    "group_by_country": group_by_country,
                    "compact_view": bool(form.compact_view.data),
                    "stage_order": stage_order,
                    "stage_slots": GRADE_STAGE_SLOTS,
                }
                if form.export_type.data == "xlsx":
                    return _grades_export_xlsx(students, options)
                return _grades_export_pdf(students, options)

    return render_template("admin/grades/export.html", form=form)
@web_bp.route("/admin/uploads/<int:assignment_id>/extra-reader-packet.pdf")
def admin_extra_reader_packet(assignment_id: int):
    assignment = (
        EssayAssignment.query.options(
            joinedload(EssayAssignment.group).joinedload(ReadingGroup.members).joinedload(ReadingGroupMember.teacher),
            joinedload(EssayAssignment.student),
            joinedload(EssayAssignment.extra_readers).joinedload(EssayAssignmentExtraReader.teacher),
        )
        .filter_by(id=assignment_id)
        .first()
    )
    if not assignment or not assignment.student:
        abort(404)

    group = assignment.group
    if not group:
        flash("Assignment is not linked to a group.", "error")
        return redirect(url_for("web.admin_uploads"))

    assessors = [entry for entry in _stage_assessors(assignment, Stage.STAGE_1) if entry.is_extra]
    if not assessors:
        flash("No extra readers assigned to this essay.", "warning")
        return redirect(url_for("web.admin_score_student", student_id=assignment.student_id))

    student = assignment.student
    pdf_path = _essay_pdf_path(student)
    if not pdf_path:
        flash("The student's essay has not been uploaded.", "error")
        return redirect(url_for("web.admin_score_student", student_id=assignment.student_id))

    try:
        essay_reader = PdfReader(str(pdf_path))
    except Exception:
        flash("Unable to open essay PDF.", "error")
        return redirect(url_for("web.admin_score_student", student_id=assignment.student_id))

    writer = PdfWriter()
    double_sided = _packets_double_sided()

    for assessor in assessors:
        teacher = assessor.teacher
        title_reader = _create_teacher_title_page(teacher, group, [student.id])
        writer.add_page(title_reader.pages[0])
        if double_sided:
            _add_blank_page(writer, title_reader.pages[0])

        essay_reader = PdfReader(str(pdf_path))
        overlay_reader = _create_essay_overlay_page(essay_reader.pages[0], group, assignment)
        first_page = essay_reader.pages[0]
        first_page.merge_page(overlay_reader.pages[0])
        writer.add_page(first_page)
        for page in essay_reader.pages[1:]:
            writer.add_page(page)

        if double_sided and len(essay_reader.pages) % 2 != 0:
            _add_blank_page(writer, essay_reader.pages[0])

    output = BytesIO()
    writer.write(output)
    output.seek(0)
    filename = f"extra-reader-assignment-{assignment.id}.pdf"
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name=filename)


@web_bp.route('/admin/stage2')
def admin_stage2_overview():
    active_batch = _get_active_batch(Stage.STAGE_1)
    if not active_batch:
        flash('Publish a Stage 1 batch before proceeding to Stage 2.', 'error')
        return redirect(url_for('web.admin_planner'))

    assignments = (
        EssayAssignment.query.options(joinedload(EssayAssignment.student))
        .filter(
            EssayAssignment.stage == Stage.STAGE_1,
            EssayAssignment.stage_two_candidate.is_(True),
            EssayAssignment.batch_id == active_batch.id,
        )
        .all()
    )
    candidates = [a for a in assignments if a.student and not a.student.is_absent]

    stage1_status_map: dict[int, SimpleNamespace] = {}
    stage2_status_map: dict[int, SimpleNamespace] = {}
    for assignment in candidates:
        stage1_status_map[assignment.student_id] = _compute_stage_one_status(assignment)
        stage2_status_map[assignment.student_id] = _compute_stage_two_status(assignment)

    is_htmx = _is_htmx_request()
    template_name = "admin/partials/stage2_overview_content.html" if is_htmx else "admin/stage2/overview.html"
    context = {
        "assignments": candidates,
        "stage1_status_map": stage1_status_map,
        "stage2_status_map": stage2_status_map,
        "is_htmx": is_htmx,
    }
    if not is_htmx:
        return render_template(template_name, **context)
    return _render_htmx_template(template_name, **context)


@web_bp.route('/admin/stage2/<int:student_id>', methods=['GET', 'POST'])
def admin_stage2_score_student(student_id: int):
    student = db.session.get(Student, student_id)
    if not student or student.is_absent:
        abort(404)

    assignment = _get_stage_one_assignment(student)
    if not assignment or not assignment.stage_two_candidate:
        flash('Student is not marked for Stage 2.', 'error')
        return redirect(url_for('web.admin_stage2_overview'))

    is_htmx = _is_htmx_request()

    def _load_stage2_state() -> SimpleNamespace:
        stage1_assessors = _stage_assessors(assignment, Stage.STAGE_1, include_group=True)
        stage1_scores = EssayScore.query.filter_by(
            assignment_id=assignment.id,
            stage=Stage.STAGE_1,
        ).all()
        stage1_scores_map = {score.teacher_id: score for score in stage1_scores}
        db.session.expire(assignment, ["scores"])
        stage1_status = _compute_stage_one_status(assignment)
        stage2_assessors = _stage_assessors(assignment, Stage.STAGE_2, include_group=False)
        stage2_scores = EssayScore.query.filter_by(
            assignment_id=assignment.id,
            stage=Stage.STAGE_2,
        ).all()
        stage2_scores_map = {score.teacher_id: score for score in stage2_scores}

        pdf_available = _essay_pdf_path(student) is not None
        stage1_global_lock = _is_stage_locked(Stage.STAGE_1)
        stage2_global_lock = _is_stage_locked(Stage.STAGE_2)
        assignment_stage1_locked = assignment.stage1_locked or stage1_global_lock
        assignment_stage2_locked = assignment.stage2_locked or stage2_global_lock

        stage1_entries = []
        for assessor in stage1_assessors:
            teacher = assessor.teacher
            score_record = stage1_scores_map.get(teacher.id)
            status_info = _resolve_portal_status(student, score_record, pdf_available, locked=assignment_stage1_locked)
            stage1_entries.append(
                SimpleNamespace(
                    teacher=teacher,
                    teacher_id=teacher.id,
                    source=assessor.source,
                    is_extra=assessor.is_extra,
                    score=score_record.score if score_record else None,
                    score_value=format_score_value(score_record.score) if score_record and score_record.score is not None else "",
                    updated_at=score_record.updated_at if score_record else None,
                    submitted_at=score_record.submitted_at if score_record else None,
                    finalized=score_record.teacher_finalized if score_record else False,
                    finalized_at=score_record.teacher_finalized_at if score_record else None,
                    portal_status_code=status_info.code,
                    portal_status_label=status_info.label,
                    portal_status_css=status_info.css,
                    portal_status_timestamp=status_info.timestamp,
                )
            )

        entries = []
        for assessor in stage2_assessors:
            teacher = assessor.teacher
            score_record = stage2_scores_map.get(teacher.id)
            status_info = _resolve_portal_status(student, score_record, pdf_available, locked=assignment_stage2_locked)
            entries.append(
                SimpleNamespace(
                    teacher=teacher,
                    teacher_id=teacher.id,
                    source=assessor.source,
                    is_extra=True,
                    score=score_record.score if score_record else None,
                    score_value=format_score_value(score_record.score) if score_record and score_record.score is not None else "",
                    updated_at=score_record.updated_at if score_record else None,
                    finalized=score_record.teacher_finalized if score_record else False,
                    finalized_at=score_record.teacher_finalized_at if score_record else None,
                    portal_status_code=status_info.code,
                    portal_status_label=status_info.label,
                    portal_status_css=status_info.css,
                    portal_status_timestamp=status_info.timestamp,
                )
            )

        assigned_teacher_ids = {entry.teacher_id for entry in entries}
        stage1_teacher_ids = {entry.teacher_id for entry in stage1_entries}
        available_query = _assessment_eligible_teachers_query()
        if assigned_teacher_ids or stage1_teacher_ids:
            available_query = available_query.filter(
                ~Teacher.id.in_(assigned_teacher_ids.union(stage1_teacher_ids))
            )
        available_teachers = [
            teacher
            for teacher in available_query.order_by(Teacher.last_name.asc(), Teacher.first_name.asc()).all()
            if _teacher_can_grade(student, teacher, student.essay_language)
        ]

        teacher_stats: dict[int, SimpleNamespace] = {}
        suggested_teachers: list[Teacher] = []
        if available_teachers:
            load_assignments = (
                EssayAssignment.query.options(
                    joinedload(EssayAssignment.scores),
                    joinedload(EssayAssignment.extra_readers).joinedload(EssayAssignmentExtraReader.teacher),
                    joinedload(EssayAssignment.group)
                    .joinedload(ReadingGroup.members)
                    .joinedload(ReadingGroupMember.teacher),
                )
                .filter(
                    EssayAssignment.stage == Stage.STAGE_1,
                    EssayAssignment.batch_id == assignment.batch_id,
                )
                .all()
            )
            load_snapshot = _teacher_load_snapshot(load_assignments)
            for teacher in available_teachers:
                entry = load_snapshot.get(teacher.id, _empty_load_entry())
                stage1_pending = max(entry.stage1_expected - entry.stage1_completed, 0)
                stage2_pending = max(entry.stage2_assigned - entry.stage2_completed, 0)
                teacher_stats[teacher.id] = SimpleNamespace(
                    stage1_expected=entry.stage1_expected,
                    stage1_completed=entry.stage1_completed,
                    stage1_pending=stage1_pending,
                    stage2_assigned=entry.stage2_assigned,
                    stage2_completed=entry.stage2_completed,
                    stage2_pending=stage2_pending,
                    last_stage2_activity=entry.last_stage2_activity,
                )

            if teacher_stats:
                def _suggestion_key(t: Teacher):
                    stats_entry = teacher_stats[t.id]
                    last_activity = stats_entry.last_stage2_activity or datetime.min
                    return (
                        1 if stats_entry.stage1_pending > 0 else 0,
                        stats_entry.stage1_pending,
                        stats_entry.stage2_pending,
                        stats_entry.stage2_assigned,
                        last_activity,
                        t.last_name,
                        t.first_name,
                    )

                suggested_teachers = sorted(available_teachers, key=_suggestion_key)[:5]

        stage2_status = _compute_stage_two_status(assignment)
        combined_status = _compute_combined_status(assignment)
        return SimpleNamespace(
            stage1_scores_map=stage1_scores_map,
            stage2_scores_map=stage2_scores_map,
            stage1_status=stage1_status,
            stage1_entries=stage1_entries,
            entries=entries,
            available_teachers=available_teachers,
            teacher_stats=teacher_stats,
            suggested_teachers=suggested_teachers,
            assigned_teacher_ids=assigned_teacher_ids,
            stage1_teacher_ids=stage1_teacher_ids,
            stage1_global_lock=stage1_global_lock,
            stage2_global_lock=stage2_global_lock,
            stage2_status=stage2_status,
            combined_status=combined_status,
        )

    def _render_stage2_response():
        state = _load_stage2_state()
        template_name = "admin/partials/stage2_scoring_content.html" if is_htmx else "admin/stage2/scoring.html"
        context = {
            "student": student,
            "assignment": assignment,
            "stage1_status": state.stage1_status,
            "stage1_entries": state.stage1_entries,
            "stage2_status": state.stage2_status,
            "combined_status": state.combined_status,
            "entries": state.entries,
            "available_teachers": state.available_teachers,
            "teacher_stats": state.teacher_stats,
            "suggested_teachers": state.suggested_teachers,
            "stage1_global_lock": state.stage1_global_lock,
            "stage2_global_lock": state.stage2_global_lock,
            "score_choices": score_choices(include_blank=False),
            "is_htmx": is_htmx,
        }
        if not is_htmx:
            context.update(_quick_message_context())
            return render_template(template_name, **context)
        return _render_htmx_template(template_name, **context)

    def _post_response():
        if is_htmx:
            return _render_stage2_response()
        return redirect(url_for('web.admin_stage2_score_student', student_id=student.id))

    state = _load_stage2_state()

    action = request.form.get('action')
    remove_id = request.form.get('remove_reader')
    unlock_id = request.form.get('unlock_reader')

    if request.method == 'POST' and action == 'toggle_stage2_lock':
        assignment.stage2_locked = not assignment.stage2_locked
        now = datetime.utcnow()
        for score in EssayScore.query.filter_by(
            assignment_id=assignment.id,
            stage=Stage.STAGE_2,
        ).all():
            score.locked_by_admin = assignment.stage2_locked
            if assignment.stage2_locked:
                score.admin_updated_at = now
        log_activity(
            'stage2_grading_lock',
            f'Stage 2 grading lock set to {assignment.stage2_locked} for student {student.id}',
            object_type='student',
            object_id=student.id,
            details={'locked': assignment.stage2_locked},
        )
        db.session.commit()
        flash('Stage 2 grading locked.' if assignment.stage2_locked else 'Stage 2 grading unlocked.', 'success')
        return _post_response()

    if request.method == 'POST' and unlock_id:
        try:
            teacher_id = int(unlock_id)
        except ValueError:
            flash('Invalid teacher selection.', 'error')
            return _post_response()
        score_record = state.stage2_scores_map.get(teacher_id)
        if not score_record or not score_record.teacher_finalized:
            flash('Score is not finalized for this teacher.', 'warning')
            return _post_response()
        score_record.teacher_finalized = False
        score_record.teacher_finalized_at = None
        score_record.admin_updated_at = datetime.utcnow()
        log_activity(
            'stage2_score_unlocked',
            f'Unlocked Stage 2 score for teacher {teacher_id} on student {student.id}',
            object_type='student',
            object_id=student.id,
            details={'teacher_id': teacher_id},
        )
        db.session.commit()
        flash('Teacher score unlocked for editing.', 'success')
        return _post_response()

    if request.method == 'POST' and remove_id:
        try:
            teacher_id = int(remove_id)
        except ValueError:
            flash('Invalid teacher selection.', 'error')
            return _post_response()
        extra = EssayAssignmentExtraReader.query.filter_by(assignment_id=assignment.id, teacher_id=teacher_id, stage=Stage.STAGE_2).first()
        if not extra:
            flash('Extra reader not found for this essay.', 'error')
            return _post_response()
        score_record = state.stage2_scores_map.get(teacher_id)
        if score_record and score_record.stage == Stage.STAGE_2:
            db.session.delete(score_record)
            state.stage2_scores_map.pop(teacher_id, None)
        db.session.delete(extra)
        log_activity(
            'stage2_extra_reader_removed',
            f'Stage 2 reader removed: {teacher_id}',
            object_type='student',
            object_id=student.id,
            details={'teacher_id': teacher_id},
        )
        db.session.commit()
        flash('Stage 2 reader removed.', 'success')
        return _post_response()

    if request.method == 'POST' and action == 'add_reader':
        teacher_id = request.form.get('teacher_id')
        if not teacher_id:
            flash('Select a teacher to add.', 'error')
            return _post_response()
        try:
            teacher_id_int = int(teacher_id)
        except ValueError:
            flash('Invalid teacher selection.', 'error')
            return _post_response()
        if teacher_id_int in state.assigned_teacher_ids or teacher_id_int in state.stage1_teacher_ids:
            flash('Teacher already assigned to this essay.', 'warning')
            return _post_response()
        teacher = db.session.get(Teacher, teacher_id_int)
        if not teacher or teacher.is_absent:
            flash('Teacher not available.', 'error')
            return _post_response()
        if not _teacher_can_grade(student, teacher, student.essay_language):
            flash('Teacher is not eligible to grade this essay.', 'error')
            return _post_response()
        db.session.add(
            EssayAssignmentExtraReader(
                assignment_id=assignment.id,
                teacher_id=teacher_id_int,
                stage=Stage.STAGE_2,
            )
        )
        log_activity(
            'stage2_extra_reader_added',
            f'Added Stage 2 reader {teacher_id_int} for student {student.id}',
            object_type='student',
            object_id=student.id,
            details={'teacher_id': teacher_id_int},
        )
        db.session.commit()
        flash('Stage 2 reader assigned.', 'success')
        return _post_response()

    if request.method == 'POST' and action == 'update_scores':
        updates = {}
        now = datetime.utcnow()
        for entry in state.entries:
            field_name = f'score_{entry.teacher_id}'
            value = request.form.get(field_name, '').strip()
            if not value:
                continue
            score_value, score_error = parse_score_value(value)
            if score_error or score_value is None:
                flash(
                    f"Invalid score for {entry.teacher.first_name} {entry.teacher.last_name}. "
                    f"{SCORE_VALIDATION_MESSAGE}",
                    'error',
                )
                return _post_response()
            numeric = float(score_value)
            score_record = state.stage2_scores_map.get(entry.teacher_id)
            if not score_record:
                score_record = EssayScore(
                    student_id=student.id,
                    teacher_id=entry.teacher_id,
                    assignment_id=assignment.id,
                    stage=Stage.STAGE_2,
                )
                db.session.add(score_record)
                state.stage2_scores_map[entry.teacher_id] = score_record
            prev_admin_entered = score_record.admin_entered if score_record else False
            prev_admin_override = score_record.admin_override if score_record else False
            had_teacher_input = bool(score_record.submitted_at) or bool(score_record.teacher_finalized)
            existing_score = score_record.score if score_record and score_record.score is not None else None

            if existing_score is not None and math.isclose(existing_score, numeric, abs_tol=1e-9):
                continue

            score_record.score = numeric
            score_record.admin_updated_at = now
            if had_teacher_input and not prev_admin_override and not prev_admin_entered:
                score_record.admin_override = True
                score_record.admin_entered = False
            elif prev_admin_override:
                score_record.admin_override = True
                score_record.admin_entered = False
            else:
                score_record.admin_entered = True
                score_record.admin_override = False

            score_record.teacher_finalized = False
            score_record.teacher_finalized_at = None
            updates[entry.teacher_id] = numeric
        if not updates:
            flash('No score changes detected.', 'info')
            return _post_response()

        log_activity(
            'stage2_scores_updated',
            f'Stage 2 scores updated for student {student.id}',
            object_type='student',
            object_id=student.id,
            details={'updates': updates},
        )
        db.session.commit()
        flash('Stage 2 scores saved.', 'success')
        return _post_response()

    if request.method == 'POST' and action == 'mark_stage3':
        assignment.stage_three_recommended = not assignment.stage_three_recommended
        log_activity(
            'stage2_mark_stage3',
            f'Stage 3 recommendation set to {assignment.stage_three_recommended} for student {student.id}',
            object_type='student',
            object_id=student.id,
        )
        db.session.commit()
        flash(
            'Marked for Stage 3 recommendation.'
            if assignment.stage_three_recommended
            else 'Stage 3 recommendation removed.',
            'success',
        )
        return _post_response()

    return _render_stage2_response()

def _ordered_survey_questions(survey: Survey) -> list[SurveyQuestion]:
    return sorted(survey.questions, key=lambda question: (question.position or 0, question.id))


def _next_question_position(survey: Survey) -> int:
    if not survey.questions:
        return 1
    return max((question.position or 0) for question in survey.questions) + 1


def _normalize_question_positions(survey: Survey) -> None:
    for index, question in enumerate(_ordered_survey_questions(survey), start=1):
        question.position = index


def _parse_survey_options(options_text: str | None, question_type_value: str) -> list[str] | None:
    try:
        question_type = SurveyQuestionType(question_type_value)
    except ValueError:
        question_type = SurveyQuestionType.SHORT_TEXT

    if question_type in {SurveyQuestionType.RADIO, SurveyQuestionType.SELECT}:
        options = [line.strip() for line in (options_text or "").splitlines() if line.strip()]
        return options
    return None


def _create_survey_assignments(
    survey: Survey,
    *,
    teacher_ids: Iterable[int],
    student_ids: Iterable[int],
) -> list[SurveyAssignment]:
    existing_assignments = SurveyAssignment.query.filter_by(survey_id=survey.id).all()
    existing_teacher_ids = {assignment.teacher_id for assignment in existing_assignments if assignment.teacher_id}
    existing_student_ids = {assignment.student_id for assignment in existing_assignments if assignment.student_id}

    created: list[SurveyAssignment] = []
    for teacher_id in teacher_ids:
        if teacher_id in existing_teacher_ids:
            continue
        assignment = SurveyAssignment(
            survey_id=survey.id,
            teacher_id=teacher_id,
            status=SurveyAssignmentStatus.PENDING,
        )
        db.session.add(assignment)
        created.append(assignment)

    for student_id in student_ids:
        if student_id in existing_student_ids:
            continue
        assignment = SurveyAssignment(
            survey_id=survey.id,
            student_id=student_id,
            status=SurveyAssignmentStatus.PENDING,
        )
        db.session.add(assignment)
        created.append(assignment)

    return created


def _survey_accepts_students(survey: Survey) -> bool:
    return survey.audience in (SurveyAudience.STUDENT, SurveyAudience.MIXED)


def _survey_accepts_teachers(survey: Survey) -> bool:
    return survey.audience in (SurveyAudience.TEACHER, SurveyAudience.MIXED)


def _survey_recipient_update_plan(survey: Survey) -> SimpleNamespace:
    assigned_student_ids = {
        student_id
        for (student_id,) in (
            db.session.query(SurveyAssignment.student_id)
            .filter(
                SurveyAssignment.survey_id == survey.id,
                SurveyAssignment.student_id.isnot(None),
            )
            .all()
        )
    }
    assigned_teacher_ids = {
        teacher_id
        for (teacher_id,) in (
            db.session.query(SurveyAssignment.teacher_id)
            .filter(
                SurveyAssignment.survey_id == survey.id,
                SurveyAssignment.teacher_id.isnot(None),
            )
            .all()
        )
    }

    students: list[Student] = []
    teachers: list[Teacher] = []
    if _survey_accepts_students(survey):
        students = Student.query.order_by(Student.last_name.asc(), Student.first_name.asc(), Student.id.asc()).all()
    if _survey_accepts_teachers(survey):
        teachers = Teacher.query.order_by(Teacher.last_name.asc(), Teacher.first_name.asc(), Teacher.id.asc()).all()

    missing_students = [student for student in students if student.id not in assigned_student_ids]
    missing_teachers = [teacher for teacher in teachers if teacher.id not in assigned_teacher_ids]

    return SimpleNamespace(
        accepts_students=_survey_accepts_students(survey),
        accepts_teachers=_survey_accepts_teachers(survey),
        eligible_students=len(students),
        eligible_teachers=len(teachers),
        assigned_students=len(assigned_student_ids),
        assigned_teachers=len(assigned_teacher_ids),
        missing_students=missing_students,
        missing_teachers=missing_teachers,
        missing_student_ids=[student.id for student in missing_students],
        missing_teacher_ids=[teacher.id for teacher in missing_teachers],
        missing_total=len(missing_students) + len(missing_teachers),
    )


def _survey_assignment_state(assignment: SurveyAssignment) -> SimpleNamespace:
    survey = assignment.survey
    now = datetime.utcnow()
    reason = None
    locked = False

    if not survey or survey.status == SurveyStatus.DRAFT:
        locked = True
        reason = "unavailable"
    elif survey.status == SurveyStatus.CLOSED:
        locked = True
        reason = "closed"
    elif survey.open_at and now < survey.open_at:
        locked = True
        reason = "not_open"
    elif survey.close_at and now > survey.close_at:
        locked = True
        reason = "past_due"
    elif assignment.status == SurveyAssignmentStatus.COMPLETED and not survey.allow_edit:
        locked = True
        reason = "completed"

    return SimpleNamespace(
        locked=locked,
        reason=reason,
        can_submit=not locked,
        allow_edit=bool(survey.allow_edit if survey else False),
    )


def _survey_assignments_for_user(
    *,
    teacher: Optional[Teacher] = None,
    student: Optional[Student] = None,
) -> list[SurveyAssignment]:
    if teacher is None and student is None:
        raise ValueError("Either teacher or student must be provided.")
    query = (
        SurveyAssignment.query.options(
            joinedload(SurveyAssignment.survey),
            joinedload(SurveyAssignment.response),
        )
        .join(Survey)
        .filter(Survey.status != SurveyStatus.DRAFT)
    )
    if teacher is not None:
        query = query.filter(SurveyAssignment.teacher_id == teacher.id)
    if student is not None:
        query = query.filter(SurveyAssignment.student_id == student.id)
    order_clauses = [Survey.open_at.asc(), Survey.created_at.desc()]
    return query.order_by(*order_clauses).all()


def _load_survey_assignment(
    assignment_id: int,
    *,
    teacher: Optional[Teacher] = None,
    student: Optional[Student] = None,
) -> Optional[SurveyAssignment]:
    query = SurveyAssignment.query.options(
        joinedload(SurveyAssignment.survey).joinedload(Survey.questions),
        joinedload(SurveyAssignment.response),
    ).filter(SurveyAssignment.id == assignment_id)
    if teacher is not None:
        query = query.filter(SurveyAssignment.teacher_id == teacher.id)
    if student is not None:
        query = query.filter(SurveyAssignment.student_id == student.id)
    return query.first()


def _parse_survey_submission(
    questions: list[SurveyQuestion],
    form_data: dict[str, Any],
) -> tuple[dict[str, Any], dict[int, str]]:
    answers: dict[str, Any] = {}
    errors: dict[int, str] = {}

    for question in questions:
        field_name = f"question_{question.id}"
        raw_value = form_data.get(field_name)
        value: Optional[str] = None

        if question.type in {SurveyQuestionType.SHORT_TEXT, SurveyQuestionType.LONG_TEXT}:
            value = (raw_value or "").strip()
            if not value:
                value = None
        elif question.type == SurveyQuestionType.DATE:
            raw_value = (raw_value or "").strip()
            if raw_value:
                try:
                    datetime.strptime(raw_value, "%Y-%m-%d")
                    value = raw_value
                except ValueError:
                    errors[question.id] = "Enter a valid date (YYYY-MM-DD)."
            else:
                value = None
        elif question.type in {SurveyQuestionType.RADIO, SurveyQuestionType.SELECT}:
            raw_value = (raw_value or "").strip()
            options = question.options or []
            if raw_value:
                if raw_value not in options:
                    errors[question.id] = "Select a valid option."
                else:
                    value = raw_value
            else:
                value = None
        else:
            value = (raw_value or "").strip() or None

        if question.required and value is None and question.id not in errors:
            errors[question.id] = "This question is required."

        answers[str(question.id)] = value

    return answers, errors


def _save_survey_response(
    assignment: SurveyAssignment,
    answers: dict[str, Any],
    *,
    respondent_type: str,
    respondent_id: int,
    actor: Optional[str] = None,
) -> None:
    now = datetime.utcnow()
    existing_response = assignment.response
    is_new_response = existing_response is None

    if existing_response is None:
        existing_response = SurveyResponse(
            survey_id=assignment.survey_id,
            assignment_id=assignment.id,
            response_data=answers,
        )
        assignment.response = existing_response
        db.session.add(existing_response)
    else:
        existing_response.response_data = answers

    assignment.status = SurveyAssignmentStatus.COMPLETED
    assignment.completed_at = now

    db.session.flush()

    log_activity(
        "survey_response_submitted" if is_new_response else "survey_response_updated",
        f"Survey '{assignment.survey.title}' response {'submitted' if is_new_response else 'updated'}",
        actor=actor,
        object_type="survey_assignment",
        object_id=assignment.id,
        details={
            "survey_id": assignment.survey_id,
            "respondent_type": respondent_type,
            "respondent_id": respondent_id,
        },
    )


def _survey_display_data(raw: dict[str, Any] | None) -> dict[str, Any]:
    display: dict[str, Any] = {}
    if not raw:
        return display
    for key, value in raw.items():
        if isinstance(key, str) and key.startswith("question_"):
            display[key] = value
        else:
            display[f"question_{key}"] = value
    return display


def _prepare_survey_assignment_view(
    assignments: list[SurveyAssignment],
    detail_url_builder: Callable[[int], str],
) -> tuple[list[SimpleNamespace], SimpleNamespace]:
    items: list[SimpleNamespace] = []
    total = len(assignments)
    completed = 0
    outstanding = 0
    closed = 0

    for assignment in assignments:
        state = _survey_assignment_state(assignment)
        is_completed = assignment.status == SurveyAssignmentStatus.COMPLETED
        if is_completed:
            completed += 1
        elif state.locked:
            closed += 1
        else:
            outstanding += 1

        response = assignment.response
        items.append(
            SimpleNamespace(
                assignment=assignment,
                survey=assignment.survey,
                state=state,
                is_completed=is_completed,
                submitted_at=assignment.completed_at,
                updated_at=response.updated_at if response else None,
                detail_url=detail_url_builder(assignment.id),
            )
        )

    items.sort(
        key=lambda item: (
            item.is_completed,
            item.state.locked,
            item.survey.close_at or datetime.max,
            item.survey.open_at or datetime.max,
            item.survey.title.lower(),
        )
    )

    summary = SimpleNamespace(
        total=total,
        completed=completed,
        outstanding=outstanding,
        closed=closed,
    )
    return items, summary


def _survey_assignment_rows(survey: Survey) -> list[SimpleNamespace]:
    assignments = (
        SurveyAssignment.query.options(
            joinedload(SurveyAssignment.response),
            joinedload(SurveyAssignment.teacher),
            joinedload(SurveyAssignment.student),
        )
        .filter(SurveyAssignment.survey_id == survey.id)
        .order_by(SurveyAssignment.created_at.asc())
        .all()
    )

    rows: list[SimpleNamespace] = []
    for assignment in assignments:
        respondent_type = "teacher" if assignment.teacher_id else "student"
        if assignment.teacher:
            respondent_name = f"{assignment.teacher.first_name} {assignment.teacher.last_name}"
            respondent_email = assignment.teacher.email
            respondent_country = assignment.teacher.country
        elif assignment.student:
            respondent_name = f"{assignment.student.first_name} {assignment.student.last_name}"
            respondent_email = assignment.student.email
            respondent_country = assignment.student.country
        else:
            respondent_name = "—"
            respondent_email = None
            respondent_country = None

        answers = _survey_display_data(assignment.response.response_data if assignment.response else {})
        state = _survey_assignment_state(assignment)

        rows.append(
            SimpleNamespace(
                assignment=assignment,
                respondent_type=respondent_type,
                respondent_name=respondent_name,
                respondent_email=respondent_email,
                respondent_country=respondent_country,
                answers=answers,
                state=state,
            )
        )

    return rows


def _survey_responses_workbook(survey: Survey, assignments: list[SimpleNamespace], questions: list[SurveyQuestion]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"

    header = [
        "Assignment ID",
        "Respondent type",
        "Respondent name",
        "Email",
        "Country",
        "Status",
        "Submitted at",
        "Updated at",
    ]
    for question in questions:
        header.append(question.prompt)
    ws.append(header)

    for row in assignments:
        assignment = row.assignment
        submitted_at = assignment.completed_at.strftime("%Y-%m-%d %H:%M") if assignment.completed_at else ""
        updated_at = (
            assignment.response.updated_at.strftime("%Y-%m-%d %H:%M")
            if assignment.response and assignment.response.updated_at
            else ""
        )
        excel_row = [
            assignment.id,
            row.respondent_type,
            row.respondent_name,
            row.respondent_email or "",
            row.respondent_country or "",
            assignment.status.value,
            submitted_at,
            updated_at,
        ]
        for question in questions:
            value = row.answers.get(f"question_{question.id}")
            excel_row.append(value if value is not None else "")
        ws.append(excel_row)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _filter_survey_rows(
    rows: list[SimpleNamespace],
    audience_filter: str,
    status_filter: str,
) -> list[SimpleNamespace]:
    filtered: list[SimpleNamespace] = []
    for row in rows:
        if audience_filter == "teacher" and row.respondent_type != "teacher":
            continue
        if audience_filter == "student" and row.respondent_type != "student":
            continue

        if status_filter == "completed" and row.assignment.status != SurveyAssignmentStatus.COMPLETED:
            continue
        if status_filter == "outstanding" and (
            row.assignment.status == SurveyAssignmentStatus.COMPLETED or row.state.locked
        ):
            continue

        filtered.append(row)
    return filtered


def _survey_summary_from_rows(rows: list[SimpleNamespace]) -> SimpleNamespace:
    total = len(rows)
    completed = sum(1 for row in rows if row.assignment.status == SurveyAssignmentStatus.COMPLETED)
    outstanding = sum(
        1
        for row in rows
        if row.assignment.status != SurveyAssignmentStatus.COMPLETED and not row.state.locked
    )
    closed = total - completed - outstanding
    return SimpleNamespace(total=total, completed=completed, outstanding=outstanding, closed=closed)


def _format_datetime(dt: datetime | None) -> str | None:
    if not dt:
        return None
    return dt.strftime("%Y-%m-%d %H:%M UTC")


def _survey_assignment_link(assignment: SurveyAssignment) -> str:
    if assignment.teacher_id:
        return url_for("teacher_portal.survey_detail", assignment_id=assignment.id, _external=True)
    if assignment.student_id:
        return url_for("web.student_survey_detail", assignment_id=assignment.id, _external=True)
    return url_for("web.admin_surveys", _external=True)


def _survey_portal_link(assignment: SurveyAssignment) -> str:
    if assignment.teacher_id:
        return url_for("teacher_portal.surveys", _external=True)
    if assignment.student_id:
        return url_for("web.student_portal", _external=True)
    return url_for("web.index", _external=True)


def _survey_email_context(assignment: SurveyAssignment) -> tuple[str | None, dict]:
    survey = assignment.survey
    if not survey:
        return None, {}

    if assignment.teacher:
        recipient_first_name = assignment.teacher.first_name
        recipient_last_name = assignment.teacher.last_name
        recipient_email = assignment.teacher.email
        respondent_type = "teacher"
    elif assignment.student:
        recipient_first_name = assignment.student.first_name
        recipient_last_name = assignment.student.last_name
        recipient_email = assignment.student.email
        respondent_type = "student"
    else:
        recipient_first_name = ""
        recipient_last_name = ""
        recipient_email = None
        respondent_type = "unknown"

    recipient_name = (recipient_first_name or "").strip()
    if recipient_last_name:
        recipient_name = f"{recipient_name} {recipient_last_name}".strip()
    if not recipient_name:
        recipient_name = "participant"

    context = {
        "recipient_first_name": recipient_first_name,
        "recipient_last_name": recipient_last_name,
        "recipient_name": recipient_name,
        "survey_title": survey.title,
        "survey_description": survey.description or "",
        "survey_deadline": _format_datetime(survey.close_at),
        "survey_open": _format_datetime(survey.open_at),
        "survey_link": _survey_assignment_link(assignment),
        "portal_link": _survey_portal_link(assignment),
        "respondent_type": respondent_type,
    }
    return recipient_email, context


def _send_survey_emails(
    assignments: list[SurveyAssignment],
    subject_key: str,
    body_key: str,
) -> int:
    if not assignments:
        return 0
    if not _auto_email_enabled():
        return -1

    sent = 0
    for assignment in assignments:
        recipient_email, context = _survey_email_context(assignment)
        if not recipient_email:
            continue
        _send_templated_email(recipient_email, subject_key, body_key, context)
        sent += 1
    return sent


@web_bp.route("/admin/students/export")
def export_students():
    students = Student.query.order_by(Student.id.asc()).all()
    workbook = _students_workbook(students)
    return send_file(
        workbook,
        as_attachment=True,
        download_name="ipo-students.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@web_bp.route("/admin/teachers/export")
def export_teachers():
    teachers = Teacher.query.order_by(Teacher.id.asc()).all()
    workbook = _teachers_workbook(teachers)
    return send_file(
        workbook,
        as_attachment=True,
        download_name="ipo-teachers.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def _students_workbook(students: list[Student]) -> io.BytesIO:
    room_sharing_labels = {
        "same_sex": "a room only with persons of this same sex as in official documents",
        "any_sex": "a room with persons of any sex",
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(
        [
            "ID",
            "Country",
            "First name",
            "Last name",
            "Email",
            "School",
            "Sex (official document)",
            "Accommodation",
            "Date of birth",
            "Essay language",
            "Keyboard preference",
            "Meal preference",
            "Special needs",
            "Created at",
            "Updated at",
        ]
    )

    for student in students:
        ws.append(
            [
                student.id,
                student.country,
                student.first_name,
                student.last_name,
                student.email,
                student.school,
                student.sex_document or "",
                room_sharing_labels.get(student.accommodation_room_sharing, ""),
                student.date_of_birth.isoformat() if student.date_of_birth else "",
                student.essay_language,
                student.keyboard_preference,
                student.meal_preference,
                student.special_needs or "",
                student.created_at.strftime("%Y-%m-%d %H:%M") if student.created_at else "",
                student.updated_at.strftime("%Y-%m-%d %H:%M") if student.updated_at else "",
            ]
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _teachers_workbook(teachers: list[Teacher]) -> io.BytesIO:
    room_sharing_labels = {
        "same_sex": "a room only with persons of this same sex as in official documents",
        "any_sex": "a room with persons of any sex",
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Teachers"
    ws.append(
        [
            "ID",
            "Country",
            "First name",
            "Last name",
            "Email",
            "School",
            "Sex (official document)",
            "Accommodation",
            "Languages",
            "Meal preference",
            "T-shirt size",
            "Delegation leader",
            "Special needs",
            "Created at",
            "Updated at",
        ]
    )

    for teacher in teachers:
        ws.append(
            [
                teacher.id,
                teacher.country,
                teacher.first_name,
                teacher.last_name,
                teacher.email,
                teacher.school,
                teacher.sex_document or "",
                room_sharing_labels.get(teacher.accommodation_room_sharing, ""),
                teacher.languages,
                teacher.meal_preference,
                teacher.tshirt_size,
                "Yes" if teacher.delegation_leader else "No",
                teacher.special_needs or "",
                teacher.created_at.strftime("%Y-%m-%d %H:%M") if teacher.created_at else "",
                teacher.updated_at.strftime("%Y-%m-%d %H:%M") if teacher.updated_at else "",
            ]
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
