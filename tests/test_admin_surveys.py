import io
from datetime import date, datetime

from openpyxl import load_workbook

import ipo_app.views as views
from ipo_app.extensions import db
from ipo_app.models import (
    ActivityLog,
    AppSetting,
    Student,
    Survey,
    SurveyAssignment,
    SurveyAssignmentStatus,
    SurveyAudience,
    SurveyQuestion,
    SurveyQuestionType,
    SurveyResponse,
    SurveyStatus,
    Teacher,
)


def _set_setting(key: str, value: str) -> None:
    existing = AppSetting.query.filter_by(key=key).first()
    if existing:
        existing.value = value
    else:
        db.session.add(AppSetting(key=key, value=value))
    db.session.commit()


def _create_student(*, email: str, country: str = "Testland") -> int:
    student = Student(
        country=country,
        first_name="Student",
        last_name=email.split("@")[0],
        email=email,
        school="Test School",
        date_of_birth=date(2000, 1, 1),
        essay_language="English",
        keyboard_preference="QWERTY",
        meal_preference="No preference",
    )
    db.session.add(student)
    db.session.commit()
    return student.id


def _create_teacher(*, email: str, country: str = "Testland") -> int:
    teacher = Teacher(
        country=country,
        first_name="Teacher",
        last_name=email.split("@")[0],
        email=email,
        school="Test School",
        languages="English",
        meal_preference="No preference",
        tshirt_size="M",
    )
    db.session.add(teacher)
    db.session.commit()
    return teacher.id


def _create_draft_student_survey(*, title: str = "Student survey") -> tuple[int, int]:
    survey = Survey(
        title=title,
        description=None,
        status=SurveyStatus.DRAFT,
        audience=SurveyAudience.STUDENT,
        allow_edit=False,
        open_at=None,
        close_at=None,
        created_by="admin",
    )
    db.session.add(survey)
    db.session.flush()
    question = SurveyQuestion(
        survey_id=survey.id,
        prompt="What did you like most?",
        help_text=None,
        type=SurveyQuestionType.SHORT_TEXT,
        required=True,
        options=None,
        position=1,
    )
    db.session.add(question)
    db.session.commit()
    return survey.id, question.id


def test_admin_survey_publish_creates_assignments_and_sets_status(client, csrf_token):
    with client.application.app_context():
        student_ids = [
            _create_student(email="student.one@example.com"),
            _create_student(email="student.two@example.com"),
        ]
        survey_id, _question_id = _create_draft_student_survey()

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    publish_page = client.get(f"/admin/surveys/{survey_id}/publish")
    token = csrf_token(publish_page)

    response = client.post(
        f"/admin/surveys/{survey_id}/publish",
        data={
            "csrf_token": token,
            "send_to_all_students": "y",
            "notify_recipients": "y",
            "submit": "Publish survey",
        },
        follow_redirects=False,
    )

    assert response.status_code == 302
    assert response.headers["Location"] == "/admin/surveys"

    with client.application.app_context():
        survey = db.session.get(Survey, survey_id)
        assert survey is not None
        assert survey.status == SurveyStatus.PUBLISHED
        assert survey.open_at is not None

        assignments = SurveyAssignment.query.filter_by(survey_id=survey_id).all()
        assert len(assignments) == len(student_ids)
        assert {assignment.student_id for assignment in assignments} == set(student_ids)
        assert {assignment.teacher_id for assignment in assignments} == {None}
        assert {assignment.status for assignment in assignments} == {SurveyAssignmentStatus.PENDING}


def test_admin_survey_publish_with_email_enabled_sends_notifications(client, csrf_token, monkeypatch):
    sent_to: list[str] = []

    def _fake_send_email(to_email: str, subject: str, html_content: str, plain_text_content=None):
        sent_to.append(to_email)
        return True, "msg-id", None

    monkeypatch.setattr(views, "_send_email", _fake_send_email)

    client.application.config["SENDGRID_API_KEY"] = "dummy"
    client.application.config["SENDGRID_FROM_EMAIL"] = "noreply@example.com"

    with client.application.app_context():
        _set_setting("email_auto_enabled", "true")
        student_emails = ["survey.one@example.com", "survey.two@example.com"]
        student_ids = [_create_student(email=email) for email in student_emails]
        survey_id, _question_id = _create_draft_student_survey(title="Email survey")

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    publish_page = client.get(f"/admin/surveys/{survey_id}/publish")
    token = csrf_token(publish_page)

    response = client.post(
        f"/admin/surveys/{survey_id}/publish",
        data={
            "csrf_token": token,
            "send_to_all_students": "y",
            "notify_recipients": "y",
            "submit": "Publish survey",
        },
        follow_redirects=False,
    )

    assert response.status_code == 302
    assert response.headers["Location"] == "/admin/surveys"
    assert set(sent_to) == set(student_emails)

    with client.application.app_context():
        assignments = SurveyAssignment.query.filter_by(survey_id=survey_id).all()
        assert {assignment.student_id for assignment in assignments} == set(student_ids)

        log = (
            ActivityLog.query
            .filter_by(action="survey_notification_sent", object_type="survey", object_id=str(survey_id))
            .first()
        )
        assert log is not None


def test_admin_update_survey_recipients_adds_late_student_without_duplicates(client, csrf_token):
    with client.application.app_context():
        existing_student_id = _create_student(email="survey.existing@example.com")
        late_student_id = _create_student(email="survey.late@example.com")
        survey_id, _question_id = _create_draft_student_survey(title="Late student survey")
        survey = db.session.get(Survey, survey_id)
        survey.status = SurveyStatus.PUBLISHED
        survey.open_at = datetime.utcnow()
        db.session.add(
            SurveyAssignment(
                survey_id=survey_id,
                student_id=existing_student_id,
                status=SurveyAssignmentStatus.PENDING,
            )
        )
        db.session.commit()

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    update_page = client.get(f"/admin/surveys/{survey_id}/recipients/update")
    assert update_page.status_code == 200
    html = update_page.get_data(as_text=True)
    assert "1 missing" in html
    assert "survey.late@example.com" in html
    token = csrf_token(update_page)

    response = client.post(
        f"/admin/surveys/{survey_id}/recipients/update",
        data={
            "csrf_token": token,
            "submit": "Add missing recipients",
        },
        follow_redirects=False,
    )

    assert response.status_code == 302
    assert response.headers["Location"] == f"/admin/surveys/{survey_id}/responses"

    with client.application.app_context():
        assignments = SurveyAssignment.query.filter_by(survey_id=survey_id).all()
        assert len(assignments) == 2
        assert {assignment.student_id for assignment in assignments} == {existing_student_id, late_student_id}
        assert {assignment.teacher_id for assignment in assignments} == {None}

        log = (
            ActivityLog.query
            .filter_by(action="survey_recipients_updated", object_type="survey", object_id=str(survey_id))
            .first()
        )
        assert log is not None
        assert log.details["new_assignments"] == 1
        assert log.details["students"] == 1
        assert log.details["teachers"] == 0

    with client.session_transaction() as session:
        session.clear()
        session["student_upload_id"] = late_student_id

    surveys_page = client.get("/participant/surveys")
    assert surveys_page.status_code == 200
    assert "Late student survey" in surveys_page.get_data(as_text=True)


def test_admin_update_survey_recipients_adds_missing_mixed_audience(client, csrf_token):
    with client.application.app_context():
        existing_student_id = _create_student(email="mixed.student.existing@example.com")
        late_student_id = _create_student(email="mixed.student.late@example.com")
        existing_teacher_id = _create_teacher(email="mixed.teacher.existing@example.com")
        late_teacher_id = _create_teacher(email="mixed.teacher.late@example.com")
        survey_id, _question_id = _create_draft_student_survey(title="Mixed audience survey")
        survey = db.session.get(Survey, survey_id)
        survey.audience = SurveyAudience.MIXED
        survey.status = SurveyStatus.PUBLISHED
        survey.open_at = datetime.utcnow()
        db.session.add_all(
            [
                SurveyAssignment(
                    survey_id=survey_id,
                    student_id=existing_student_id,
                    status=SurveyAssignmentStatus.PENDING,
                ),
                SurveyAssignment(
                    survey_id=survey_id,
                    teacher_id=existing_teacher_id,
                    status=SurveyAssignmentStatus.PENDING,
                ),
            ]
        )
        db.session.commit()

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    update_page = client.get(f"/admin/surveys/{survey_id}/recipients/update")
    token = csrf_token(update_page)
    response = client.post(
        f"/admin/surveys/{survey_id}/recipients/update",
        data={
            "csrf_token": token,
            "submit": "Add missing recipients",
        },
        follow_redirects=False,
    )

    assert response.status_code == 302

    with client.application.app_context():
        assignments = SurveyAssignment.query.filter_by(survey_id=survey_id).all()
        assert len(assignments) == 4
        assert {assignment.student_id for assignment in assignments if assignment.student_id} == {
            existing_student_id,
            late_student_id,
        }
        assert {assignment.teacher_id for assignment in assignments if assignment.teacher_id} == {
            existing_teacher_id,
            late_teacher_id,
        }


def test_admin_update_survey_recipients_notifies_only_new_assignments(client, csrf_token, monkeypatch):
    sent_to: list[str] = []

    def _fake_send_email(to_email: str, subject: str, html_content: str, plain_text_content=None):
        sent_to.append(to_email)
        return True, "msg-id", None

    monkeypatch.setattr(views, "_send_email", _fake_send_email)
    client.application.config["SENDGRID_API_KEY"] = "dummy"
    client.application.config["SENDGRID_FROM_EMAIL"] = "noreply@example.com"

    with client.application.app_context():
        _set_setting("email_auto_enabled", "true")
        existing_student_id = _create_student(email="notify.existing@example.com")
        late_student_id = _create_student(email="notify.late@example.com")
        survey_id, _question_id = _create_draft_student_survey(title="Notify late student")
        survey = db.session.get(Survey, survey_id)
        survey.status = SurveyStatus.PUBLISHED
        survey.open_at = datetime.utcnow()
        db.session.add(
            SurveyAssignment(
                survey_id=survey_id,
                student_id=existing_student_id,
                status=SurveyAssignmentStatus.PENDING,
            )
        )
        db.session.commit()

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    update_page = client.get(f"/admin/surveys/{survey_id}/recipients/update")
    token = csrf_token(update_page)

    response = client.post(
        f"/admin/surveys/{survey_id}/recipients/update",
        data={
            "csrf_token": token,
            "notify_recipients": "y",
            "submit": "Add missing recipients",
        },
        follow_redirects=False,
    )

    assert response.status_code == 302
    assert sent_to == ["notify.late@example.com"]

    with client.application.app_context():
        assignments = SurveyAssignment.query.filter_by(survey_id=survey_id).all()
        assert len(assignments) == 2
        assert {assignment.student_id for assignment in assignments} == {existing_student_id, late_student_id}

        log = (
            ActivityLog.query
            .filter_by(
                action="survey_recipient_update_notification_sent",
                object_type="survey",
                object_id=str(survey_id),
            )
            .first()
        )
        assert log is not None
        assert log.details["recipients"] == 1


def test_admin_update_survey_recipients_rejects_draft_survey(client):
    with client.application.app_context():
        survey_id, _question_id = _create_draft_student_survey(title="Draft recipient update")

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    response = client.get(f"/admin/surveys/{survey_id}/recipients/update", follow_redirects=False)

    assert response.status_code == 302
    assert response.headers["Location"] == "/admin/surveys"


def test_admin_survey_responses_export_xlsx_contains_question_columns_and_answers(client):
    with client.application.app_context():
        survey_id, question_id = _create_draft_student_survey(title="Export survey")
        student_id = _create_student(email="export.survey@example.com")

        assignment = SurveyAssignment(
            survey_id=survey_id,
            student_id=student_id,
            status=SurveyAssignmentStatus.COMPLETED,
            completed_at=datetime.utcnow(),
        )
        db.session.add(assignment)
        db.session.flush()
        db.session.add(
            SurveyResponse(
                survey_id=survey_id,
                assignment_id=assignment.id,
                response_data={f"question_{question_id}": "Great organization."},
            )
        )
        db.session.commit()

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    response = client.get(f"/admin/surveys/{survey_id}/responses.xlsx?audience=all&status=all", follow_redirects=False)
    assert response.status_code == 200
    assert (
        response.headers.get("Content-Type", "")
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(io.BytesIO(response.data))
    assert "Responses" in workbook.sheetnames
    sheet = workbook["Responses"]

    assert sheet["A1"].value == "Assignment ID"
    assert sheet["B1"].value == "Respondent type"
    assert sheet["I1"].value == "What did you like most?"

    assert sheet["B2"].value == "student"
    assert sheet["I2"].value == "Great organization."


def test_admin_survey_publish_requires_at_least_one_question(client):
    with client.application.app_context():
        survey = Survey(
            title="Empty survey",
            description=None,
            status=SurveyStatus.DRAFT,
            audience=SurveyAudience.STUDENT,
            allow_edit=False,
            open_at=None,
            close_at=None,
            created_by="admin",
        )
        db.session.add(survey)
        db.session.commit()
        survey_id = survey.id

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    response = client.get(f"/admin/surveys/{survey_id}/publish", follow_redirects=False)
    assert response.status_code == 302
    assert response.headers["Location"] == f"/admin/surveys/{survey_id}/edit"

    with client.application.app_context():
        refreshed = db.session.get(Survey, survey_id)
        assert refreshed is not None
        assert refreshed.status == SurveyStatus.DRAFT
        assert SurveyAssignment.query.filter_by(survey_id=survey_id).all() == []
