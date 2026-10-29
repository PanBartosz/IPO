import csv
import io
import zipfile
from datetime import date, datetime
from pathlib import Path
import re
import sqlite3

from openpyxl import load_workbook
import pytest

import ipo_app.views as views
from ipo_app.extensions import db
from ipo_app.models import ActivityLog, Student, Teacher


def test_admin_export_essays_zip_contains_manifest(client, csrf_token):
    with client.application.app_context():
        student = Student(
            country="Testland",
            first_name="Student",
            last_name="Export",
            email="student.export@example.test",
            school="Test School",
            date_of_birth=date(2000, 1, 1),
            essay_language="English",
            keyboard_preference="QWERTY",
            meal_preference="No preference",
        )
        db.session.add(student)
        db.session.commit()

        student.uploaded_filename = f"student-{student.id}-export.pdf"
        student.uploaded_at = datetime.utcnow()
        db.session.commit()

        upload_folder = Path(client.application.config["UPLOAD_FOLDER"])
        (upload_folder / student.uploaded_filename).write_bytes(b"%PDF-1.4\n% export test\n%%EOF\n")

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    exports_page = client.get("/admin/exports")
    token = csrf_token(exports_page)

    response = client.post(
        "/admin/exports/essays",
        data={"csrf_token": token},
        follow_redirects=False,
    )

    assert response.status_code == 200
    assert "application/zip" in response.headers.get("Content-Type", "")

    archive = zipfile.ZipFile(io.BytesIO(response.data))
    names = set(archive.namelist())
    assert "manifest.csv" in names

    expected_name = f"essays/{student.id}-{student.last_name}_{student.first_name}.pdf"
    assert expected_name in names

    manifest = archive.read("manifest.csv").decode("utf-8")
    assert str(student.id) in manifest
    assert student.uploaded_filename in manifest
    with client.application.app_context():
        assert ActivityLog.query.filter_by(action="essays_exported").count() == 1


def test_admin_exports_redirects_to_login_for_non_admin(client):
    response = client.get("/admin/exports", follow_redirects=False)
    assert response.status_code == 302
    assert response.headers["Location"].startswith("/admin/login")


def test_admin_export_essays_finalized_only_filters_files(client, csrf_token):
    with client.application.app_context():
        upload_folder = Path(client.application.config["UPLOAD_FOLDER"])

        finalized = Student(
            country="Testland",
            first_name="Student",
            last_name="Final",
            email="student.final@example.test",
            school="Test School",
            date_of_birth=date(2000, 1, 1),
            essay_language="English",
            keyboard_preference="QWERTY",
            meal_preference="No preference",
            uploaded_filename="final.pdf",
            uploaded_at=datetime.utcnow(),
            essay_finalized=True,
            essay_finalized_at=datetime.utcnow(),
        )
        draft = Student(
            country="Testland",
            first_name="Student",
            last_name="Draft",
            email="student.draft@example.test",
            school="Test School",
            date_of_birth=date(2000, 1, 1),
            essay_language="English",
            keyboard_preference="QWERTY",
            meal_preference="No preference",
            uploaded_filename="draft.pdf",
            uploaded_at=datetime.utcnow(),
            essay_finalized=False,
        )
        db.session.add_all([finalized, draft])
        db.session.commit()

        (upload_folder / finalized.uploaded_filename).write_bytes(b"%PDF-1.4\n% finalized\n%%EOF\n")
        (upload_folder / draft.uploaded_filename).write_bytes(b"%PDF-1.4\n% draft\n%%EOF\n")

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    token = csrf_token(client.get("/admin/exports"))
    response = client.post(
        "/admin/exports/essays",
        data={"csrf_token": token, "finalized_only": "1"},
        follow_redirects=False,
    )

    assert response.status_code == 200
    archive = zipfile.ZipFile(io.BytesIO(response.data))
    names = set(archive.namelist())

    assert f"essays/{finalized.id}-{finalized.last_name}_{finalized.first_name}.pdf" in names
    assert f"essays/{draft.id}-{draft.last_name}_{draft.first_name}.pdf" not in names

    manifest = archive.read("manifest.csv").decode("utf-8")
    reader = csv.DictReader(io.StringIO(manifest))
    exported_ids = [int(row["student_id"]) for row in reader if row.get("student_id")]
    assert exported_ids == [finalized.id]


def test_admin_export_essays_redirects_when_no_essays_available(client, csrf_token):
    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    token = csrf_token(client.get("/admin/exports"))
    response = client.post(
        "/admin/exports/essays",
        data={"csrf_token": token},
        follow_redirects=True,
    )

    assert response.status_code == 200
    assert "No essays available to export." in response.get_data(as_text=True)


def test_admin_export_database_sqlite_creates_backup(client, csrf_token):
    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    token = csrf_token(client.get("/admin/exports"))
    response = client.post(
        "/admin/exports/database",
        data={"csrf_token": token},
        follow_redirects=False,
    )

    assert response.status_code == 200
    assert "application/sql" in response.headers.get("Content-Type", "")

    content_disposition = response.headers.get("Content-Disposition", "")
    match = re.search(r"filename=([^;]+)", content_disposition)
    assert match, "Expected filename in Content-Disposition"
    filename = match.group(1).strip('"')

    backups_folder = Path(client.application.config["BACKUP_FOLDER"])
    assert (backups_folder / filename).exists()
    with client.application.app_context():
        assert ActivityLog.query.filter_by(action="database_exported").count() == 1


def test_sqlite_backup_includes_committed_wal_data(app):
    with app.app_context():
        connection = db.engine.raw_connection()
        try:
            connection.execute("PRAGMA journal_mode=WAL")
            connection.execute("PRAGMA wal_autocheckpoint=0")
            connection.execute("CREATE TABLE wal_backup_probe (value TEXT)")
            connection.commit()
            connection.execute("PRAGMA wal_checkpoint(TRUNCATE)")
            connection.execute("INSERT INTO wal_backup_probe VALUES ('committed in WAL')")
            connection.commit()

            path, _ = views._export_database()
            with sqlite3.connect(path) as backup:
                assert backup.execute("SELECT value FROM wal_backup_probe").fetchall() == [("committed in WAL",)]
        finally:
            connection.close()


def test_sqlite_backup_uses_engine_for_relative_database_path(app):
    with app.app_context():
        # Flask resolves relative SQLite paths against instance_path, not cwd.
        app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///relative.db"
        db.session.add(ActivityLog(action="relative-backup", description="keep this row"))
        db.session.commit()
        path, _ = views._export_database()
        with sqlite3.connect(path) as backup:
            assert backup.execute("SELECT action FROM activity_logs").fetchall() == [("relative-backup",)]


@pytest.mark.parametrize("role", ["student", "teacher"])
def test_registration_workbooks_preserve_formula_like_input_as_text(client, role):
    with client.application.app_context():
        common = dict(
            country="Testland", first_name='=HYPERLINK("https://evil.example","click")', last_name="Export",
            email=f"{role}.formula@example.test", school="=1+1", meal_preference="No preference",
        )
        if role == "student":
            record = Student(**common, date_of_birth=date(2000, 1, 1), essay_language="English", keyboard_preference="QWERTY")
        else:
            record = Teacher(**common, languages="English", tshirt_size="M")
        db.session.add(record)
        db.session.commit()
    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    response = client.get(f"/admin/{role}s/export")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.data)).active
    assert sheet["A2"].data_type == "n"
    for address, expected in [("C2", common["first_name"]), ("F2", common["school"])]:
        assert sheet[address].value == expected
        assert sheet[address].data_type == "s"


def test_essay_archive_sanitizes_paths_and_csv_formula_values(app):
    with app.app_context():
        student = Student(
            country="@SUM(1)", first_name="=1+1", last_name="../../outside",
            email="archive.paths@example.test", school="Test School", date_of_birth=date(2000, 1, 1),
            essay_language="English", keyboard_preference="QWERTY", meal_preference="No preference",
            uploaded_filename="archive.pdf",
        )
        db.session.add(student)
        db.session.commit()
        Path(app.config["UPLOAD_FOLDER"], "archive.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
        stream, _ = views._zip_essays()

    with zipfile.ZipFile(stream) as archive:
        essay_names = [name for name in archive.namelist() if name.endswith(".pdf")]
        assert len(essay_names) == 1
        assert len(Path(essay_names[0]).parts) == 2
        assert Path(essay_names[0]).parts[0] == "essays"
        manifest = list(csv.DictReader(io.StringIO(archive.read("manifest.csv").decode())))
        assert manifest[0]["first_name"] == "'=1+1"
        assert manifest[0]["country"] == "'@SUM(1)"


def test_log_csv_export_neutralizes_formula_values(client):
    with client.application.app_context():
        db.session.add(ActivityLog(action="review", actor="=1+1", description="@SUM(1)"))
        db.session.commit()
    with client.session_transaction() as session:
        session["admin_user"] = "admin"
    response = client.get("/admin/logs/export")
    row = next(csv.DictReader(io.StringIO(response.get_data(as_text=True))))
    assert row["actor"] == "'=1+1"
    assert row["description"] == "'@SUM(1)"


def test_admin_logs_export_csv_respects_limit(client):
    with client.application.app_context():
        db.session.add_all(
            [
                ActivityLog(action="a1", description="one"),
                ActivityLog(action="a2", description="two"),
                ActivityLog(action="a3", description="three"),
            ]
        )
        db.session.commit()

    client.application.config["LOG_EXPORT_LIMIT"] = 2
    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    response = client.get("/admin/logs/export", follow_redirects=False)
    assert response.status_code == 200
    assert response.headers.get("Content-Type") == "text/csv; charset=utf-8"

    lines = response.get_data(as_text=True).strip().splitlines()
    assert lines[0] == "timestamp,actor,action,object_type,object_id,description,details"
    assert len(lines) == 1 + 2


def test_admin_students_export_xlsx_contains_header_row(client):
    with client.application.app_context():
        student = Student(
            country="Testland",
            first_name="Student",
            last_name="Xlsx",
            email="student.xlsx@example.test",
            school="Test School",
            date_of_birth=date(2000, 1, 1),
            essay_language="English",
            keyboard_preference="QWERTY",
            meal_preference="No preference",
        )
        db.session.add(student)
        db.session.commit()

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    response = client.get("/admin/students/export", follow_redirects=False)
    assert response.status_code == 200
    assert (
        response.headers.get("Content-Type", "")
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(io.BytesIO(response.data))
    assert "Students" in workbook.sheetnames
    sheet = workbook["Students"]
    assert sheet["A1"].value == "ID"
    assert sheet["B1"].value == "Country"


def test_admin_teachers_export_xlsx_contains_header_row(client):
    with client.application.app_context():
        teacher = Teacher(
            country="Testland",
            first_name="Teacher",
            last_name="Xlsx",
            email="teacher.xlsx@example.test",
            school="Test School",
            languages="English",
            meal_preference="No preference",
            tshirt_size="M",
        )
        db.session.add(teacher)
        db.session.commit()

    with client.session_transaction() as session:
        session["admin_user"] = "admin"

    response = client.get("/admin/teachers/export", follow_redirects=False)
    assert response.status_code == 200
    assert (
        response.headers.get("Content-Type", "")
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(io.BytesIO(response.data))
    assert "Teachers" in workbook.sheetnames
    sheet = workbook["Teachers"]
    assert sheet["A1"].value == "ID"
    assert sheet["B1"].value == "Country"
